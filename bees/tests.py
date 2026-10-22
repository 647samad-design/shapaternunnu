"""
Automated tests for the marketplace.

These aren't exhaustive - they focus on the business logic that's most
likely to break silently: money calculations (order totals, commission),
permission checks (who can access what), and the bugs found and fixed
during development (display_name, N+1 query annotations).

Run with: python manage.py test bees
"""
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import Client
from django.test import TestCase as DjangoTestCase


class TestCase(DjangoTestCase):
    """Clears the cache before every test: site settings and rate-limit
    counters are cached, and the test database is rolled back between
    tests while the in-memory cache is not."""

    def _pre_setup(self):
        super()._pre_setup()
        cache.clear()
        from django.core.cache import caches
        caches["live"].clear()
        # Never call the real Stripe API from tests.
        from unittest import mock as _mock
        self._stripe_customer = _mock.patch("stripe.Customer.create", return_value=_mock.MagicMock(id="cus_test"))
        self._stripe_customer.start()
        # Older admin tests use staff without an authenticator app; the
        # two-step tests switch this back on.
        from .models import SiteSettings as _SS
        _SS.objects.update_or_create(pk=1, defaults={"require_staff_2fa": False})
        # Older checkout tests treat every buyer as having a confirmed email;
        # the cash-on-delivery confirmation tests switch this back on.
        self._cod_patch = None
        if not getattr(self, "real_cod_confirmation", False):
            self._cod_patch = _mock.patch("bees.views._email_verified", return_value=True)
            self._cod_patch.start()

    def _post_teardown(self):
        self._stripe_customer.stop()
        if self._cod_patch:
            self._cod_patch.stop()
        super()._post_teardown()
from django.urls import reverse

from .models import (
    Product, Order, OrderItem, SellerAccount, OrganizationMember,
    Review, Notification, AuditLog, Claim, ReturnRequest, Shipment,
)
from .views import with_ratings, get_seller_account_for_user


def make_product(**kwargs):
    defaults = dict(
        name="Test Product", category="skincare", price=Decimal("500.00"),
        stock=10, image_url="https://example.com/img.jpg",
    )
    defaults.update(kwargs)
    return Product.objects.create(**defaults)


class ProductRatingTests(TestCase):
    """The average_rating/rating_count properties used to run a fresh query
    per product (N+1). They now use a DB annotation when present, and fall
    back to the old per-query behaviour otherwise. Both paths must agree."""

    def setUp(self):
        self.product = make_product()

    def test_no_reviews_gives_zero(self):
        self.assertEqual(self.product.average_rating, 0)
        self.assertEqual(self.product.rating_count, 0)

    def test_average_rating_matches_manual_calculation(self):
        Review.objects.create(product=self.product, username="a", rating=4, comment="Good")
        Review.objects.create(product=self.product, username="b", rating=5, comment="Great")
        self.product.refresh_from_db()
        self.assertEqual(self.product.average_rating, 4.5)
        self.assertEqual(self.product.rating_count, 2)

    def test_annotated_queryset_matches_property_fallback(self):
        Review.objects.create(product=self.product, username="a", rating=4, comment="Good")
        Review.objects.create(product=self.product, username="b", rating=5, comment="Great")
        annotated = with_ratings(Product.objects.filter(pk=self.product.pk)).first()
        self.assertEqual(annotated.average_rating, 4.5)
        self.assertEqual(annotated.rating_count, 2)


class ProductStockNotificationTests(TestCase):
    """Product.save() notifies the seller once when stock crosses into the
    low-stock/out-of-stock zone. Must fire exactly once per crossing, not
    on every save while it stays low (that would spam the seller)."""

    def setUp(self):
        self.user = User.objects.create_user("seller1", "s1@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(
            user=self.user, account_type="individual", status="approved",
            business_name="Test Shop",
        )
        self.product = make_product(stock=20, seller_account=self.seller)

    def test_low_stock_notification_fires_once(self):
        self.product.stock = 4
        self.product.save()
        self.assertEqual(
            Notification.objects.filter(user=self.user, message__icontains="running low").count(), 1
        )
        self.product.name = self.product.name
        self.product.save()
        self.assertEqual(
            Notification.objects.filter(user=self.user, message__icontains="running low").count(), 1
        )

    def test_out_of_stock_notification(self):
        self.product.stock = 0
        self.product.save()
        self.assertTrue(
            Notification.objects.filter(user=self.user, message__icontains="out of stock").exists()
        )


class OrderTotalTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("buyer1", "b1@example.com", "pass12345")
        self.product = make_product(price=Decimal("250.00"))

    def test_total_sums_items_and_subtracts_discount(self):
        order = Order.objects.create(
            user=self.user, full_name="Buyer", address="St", city="Karachi",
            phone="0300", payment_method="cod", discount_amount=Decimal("50.00"),
        )
        OrderItem.objects.create(order=order, product=self.product, product_name=self.product.name,
                                  price=Decimal("250.00"), quantity=2)
        self.assertEqual(order.total, Decimal("450.00"))

    def test_total_never_goes_negative(self):
        order = Order.objects.create(
            user=self.user, full_name="Buyer", address="St", city="Karachi",
            phone="0300", payment_method="cod", discount_amount=Decimal("999.00"),
        )
        OrderItem.objects.create(order=order, product=self.product, product_name=self.product.name,
                                  price=Decimal("250.00"), quantity=1)
        self.assertEqual(order.total, 0)

    def test_is_cancellable_only_for_pending_or_confirmed(self):
        order = Order.objects.create(
            user=self.user, full_name="Buyer", address="St", city="Karachi",
            phone="0300", payment_method="cod", status="pending",
        )
        self.assertTrue(order.is_cancellable)
        order.status = "shipped"
        self.assertFalse(order.is_cancellable)
        order.status = "delivered"
        self.assertFalse(order.is_cancellable)


class SellerAccountTests(TestCase):
    """display_name used to silently ignore organization_name - every org
    account's name fell back to the raw username everywhere it was shown."""

    def test_display_name_prefers_business_name(self):
        user = User.objects.create_user("u1", "u1@example.com", "pass12345")
        seller = SellerAccount.objects.create(
            user=user, account_type="individual", business_name="Ali's Shop",
        )
        self.assertEqual(seller.display_name, "Ali's Shop")

    def test_display_name_falls_back_to_organization_name(self):
        user = User.objects.create_user("u2", "u2@example.com", "pass12345")
        org = SellerAccount.objects.create(
            user=user, account_type="organization", organization_name="Acme Traders",
        )
        self.assertEqual(org.display_name, "Acme Traders")

    def test_display_name_falls_back_to_username_last(self):
        user = User.objects.create_user("plainuser", "u3@example.com", "pass12345")
        seller = SellerAccount.objects.create(user=user, account_type="individual")
        self.assertEqual(seller.display_name, "plainuser")

    def test_lifetime_sales_aggregate(self):
        user = User.objects.create_user("u4", "u4@example.com", "pass12345")
        seller = SellerAccount.objects.create(user=user, account_type="individual", status="approved")
        product = make_product(price=Decimal("100.00"), seller_account=seller)
        buyer = User.objects.create_user("buyer2", "b2@example.com", "pass12345")
        order = Order.objects.create(user=buyer, full_name="B", address="St", city="Karachi",
                                      phone="0300", payment_method="cod")
        item = OrderItem(order=order, product=product, product_name=product.name,
                         price=Decimal("100.00"), quantity=3)
        item.apply_commission(seller)
        item.save()
        self.assertEqual(seller.lifetime_sales, Decimal("300.00"))
        self.assertEqual(item.commission_amount, Decimal("30.00"))

    def test_commission_rate_reduces_at_volume_thresholds(self):
        user = User.objects.create_user("u5", "u5@example.com", "pass12345")
        seller = SellerAccount.objects.create(
            user=user, account_type="individual", commission_rate=Decimal("10.00"),
        )
        self.assertEqual(seller.effective_commission_rate, 10.0)


class OrganizationTeamAccessTests(TestCase):
    """The core organization feature: team members can act on behalf of the
    org account without the owner's credentials."""

    def setUp(self):
        self.owner = User.objects.create_user("owner1", "o1@example.com", "pass12345")
        self.staff = User.objects.create_user("staff1", "s1@example.com", "pass12345")
        self.outsider = User.objects.create_user("outsider1", "out1@example.com", "pass12345")
        self.org = SellerAccount.objects.create(
            user=self.owner, account_type="organization", status="approved",
            organization_name="Test Org",
        )
        OrganizationMember.objects.create(organization=self.org, user=self.staff, role="staff")

    def test_owner_resolves_to_their_own_account(self):
        account, role = get_seller_account_for_user(self.owner)
        self.assertEqual(account, self.org)
        self.assertEqual(role, "owner")

    def test_team_member_resolves_to_organization_account(self):
        account, role = get_seller_account_for_user(self.staff)
        self.assertEqual(account, self.org)
        self.assertEqual(role, "staff")

    def test_unrelated_user_has_no_seller_account(self):
        account, role = get_seller_account_for_user(self.outsider)
        self.assertIsNone(account)
        self.assertIsNone(role)

    def test_staff_can_load_seller_dashboard(self):
        client = Client()
        client.force_login(self.staff)
        response = client.get(reverse("seller_dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Test Org")

    def test_outsider_is_sent_to_seller_signup(self):
        client = Client()
        client.force_login(self.outsider)
        response = client.get(reverse("seller_dashboard"))
        self.assertRedirects(response, reverse("become_seller"), fetch_redirect_response=False)
        self.assertEqual(client.get(reverse("seller_order", args=[999])).status_code, 302)

    def test_only_owner_can_add_team_members(self):
        client = Client()
        client.force_login(self.staff)
        new_user = User.objects.create_user("newperson", "np@example.com", "pass12345")
        client.post(reverse("add_team_member"), {"username_or_email": "newperson", "role": "staff"})
        self.assertFalse(OrganizationMember.objects.filter(user=new_user).exists())


class CorePageLoadTests(TestCase):
    """Smoke tests: the most-visited pages should always return 200."""

    def setUp(self):
        make_product(name="Homepage Product", stock=5)

    def test_homepage_loads(self):
        self.assertEqual(self.client.get(reverse("home")).status_code, 200)

    def test_all_products_loads(self):
        self.assertEqual(self.client.get(reverse("all_products")).status_code, 200)

    def test_category_page_loads(self):
        self.assertEqual(self.client.get(reverse("category_products", args=["skincare"])).status_code, 200)

    def test_product_detail_loads(self):
        product = Product.objects.first()
        self.assertEqual(self.client.get(reverse("product_detail", args=[product.id])).status_code, 200)


class CartAndCheckoutTests(TestCase):
    def setUp(self):
        self.product = make_product(stock=5)
        self.user = User.objects.create_user("shopper", "sh@example.com", "pass12345")

    def test_add_to_cart_updates_session(self):
        response = self.client.post(
            reverse("add_to_cart", args=[self.product.id]),
            {"quantity": 2},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.session["cart"], {str(self.product.id): 2})

    def test_add_to_cart_cannot_exceed_stock(self):
        self.client.post(
            reverse("add_to_cart", args=[self.product.id]),
            {"quantity": 999},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(self.client.session["cart"][str(self.product.id)], self.product.stock)

    def test_empty_checkout_goes_back_to_cart(self):
        response = self.client.get(reverse("checkout"))
        self.assertRedirects(response, reverse("cart"))

    def test_cancel_order_only_works_for_own_pending_order(self):
        other_user = User.objects.create_user("someone_else", "oe@example.com", "pass12345")
        order = Order.objects.create(
            user=other_user, full_name="Someone Else", address="St", city="Karachi",
            phone="0300", payment_method="cod", status="pending",
        )
        self.client.force_login(self.user)
        response = self.client.post(reverse("cancel_order", args=[order.id]))
        self.assertEqual(response.status_code, 404)
        order.refresh_from_db()
        self.assertEqual(order.status, "pending")


class CouponValidationTests(TestCase):
    """Coupons used to have no expiry, no usage limit, no per-user limit,
    and no minimum order check - any active code worked forever, for
    anyone, any number of times."""

    def setUp(self):
        self.user = User.objects.create_user("shopper1", "sh1@example.com", "pass12345")
        self.other_user = User.objects.create_user("shopper2", "sh2@example.com", "pass12345")
        self.product = make_product(price=Decimal("500.00"))

    def _make_order(self, user, coupon_code, status="confirmed"):
        order = Order.objects.create(
            user=user, full_name="Buyer", address="St", city="Karachi",
            phone="0300", payment_method="cod", coupon_code=coupon_code, status=status,
        )
        OrderItem.objects.create(order=order, product=self.product, product_name=self.product.name,
                                  price=Decimal("500.00"), quantity=1)
        return order

    def test_inactive_coupon_is_invalid(self):
        from .models import Coupon
        coupon = Coupon.objects.create(code="OFF10", percent_off=10, active=False)
        is_valid, error = coupon.is_valid_for(self.user, Decimal("1000"))
        self.assertFalse(is_valid)
        self.assertIn("no longer active", error)

    def test_expired_coupon_is_invalid(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models import Coupon
        coupon = Coupon.objects.create(
            code="OLD10", percent_off=10, expiry_date=timezone.localdate() - timedelta(days=1),
        )
        is_valid, error = coupon.is_valid_for(self.user, Decimal("1000"))
        self.assertFalse(is_valid)
        self.assertIn("expired", error)

    def test_future_expiry_is_still_valid(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models import Coupon
        coupon = Coupon.objects.create(
            code="NEW10", percent_off=10, expiry_date=timezone.localdate() + timedelta(days=1),
        )
        is_valid, error = coupon.is_valid_for(self.user, Decimal("1000"))
        self.assertTrue(is_valid)

    def test_minimum_order_value_enforced(self):
        from .models import Coupon
        coupon = Coupon.objects.create(code="BIG50", percent_off=50, min_order_value=Decimal("2000"))
        is_valid, error = coupon.is_valid_for(self.user, Decimal("500"))
        self.assertFalse(is_valid)
        self.assertIn("minimum order", error)
        is_valid, error = coupon.is_valid_for(self.user, Decimal("2500"))
        self.assertTrue(is_valid)

    def test_global_usage_limit_enforced(self):
        from .models import Coupon
        coupon = Coupon.objects.create(code="LIMITED", percent_off=10, usage_limit=1)
        self._make_order(self.user, "LIMITED")
        # Already used once globally, limit is 1 - a different user should now be blocked too.
        is_valid, error = coupon.is_valid_for(self.other_user, Decimal("1000"))
        self.assertFalse(is_valid)
        self.assertIn("usage limit", error)

    def test_cancelled_orders_dont_count_against_usage_limit(self):
        from .models import Coupon
        coupon = Coupon.objects.create(code="LIMITED2", percent_off=10, usage_limit=1)
        self._make_order(self.user, "LIMITED2", status="cancelled")
        is_valid, error = coupon.is_valid_for(self.other_user, Decimal("1000"))
        self.assertTrue(is_valid)

    def test_per_user_limit_enforced(self):
        from .models import Coupon
        coupon = Coupon.objects.create(code="ONEUSE", percent_off=10, per_user_limit=1)
        self._make_order(self.user, "ONEUSE")
        # This user already used it once - should now be blocked for them...
        is_valid, error = coupon.is_valid_for(self.user, Decimal("1000"))
        self.assertFalse(is_valid)
        self.assertIn("maximum number of times", error)
        # ...but a different user should still be able to use it.
        is_valid, error = coupon.is_valid_for(self.other_user, Decimal("1000"))
        self.assertTrue(is_valid)

    def test_apply_coupon_view_rejects_invalid_code(self):
        self.client.force_login(self.user)
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 1})
        response = self.client.post(reverse("apply_coupon"), {"coupon_code": "DOESNOTEXIST"}, follow=True)
        self.assertContains(response, "coupon code isn")

    def test_apply_coupon_view_accepts_valid_code(self):
        from .models import Coupon
        Coupon.objects.create(code="WORKS10", percent_off=10)
        self.client.force_login(self.user)
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 1})
        response = self.client.post(reverse("apply_coupon"), {"coupon_code": "WORKS10"}, follow=True)
        self.assertContains(response, "Coupon applied")


class ProductFilterTests(TestCase):
    def setUp(self):
        make_product(name="Cheap Item", price=Decimal("100.00"), stock=5, seller_name="SellerA")
        make_product(name="Mid Item", price=Decimal("500.00"), stock=0, seller_name="SellerB")
        make_product(name="Expensive Item", price=Decimal("2000.00"), stock=10, seller_name="SellerA")

    def test_price_range_filter(self):
        response = self.client.get(reverse("all_products"), {"min_price": "200", "max_price": "1000"})
        self.assertContains(response, "Mid Item")
        self.assertNotContains(response, "Cheap Item")
        self.assertNotContains(response, "Expensive Item")

    def test_in_stock_filter_excludes_zero_stock(self):
        response = self.client.get(reverse("all_products"), {"in_stock": "1"})
        self.assertNotContains(response, "Mid Item")
        self.assertContains(response, "Cheap Item")

    def test_seller_filter(self):
        response = self.client.get(reverse("all_products"), {"seller": "SellerB"})
        self.assertContains(response, "Mid Item")
        self.assertNotContains(response, "Cheap Item")

    def test_filters_combine_with_category_page(self):
        response = self.client.get(reverse("category_products", args=["skincare"]), {"min_price": "1000"})
        self.assertContains(response, "Expensive Item")
        self.assertNotContains(response, "Cheap Item")


class ChatTests(TestCase):
    def test_guest_can_send_and_read_chat_messages(self):
        response = self.client.post(reverse("chat_send"), {"message": "Hello, is anyone there?"})
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertIn("reply", data)
        self.assertTrue(data["reply"]["message"])

        history = self.client.get(reverse("chat_messages"))
        messages = history.json()["messages"]
        self.assertEqual(len(messages), 2)
        self.assertEqual(messages[0]["sender"], "user")
        self.assertEqual(messages[0]["message"], "Hello, is anyone there?")
        self.assertEqual(messages[1]["sender"], "support")

    def test_empty_chat_message_rejected(self):
        response = self.client.post(reverse("chat_send"), {"message": "   "})
        self.assertEqual(response.status_code, 400)

    def test_order_keyword_triggers_relevant_auto_reply(self):
        response = self.client.post(reverse("chat_send"), {"message": "where is my order tracking"})
        self.assertIn("My orders", response.json()["reply"]["message"])

    def test_logged_in_user_thread_persists_across_requests(self):
        user = User.objects.create_user(username="chatuser", password="pass12345")
        self.client.force_login(user)
        self.client.post(reverse("chat_send"), {"message": "First message"})
        self.client.post(reverse("chat_send"), {"message": "Second message"})

        from .models import ChatThread
        self.assertEqual(ChatThread.objects.filter(user=user).count(), 1)
        history = self.client.get(reverse("chat_messages")).json()["messages"]
        user_messages = [m for m in history if m["sender"] == "user"]
        self.assertEqual(len(user_messages), 2)


# ---------------------------------------------------------------------------
# Checkout, Stripe and security regression tests
# ---------------------------------------------------------------------------
import json
import os
from unittest import mock

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings

from .models import Address, Coupon, NewsletterSubscriber, ProductImage, SiteSettings


CHECKOUT_FORM = {
    "full_name": "Jane Doe", "email": "jane@example.com", "phone": "+1 555 0100",
    "address": "1 Main St", "city": "Austin", "state": "TX", "postal_code": "73301", "country": "us",
}


class CheckoutTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("jane", "jane@example.com", "pass12345")
        self.product = make_product(price=Decimal("20.00"), stock=3)
        self.client.force_login(self.user)

    def _add(self, qty=1, product=None):
        self.client.post(reverse("add_to_cart", args=[(product or self.product).id]), {"quantity": qty})

    def test_cod_checkout_creates_order_and_decrements_stock(self):
        self._add(2)
        response = self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get(user=self.user)
        self.assertRedirects(response, reverse("order_success", args=[order.id]))
        self.assertEqual(order.country, "US")
        self.assertEqual(order.total, Decimal("40.00"))
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock, 1)
        self.assertEqual(self.client.session["cart"], {})

    def test_checkout_blocks_when_stock_ran_out(self):
        self._add(3)
        Product.objects.filter(pk=self.product.pk).update(stock=1)  # someone else bought it
        response = self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertRedirects(response, reverse("cart"))
        self.assertFalse(Order.objects.exists())
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock, 1)

    def test_missing_address_fields_rejected(self):
        self._add()
        form = {**CHECKOUT_FORM, "postal_code": "", "payment_method": "cod"}
        response = self.client.post(reverse("checkout"), form)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Order.objects.exists())

    def test_shipping_and_tax_applied(self):
        settings_obj = SiteSettings.load()
        settings_obj.shipping_flat_fee = Decimal("5.00")
        settings_obj.tax_percent = Decimal("10")
        settings_obj.save()
        self._add(1)
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get()
        self.assertEqual(order.shipping_amount, Decimal("5.00"))
        self.assertEqual(order.tax_amount, Decimal("2.00"))
        self.assertEqual(order.total, Decimal("27.00"))

    def test_free_shipping_threshold(self):
        settings_obj = SiteSettings.load()
        settings_obj.shipping_flat_fee = Decimal("5.00")
        settings_obj.free_shipping_threshold = Decimal("30.00")
        settings_obj.save()
        self._add(2)
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertEqual(Order.objects.get().shipping_amount, Decimal("0"))

    def test_cod_disabled_hides_method(self):
        settings_obj = SiteSettings.load()
        settings_obj.allow_cash_on_delivery = False
        settings_obj.save()
        self._add()
        response = self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertRedirects(response, reverse("checkout"), fetch_redirect_response=False)
        self.assertFalse(Order.objects.exists())

    def test_coupon_applied_to_order(self):
        Coupon.objects.create(code="TEN", percent_off=10)
        self._add(1)
        self.client.post(reverse("apply_coupon"), {"coupon_code": "ten"})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get()
        self.assertEqual(order.coupon_code, "TEN")
        self.assertEqual(order.discount_amount, Decimal("2.00"))

    def test_negative_quantity_cannot_lower_cart_total(self):
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": -5})
        self.assertEqual(self.client.session["cart"][str(self.product.id)], 1)

    def test_cannot_add_unapproved_product(self):
        hidden = make_product(name="Hidden", approval_status="pending")
        response = self.client.post(reverse("add_to_cart", args=[hidden.id]))
        self.assertEqual(response.status_code, 404)

    def test_cancel_cod_order_restocks(self):
        self._add(2)
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get()
        self.client.post(reverse("cancel_order", args=[order.id]))
        order.refresh_from_db()
        self.product.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.assertEqual(self.product.stock, 3)

    def test_cancel_requires_post(self):
        self._add(1)
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get()
        response = self.client.get(reverse("cancel_order", args=[order.id]))
        self.assertEqual(response.status_code, 405)


def _fake_session(order, **overrides):
    data = {
        "id": order.stripe_session_id or "cs_test_123",
        "payment_status": "paid",
        "status": "complete",
        "amount_total": order.total_cents,
        "currency": order.currency,
        "payment_intent": "pi_test_123",
        "metadata": {"order_id": str(order.id)},
    }
    data.update(overrides)
    return data


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy")
class StripeTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("payer", "payer@example.com", "pass12345")
        self.product = make_product(price=Decimal("25.00"), stock=5)
        self.client.force_login(self.user)

    def _checkout_card(self, qty=2):
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": qty})
        fake = mock.MagicMock()
        fake.id = "cs_test_123"
        fake.url = "https://checkout.stripe.com/c/pay/cs_test_123"
        with mock.patch("stripe.checkout.Session.create", return_value=fake) as create:
            response = self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card"})
        return response, create

    def test_card_checkout_redirects_to_stripe_with_correct_amounts(self):
        response, create = self._checkout_card()
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith("https://checkout.stripe.com/"))
        kwargs = create.call_args.kwargs
        self.assertEqual(kwargs["line_items"][0]["price_data"]["unit_amount"], 2500)
        self.assertEqual(kwargs["line_items"][0]["quantity"], 2)
        self.assertEqual(kwargs["metadata"]["order_id"], str(Order.objects.get().id))
        order = Order.objects.get()
        self.assertEqual(order.payment_status, "pending")
        self.assertEqual(order.stripe_session_id, "cs_test_123")
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock, 3)  # reserved while paying

    def test_stripe_failure_releases_stock_and_restores_cart(self):
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 2})
        with mock.patch("stripe.checkout.Session.create", side_effect=Exception("boom")):
            response = self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card"})
        self.assertRedirects(response, reverse("cart"))
        order = Order.objects.get()
        self.assertEqual(order.status, "cancelled")
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock, 5)
        self.assertEqual(self.client.session["cart"], {str(self.product.id): 2})

    def _webhook(self, event_type, session):
        payload = json.dumps({"type": event_type, "data": {"object": session}})
        with mock.patch("stripe.Webhook.construct_event", return_value={}):
            return self.client.post(
                reverse("stripe_webhook"), data=payload, content_type="application/json",
                HTTP_STRIPE_SIGNATURE="t=1,v1=fake",
            )

    def test_webhook_marks_order_paid(self):
        self._checkout_card()
        order = Order.objects.get()
        response = self._webhook("checkout.session.completed", _fake_session(order))
        self.assertEqual(response.status_code, 200)
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "paid")
        self.assertEqual(order.status, "confirmed")
        self.assertEqual(order.stripe_payment_intent, "pi_test_123")

    def test_webhook_rejects_amount_mismatch(self):
        self._checkout_card()
        order = Order.objects.get()
        self._webhook("checkout.session.completed", _fake_session(order, amount_total=100))
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "pending")

    def test_webhook_bad_signature_rejected(self):
        with mock.patch("stripe.Webhook.construct_event", side_effect=ValueError("bad")):
            response = self.client.post(reverse("stripe_webhook"), data="{}", content_type="application/json")
        self.assertEqual(response.status_code, 400)

    def test_expired_session_cancels_and_restocks(self):
        self._checkout_card()
        order = Order.objects.get()
        self._webhook("checkout.session.expired", _fake_session(order, payment_status="unpaid", status="expired"))
        order.refresh_from_db()
        self.product.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.assertEqual(self.product.stock, 5)

    def test_webhook_is_idempotent(self):
        self._checkout_card()
        order = Order.objects.get()
        self._webhook("checkout.session.completed", _fake_session(order))
        self._webhook("checkout.session.completed", _fake_session(order))
        self._webhook("checkout.session.expired", _fake_session(order, status="expired"))
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "paid")
        self.assertEqual(order.status, "confirmed")

    def test_success_page_confirms_payment(self):
        self._checkout_card()
        order = Order.objects.get()
        fake = mock.MagicMock()
        fake.to_dict.return_value = _fake_session(order)
        with mock.patch("stripe.checkout.Session.retrieve", return_value=fake):
            response = self.client.get(reverse("payment_success"), {"session_id": "cs_test_123"})
        self.assertRedirects(response, reverse("order_success", args=[order.id]))
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "paid")

    def test_cancel_page_restores_cart(self):
        self._checkout_card()
        order = Order.objects.get()
        fake = mock.MagicMock()
        fake.to_dict.return_value = _fake_session(order, payment_status="unpaid", status="open")
        with mock.patch("stripe.checkout.Session.retrieve", return_value=fake), \
                mock.patch("stripe.checkout.Session.expire"):
            response = self.client.get(reverse("payment_cancel", args=[order.id]))
        self.assertRedirects(response, reverse("cart"))
        order.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.assertEqual(self.client.session["cart"], {str(self.product.id): 2})

    def test_cancelling_paid_order_refunds(self):
        self._checkout_card()
        order = Order.objects.get()
        self._webhook("checkout.session.completed", _fake_session(order))
        with mock.patch("stripe.Refund.create") as refund:
            self.client.post(reverse("cancel_order", args=[order.id]))
        refund.assert_called_once()
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "refunded")
        self.assertEqual(order.status, "cancelled")

    def test_refund_failure_keeps_order_active(self):
        self._checkout_card()
        order = Order.objects.get()
        self._webhook("checkout.session.completed", _fake_session(order))
        with mock.patch("stripe.Refund.create", side_effect=Exception("declined")):
            self.client.post(reverse("cancel_order", args=[order.id]))
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "paid")
        self.assertEqual(order.status, "confirmed")


class SecurityTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("alice", "alice@example.com", "Sup3r-secret-pass")

    def test_login_ignores_external_next_url(self):
        response = self.client.post(
            reverse("login") + "?next=https://evil.example/phish",
            {"username": "alice", "password": "Sup3r-secret-pass"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, reverse("home"))

    def test_login_allows_local_next_url(self):
        response = self.client.post(
            reverse("login"), {"username": "alice", "password": "Sup3r-secret-pass", "next": "/my-orders/"},
        )
        self.assertEqual(response.url, "/my-orders/")

    def test_login_with_email(self):
        response = self.client.post(reverse("login"), {"username": "ALICE@example.com", "password": "Sup3r-secret-pass"})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(int(self.client.session["_auth_user_id"]), self.user.id)

    def test_spoofed_forwarded_for_does_not_bypass_rate_limit(self):
        from django.core.cache import cache
        cache.clear()
        for i in range(8):
            self.client.post(reverse("login"), {"username": "alice", "password": "wrong"},
                             HTTP_X_FORWARDED_FOR=f"10.0.0.{i}")
        response = self.client.post(reverse("login"), {"username": "alice", "password": "Sup3r-secret-pass"},
                                    HTTP_X_FORWARDED_FOR="10.0.0.99")
        self.assertNotIn("_auth_user_id", self.client.session)
        self.assertRedirects(response, reverse("login"), fetch_redirect_response=False)
        cache.clear()

    def test_google_login_route_removed(self):
        self.assertEqual(self.client.post("/auth/google/").status_code, 404)

    def test_signup_rejects_weak_password_and_duplicate_email(self):
        self.client.post(reverse("signup"), {
            "username": "bob", "email": "bob@example.com", "password": "123", "confirm_password": "123",
        })
        self.assertFalse(User.objects.filter(username="bob").exists())
        self.client.post(reverse("signup"), {
            "username": "bob2", "email": "ALICE@example.com",
            "password": "Another-strong-pass1", "confirm_password": "Another-strong-pass1",
        })
        self.assertFalse(User.objects.filter(username="bob2").exists())

    def test_signup_rejects_email_as_username(self):
        self.client.post(reverse("signup"), {
            "username": "victim@example.com", "email": "attacker@example.com",
            "password": "Another-strong-pass1", "confirm_password": "Another-strong-pass1",
        })
        self.assertFalse(User.objects.filter(username="victim@example.com").exists())

    def test_logout_requires_post(self):
        self.client.force_login(self.user)
        self.client.get(reverse("logout"))
        self.assertIn("_auth_user_id", self.client.session)
        self.client.post(reverse("logout"))
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_invoice_of_guest_order_not_public(self):
        order = Order.objects.create(user=None, full_name="Guest", address="x", city="y", phone="1")
        self.assertEqual(self.client.get(reverse("invoice_pdf", args=[order.id])).status_code, 404)
        self.assertEqual(self.client.get(reverse("invoice_pdf", args=[order.id]) + "?t=wrong").status_code, 404)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("invoice_pdf", args=[order.id])).status_code, 404)
        self.assertEqual(Client().get(reverse("invoice_pdf", args=[order.id]) + f"?t={order.access_token}").status_code, 200)

    def test_email_code_locks_after_too_many_attempts(self):
        from django.core.cache import cache
        self.client.force_login(self.user)
        Profile_ = __import__("bees.models", fromlist=["Profile"]).Profile
        Profile_.objects.create(user=self.user, referral_code="ABC123")
        cache.set(f"email_verify_code:{self.user.id}", "123456", 900)
        for _ in range(5):
            self.client.post(reverse("verify_email"), {"code": "000000"})
        self.client.post(reverse("verify_email"), {"code": "123456"})
        self.user.profile.refresh_from_db()
        self.assertFalse(self.user.profile.email_verified)

    def test_review_requires_login_and_clamps_rating(self):
        product = make_product()
        self.client.post(reverse("product_detail", args=[product.id]), {"rating": 5, "comment": "Spam"})
        self.assertFalse(Review.objects.exists())
        self.client.force_login(self.user)
        self.client.post(reverse("product_detail", args=[product.id]), {"rating": 4, "comment": "Never bought it"})
        self.assertFalse(Review.objects.exists())  # only customers who received it can review
        order = Order.objects.create(user=self.user, full_name="R", address="x", city="y", phone="1", status="delivered")
        OrderItem.objects.create(order=order, product=product, product_name=product.name, price=product.price, quantity=1)
        self.client.post(reverse("product_detail", args=[product.id]), {"rating": 999, "comment": "Great"})
        self.assertFalse(Review.objects.exists())
        self.client.post(reverse("product_detail", args=[product.id]), {"rating": 4, "comment": "Great"})
        self.client.post(reverse("product_detail", args=[product.id]), {"rating": 5, "comment": "Even better"})
        self.assertEqual(Review.objects.get().rating, 5)


class SellerProductSecurityTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("maker", "maker@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.user, status="approved", business_name="Maker Co")
        self.product = make_product(seller_account=self.seller, seller_name="Maker Co", approval_status="approved")
        self.client.force_login(self.user)

    def _form(self, **overrides):
        data = {
            "name": self.product.name, "category": "skincare", "price": "10.00", "stock": "4",
            "image_url": self.product.image_url, "description": self.product.description,
        }
        data.update(overrides)
        return data

    def test_editing_listing_details_requires_re_approval(self):
        self.client.post(reverse("seller_edit_product", args=[self.product.id]), self._form(name="Totally different"))
        self.product.refresh_from_db()
        self.assertEqual(self.product.approval_status, "pending")

    def test_price_and_stock_changes_stay_live(self):
        self.client.post(reverse("seller_edit_product", args=[self.product.id]), self._form(price="12.50", stock="9"))
        self.product.refresh_from_db()
        self.assertEqual(self.product.approval_status, "approved")
        self.assertEqual(self.product.price, Decimal("12.50"))

    def test_html_upload_rejected(self):
        evil = SimpleUploadedFile("x.html", b"<script>alert(1)</script>", content_type="text/html")
        self.client.post(reverse("seller_add_product"), {**self._form(name="New", image_url=""), "image_file": evil})
        self.assertFalse(Product.objects.filter(name="New").exists())

    def test_fake_image_rejected(self):
        evil = SimpleUploadedFile("x.png", b"<svg onload=alert(1)>", content_type="image/png")
        self.client.post(reverse("seller_add_product"), {**self._form(name="New2", image_url=""), "image_file": evil})
        self.assertFalse(Product.objects.filter(name="New2").exists())

    def test_invalid_price_does_not_crash(self):
        response = self.client.post(reverse("seller_add_product"), self._form(name="Bad", price="abc"))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Product.objects.filter(name="Bad").exists())

    def test_delete_requires_post(self):
        response = self.client.get(reverse("seller_delete_product", args=[self.product.id]))
        self.assertEqual(response.status_code, 405)
        self.assertTrue(Product.objects.filter(pk=self.product.pk).exists())

    def test_seller_documents_are_staff_only(self):
        response = self.client.get(reverse("seller_document", args=[self.seller.id, "id_document"]))
        self.assertEqual(response.status_code, 403)


class WhiteLabelTests(TestCase):
    def test_store_name_comes_from_settings(self):
        make_product(name="Branded Product")
        settings_obj = SiteSettings.load()
        settings_obj.site_name = "Acme Goods"
        settings_obj.save()
        response = self.client.get(reverse("home"))
        self.assertContains(response, "Acme Goods")
        self.assertNotContains(response, "19Bees")

    @override_settings(STORE_CURRENCY="eur")
    def test_prices_use_store_currency(self):
        make_product(name="Euro Product", price=Decimal("1234.50"))
        response = self.client.get(reverse("all_products"))
        self.assertContains(response, "€1,234.50")


# ---------------------------------------------------------------------------
# Store admin (/manage/)
# ---------------------------------------------------------------------------
from django.core import mail

from .models import ChatMessage, ChatThread, Profile, ProductVariant, Question, ReturnRequest


class ManageAccessTests(TestCase):
    PAGES = ["manage_dashboard", "manage_orders", "manage_products", "manage_product_new", "manage_sellers",
             "manage_customers", "manage_coupons", "manage_reviews", "manage_returns", "manage_support", "manage_settings"]

    def setUp(self):
        self.staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        self.shopper = User.objects.create_user("shop", "shop@example.com", "pass12345")
        make_product(name="Admin Visible Product")

    def test_anonymous_redirected_to_login(self):
        for name in self.PAGES:
            r = self.client.get(reverse(name))
            self.assertEqual(r.status_code, 302, name)
            self.assertIn("/login/", r.url)

    def test_customers_are_forbidden(self):
        self.client.force_login(self.shopper)
        for name in self.PAGES:
            self.assertEqual(self.client.get(reverse(name)).status_code, 403, name)

    def test_staff_can_open_every_page(self):
        self.client.force_login(self.staff)
        for name in self.PAGES:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)
        self.assertEqual(self.client.get(reverse("manage_reviews") + "?tab=questions").status_code, 200)
        self.assertEqual(self.client.get(reverse("manage_dashboard") + "?range=30").status_code, 200)

    def test_admin_icon_only_for_staff(self):
        self.client.force_login(self.shopper)
        self.assertNotContains(self.client.get(reverse("home")), f'href="{reverse("manage_dashboard")}"')
        self.client.force_login(self.staff)
        self.assertContains(self.client.get(reverse("home")), f'href="{reverse("manage_dashboard")}"')

    def test_old_dashboard_url_redirects(self):
        self.client.force_login(self.staff)
        self.assertRedirects(self.client.get(reverse("owner_dashboard")), reverse("manage_dashboard"))


class ManageActionTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        self.buyer = User.objects.create_user("buyer", "buyer@example.com", "pass12345")
        self.product = make_product(price=Decimal("30.00"), stock=5)
        self.client.force_login(self.staff)

    def _order(self, **kw):
        order = Order.objects.create(user=self.buyer, email="buyer@example.com", full_name="Buy Er", address="1 St",
                                     city="Austin", postal_code="1", country="US", phone="1", **kw)
        OrderItem.objects.create(order=order, product=self.product, product_name=self.product.name, price=Decimal("30.00"), quantity=2)
        return order

    def test_marking_shipped_emails_customer_and_notifies(self):
        order = self._order(status="confirmed")
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("manage_order", args=[order.id]), {
                "action": "update", "status": "shipped", "tracking_number": "1Z999", "courier_name": "UPS", "estimated_delivery": "2026-12-01",
            })
        order.refresh_from_db()
        self.assertEqual(order.status, "shipped")
        self.assertEqual(order.tracking_number, "1Z999")
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("1Z999", mail.outbox[0].alternatives[0][0])
        self.assertTrue(Notification.objects.filter(user=self.buyer, message__icontains="Shipped").exists())

    def test_cancel_restocks(self):
        order = self._order(status="pending")
        Product.objects.filter(pk=self.product.pk).update(stock=3)
        self.client.post(reverse("manage_order", args=[order.id]), {"action": "cancel"})
        order.refresh_from_db()
        self.product.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.assertEqual(self.product.stock, 5)

    def test_orders_csv_export(self):
        self._order()
        r = self.client.get(reverse("manage_orders") + "?export=csv")
        self.assertEqual(r["Content-Type"], "text/csv")
        self.assertIn(b"Buy Er", r.content)

    def test_create_product(self):
        r = self.client.post(reverse("manage_product_new"), {
            "name": "Staff Made", "category": "skincare", "price": "12.50", "stock": "7",
            "image_url": "https://example.com/a.jpg", "approval_status": "approved", "is_flash_sale": "1",
            "extra_images": "https://example.com/b.jpg\nnot-a-url",
        })
        p = Product.objects.get(name="Staff Made")
        self.assertRedirects(r, reverse("manage_product_edit", args=[p.id]))
        self.assertTrue(p.is_flash_sale)
        self.assertEqual(p.extra_images.count(), 1)

    def test_bulk_approve_notifies_seller(self):
        seller_user = User.objects.create_user("sel", "sel@example.com", "pass12345")
        seller = SellerAccount.objects.create(user=seller_user, status="approved", business_name="Sel Co")
        pending = make_product(name="Pending One", seller_account=seller, approval_status="pending")
        self.client.post(reverse("manage_products_bulk"), {"ids": [pending.id], "action": "approve"})
        pending.refresh_from_db()
        self.assertEqual(pending.approval_status, "approved")
        self.assertTrue(Notification.objects.filter(user=seller_user, message__icontains="approved").exists())

    def test_seller_approve_and_payout(self):
        u = User.objects.create_user("app", "app@example.com", "pass12345")
        s = SellerAccount.objects.create(user=u, business_name="App Shop")
        self.client.post(reverse("manage_seller", args=[s.id]), {"action": "approved"})
        s.refresh_from_db()
        self.assertEqual(s.status, "approved")
        self.client.post(reverse("manage_seller", args=[s.id]), {"action": "payout", "amount": "25.50"})
        s.refresh_from_db()
        self.assertEqual(s.total_paid_out, Decimal("25.50"))
        self.client.post(reverse("manage_seller", args=[s.id]), {"action": "payout", "amount": "-4"})
        s.refresh_from_db()
        self.assertEqual(s.total_paid_out, Decimal("25.50"))

    def test_coupon_create_toggle_delete(self):
        self.client.post(reverse("manage_coupons"), {"action": "create", "code": "spring", "percent_off": "15",
                                                      "min_order_value": "0", "per_user_limit": "1", "active": "on"})
        c = Coupon.objects.get(code="SPRING")
        self.client.post(reverse("manage_coupons"), {"action": "toggle", "id": c.id})
        c.refresh_from_db()
        self.assertFalse(c.active)
        self.client.post(reverse("manage_coupons"), {"action": "create", "code": "bad", "percent_off": "500",
                                                      "min_order_value": "0", "per_user_limit": "1"})
        self.assertFalse(Coupon.objects.filter(code="BAD").exists())
        self.client.post(reverse("manage_coupons"), {"action": "delete", "id": c.id})
        self.assertFalse(Coupon.objects.filter(pk=c.id).exists())

    def test_answer_question(self):
        qn = Question.objects.create(product=self.product, username="q", question="Waterproof?")
        self.client.post(reverse("manage_reviews") + "?tab=questions", {"action": "answer", "id": qn.id, "answer": "Yes."})
        qn.refresh_from_db()
        self.assertEqual(qn.answer, "Yes.")

    @override_settings(STRIPE_SECRET_KEY="sk_test_dummy")
    def test_return_refund_to_card_is_partial(self):
        order = self._order(status="delivered", payment_method="card", payment_status="paid", stripe_payment_intent="pi_1")
        item = order.items.get()
        rr = ReturnRequest.objects.create(order_item=item, user=self.buyer, reason="Broken")
        with mock.patch("stripe.Refund.create") as refund:
            self.client.post(reverse("manage_returns"), {"id": rr.id, "action": "refund"})
        self.assertEqual(refund.call_args.kwargs["amount"], 6000)
        rr.refresh_from_db()
        self.product.refresh_from_db()
        self.assertEqual(rr.status, "refunded")
        self.assertEqual(self.product.stock, 7)

    def test_support_reply(self):
        thread = ChatThread.objects.create(user=self.buyer)
        ChatMessage.objects.create(thread=thread, sender="user", message="Where is my order?")
        self.client.get(reverse("manage_support_thread", args=[thread.id]))
        self.assertFalse(thread.messages.filter(is_read=False).exists())
        self.client.post(reverse("manage_support_thread", args=[thread.id]), {"action": "reply", "message": "Shipping today!"})
        self.assertTrue(thread.messages.filter(sender="support", message="Shipping today!").exists())
        self.client.logout()
        self.client.force_login(self.buyer)
        msgs = self.client.get(reverse("chat_messages")).json()["messages"]
        self.assertEqual(msgs[-1]["message"], "Shipping today!")

    def test_settings_update_rebrands_store(self):
        data = {f: v for f, v in {
            "site_name": "Nova Goods", "tagline": "Good things", "primary_color": "#112233", "accent_color": "#FFCC00",
            "hero_title": "Hello", "hero_subtitle": "World", "tax_percent": "0", "shipping_flat_fee": "0",
            "free_shipping_threshold": "0", "delivery_days": "5", "return_days": "14", "allow_cash_on_delivery": "on",
        }.items()}
        self.client.post(reverse("manage_settings"), data)
        self.assertContains(self.client.get(reverse("home")), "Nova Goods")
        bad = dict(data, primary_color="red")
        self.client.post(reverse("manage_settings"), bad)
        self.assertEqual(SiteSettings.load().primary_color, "#112233")


class SpeculativeRequestTests(TestCase):
    def test_prefetch_does_not_count_as_search(self):
        from .models import SearchLog
        make_product(name="Oil Lamp")
        self.client.get(reverse("search_products"), {"q": "oil"}, HTTP_SEC_PURPOSE="prefetch;prerender")
        self.assertFalse(SearchLog.objects.exists())
        self.client.get(reverse("search_products"), {"q": "oil"})
        self.assertTrue(SearchLog.objects.exists())


class MarketplaceFlowTests(TestCase):
    """End to end: a business signs up to sell, gets approved, lists a
    product, a buyer orders it, and the commission is recorded."""

    def test_full_seller_to_buyer_flow(self):
        staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        # 1. A business applies to sell
        self.client.post(reverse("signup"), {
            "username": "acme", "email": "acme@example.com", "password": "Strong-pass-123", "confirm_password": "Strong-pass-123",
            "user_type": "organization", "organization_name": "Acme Ltd", "phone": "+1 555 0101", "country": "US",
        })
        seller = SellerAccount.objects.get(user__username="acme")
        self.assertEqual(seller.status, "pending")
        self.assertEqual(seller.commission_rate, Decimal("20"))
        # 2. Staff approve the seller
        self.client.force_login(staff)
        self.client.post(reverse("manage_seller", args=[seller.id]), {"action": "approved"})
        # 3. Seller lists a product; it waits for review
        self.client.force_login(seller.user)
        self.client.post(reverse("seller_add_product"), {
            "name": "Acme Lamp", "category": "table-lamp", "price": "50.00", "stock": "4",
            "image_url": "https://example.com/lamp.jpg", "description": "A lamp",
        })
        product = Product.objects.get(name="Acme Lamp")
        self.assertEqual(product.approval_status, "pending")
        self.assertEqual(self.client.get(reverse("product_detail", args=[product.id])).status_code, 200)  # owner preview
        self.client.logout()
        self.assertEqual(self.client.get(reverse("product_detail", args=[product.id])).status_code, 404)  # hidden from public
        # 4. Staff approve the product
        self.client.force_login(staff)
        self.client.post(reverse("manage_products_bulk"), {"ids": [product.id], "action": "approve"})
        # 5. A buyer orders 2 with cash on delivery
        buyer = User.objects.create_user("buyer", "buyer@example.com", "pass12345")
        self.client.force_login(buyer)
        self.client.post(reverse("add_to_cart", args=[product.id]), {"quantity": 2})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        item = OrderItem.objects.get(product=product)
        self.assertEqual(item.seller_account, seller)
        self.assertEqual(item.commission_rate, Decimal("20"))
        self.assertEqual(item.commission_amount, Decimal("20.00"))
        seller.refresh_from_db()
        self.assertEqual(seller.lifetime_sales, Decimal("100.00"))
        self.assertEqual(seller.net_earnings, 80.0)
        # 6. The seller is alerted and sees the sale, the address and their earning
        self.assertTrue(Notification.objects.filter(user=seller.user, message__icontains=f"New order #{item.order_id}").exists())
        self.client.force_login(seller.user)
        page = self.client.get(reverse("seller_dashboard"))
        self.assertContains(page, "Acme Lamp")
        self.assertContains(page, "$80.00")
        page = self.client.get(reverse("seller_orders"))
        self.assertContains(page, "Austin")
        page = self.client.get(reverse("seller_order", args=[item.order_id]))
        self.assertContains(page, "Austin")
        self.assertContains(page, "$80.00")
        # 7. Cancelled orders don't count toward earnings
        self.client.force_login(buyer)
        self.client.post(reverse("cancel_order", args=[item.order.id]))
        seller = SellerAccount.objects.get(pk=seller.pk)
        self.assertEqual(seller.lifetime_sales, 0)
        product.refresh_from_db()
        self.assertEqual(product.stock, 4)

    def test_store_products_have_no_commission(self):
        buyer = User.objects.create_user("b2", "b2@example.com", "pass12345")
        p = make_product(price=Decimal("10.00"))
        self.client.force_login(buyer)
        self.client.post(reverse("add_to_cart", args=[p.id]))
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        item = OrderItem.objects.get()
        self.assertIsNone(item.seller_account)
        self.assertEqual(item.commission_amount, 0)


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy")
class PayCodOrderOnlineTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("payer", "payer@example.com", "pass12345")
        self.product = make_product(price=Decimal("25.00"), stock=5)
        self.client.force_login(self.user)
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.order = Order.objects.get()

    def _start(self):
        fake = mock.MagicMock()
        fake.id = "cs_cod_1"
        fake.url = "https://checkout.stripe.com/c/pay/cs_cod_1"
        with mock.patch("stripe.checkout.Session.create", return_value=fake):
            return self.client.post(reverse("pay_online", args=[self.order.id]))

    def test_cod_order_can_be_paid_in_advance(self):
        r = self._start()
        self.assertTrue(r.url.startswith("https://checkout.stripe.com/"))
        self.order.refresh_from_db()
        payload = json.dumps({"type": "checkout.session.completed", "data": {"object": _fake_session(self.order)}})
        with mock.patch("stripe.Webhook.construct_event", return_value={}):
            self.client.post(reverse("stripe_webhook"), data=payload, content_type="application/json", HTTP_STRIPE_SIGNATURE="x")
        self.order.refresh_from_db()
        self.assertEqual(self.order.payment_status, "paid")
        self.assertEqual(self.order.payment_method, "card")
        self.assertFalse(self.order.cod_fallback)

    def test_abandoned_advance_payment_keeps_cod_order(self):
        self._start()
        self.order.refresh_from_db()
        payload = json.dumps({"type": "checkout.session.expired", "data": {"object": _fake_session(self.order, payment_status="unpaid", status="expired")}})
        with mock.patch("stripe.Webhook.construct_event", return_value={}):
            self.client.post(reverse("stripe_webhook"), data=payload, content_type="application/json", HTTP_STRIPE_SIGNATURE="x")
        self.order.refresh_from_db()
        self.product.refresh_from_db()
        self.assertEqual(self.order.status, "pending")
        self.assertEqual(self.order.payment_method, "cod")
        self.assertEqual(self.order.payment_status, "not_applicable")
        self.assertEqual(self.product.stock, 4)  # still reserved for the COD order

    def test_cancel_page_returns_to_cod(self):
        self._start()
        fake = mock.MagicMock()
        self.order.refresh_from_db()
        fake.to_dict.return_value = _fake_session(self.order, payment_status="unpaid", status="open")
        with mock.patch("stripe.checkout.Session.retrieve", return_value=fake), mock.patch("stripe.checkout.Session.expire"):
            r = self.client.get(reverse("payment_cancel", args=[self.order.id]))
        self.assertRedirects(r, reverse("my_orders"))
        self.order.refresh_from_db()
        self.assertEqual(self.order.payment_method, "cod")
        self.assertNotEqual(self.order.status, "cancelled")

    def test_paid_or_shipped_orders_cannot_switch(self):
        Order.objects.filter(pk=self.order.pk).update(status="shipped")
        self._start()
        self.order.refresh_from_db()
        self.assertEqual(self.order.payment_method, "cod")


class DemoCatalogueTests(TestCase):
    def test_seed_gives_every_product_a_matching_image(self):
        import tempfile
        from django.core.management import call_command
        with tempfile.TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp):
            Product.objects.create(name="Charcoal Face Wash", category="skincare", price=5,
                                   image_url="https://picsum.photos/id/10/400/400")
            call_command("seed_data", stdout=open(os.devnull, "w"))
            fw = Product.objects.get(name="Charcoal Face Wash")
            self.assertIn("charcoal-face-wash", fw.image_url)
            self.assertFalse(Product.objects.filter(image_url__contains="picsum").exists())
            self.assertFalse(ProductImage.objects.filter(image_url__contains="picsum").exists())


class StripeRelayTests(TestCase):
    def tearDown(self):
        import stripe
        stripe.api_base = "https://api.stripe.com"

    def test_default_api_base(self):
        from . import payments
        with self.settings(STRIPE_SECRET_KEY="sk_test_x", STRIPE_API_BASE=""):
            self.assertEqual(payments._stripe().api_base, "https://api.stripe.com")

    def test_relay_api_base(self):
        from . import payments
        relay = "https://ref.supabase.co/functions/v1/stripe-relay/secret"
        with self.settings(STRIPE_SECRET_KEY="sk_test_x", STRIPE_API_BASE=relay):
            self.assertEqual(payments._stripe().api_base, relay)
@override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy",
                   ALLOWED_HOSTS=["shop.example.com", "testserver"], SITE_URL="")
class OrderJourneyTests(TestCase):
    """The customer is kept informed at every step, and cancelling a paid
    order refunds it."""

    def setUp(self):
        self.staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        self.buyer = User.objects.create_user("buyer", "buyer@example.com", "pass12345")
        self.product = make_product(name="Charcoal Face Wash", price=Decimal("20.00"), stock=10)

    def _cod_order(self):
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 2})
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        return Order.objects.get()

    def _admin_set(self, order, status, **extra):
        self.client.force_login(self.staff)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("manage_order", args=[order.id]), {
                "action": "update", "status": status, "tracking_number": extra.get("tracking", ""),
                "courier_name": extra.get("courier", ""),
                "estimated_delivery": order.estimated_delivery.isoformat() if order.estimated_delivery else "",
            })
        order.refresh_from_db()
        return order

    def test_full_cod_journey_emails_every_step(self):
        order = self._cod_order()
        # Placed: estimated delivery set, "received" email with the date
        self.assertIsNotNone(order.estimated_delivery)
        self.assertGreater(order.estimated_delivery, order.created_at.date())
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("received", mail.outbox[0].subject)
        self.assertIn("Estimated delivery", mail.outbox[0].alternatives[0][0])

        order = self._admin_set(order, "confirmed")
        self.assertEqual(len(mail.outbox), 2)
        self.assertIn("is confirmed", mail.outbox[1].subject)
        self.assertIn(order.estimated_delivery.strftime("%B"), mail.outbox[1].alternatives[0][0])

        order = self._admin_set(order, "shipped", tracking="TRK123", courier="DHL")
        self.assertEqual(len(mail.outbox), 3)
        self.assertIn("on its way", mail.outbox[2].subject)
        self.assertIn("TRK123", mail.outbox[2].alternatives[0][0])
        self.assertIn("https://shop.example.com/my-orders/", mail.outbox[2].alternatives[0][0])
        self.assertEqual(set(order.items.values_list("fulfillment_status", flat=True)), {"handed_to_courier"})

        order = self._admin_set(order, "delivered")
        self.assertEqual(len(mail.outbox), 4)
        self.assertIn("delivered", mail.outbox[3].subject)
        self.assertIn("Write a review", mail.outbox[3].alternatives[0][0])
        self.assertTrue(all(m.to == ["jane@example.com"] for m in mail.outbox))  # the email given at checkout

    def test_customer_cancels_cod_order(self):
        order = self._cod_order()
        mail.outbox.clear()
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("cancel_order", args=[order.id]))
        order.refresh_from_db()
        self.product.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.assertEqual(self.product.stock, 10)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("cancelled", mail.outbox[0].subject)
        self.assertIn("not been charged", mail.outbox[0].alternatives[0][0])

    def test_customer_cancels_paid_order_gets_refund_and_email(self):
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 1})
        fake = mock.MagicMock(id="cs_test_9", url="https://checkout.stripe.com/x")
        with mock.patch("stripe.checkout.Session.create", return_value=fake):
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card"})
        order = Order.objects.get()
        payload = json.dumps({"type": "checkout.session.completed", "data": {"object": _fake_session(order)}})
        with mock.patch("stripe.Webhook.construct_event", return_value={}), self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("stripe_webhook"), data=payload, content_type="application/json", HTTP_STRIPE_SIGNATURE="t=1,v1=x")
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "paid")
        self.assertIsNotNone(order.estimated_delivery)
        self.assertEqual(len(mail.outbox), 1)  # one "confirmed + paid" email, not two
        mail.outbox.clear()

        with mock.patch("stripe.Refund.create") as refund, self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("cancel_order", args=[order.id]))
        refund.assert_called_once()
        self.assertEqual(refund.call_args.kwargs["payment_intent"], "pi_test_123")
        order.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.assertEqual(order.payment_status, "refunded")
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("full refund", mail.outbox[0].alternatives[0][0])
        self.assertIn("$20.00", mail.outbox[0].alternatives[0][0])
        self.assertContains(self.client.get(reverse("my_orders")), "refunded to your card")

    def test_failed_refund_keeps_order_and_warns(self):
        order = self._cod_order()
        Order.objects.filter(pk=order.pk).update(payment_status="paid", payment_method="card", stripe_payment_intent="pi_x")
        with mock.patch("stripe.Refund.create", side_effect=Exception("down")):
            r = self.client.post(reverse("cancel_order", args=[order.id]), follow=True)
        order.refresh_from_db()
        self.assertEqual(order.status, "pending")
        self.assertContains(r, "couldn&#x27;t process the refund")

    def test_shipped_order_cannot_be_cancelled_by_customer(self):
        order = self._cod_order()
        self._admin_set(order, "shipped")
        self.client.force_login(self.buyer)
        self.client.post(reverse("cancel_order", args=[order.id]))
        order.refresh_from_db()
        self.assertEqual(order.status, "shipped")

    def test_seller_fulfilment_moves_order_forward(self):
        seller_user = User.objects.create_user("sel", "sel@example.com", "pass12345")
        seller = SellerAccount.objects.create(user=seller_user, status="approved", business_name="Sel Co")
        Product.objects.filter(pk=self.product.pk).update(seller_account=seller)
        order = self._cod_order()
        item = order.items.get()
        mail.outbox.clear()
        self.client.force_login(seller_user)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("update_fulfillment_status", args=[item.id]), {"fulfillment_status": "handed_to_courier"})
        order.refresh_from_db()
        self.assertEqual(order.status, "shipped")
        self.assertIn("on its way", mail.outbox[-1].subject)
        with self.captureOnCommitCallbacks(execute=True):
            # Cash orders: the customer's doorstep code is needed.
            self.client.post(reverse("update_fulfillment_status", args=[item.id]), {"fulfillment_status": "delivered",
                                                                                    "delivery_code": order.delivery_code})
        order.refresh_from_db()
        self.assertEqual(order.status, "delivered")
        self.assertIn("delivered", mail.outbox[-1].subject)

    def test_seller_cannot_ship_cancelled_order(self):
        seller_user = User.objects.create_user("sel", "sel@example.com", "pass12345")
        seller = SellerAccount.objects.create(user=seller_user, status="approved", business_name="Sel Co")
        Product.objects.filter(pk=self.product.pk).update(seller_account=seller)
        order = self._cod_order()
        self.client.post(reverse("cancel_order", args=[order.id]))
        item = order.items.get()
        self.client.force_login(seller_user)
        self.client.post(reverse("update_fulfillment_status", args=[item.id]), {"fulfillment_status": "handed_to_courier"})
        item.refresh_from_db()
        self.assertEqual(item.fulfillment_status, "pending")

    def test_delivery_days_setting_controls_date(self):
        from .order_emails import add_business_days
        from datetime import date
        self.assertEqual(add_business_days(date(2026, 10, 9), 1), date(2026, 10, 12))  # Fri -> Mon
        SiteSettings.objects.update_or_create(pk=1, defaults={"delivery_days": 10})
        cache.clear()
        order = self._cod_order()
        from django.utils import timezone
        self.assertEqual(order.estimated_delivery, add_business_days(timezone.localdate(), 10))


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy")
class AuditFixTests(TestCase):
    """Regression tests for the problems found in the full-site review."""

    def setUp(self):
        self.buyer = User.objects.create_user("buyer", "buyer@example.com", "pass12345")
        self.staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        self.product = make_product(name="Face Wash", price=Decimal("20.00"), stock=10)

    # --- helpers -----------------------------------------------------------
    def _seller(self, name="Sel Co", username="sel", **kw):
        u = User.objects.create_user(username, f"{username}@example.com", "pass12345")
        return SellerAccount.objects.create(user=u, status="approved", business_name=name, **kw)

    def _delivered_order(self, user=None, product=None, qty=1, **kw):
        product = product or self.product
        order = Order.objects.create(user=user or self.buyer, email="buyer@example.com", full_name="B", address="1 St",
                                     city="Austin", country="US", phone="1", status="delivered", **kw)
        OrderItem.objects.create(order=order, product=product, product_name=product.name, price=product.price, quantity=qty)
        return order

    def _webhook(self, session):
        payload = json.dumps({"type": "checkout.session.completed", "data": {"object": session}})
        with mock.patch("stripe.Webhook.construct_event", return_value={}):
            return self.client.post(reverse("stripe_webhook"), data=payload, content_type="application/json", HTTP_STRIPE_SIGNATURE="t=1,v1=x")

    def _card_checkout(self, qty=1, **post):
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": qty})
        fake = mock.MagicMock(id="cs_test_1", url="https://checkout.stripe.com/x")
        with mock.patch("stripe.checkout.Session.create", return_value=fake) as create, \
             mock.patch("stripe.Coupon.create", return_value=mock.MagicMock(id="co_1")) as coupon:
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card", **post})
        return Order.objects.latest("id"), create, coupon

    # --- sellers -----------------------------------------------------------
    def test_seller_product_pages_open_and_keep_zero_stock(self):
        seller = self._seller()
        p = make_product(name="Sold Out", seller_account=seller, seller_name="Sel Co", stock=0)
        self.client.force_login(seller.user)
        self.assertEqual(self.client.get(reverse("seller_add_product")).status_code, 200)
        page = self.client.get(reverse("seller_edit_product", args=[p.id]))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, 'name="stock" min="0" step="1" value="0"')

    def test_seller_can_keep_uploaded_image_path(self):
        seller = self._seller()
        p = make_product(name="Local", seller_account=seller, seller_name="Sel Co", image_url="/media/products/a.png")
        self.client.force_login(seller.user)
        self.client.post(reverse("seller_edit_product", args=[p.id]), {
            "name": "Local", "category": "skincare", "price": "11.00", "stock": "3", "image_url": "/media/products/a.png",
        })
        p.refresh_from_db()
        self.assertEqual(p.price, Decimal("11.00"))

    def test_store_name_with_slash_works(self):
        seller = self._seller(name="A/B Goods")
        p = make_product(name="Slashy", seller_account=seller, seller_name="A/B Goods")
        self.assertEqual(self.client.get(reverse("product_detail", args=[p.id])).status_code, 200)
        self.assertContains(self.client.get(reverse("store_page", args=["A/B Goods"])), "Slashy")

    def test_suspended_seller_products_hidden_and_unbuyable(self):
        seller = self._seller()
        p = make_product(name="Gone", seller_account=seller, seller_name="Sel Co")
        SellerAccount.objects.filter(pk=seller.pk).update(status="suspended")
        self.assertEqual(self.client.get(reverse("product_detail", args=[p.id])).status_code, 404)
        self.assertNotContains(self.client.get(reverse("all_products")), "Gone")
        self.client.post(reverse("add_to_cart", args=[p.id]))
        self.assertEqual(self.client.session.get("cart", {}), {})

    def test_team_admin_can_manage_team(self):
        org = self._seller(name="Org", username="orgo", account_type="organization")
        admin = User.objects.create_user("adm", "adm@example.com", "pass12345")
        OrganizationMember.objects.create(organization=org, user=admin, role="admin")
        User.objects.create_user("newbie", "newbie@example.com", "pass12345")
        self.client.force_login(admin)
        self.client.post(reverse("add_team_member"), {"username_or_email": "newbie", "role": "staff"})
        self.assertTrue(OrganizationMember.objects.filter(organization=org, user__username="newbie").exists())

    def test_seller_answers_question_and_asker_is_notified(self):
        seller = self._seller()
        p = make_product(name="Q Product", seller_account=seller, seller_name="Sel Co")
        self.client.force_login(self.buyer)
        self.client.post(reverse("ask_question", args=[p.id]), {"question": "Is it vegan?"})
        self.assertTrue(Notification.objects.filter(user=seller.user, message__icontains="question").exists())
        q = Question.objects.get()
        self.client.force_login(seller.user)
        self.assertContains(self.client.get(reverse("seller_reviews")), "Is it vegan?")
        self.client.post(reverse("seller_answer_question", args=[q.id]), {"answer": "Yes, 100%."})
        q.refresh_from_db()
        self.assertEqual(q.answer, "Yes, 100%.")
        self.assertTrue(Notification.objects.filter(user=self.buyer, message__icontains="answered").exists())

    # --- payments ----------------------------------------------------------
    def test_cancelled_card_order_cannot_be_charged(self):
        order, _, _ = self._card_checkout()
        with mock.patch("stripe.checkout.Session.expire") as expire:
            self.client.post(reverse("cancel_order", args=[order.id]))
        expire.assert_called_once_with("cs_test_1")
        order.refresh_from_db()
        self.assertEqual((order.status, order.payment_status), ("cancelled", "failed"))
        # Even if the old Stripe tab is paid anyway, the money goes straight back.
        order.payment_status = "pending"
        order.save(update_fields=["payment_status"])
        with mock.patch("stripe.Refund.create") as refund:
            self._webhook(_fake_session(order))
        refund.assert_called_once()
        order.refresh_from_db()
        self.assertEqual((order.status, order.payment_status), ("cancelled", "refunded"))

    def test_zero_decimal_currency_amounts(self):
        from .payments import to_cents
        self.assertEqual(to_cents(Decimal("1000"), "jpy"), 1000)
        self.assertEqual(to_cents(Decimal("10.50"), "usd"), 1050)

    def test_coupon_kept_when_payment_cancelled(self):
        Coupon.objects.create(code="SAVE10", percent_off=10)
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        self.client.post(reverse("apply_coupon"), {"coupon_code": "SAVE10"})
        fake = mock.MagicMock(id="cs_test_2", url="https://checkout.stripe.com/x")
        with mock.patch("stripe.checkout.Session.create", return_value=fake), \
             mock.patch("stripe.Coupon.create", return_value=mock.MagicMock(id="co")):
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card"})
        order = Order.objects.get()
        self.assertEqual(self.client.session["coupon_code"], "")
        with mock.patch("stripe.checkout.Session.retrieve", side_effect=Exception("offline")), \
             mock.patch("stripe.checkout.Session.expire"):
            self.client.get(reverse("payment_cancel", args=[order.id]))
        self.assertEqual(self.client.session["coupon_code"], "SAVE10")

    def test_free_shipping_uses_price_after_coupon(self):
        SiteSettings.objects.update_or_create(pk=1, defaults={"shipping_flat_fee": Decimal("5"), "free_shipping_threshold": Decimal("50")})
        cache.clear()
        from .views import _price_cart
        coupon = Coupon(code="HALF", percent_off=50)
        self.assertEqual(_price_cart(Decimal("60"), coupon)["shipping"], Decimal("5"))
        self.assertEqual(_price_cart(Decimal("120"), coupon)["shipping"], Decimal("0"))

    # --- store credit & points ---------------------------------------------
    def test_store_credit_partly_pays_card_order_and_returns_on_cancel(self):
        Profile.objects.create(user=self.buyer, referral_code="B1", store_credit=Decimal("5.00"))
        order, create, coupon = self._card_checkout(use_credit="1")
        self.assertEqual(order.credit_used, Decimal("5.00"))
        self.assertEqual(order.total, Decimal("15.00"))
        self.assertEqual(coupon.call_args.kwargs["amount_off"], 500)
        self.assertEqual(Profile.objects.get(user=self.buyer).store_credit, 0)
        with mock.patch("stripe.checkout.Session.expire"):
            self.client.post(reverse("cancel_order", args=[order.id]))
        self.assertEqual(Profile.objects.get(user=self.buyer).store_credit, Decimal("5.00"))

    def test_store_credit_can_cover_whole_order(self):
        Profile.objects.create(user=self.buyer, referral_code="B1", store_credit=Decimal("50.00"))
        order, create, _ = self._card_checkout(use_credit="1")
        create.assert_not_called()
        self.assertEqual((order.payment_status, order.status), ("paid", "confirmed"))
        self.assertEqual(Profile.objects.get(user=self.buyer).store_credit, Decimal("30.00"))
        self.client.post(reverse("cancel_order", args=[order.id]))  # nothing on a card; credit comes back
        order.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.assertEqual(Profile.objects.get(user=self.buyer).store_credit, Decimal("50.00"))

    def test_credit_not_used_unless_ticked(self):
        Profile.objects.create(user=self.buyer, referral_code="B1", store_credit=Decimal("5.00"))
        order, _, _ = self._card_checkout()
        self.assertEqual(order.credit_used, 0)

    def test_redeem_points(self):
        Profile.objects.create(user=self.buyer, referral_code="B1", loyalty_points=250)
        self.client.force_login(self.buyer)
        self.assertContains(self.client.get(reverse("profile")), "Turn 200 points into $2.00")
        self.client.post(reverse("redeem_points"))
        prof = Profile.objects.get(user=self.buyer)
        self.assertEqual((prof.loyalty_points, prof.store_credit), (50, Decimal("2.00")))

    # --- returns -----------------------------------------------------------
    def _return(self, order, method="original_payment"):
        return ReturnRequest.objects.create(order_item=order.items.get(), user=self.buyer, reason="Broken", refund_method=method)

    def test_return_refund_respects_coupon_and_only_once(self):
        order = self._delivered_order(qty=1, discount_amount=Decimal("10.00"), payment_status="paid",
                                      payment_method="card", stripe_payment_intent="pi_1")
        rr = self._return(order)
        self.client.force_login(self.staff)
        with mock.patch("stripe.Refund.create") as refund:
            self.client.post(reverse("manage_returns"), {"id": rr.id, "action": "refund"})
            self.client.post(reverse("manage_returns"), {"id": rr.id, "action": "refund"})
        refund.assert_called_once()
        self.assertEqual(refund.call_args.kwargs["amount"], 1000)  # paid $10 after coupon, not $20
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock, 11)  # restocked once

    def test_store_credit_refund_and_cod_refund(self):
        order = self._delivered_order(payment_method="cod")
        rr = self._return(order, method="store_credit")
        self.client.force_login(self.staff)
        self.client.post(reverse("manage_returns"), {"id": rr.id, "action": "refund"})
        self.assertEqual(Profile.objects.get(user=self.buyer).store_credit, Decimal("20.00"))
        order2 = self._delivered_order(payment_method="cod")
        rr2 = self._return(order2)
        r = self.client.post(reverse("manage_returns"), {"id": rr2.id, "action": "refund"}, follow=True)
        self.assertContains(r, "pay the customer $20.00 yourself")
        rr2.refresh_from_db()
        self.assertEqual(rr2.status, "refunded")

    def test_return_window_enforced(self):
        from datetime import timedelta
        from django.utils import timezone
        order = self._delivered_order()
        Order.objects.filter(pk=order.pk).update(delivered_at=timezone.now() - timedelta(days=20))
        self.client.force_login(self.buyer)
        item = order.items.get()
        self.client.post(reverse("request_return", args=[item.id]), {"reason": "late"})
        self.assertFalse(ReturnRequest.objects.exists())
        self.assertContains(self.client.get(reverse("my_orders")), "Return window closed")

    def test_delivered_at_recorded(self):
        order = self._delivered_order()
        order.status = "confirmed"; order.save()
        order.status = "delivered"; order.save(update_fields=["status"])
        order.refresh_from_db()
        self.assertIsNotNone(order.delivered_at)

    # --- accounts, reviews, misc -------------------------------------------
    def test_seller_cannot_review_own_product(self):
        seller = self._seller()
        p = make_product(name="Mine", seller_account=seller, seller_name="Sel Co")
        self._delivered_order(user=seller.user, product=p)
        self.client.force_login(seller.user)
        self.client.post(reverse("product_detail", args=[p.id]), {"rating": 5, "comment": "Best ever"})
        self.assertFalse(Review.objects.exists())

    def test_referral_rewarded_once_after_first_delivery(self):
        referrer = Profile.objects.create(user=self.staff, referral_code="REFCODE1")
        self.client.post(reverse("signup") + "?ref=REFCODE1", {
            "username": "friend", "email": "friend@example.com", "password": "Strong-pass-123",
            "confirm_password": "Strong-pass-123", "ref": "REFCODE1",
        })
        friend = User.objects.get(username="friend")
        self.assertFalse(Coupon.objects.filter(code__startswith="REF-").exists())  # nothing for just signing up
        welcome = Coupon.objects.get(code__startswith="WELCOME-")
        self.assertEqual(welcome.usage_limit, 1)
        order = self._delivered_order(user=friend)
        order.status = "confirmed"; order.save()
        order.status = "delivered"; order.save()
        order2 = self._delivered_order(user=friend)
        order2.status = "confirmed"; order2.save(); order2.status = "delivered"; order2.save()
        self.assertEqual(Coupon.objects.filter(code__startswith="REF-").count(), 1)
        self.assertTrue(Notification.objects.filter(user=referrer.user, message__icontains="first purchase").exists())

    def test_signup_returns_to_next_page(self):
        r = self.client.post(reverse("signup"), {
            "username": "newone", "email": "newone@example.com", "password": "Strong-pass-123",
            "confirm_password": "Strong-pass-123", "next": "/checkout/",
        })
        self.assertEqual(r.url, "/checkout/")
        r = self.client.post(reverse("logout"))
        r = self.client.post(reverse("signup"), {
            "username": "evil", "email": "evil@example.com", "password": "Strong-pass-123",
            "confirm_password": "Strong-pass-123", "next": "https://evil.example/",
        })
        self.assertEqual(r.url, "/")

    def test_notifications_not_marked_read_by_prefetch(self):
        Notification.objects.create(user=self.buyer, message="hi", link="/")
        self.client.force_login(self.buyer)
        self.client.get(reverse("notifications_list"), HTTP_SEC_PURPOSE="prefetch;prerender")
        self.assertTrue(Notification.objects.filter(user=self.buyer, is_read=False).exists())
        self.client.get(reverse("notifications_list"))
        self.assertFalse(Notification.objects.filter(user=self.buyer, is_read=False).exists())

    def test_add_to_cart_requires_post(self):
        self.assertEqual(self.client.get(reverse("add_to_cart", args=[self.product.id])).status_code, 405)

    def test_cart_badge_ignores_removed_products(self):
        gone = make_product(name="Temp")
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        self.client.post(reverse("add_to_cart", args=[gone.id]))
        gone.delete()
        self.assertEqual(self.client.get(reverse("home")).context["cart_count"], 1)

    def test_cancelled_order_success_page(self):
        self.client.force_login(self.buyer)
        order = self._delivered_order()
        Order.objects.filter(pk=order.pk).update(status="cancelled")
        page = self.client.get(reverse("order_success", args=[order.id]))
        self.assertContains(page, "was cancelled")
        self.assertNotContains(page, "your order is confirmed")

    def test_guest_chat_kept_after_login(self):
        self.client.get(reverse("chat_messages"))
        self.client.post(reverse("chat_send"), {"message": "Where is my parcel?"})
        self.client.post(reverse("login"), {"username": "buyer", "password": "pass12345"})
        msgs = self.client.get(reverse("chat_messages")).json()["messages"]
        self.assertTrue(any("parcel" in m["message"] for m in msgs))
        self.assertEqual(ChatThread.objects.get(user=self.buyer).messages.filter(sender="user").count(), 1)


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy")
class AccountSettingsTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("acct", "acct@example.com", "Old-pass-123")
        self.client.force_login(self.user)

    # --- password & devices -------------------------------------------------
    def test_change_password(self):
        other = Client()
        other.force_login(self.user)
        r = self.client.post(reverse("account_security"), {"old_password": "wrong", "new_password1": "New-pass-456!", "new_password2": "New-pass-456!"})
        self.assertContains(r, "incorrectly")
        self.client.post(reverse("account_security"), {"old_password": "Old-pass-123", "new_password1": "New-pass-456!", "new_password2": "New-pass-456!"})
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password("New-pass-456!"))
        self.assertEqual(self.client.get(reverse("profile")).status_code, 200)  # still signed in here
        self.assertEqual(other.get(reverse("profile")).status_code, 302)  # signed out elsewhere
        self.assertTrue(any("password was changed" in m.subject.lower() for m in mail.outbox))

    def test_sign_out_other_devices(self):
        self.client.logout()
        self.client.post(reverse("login"), {"username": "acct", "password": "Old-pass-123", "remember": "1"})
        other = Client()
        other.post(reverse("login"), {"username": "acct", "password": "Old-pass-123", "remember": "1"})
        self.assertContains(self.client.get(reverse("account_security")), "1 other device")
        self.client.post(reverse("account_security"), {"action": "signout_others"})
        self.assertEqual(other.get(reverse("profile")).status_code, 302)
        self.assertEqual(self.client.get(reverse("profile")).status_code, 200)

    def test_remember_me(self):
        self.client.logout()
        self.client.post(reverse("login"), {"username": "acct", "password": "Old-pass-123"})
        self.assertTrue(self.client.session.get_expire_at_browser_close())
        self.client.logout()
        self.client.post(reverse("login"), {"username": "acct", "password": "Old-pass-123", "remember": "1"})
        self.assertFalse(self.client.session.get_expire_at_browser_close())

    def test_email_change_alerts_old_address(self):
        self.client.post(reverse("profile"), {"first_name": "A", "email": "new@example.com", "phone": ""})
        self.user.refresh_from_db()
        self.assertEqual(self.user.email, "acct@example.com")  # no password: not changed
        self.client.post(reverse("profile"), {"first_name": "A", "email": "new@example.com", "phone": "", "current_password": "Old-pass-123"})
        self.user.refresh_from_db()
        self.assertEqual(self.user.email, "new@example.com")
        self.assertTrue(any(m.to == ["acct@example.com"] and "email address was changed" in m.subject.lower() for m in mail.outbox))

    # --- saved cards ----------------------------------------------------------
    def test_cards_listed_added_and_removed(self):
        Profile.objects.create(user=self.user, referral_code="AC1", stripe_customer_id="cus_1")
        listing = {"data": [{"id": "pm_1", "card": {"brand": "visa", "last4": "4242", "exp_month": 4, "exp_year": 2030}}]}
        with mock.patch("stripe.Customer.list_payment_methods", return_value=listing):
            page = self.client.get(reverse("payment_methods"))
        self.assertContains(page, "Visa ending in 4242")
        self.assertContains(page, "Expires 04/2030")
        with mock.patch("stripe.checkout.Session.create", return_value=mock.MagicMock(url="https://checkout.stripe.com/setup")) as create:
            r = self.client.post(reverse("add_card"))
        self.assertEqual(r.url, "https://checkout.stripe.com/setup")
        self.assertEqual(create.call_args.kwargs["mode"], "setup")
        self.assertEqual(create.call_args.kwargs["customer"], "cus_1")
        # Someone else's card can't be removed
        with mock.patch("stripe.PaymentMethod.retrieve", return_value={"customer": "cus_OTHER"}), \
             mock.patch("stripe.PaymentMethod.detach") as detach:
            self.client.post(reverse("remove_card"), {"card": "pm_x"})
        detach.assert_not_called()
        with mock.patch("stripe.PaymentMethod.retrieve", return_value={"customer": "cus_1"}), \
             mock.patch("stripe.PaymentMethod.detach") as detach:
            self.client.post(reverse("remove_card"), {"card": "pm_1"})
        detach.assert_called_once_with("pm_1")

    def test_card_added_on_stripe_is_offered_again(self):
        Profile.objects.create(user=self.user, referral_code="AC1", stripe_customer_id="cus_1")
        session = {"customer": "cus_1", "status": "complete", "setup_intent": {"payment_method": "pm_9"}}
        with mock.patch("stripe.checkout.Session.retrieve", return_value=session), \
             mock.patch("stripe.Customer.list_payment_methods", return_value={"data": []}), \
             mock.patch("stripe.PaymentMethod.modify") as modify:
            r = self.client.get(reverse("payment_methods") + "?added=cs_setup_1", follow=True)
        modify.assert_called_once_with("pm_9", allow_redisplay="always")
        self.assertContains(r, "Card saved")

    def test_checkout_uses_saved_customer(self):
        product = make_product(price=Decimal("10.00"))
        self.client.post(reverse("add_to_cart", args=[product.id]))
        fake = mock.MagicMock(id="cs_1", url="https://checkout.stripe.com/pay")
        with mock.patch("stripe.checkout.Session.create", return_value=fake) as create:
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card"})
        kwargs = create.call_args.kwargs
        self.assertEqual(kwargs["customer"], "cus_test")
        self.assertEqual(kwargs["saved_payment_method_options"], {"payment_method_save": "enabled"})
        self.assertNotIn("customer_email", kwargs)
        self.assertEqual(Profile.objects.get(user=self.user).stripe_customer_id, "cus_test")

    @override_settings(STRIPE_SECRET_KEY="")
    def test_cards_page_without_stripe(self):
        self.assertContains(self.client.get(reverse("payment_methods")), "aren't switched on")

    # --- addresses ------------------------------------------------------------
    def test_edit_and_default_address_prefills_checkout(self):
        self.client.post(reverse("add_address"), {"label": "Home", "full_name": "A One", "address": "1 Road", "city": "Austin", "country": "US", "phone": "1"})
        self.client.post(reverse("add_address"), {"label": "Work", "full_name": "A Two", "address": "2 Office", "city": "Boston", "country": "US", "phone": "2"})
        home, work = Address.objects.order_by("id")
        self.assertTrue(home.is_default)
        self.client.post(reverse("edit_address", args=[work.id]), {"label": "Office", "full_name": "A Two", "address": "22 Office", "city": "Boston", "country": "US", "phone": "2", "is_default": "1"})
        work.refresh_from_db(); home.refresh_from_db()
        self.assertEqual((work.label, work.address, work.is_default, home.is_default), ("Office", "22 Office", True, False))
        product = make_product()
        self.client.post(reverse("add_to_cart", args=[product.id]))
        self.assertContains(self.client.get(reverse("checkout")), 'value="22 Office"')
        other = User.objects.create_user("other", "o@example.com", "pass12345")
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse("edit_address", args=[work.id])).status_code, 404)

    # --- privacy --------------------------------------------------------------
    def test_email_preferences(self):
        self.client.post(reverse("account_privacy"), {"action": "emails", "marketing": "1"})
        self.assertTrue(NewsletterSubscriber.objects.filter(email="acct@example.com").exists())
        self.client.post(reverse("account_privacy"), {"action": "emails"})
        self.assertFalse(NewsletterSubscriber.objects.exists())

    def test_download_data(self):
        order = Order.objects.create(user=self.user, full_name="A", address="x", city="y", phone="1", status="delivered")
        OrderItem.objects.create(order=order, product_name="Face Wash", price=Decimal("5"), quantity=1)
        r = self.client.get(reverse("download_data"))
        data = json.loads(r.content)
        self.assertEqual(data["account"]["email"], "acct@example.com")
        self.assertEqual(data["orders"][0]["items"][0]["product"], "Face Wash")

    def test_delete_account(self):
        r = self.client.post(reverse("delete_account"), {"password": "nope"}, follow=True)
        self.assertContains(r, "isn&#x27;t right")
        open_order = Order.objects.create(user=self.user, full_name="A", address="x", city="y", phone="1", status="shipped")
        self.client.post(reverse("delete_account"), {"password": "Old-pass-123"})
        self.assertTrue(User.objects.get(pk=self.user.pk).is_active)
        Order.objects.filter(pk=open_order.pk).update(status="delivered")
        Address.objects.create(user=self.user, full_name="A", phone="1", address="x", city="y", country="US")
        Profile.objects.update_or_create(user=self.user, defaults={"referral_code": "AC1", "stripe_customer_id": "cus_1"})
        with mock.patch("stripe.Customer.delete") as delete:
            self.client.post(reverse("delete_account"), {"password": "Old-pass-123"})
        delete.assert_called_once_with("cus_1")
        u = User.objects.get(pk=self.user.pk)
        self.assertFalse(u.is_active)
        self.assertEqual(u.email, "")
        self.assertFalse(Address.objects.filter(user=u).exists())
        self.assertTrue(Order.objects.filter(pk=open_order.pk).exists())  # sales records kept
        self.assertFalse(Client().login(username="acct", password="Old-pass-123"))
        self.assertEqual(self.client.get(reverse("profile")).status_code, 302)  # signed out

    def test_sellers_cannot_self_delete(self):
        SellerAccount.objects.create(user=self.user, status="approved", business_name="S")
        self.client.post(reverse("delete_account"), {"password": "Old-pass-123"})
        self.assertTrue(User.objects.get(pk=self.user.pk).is_active)

    # --- seller store settings ------------------------------------------------
    def test_store_settings_rename_updates_products(self):
        seller = SellerAccount.objects.create(user=self.user, status="approved", business_name="Old Name")
        p = make_product(seller_account=seller, seller_name="Old Name")
        self.client.post(reverse("seller_store_settings"), {"store_name": "New Name", "store_description": "Hello", "bank_details": "IBAN 123"})
        seller.refresh_from_db(); p.refresh_from_db()
        self.assertEqual((seller.business_name, seller.bank_details, p.seller_name), ("New Name", "IBAN 123", "New Name"))
        SellerAccount.objects.create(user=User.objects.create_user("z", "z@example.com", "x"), business_name="Taken")
        self.client.post(reverse("seller_store_settings"), {"store_name": "taken"})
        seller.refresh_from_db()
        self.assertEqual(seller.business_name, "New Name")

    def test_team_staff_cannot_change_store_settings(self):
        owner = User.objects.create_user("own", "own@example.com", "x")
        org = SellerAccount.objects.create(user=owner, status="approved", business_name="Org", account_type="organization")
        OrganizationMember.objects.create(organization=org, user=self.user, role="staff")
        self.client.post(reverse("seller_store_settings"), {"store_name": "Hijack"})
        org.refresh_from_db()
        self.assertEqual(org.business_name, "Org")


class ShippingZoneTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("ship", "ship@example.com", "pass12345")
        self.product = make_product(price=Decimal("20.00"), stock=10)
        self.client.force_login(self.user)

    def _checkout(self, country):
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        return self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "country": country, "payment_method": "cod"}, follow=True)

    def test_flat_fee_without_zones(self):
        SiteSettings.objects.update_or_create(pk=1, defaults={"shipping_flat_fee": Decimal("7")})
        cache.clear()
        self._checkout("DE")
        self.assertEqual(Order.objects.get().shipping_amount, Decimal("7"))

    def test_fee_by_country_and_unsupported_country(self):
        from .models import ShippingZone
        ShippingZone.objects.create(name="US", countries="US", fee=Decimal("5"), free_over=Decimal("50"), delivery_days=3)
        rest = ShippingZone.objects.create(name="World", countries="*", fee=Decimal("20"), delivery_days=12)
        self._checkout("US")
        us = Order.objects.get()
        self.assertEqual(us.shipping_amount, Decimal("5"))
        from .order_emails import add_business_days
        from django.utils import timezone
        self.assertEqual(us.estimated_delivery, add_business_days(timezone.localdate(), 3))
        self._checkout("DE")
        self.assertEqual(Order.objects.latest("id").shipping_amount, Decimal("20"))
        rest.delete()
        r = self._checkout("DE")
        self.assertContains(r, "don&#x27;t deliver to Germany")
        self.assertEqual(Order.objects.count(), 2)

    def test_free_over_threshold(self):
        from .models import ShippingZone
        ShippingZone.objects.create(name="US", countries="US", fee=Decimal("5"), free_over=Decimal("50"))
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 3})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "country": "US", "payment_method": "cod"})
        self.assertEqual(Order.objects.get().shipping_amount, 0)

    def test_checkout_page_has_shipping_table(self):
        from .models import ShippingZone
        ShippingZone.objects.create(name="US", countries="US, CA", fee=Decimal("5"))
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        page = self.client.get(reverse("checkout"))
        self.assertContains(page, '"CA": {"fee": "5.00"')
        self.assertContains(page, "Choose a country")

    def test_admin_manages_zones(self):
        from .models import ShippingZone
        staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        self.client.force_login(staff)
        r = self.client.post(reverse("manage_shipping"), {"action": "save", "name": "Bad", "countries": "US, XX", "fee": "5"})
        self.assertContains(r, "Unknown country code(s): XX")
        self.client.post(reverse("manage_shipping"), {"action": "save", "name": "North America", "countries": "us,ca, us", "fee": "6", "active": "on"})
        self.assertEqual(ShippingZone.objects.get().countries, "US, CA")
        ShippingZone.objects.all().delete()
        self.client.post(reverse("manage_shipping"), {"action": "starter"})
        self.assertEqual(ShippingZone.objects.count(), 3)
        self.assertEqual(self.client.get(reverse("manage_shipping")).status_code, 200)


class VariantTests(TestCase):
    def setUp(self):
        self.seller_user = User.objects.create_user("vs", "vs@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.seller_user, status="approved", business_name="Hood Co")
        self.buyer = User.objects.create_user("vb", "vb@example.com", "pass12345")

    def _form(self, **extra):
        data = {"name": "Hoodie", "category": "hoodies", "price": "40.00", "stock": "0", "image_url": "https://example.com/h.jpg",
                "variant_id": ["", ""], "variant_size": ["M", "L"], "variant_color": ["Black", "Black"],
                "variant_price": ["", "45.00"], "variant_stock": ["3", "0"], "variant_sku": ["H-M", "H-L"]}
        data.update(extra)
        return data

    def _product(self):
        self.client.force_login(self.seller_user)
        self.client.post(reverse("seller_add_product"), self._form())
        p = Product.objects.get(name="Hoodie")
        Product.objects.filter(pk=p.pk).update(approval_status="approved")
        p.refresh_from_db()
        return p

    def test_seller_creates_variants(self):
        p = self._product()
        self.assertTrue(p.has_variants)
        self.assertEqual(p.stock, 3)
        self.assertEqual([v.label for v in p.variants.all()], ["M / Black", "L / Black"])
        self.assertEqual(self.client.get(reverse("seller_edit_product", args=[p.id])).status_code, 200)

    def test_duplicate_options_rejected(self):
        self.client.force_login(self.seller_user)
        r = self.client.post(reverse("seller_add_product"), self._form(variant_size=["M", "m"]))
        self.assertContains(r, "listed twice")
        self.assertFalse(Product.objects.exists())

    def test_buy_variant_and_cancel_restocks(self):
        p = self._product()
        m, large = p.variants.all()
        self.client.force_login(self.buyer)
        page = self.client.get(reverse("product_detail", args=[p.id]))
        self.assertContains(page, 'data-value="M"')
        self.client.post(reverse("add_to_cart", args=[p.id]))  # no size chosen
        self.assertEqual(self.client.session.get("cart", {}), {})
        self.client.post(reverse("add_to_cart", args=[p.id]), {"variant": large.id})  # sold out
        self.assertEqual(self.client.session.get("cart", {}), {})
        self.client.post(reverse("add_to_cart", args=[p.id]), {"variant": m.id, "quantity": 5})
        self.assertEqual(self.client.session["cart"], {f"{p.id}-{m.id}": 3})  # capped at stock
        self.assertContains(self.client.get(reverse("cart")), "M / Black")
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        item = OrderItem.objects.get()
        self.assertEqual((item.product_name, item.variant, item.quantity, item.price), ("Hoodie (M / Black)", m, 3, Decimal("40.00")))
        m.refresh_from_db(); p.refresh_from_db()
        self.assertEqual((m.stock, p.stock), (0, 0))
        self.client.post(reverse("cancel_order", args=[item.order_id]))
        m.refresh_from_db(); p.refresh_from_db()
        self.assertEqual((m.stock, p.stock), (3, 3))

    def test_variant_price_used(self):
        p = self._product()
        large = p.variants.get(size="L")
        ProductVariant.objects.filter(pk=large.pk).update(stock=2)
        p.sync_variants()
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[p.id]), {"variant": large.id})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertEqual(OrderItem.objects.get().price, Decimal("45.00"))

    def test_admin_removes_variant(self):
        p = self._product()
        staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        self.client.force_login(staff)
        m = p.variants.get(size="M")
        self.client.post(reverse("manage_product_edit", args=[p.id]), {
            "name": "Hoodie", "category": "hoodies", "price": "40", "stock": "0", "image_url": p.image_url,
            "approval_status": "approved", "variant_id": [m.id], "variant_size": ["M"], "variant_color": ["Black"],
            "variant_price": [""], "variant_stock": ["7"], "variant_sku": [""],
        })
        p.refresh_from_db()
        self.assertEqual(p.variants.count(), 1)
        self.assertEqual(p.stock, 7)

    def test_buy_again_keeps_size(self):
        p = self._product()
        m = p.variants.get(size="M")
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[p.id]), {"variant": m.id})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.client.post(reverse("buy_again", args=[Order.objects.get().id]))
        self.assertEqual(list(self.client.session["cart"]), [f"{p.id}-{m.id}"])


class StockAlertTests(TestCase):
    def test_notified_once_when_back_in_stock(self):
        from .models import StockAlert
        product = make_product(name="Rare Serum", stock=0)
        self.client.post(reverse("stock_alert", args=[product.id]), {"email": "Fan@Example.com"})
        self.client.post(reverse("stock_alert", args=[product.id]), {"email": "fan@example.com"})
        self.assertEqual(StockAlert.objects.count(), 1)
        self.assertContains(self.client.get(reverse("product_detail", args=[product.id])), "Notify me")
        with self.captureOnCommitCallbacks(execute=True):
            product.stock = 4
            product.save()
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("back in stock", mail.outbox[0].subject)
        self.assertEqual(mail.outbox[0].to, ["fan@example.com"])
        with self.captureOnCommitCallbacks(execute=True):
            product.stock = 5
            product.save()
        self.assertEqual(len(mail.outbox), 1)

    def test_variant_alert_waits_for_that_variant(self):
        product = make_product(name="Tee", stock=0)
        small = ProductVariant.objects.create(product=product, size="S", stock=0)
        large = ProductVariant.objects.create(product=product, size="L", stock=0)
        product.sync_variants()
        self.client.post(reverse("stock_alert", args=[product.id]), {"email": "a@example.com", "variant": small.id})
        with self.captureOnCommitCallbacks(execute=True):
            large.stock = 3
            large.save()
        self.assertEqual(len(mail.outbox), 0)
        with self.captureOnCommitCallbacks(execute=True):
            small.stock = 1
            small.save()
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("Tee (S)", mail.outbox[0].subject)


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy")
class GuestCheckoutTests(TestCase):
    """Guest checkout is off by default; these cover stores that switch it on."""
    def setUp(self):
        SiteSettings.objects.update_or_create(pk=1, defaults={"allow_guest_checkout": True})
        cache.clear()
        self.product = make_product(name="Guest Serum", price=Decimal("15.00"), stock=5)

    def test_guest_must_sign_in_when_switched_off(self):
        SiteSettings.objects.filter(pk=1).update(allow_guest_checkout=False)
        cache.clear()
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        r = self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertTrue(r.url.startswith(reverse("login")))
        self.assertFalse(Order.objects.exists())
        self.assertContains(self.client.get(reverse("cart")), "Sign in to check out")
        user = User.objects.create_user("later", "later@example.com", "pass12345")
        self.client.post(reverse("login"), {"username": "later", "password": "pass12345"})
        self.assertContains(self.client.get(reverse("cart")), "Guest Serum")  # cart kept through sign-in

    def _guest_cod(self, client=None):
        c = client or self.client
        c.post(reverse("add_to_cart", args=[self.product.id]))
        with self.captureOnCommitCallbacks(execute=True):
            r = c.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        return Order.objects.latest("id"), r

    def test_guest_places_and_tracks_order(self):
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        self.assertContains(self.client.get(reverse("checkout")), "Checking out as a guest")
        order, r = self._guest_cod()
        self.assertIsNone(order.user)
        self.assertEqual(order.guest_email, "jane@example.com")
        self.assertRedirects(r, reverse("order_success", args=[order.id]))
        self.assertContains(self.client.get(r.url), "Create an account")
        email = mail.outbox[0].alternatives[0][0]
        self.assertIn(order.tracking_path(), email)
        # Anyone else needs the private link
        stranger = Client()
        self.assertEqual(stranger.get(reverse("order_success", args=[order.id])).status_code, 404)
        self.assertEqual(stranger.get(reverse("order_track", args=[order.id, "nope"])).status_code, 404)
        page = stranger.get(order.tracking_path())
        self.assertContains(page, "Guest Serum")
        # ... and can cancel from it
        stranger.post(reverse("cancel_order", args=[order.id]))
        order.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock, 5)

    def test_guest_cannot_touch_member_orders(self):
        member = User.objects.create_user("mem", "mem@example.com", "pass12345")
        order = Order.objects.create(user=member, full_name="M", address="x", city="y", phone="1")
        self.client.post(reverse("cancel_order", args=[order.id]))
        order.refresh_from_db()
        self.assertEqual(order.status, "pending")
        r = self.client.get(order.tracking_path())
        self.assertEqual(r.status_code, 302)

    def test_guest_card_payment_uses_email_not_customer(self):
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        fake = mock.MagicMock(id="cs_g", url="https://checkout.stripe.com/g")
        with mock.patch("stripe.checkout.Session.create", return_value=fake) as create:
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card"})
        kwargs = create.call_args.kwargs
        self.assertEqual(kwargs["customer_email"], "jane@example.com")
        self.assertNotIn("customer", kwargs)
        order = Order.objects.get()
        session = {"id": "cs_g", "payment_status": "paid", "status": "complete", "amount_total": order.total_cents,
                   "currency": order.currency, "payment_intent": "pi_g", "metadata": {"order_id": str(order.id)}}
        with mock.patch("stripe.checkout.Session.retrieve", return_value=mock.MagicMock(to_dict=lambda: session)):
            r = self.client.get(reverse("payment_success") + "?session_id=cs_g")
        self.assertRedirects(r, reverse("order_success", args=[order.id]), fetch_redirect_response=False)
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "paid")

    def test_guest_coupon_limit_by_email(self):
        Coupon.objects.create(code="ONCE", percent_off=10, per_user_limit=1)
        self.client.post(reverse("apply_coupon"), {"coupon_code": "ONCE"})
        self._guest_cod()
        other = Client()
        other.post(reverse("add_to_cart", args=[self.product.id]))
        other.post(reverse("apply_coupon"), {"coupon_code": "ONCE"})
        other.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertEqual(Order.objects.count(), 1)  # same email can't reuse it

    def test_guest_orders_join_account_after_email_verified(self):
        order, _ = self._guest_cod()
        user = User.objects.create_user("jane", "jane@example.com", "pass12345")
        self.client.force_login(user)
        from django.core.cache import cache as _cache
        Profile.objects.create(user=user, referral_code="JN1")
        _cache.set(f"email_verify_code:{user.id}", "123456", 900)
        self.client.post(reverse("verify_email"), {"code": "123456"})
        order.refresh_from_db()
        self.assertEqual(order.user, user)
        self.assertContains(self.client.get(reverse("my_orders")), "Guest Serum")


class SavedCartTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("cart", "cart@example.com", "pass12345")
        self.product = make_product(name="Lip Balm", price=Decimal("4.00"), stock=9)

    def test_cart_follows_account_to_another_device(self):
        phone = Client()
        phone.post(reverse("login"), {"username": "cart", "password": "pass12345"})
        phone.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 2})
        laptop = Client()
        laptop.post(reverse("add_to_cart", args=[make_product(name="Other").id]))  # browsing as a visitor
        laptop.post(reverse("login"), {"username": "cart", "password": "pass12345"})
        self.assertEqual(laptop.session["cart"][str(self.product.id)], 2)
        self.assertEqual(len(laptop.session["cart"]), 2)
        self.assertEqual(len(Profile.objects.get(user=self.user).saved_cart), 2)

    def test_reminder_sent_once_after_a_day(self):
        from datetime import timedelta
        from django.core.management import call_command
        from django.utils import timezone
        self.client.force_login(self.user)
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        call_command("send_cart_reminders", stdout=open(os.devnull, "w"))
        self.assertEqual(len(mail.outbox), 0)  # too soon
        Profile.objects.filter(user=self.user).update(cart_updated_at=timezone.now() - timedelta(hours=25))
        call_command("send_cart_reminders", stdout=open(os.devnull, "w"))
        call_command("send_cart_reminders", stdout=open(os.devnull, "w"))
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn("Lip Balm", mail.outbox[0].alternatives[0][0])

    def test_no_reminder_when_opted_out_or_ordered(self):
        from datetime import timedelta
        from django.core.management import call_command
        from django.utils import timezone
        self.client.force_login(self.user)
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        self.client.post(reverse("account_privacy"), {"action": "emails"})  # both boxes unticked
        Profile.objects.filter(user=self.user).update(cart_updated_at=timezone.now() - timedelta(hours=30))
        call_command("send_cart_reminders", stdout=open(os.devnull, "w"))
        self.assertEqual(len(mail.outbox), 0)
        Profile.objects.filter(user=self.user).update(cart_reminders=True)
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertEqual(Profile.objects.get(user=self.user).saved_cart, {})


from django.test import TransactionTestCase  # noqa: E402


class BackupTests(TransactionTestCase):
    """Not wrapped in a transaction: SQLite can't copy a database while the
    test's own open transaction holds its write lock."""

    def test_backup_and_list(self):
        import tempfile
        from django.core.management import call_command
        with tempfile.TemporaryDirectory() as tmp, override_settings(BASE_DIR=tmp, MEDIA_ROOT=os.path.join(tmp, "m")):
            for _ in range(3):
                call_command("backup_data", "--keep", "2", stdout=open(os.devnull, "w"))
            files = os.listdir(os.path.join(tmp, "backups"))
            self.assertTrue(1 <= len(files) <= 2)
            self.assertTrue(all(f.startswith("backup-") and f.endswith(".tar.gz") for f in files))


class TwoStepSignInTests(TestCase):
    def setUp(self):
        SiteSettings.objects.update_or_create(pk=1, defaults={"require_staff_2fa": True})
        cache.clear()
        self.staff = User.objects.create_user("owner", "owner@example.com", "Owner-pass-1", is_staff=True)

    def _now_code(self, secret, offset=0):
        from . import twofactor
        return twofactor.code_at(secret, twofactor.current_step() + offset)

    def _enable(self, client):
        client.get(reverse("two_factor_setup"))
        secret = client.session["totp_setup_secret"]
        r = client.post(reverse("two_factor_setup"), {"code": self._now_code(secret)})
        return secret, r

    def test_staff_must_set_up_before_admin(self):
        self.client.force_login(self.staff)
        self.assertRedirects(self.client.get(reverse("manage_dashboard")), reverse("two_factor_setup"), fetch_redirect_response=False)
        self.assertRedirects(self.client.get("/admin/"), reverse("two_factor_setup"), fetch_redirect_response=False)
        page = self.client.get(reverse("two_factor_setup"))
        self.assertContains(page, "<svg")
        r = self.client.post(reverse("two_factor_setup"), {"code": "000000"})
        self.assertContains(r, "didn&#x27;t match")
        secret, r = self._enable(self.client)
        self.assertContains(r, "Save your backup codes")
        self.assertEqual(self.client.get(reverse("manage_dashboard")).status_code, 200)
        self.assertTrue(Profile.objects.get(user=self.staff).totp_enabled)

    def test_sign_in_needs_code(self):
        self.client.force_login(self.staff)
        secret, r = self._enable(self.client)
        codes = r.context["codes"]
        self.client.logout()
        r = self.client.post(reverse("login"), {"username": "owner", "password": "Owner-pass-1", "next": "/manage/"})
        self.assertRedirects(r, reverse("login_code"), fetch_redirect_response=False)
        self.assertNotIn("_auth_user_id", self.client.session)  # not signed in yet
        self.client.post(reverse("login_code"), {"code": "123456"})
        self.assertNotIn("_auth_user_id", self.client.session)
        r = self.client.post(reverse("login_code"), {"code": self._now_code(secret, 1)})
        self.assertRedirects(r, "/manage/", fetch_redirect_response=False)
        self.assertEqual(self.client.get("/manage/").status_code, 200)
        # The same code can't be used again
        other = Client()
        other.post(reverse("login"), {"username": "owner", "password": "Owner-pass-1"})
        other.post(reverse("login_code"), {"code": self._now_code(secret, 1)})
        self.assertNotIn("_auth_user_id", other.session)
        # A backup code works once
        other.post(reverse("login"), {"username": "owner", "password": "Owner-pass-1"})
        other.post(reverse("login_code"), {"code": codes[0]})
        self.assertIn("_auth_user_id", other.session)
        third = Client()
        third.post(reverse("login"), {"username": "owner", "password": "Owner-pass-1"})
        third.post(reverse("login_code"), {"code": codes[0]})
        self.assertNotIn("_auth_user_id", third.session)

    def test_django_admin_login_goes_through_store_login(self):
        r = self.client.get("/admin/login/?next=/admin/")
        self.assertEqual(r.status_code, 302)
        self.assertTrue(r.url.startswith("/login/"))

    def test_reset_command(self):
        from django.core.management import call_command
        self.client.force_login(self.staff)
        self._enable(self.client)
        call_command("reset_two_factor", "owner@example.com", stdout=open(os.devnull, "w"))
        self.assertFalse(Profile.objects.get(user=self.staff).totp_enabled)

    def test_customers_can_opt_in_but_are_not_forced(self):
        buyer = User.objects.create_user("cust", "cust@example.com", "Cust-pass-1")
        self.client.force_login(buyer)
        self.assertContains(self.client.get(reverse("account_security")), "Set up")
        secret, _ = self._enable(self.client)
        self.client.post(reverse("two_factor_disable"), {"password": "Cust-pass-1", "code": self._now_code(secret, 1)})
        self.assertFalse(Profile.objects.get(user=buyer).totp_enabled)


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy")
class SystemCheckTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        self.client.force_login(self.staff)

    def test_page_and_live_tests(self):
        page = self.client.get(reverse("manage_system"))
        self.assertContains(page, "System check")
        self.assertContains(page, "Database")
        r = self.client.post(reverse("manage_system"), {"action": "email", "to": "me@example.com"}, follow=True)
        self.assertContains(r, "Test email sent to me@example.com")
        self.assertEqual(mail.outbox[-1].to, ["me@example.com"])
        with mock.patch("stripe.Balance.retrieve", return_value={}):
            r = self.client.post(reverse("manage_system"), {"action": "stripe"}, follow=True)
        self.assertContains(r, "Stripe connection works")
        import stripe as _stripe
        err = _stripe.error.AuthenticationError("Invalid API Key provided", http_status=401)
        with mock.patch("stripe.Balance.retrieve", side_effect=err):
            r = self.client.post(reverse("manage_system"), {"action": "stripe"}, follow=True)
        self.assertContains(r, "rejected the secret key")
        r = self.client.post(reverse("manage_system"), {"action": "storage"}, follow=True)
        self.assertContains(r, "Image storage works")

    def test_email_failures_are_recorded(self):
        from .views import _send_html_email
        with mock.patch("django.core.mail.EmailMultiAlternatives.send", side_effect=OSError("SMTP down")):
            _send_html_email("Hello", "bees/emails/security_notice.html", {"user": self.staff, "event": "x", "detail": "y", "reset_url": "/"}, "a@example.com")
        self.assertTrue(AuditLog.objects.filter(action__contains="SMTP down").exists())
        self.assertContains(self.client.get(reverse("manage_system")), "SMTP down")

    def _hook(self, typ, obj):
        payload = json.dumps({"type": typ, "data": {"object": obj}})
        with mock.patch("stripe.Webhook.construct_event", return_value={}):
            return self.client.post(reverse("stripe_webhook"), data=payload, content_type="application/json", HTTP_STRIPE_SIGNATURE="t=1,v1=x")

    def test_refund_made_in_stripe_dashboard_updates_order(self):
        order = Order.objects.create(user=self.staff, full_name="A", address="x", city="y", phone="1", status="confirmed",
                                     payment_method="card", payment_status="paid", stripe_payment_intent="pi_dash")
        self.assertEqual(self._hook("charge.refunded", {"payment_intent": "pi_dash", "refunded": True, "amount": 1000}).status_code, 200)
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "refunded")
        self.assertTrue(Notification.objects.filter(user=self.staff, message__icontains="refunded in Stripe").exists())
        self.assertIsNotNone(SiteSettings.objects.get(pk=1).stripe_last_webhook)

    def test_dispute_alerts_staff(self):
        self._hook("charge.dispute.created", {"payment_intent": "pi_x", "amount": 2500})
        self.assertTrue(Notification.objects.filter(user=self.staff, message__icontains="disputed").exists())

    def test_failed_refund_alerts_staff(self):
        buyer = User.objects.create_user("b", "b@example.com", "pass12345")
        order = Order.objects.create(user=buyer, full_name="A", address="x", city="y", phone="1", status="confirmed",
                                     payment_method="card", payment_status="paid", stripe_payment_intent="pi_f")
        OrderItem.objects.create(order=order, product_name="X", price=Decimal("5"), quantity=1)
        c = Client(); c.force_login(buyer)
        with mock.patch("stripe.Refund.create", side_effect=Exception("card_declined")):
            r = c.post(reverse("cancel_order", args=[order.id]), follow=True)
        self.assertContains(r, "team has been alerted")
        self.assertTrue(Notification.objects.filter(user=self.staff, message__icontains="refund failed").exists())
        order.refresh_from_db()
        self.assertEqual(order.status, "confirmed")


class HelpAssistantTests(TestCase):
    def test_topics_answer_from_store_settings(self):
        from .models import ShippingZone
        ShippingZone.objects.create(name="Europe", countries="DE, FR", fee=Decimal("12"), delivery_days=8)
        SiteSettings.objects.update_or_create(pk=1, defaults={"return_days": 30})
        cache.clear()
        d = self.client.get(reverse("chat_topic") + "?key=start").json()
        self.assertIn("help assistant", d["message"])
        self.assertTrue(any(o["key"] == "human" for o in d["options"]))
        self.assertIn("Europe: $12.00", self.client.get(reverse("chat_topic") + "?key=shipping").json()["message"])
        self.assertIn("within 30 days", self.client.get(reverse("chat_topic") + "?key=returns").json()["message"])

    def test_track_shows_own_orders(self):
        user = User.objects.create_user("t", "t@example.com", "pass12345")
        Order.objects.create(user=user, full_name="T", address="x", city="y", phone="1", status="shipped", tracking_number="1Z9")
        self.client.force_login(user)
        msg = self.client.get(reverse("chat_topic") + "?key=track").json()["message"]
        self.assertIn("Shipped", msg)
        self.assertIn("1Z9", msg)

    def test_typed_question_gets_instant_answer_and_reaches_team(self):
        r = self.client.post(reverse("chat_send"), {"message": "How do I return a broken item?"}).json()
        self.assertIn("return items within", r["reply"]["message"])
        self.assertTrue(r["reply"]["auto"])
        r = self.client.post(reverse("chat_send"), {"message": "I want to talk to someone about bulk orders"}).json()
        self.assertIn("member of our team will reply", r["reply"]["message"])
        thread = ChatThread.objects.get()
        self.assertEqual(thread.messages.filter(sender="user").count(), 2)

    def test_team_reply_badge_notification_and_email(self):
        user = User.objects.create_user("c", "c@example.com", "pass12345")
        staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        self.client.force_login(user)
        self.client.post(reverse("chat_send"), {"message": "Hello?"})
        thread = ChatThread.objects.get(user=user)
        s = Client(); s.force_login(staff)
        s.post(reverse("manage_support_thread", args=[thread.id]), {"action": "reply", "message": "Hi! How can we help?"})
        self.assertEqual(self.client.get(reverse("chat_unread")).json()["unread"], 1)
        self.assertTrue(any(m.to == ["c@example.com"] and "reply" in m.subject.lower() for m in mail.outbox))
        self.client.get(reverse("chat_messages"))
        self.assertEqual(self.client.get(reverse("chat_unread")).json()["unread"], 0)


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy")
class PaymentPageExpiryTests(TestCase):
    def test_expiry_safely_above_stripe_minimum(self):
        import time as _time
        user = User.objects.create_user("exp", "exp@example.com", "pass12345")
        self.client.force_login(user)
        product = make_product(price=Decimal("10.00"))
        self.client.post(reverse("add_to_cart", args=[product.id]))
        fake = mock.MagicMock(id="cs_e", url="https://checkout.stripe.com/e")
        with mock.patch("stripe.checkout.Session.create", return_value=fake) as create:
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card"})
        # Stripe rejects sessions that expire in under 30 minutes when they arrive.
        self.assertGreaterEqual(create.call_args.kwargs["expires_at"] - int(_time.time()), 45 * 60)


class SellerCenterTests(TestCase):
    """The Seller Center: numbers, orders, products, payouts and team access."""

    def setUp(self):
        from django.core import mail as _mail
        self.mail = _mail
        owner = User.objects.create_user("shop", "shop@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, account_type="organization", business_name="Shop Co",
                                                   bank_details="IBAN GB00 TEST")
        self.seller.status = "approved"
        self.seller.save()
        self.product = make_product(name="Lamp", price=Decimal("50.00"), stock=20,
                                    seller_account=self.seller, seller_name="Shop Co")
        self.buyer = User.objects.create_user("buyer", "buyer@example.com", "pass12345")

    def _order(self, qty=2, method="cod"):
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": qty})
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": method})
        order = Order.objects.filter(user=self.buyer).latest("id")
        self.client.force_login(self.seller.user)
        return order

    def test_new_cod_order_alerts_seller_once(self):
        order = self._order()
        alerts = Notification.objects.filter(user=self.seller.user, message__icontains=f"New order #{order.id}")
        self.assertEqual(alerts.count(), 1)
        self.assertTrue(any("new order" in m.subject.lower() and "shop@example.com" in m.to for m in self.mail.outbox))
        from . import seller_center
        seller_center.notify_new_order(order)  # already done: no second alert
        self.assertEqual(alerts.count(), 1)
        order.refresh_from_db()
        self.assertTrue(order.sellers_notified)

    def test_unpaid_card_order_does_not_alert_seller_until_paid(self):
        from . import seller_center
        order = self._order()
        Notification.objects.filter(user=self.seller.user).delete()
        Order.objects.filter(pk=order.pk).update(payment_method="card", payment_status="pending", sellers_notified=False)
        order.refresh_from_db()
        seller_center.notify_new_order(order)
        self.assertFalse(Notification.objects.filter(user=self.seller.user, message__icontains="New order").exists())
        Order.objects.filter(pk=order.pk).update(payment_status="paid", status="confirmed")
        order.refresh_from_db()
        seller_center.notify_new_order(order)
        self.assertTrue(Notification.objects.filter(user=self.seller.user, message__icontains=f"New order #{order.id}").exists())

    def test_cancelling_tells_seller_not_to_ship(self):
        order = self._order()
        self.client.force_login(self.buyer)
        self.client.post(reverse("cancel_order", args=[order.id]))
        self.assertTrue(Notification.objects.filter(user=self.seller.user, message__icontains="cancelled").exists())

    def test_overview_shows_sales_commission_and_earnings(self):
        self._order(qty=2)  # $100 sale, 20% commission -> $80
        page = self.client.get(reverse("seller_dashboard") + "?range=7")
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "$100.00")
        self.assertContains(page, "$20.00")
        self.assertContains(page, "$80.00")
        self.assertContains(page, "United States")
        for key in ("30", "90", "365", "bogus"):
            self.assertEqual(self.client.get(reverse("seller_dashboard") + f"?range={key}").status_code, 200)

    def test_every_seller_center_page_loads(self):
        order = self._order()
        for name, args in [
            ("seller_dashboard", []), ("seller_orders", []), ("seller_order", [order.id]),
            ("seller_packing_slip", [order.id]), ("seller_products", []), ("seller_earnings", []),
            ("seller_returns", []), ("seller_reviews", []), ("seller_team", []), ("seller_store_settings", []),
            ("seller_add_product", []), ("seller_edit_product", [self.product.id]), ("store_page", ["Shop Co"]),
        ]:
            with self.subTest(page=name):
                self.assertEqual(self.client.get(reverse(name, args=args)).status_code, 200)
        for tab in ("to_pack", "packed", "in_transit", "delivered", "unpaid", "cancelled", "all"):
            self.assertEqual(self.client.get(reverse("seller_orders") + f"?tab={tab}&q=Jane").status_code, 200)
        for tab in ("all", "live", "review", "rejected", "low", "out"):
            self.assertEqual(self.client.get(reverse("seller_products") + f"?tab={tab}&sort=stock").status_code, 200)
        for tab in ("questions", "answered", "reviews", "store"):
            self.assertEqual(self.client.get(reverse("seller_reviews") + f"?tab={tab}").status_code, 200)

    def test_ship_order_with_tracking_moves_order_and_emails_customer(self):
        order = self._order()
        self.mail.outbox.clear()
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("seller_ship_order", args=[order.id]), {
                "status": "handed_to_courier", "courier_name": "DHL", "tracking_number": "JD0001",
            })
        order.refresh_from_db()
        self.assertEqual(order.status, "shipped")
        self.assertEqual(order.tracking_number, "JD0001")
        self.assertTrue(order.items.filter(fulfillment_status="handed_to_courier").exists())
        self.assertTrue(any("on its way" in m.subject for m in self.mail.outbox))

    def test_bulk_mark_packed(self):
        o1, o2 = self._order(qty=1), self._order(qty=1)
        self.client.post(reverse("seller_bulk_fulfilment"), {"status": "packed", "order": [o1.id, o2.id]})
        self.assertEqual(OrderItem.objects.filter(order__in=[o1, o2], fulfillment_status="packed").count(), 2)
        OrderItem.objects.filter(order=o1).update(fulfillment_status="delivered")
        self.client.post(reverse("seller_bulk_fulfilment"), {"status": "packed", "order": [o1.id]})
        self.assertTrue(OrderItem.objects.filter(order=o1, fulfillment_status="delivered").exists())  # never moved back

    def test_cannot_touch_another_sellers_order(self):
        order = self._order()
        other = User.objects.create_user("other", "o@example.com", "pass12345")
        acct = SellerAccount.objects.create(user=other, business_name="Other")
        acct.status = "approved"
        acct.save()
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse("seller_order", args=[order.id])).status_code, 404)
        self.assertEqual(self.client.get(reverse("seller_packing_slip", args=[order.id])).status_code, 404)
        self.client.post(reverse("seller_ship_order", args=[order.id]), {"status": "delivered"})
        self.client.post(reverse("seller_bulk_fulfilment"), {"status": "delivered", "order": [order.id]})
        self.assertFalse(OrderItem.objects.filter(order=order).exclude(fulfillment_status="pending").exists())
        self.client.post(reverse("seller_quick_update", args=[self.product.id]), {"price": "1", "stock": "0"})
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, Decimal("50.00"))

    def test_csv_exports(self):
        self._order()
        r = self.client.get(reverse("seller_orders") + "?tab=all&export=csv")
        self.assertEqual(r["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("Lamp", r.content.decode())
        r = self.client.get(reverse("seller_earnings") + "?export=csv")
        body = r.content.decode()
        self.assertIn("Statement for Shop Co", body)
        self.assertIn("80.00", body)

    def test_csv_neutralises_formulas(self):
        self.product.name = "=HYPERLINK(\"http://x\")"
        self.product.save()
        self._order()
        body = self.client.get(reverse("seller_orders") + "?tab=all&export=csv").content.decode()
        self.assertIn("'=HYPERLINK", body)

    def test_quick_update_price_and_stock(self):
        self.client.force_login(self.seller.user)
        self.client.post(reverse("seller_quick_update", args=[self.product.id]), {"price": "42.50", "stock": "3"})
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, Decimal("42.50"))
        self.assertEqual(self.product.stock, 3)
        self.client.post(reverse("seller_quick_update", args=[self.product.id]), {"price": "-1", "stock": "3"})
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, Decimal("42.50"))

    def test_duplicate_product_goes_to_review(self):
        self.client.force_login(self.seller.user)
        r = self.client.post(reverse("seller_duplicate_product", args=[self.product.id]))
        copy = Product.objects.exclude(pk=self.product.pk).get()
        self.assertRedirects(r, reverse("seller_edit_product", args=[copy.pk]), fetch_redirect_response=False)
        self.assertEqual(copy.approval_status, "pending")
        self.assertEqual(copy.seller_account, self.seller)

    def test_payout_request_and_staff_records_it(self):
        from .models import Payout
        order = self._order(qty=2)
        # Not delivered yet: nothing available
        self.client.post(reverse("seller_request_payout"), {"amount": "80"})
        self.assertFalse(Payout.objects.exists())
        OrderItem.objects.filter(order=order).update(fulfillment_status="delivered")
        o = Order.objects.get(pk=order.pk)
        o.status = "delivered"
        o.save()
        # Delivered but still inside the return window: held
        self.client.post(reverse("seller_request_payout"), {"amount": "80"})
        self.assertFalse(Payout.objects.exists())
        from datetime import timedelta
        Order.objects.filter(pk=order.pk).update(delivered_at=timezone.now() - timedelta(days=30))
        self.client.post(reverse("seller_request_payout"), {"amount": "500"})  # too much
        self.assertFalse(Payout.objects.exists())
        self.client.post(reverse("seller_request_payout"), {"amount": "80"})
        payout = Payout.objects.get()
        self.assertEqual(payout.status, "requested")
        self.client.post(reverse("seller_request_payout"), {"amount": "10"})  # one at a time
        self.assertEqual(Payout.objects.count(), 1)
        staff = User.objects.create_user("boss", "boss@example.com", "pass12345", is_staff=True)
        self.client.force_login(staff)
        self.assertContains(self.client.get(reverse("manage_sellers") + "?status=payout"), "Shop Co")
        self.assertContains(self.client.get(reverse("manage_seller", args=[self.seller.id])), "asked for $80.00")
        self.client.post(reverse("manage_seller", args=[self.seller.id]), {"action": "payout", "amount": "80", "method": "wise", "reference": "TX-1"})
        payout.refresh_from_db()
        self.seller.refresh_from_db()
        self.assertEqual(payout.status, "paid")
        self.assertEqual(payout.method, "wise")
        self.assertEqual(self.seller.total_paid_out, Decimal("80.00"))
        self.assertEqual(self.seller.amount_owed, 0)
        self.client.force_login(self.seller.user)
        self.assertContains(self.client.get(reverse("seller_earnings")), "TX-1")

    def test_staff_team_member_cannot_see_money_pages(self):
        helper = User.objects.create_user("helper", "h@example.com", "pass12345")
        OrganizationMember.objects.create(organization=self.seller, user=helper, role="staff")
        self.client.force_login(helper)
        self.assertEqual(self.client.get(reverse("seller_dashboard")).status_code, 200)
        self.assertEqual(self.client.get(reverse("seller_orders")).status_code, 200)
        for name in ("seller_earnings", "seller_team", "seller_store_settings"):
            self.assertRedirects(self.client.get(reverse(name)), reverse("seller_dashboard"), fetch_redirect_response=False)
        self.client.post(reverse("seller_request_payout"), {"amount": "10"})
        from .models import Payout
        self.assertFalse(Payout.objects.exists())

    def test_holiday_mode_hides_products(self):
        self.client.force_login(self.seller.user)
        self.client.post(reverse("seller_vacation"), {"vacation_mode": "on", "vacation_message": "Back Monday"})
        self.seller.refresh_from_db()
        self.assertTrue(self.seller.vacation_mode)
        self.assertFalse(Product.objects.live().filter(pk=self.product.pk).exists())
        self.client.logout()
        self.assertEqual(self.client.get(reverse("product_detail", args=[self.product.id])).status_code, 404)
        self.assertContains(self.client.get(reverse("store_page", args=["Shop Co"])), "Back Monday")
        self.client.force_login(self.seller.user)
        self.client.post(reverse("seller_vacation"), {"vacation_mode": "off"})
        self.assertTrue(Product.objects.live().filter(pk=self.product.pk).exists())

    def test_pending_seller_sees_center_but_cannot_sell(self):
        u = User.objects.create_user("newbie", "n@example.com", "pass12345")
        SellerAccount.objects.create(user=u, business_name="Newbie")
        self.client.force_login(u)
        page = self.client.get(reverse("seller_dashboard"))
        self.assertContains(page, "under review")
        self.assertEqual(self.client.get(reverse("seller_add_product")).status_code, 404)

    def test_store_page_search_and_stats(self):
        self._order()
        self.client.logout()
        page = self.client.get(reverse("store_page", args=["Shop Co"]) + "?q=lamp&sort=price_asc")
        self.assertContains(page, "Lamp")
        self.assertContains(page, "Registered business")
        self.assertContains(page, "2+ sold")


class ProfileReferralCodeTests(TestCase):
    def test_profiles_created_without_a_code_get_a_unique_one(self):
        from .models import Profile
        a = Profile.objects.create(user=User.objects.create_user("pa", "pa@example.com", "x"))
        b, _ = Profile.objects.get_or_create(user=User.objects.create_user("pb", "pb@example.com", "x"))
        self.assertTrue(a.referral_code and b.referral_code)
        self.assertNotEqual(a.referral_code, b.referral_code)


class SystemCheckPaymentHistoryTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user("boss2", "boss2@example.com", "pass12345", is_staff=True)
        self.client.force_login(self.staff)

    @override_settings(STRIPE_SECRET_KEY="sk_test_x")
    def test_old_payment_errors_clear_once_stripe_works(self):
        from .alerts import PAYMENT_FAILED, log
        log(f"{PAYMENT_FAILED}: order #1 - The Supabase relay refused the request: RELAY_SECRET mismatch")
        with mock.patch("stripe.Balance.retrieve", side_effect=Exception("still down")):
            page = self.client.get(reverse("manage_system"))
        self.assertContains(page, "RELAY_SECRET mismatch")
        with mock.patch("stripe.Balance.retrieve", return_value={}):
            self.client.post(reverse("manage_system"), {"action": "stripe"})
        page = self.client.get(reverse("manage_system"))
        self.assertNotContains(page, "RELAY_SECRET mismatch")
        self.assertContains(page, "1 earlier error was already fixed")
        # A new failure after that is reported again
        log(f"{PAYMENT_FAILED}: order #2 - Stripe rejected the secret key")
        self.assertContains(self.client.get(reverse("manage_system")), "Stripe rejected the secret key")

    def test_success_marker_is_a_single_entry(self):
        from .alerts import PAYMENT_OK, payments_working
        for _ in range(3):
            payments_working("test")
        self.assertEqual(AuditLog.objects.filter(action__startswith=PAYMENT_OK).count(), 1)


class BackupButtonTests(TransactionTestCase):
    def test_back_up_now_button(self):
        import tempfile
        staff = User.objects.create_user("boss3", "boss3@example.com", "pass12345", is_staff=True)
        from .models import SiteSettings
        SiteSettings.objects.update_or_create(pk=1, defaults={"require_staff_2fa": False})
        self.client.force_login(staff)
        with tempfile.TemporaryDirectory() as tmp, override_settings(BASE_DIR=tmp, MEDIA_ROOT=os.path.join(tmp, "m")):
            r = self.client.post(reverse("manage_system"), {"action": "backup"}, follow=True)
            self.assertContains(r, "Backup saved.")
            self.assertEqual(len(os.listdir(os.path.join(tmp, "backups"))), 1)
            self.assertNotContains(r, "No backups yet")
class StripeObjectCompatibilityTests(TestCase):
    """stripe-python v15+ returns objects that are not dicts; saved cards
    must still work (this used to crash the Saved cards page)."""

    def _obj(self, data):
        from stripe import StripeObject
        return StripeObject.construct_from(data, "sk_test")

    @override_settings(STRIPE_SECRET_KEY="sk_test_x")
    def test_saved_cards_page_with_real_stripe_objects(self):
        from .models import Profile
        user = User.objects.create_user("cardy", "c@example.com", "pass12345")
        Profile.objects.create(user=user, referral_code="CARDY1", stripe_customer_id="cus_1")
        listing = self._obj({"object": "list", "data": [
            {"id": "pm_1", "object": "payment_method", "card": {"brand": "visa", "last4": "4242", "exp_month": 4, "exp_year": 2030}},
        ]})
        self.client.force_login(user)
        with mock.patch("stripe.Customer.list_payment_methods", return_value=listing):
            r = self.client.get(reverse("payment_methods"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Visa ending in 4242")
        self.assertContains(r, "Expires 04/2030")

    @override_settings(STRIPE_SECRET_KEY="sk_test_x")
    def test_stripe_outage_shows_message_not_error_page(self):
        from .models import Profile
        user = User.objects.create_user("cardz", "z@example.com", "pass12345")
        Profile.objects.create(user=user, referral_code="CARDZ1", stripe_customer_id="cus_2")
        self.client.force_login(user)
        with mock.patch("stripe.Customer.list_payment_methods", side_effect=Exception("down")):
            r = self.client.get(reverse("payment_methods"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "couldn&#x27;t load your saved cards")

    @override_settings(STRIPE_SECRET_KEY="sk_test_x")
    def test_remove_and_finish_setup_with_real_stripe_objects(self):
        from . import payments
        with mock.patch("stripe.PaymentMethod.retrieve", return_value=self._obj({"id": "pm_1", "customer": "cus_1"})), \
                mock.patch("stripe.PaymentMethod.detach") as detach:
            self.assertTrue(payments.remove_card("cus_1", "pm_1"))
            detach.assert_called_once()
        with mock.patch("stripe.PaymentMethod.retrieve", return_value=self._obj({"id": "pm_9", "customer": "cus_other"})):
            self.assertFalse(payments.remove_card("cus_1", "pm_9"))
        session = self._obj({"id": "cs_1", "customer": "cus_1", "status": "complete", "setup_intent": {"id": "seti_1", "payment_method": "pm_1"}})
        with mock.patch("stripe.checkout.Session.retrieve", return_value=session), \
                mock.patch("stripe.PaymentMethod.modify") as modify:
            self.assertTrue(payments.finish_card_setup("cs_1", "cus_1"))
            modify.assert_called_once_with("pm_1", allow_redisplay="always")


class PaymentDiagnosticsTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user("boss4", "boss4@example.com", "pass12345", is_staff=True)
        self.client.force_login(self.staff)

    @override_settings(STRIPE_SECRET_KEY="sk_test_x")
    def test_old_errors_clear_automatically_when_stripe_works(self):
        from .alerts import PAYMENT_FAILED, log
        log(f"{PAYMENT_FAILED}: order #1 - RELAY_SECRET mismatch")
        with mock.patch("stripe.Balance.retrieve", return_value={}) as bal:
            page = self.client.get(reverse("manage_system"))
            bal.assert_called_once()
        self.assertNotContains(page, "RELAY_SECRET mismatch")
        self.assertContains(page, "already fixed")

    @override_settings(STRIPE_SECRET_KEY="sk_test_x", STRIPE_WEBHOOK_SECRET="whsec_x")
    def test_webhook_with_wrong_secret_is_explained(self):
        self.client.logout()
        r = self.client.post(reverse("stripe_webhook"), data="{}", content_type="application/json", HTTP_STRIPE_SIGNATURE="t=1,v1=bad")
        self.assertEqual(r.status_code, 400)
        self.client.force_login(self.staff)
        with mock.patch("stripe.Balance.retrieve", return_value={}):
            page = self.client.get(reverse("manage_system"))
        self.assertContains(page, "webhook secret doesn&#x27;t match")

    @override_settings(STRIPE_SECRET_KEY="sk_test_x", STRIPE_WEBHOOK_SECRET="whsec_x")
    def test_no_webhook_yet_shows_exact_url(self):
        with mock.patch("stripe.Balance.retrieve", return_value={}):
            page = self.client.get(reverse("manage_system"))
        self.assertContains(page, "/payment/stripe/webhook/")

    @override_settings(STRIPE_SECRET_KEY="sk_test_x")
    def test_add_card_recovers_from_missing_customer(self):
        import stripe as _stripe
        from .models import Profile
        Profile.objects.create(user=self.staff, referral_code="BOSS4X", stripe_customer_id="cus_old")
        missing = _stripe.error.InvalidRequestError("No such customer: 'cus_old'", "customer", code="resource_missing", http_status=400)
        ok = mock.MagicMock(url="https://checkout.stripe.com/c/pay/cs_1")
        with mock.patch("stripe.checkout.Session.create", side_effect=[missing, ok]) as create, \
                mock.patch("stripe.Customer.create", return_value=mock.MagicMock(id="cus_new")):
            r = self.client.post(reverse("add_card"))
        self.assertEqual(r.status_code, 302)
        self.assertEqual(r["Location"], "https://checkout.stripe.com/c/pay/cs_1")
        self.assertEqual(create.call_args.kwargs["customer"], "cus_new")
        self.assertEqual(Profile.objects.get(user=self.staff).stripe_customer_id, "cus_new")

    @override_settings(STRIPE_SECRET_KEY="sk_test_x")
    def test_add_card_failure_shows_reason_to_staff_and_is_logged(self):
        import stripe as _stripe
        from .models import Profile
        Profile.objects.create(user=self.staff, referral_code="BOSS4Y", stripe_customer_id="cus_1")
        err = _stripe.error.InvalidRequestError("Something specific is wrong", None, http_status=400)
        with mock.patch("stripe.checkout.Session.create", side_effect=err), \
                mock.patch("stripe.Customer.list_payment_methods", return_value={"data": []}):
            r = self.client.post(reverse("add_card"), follow=True)
        self.assertContains(r, "Admin info: Stripe error 400")
        self.assertTrue(AuditLog.objects.filter(action__contains="saved-card page").exists())


class ImageUploadTests(TestCase):
    """Uploads are turned upright, resized, stored and shown; links work."""

    def setUp(self):
        import tempfile
        self._tmp = tempfile.TemporaryDirectory()
        self._media = override_settings(MEDIA_ROOT=self._tmp.name)
        self._media.enable()
        owner = User.objects.create_user("imgshop", "imgshop@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, business_name="Img Co")
        self.seller.status = "approved"
        self.seller.save()
        self.client.force_login(owner)

    def tearDown(self):
        self._media.disable()
        self._tmp.cleanup()

    def _photo(self, name="phone.jpg", size=(4000, 3000), orientation=6, fmt="JPEG", mode="RGB"):
        from io import BytesIO
        from PIL import Image
        img = Image.new(mode, size, (200, 30, 30) if mode == "RGB" else (200, 30, 30, 120))
        buf = BytesIO()
        if fmt == "JPEG":
            exif = Image.Exif()
            exif[0x0112] = orientation  # camera says "rotate 90°"
            exif[0x8825] = {1: "N"}     # GPS info that must be stripped
            img.save(buf, fmt, exif=exif.tobytes())
        else:
            img.save(buf, fmt)
        return SimpleUploadedFile(name, buf.getvalue(), content_type="image/jpeg" if fmt == "JPEG" else "image/png")

    def _open_saved(self, url):
        from PIL import Image
        from django.conf import settings
        return Image.open(os.path.join(settings.MEDIA_ROOT, url.replace(settings.MEDIA_URL, "", 1)))

    def _form(self, **extra):
        data = {"name": "Lamp", "category": "table-lamp", "price": "20", "stock": "5", "description": "x"}
        data.update(extra)
        return data

    def test_phone_photo_is_rotated_resized_and_stripped(self):
        r = self.client.post(reverse("seller_add_product"), self._form(image_file=self._photo()))
        self.assertEqual(r.status_code, 302)
        p = Product.objects.get(name="Lamp")
        img = self._open_saved(p.image_url)
        self.assertEqual(max(img.size), 1600)
        self.assertGreater(img.height, img.width)  # turned upright (was 4000x3000 landscape + rotate)
        self.assertFalse(img.getexif())
        self.assertTrue(p.image_url.endswith(".jpg"))

    def test_transparent_png_stays_png(self):
        self.client.post(reverse("seller_add_product"), self._form(image_file=self._photo("logo.png", (300, 300), fmt="PNG", mode="RGBA")))
        self.assertTrue(Product.objects.get(name="Lamp").image_url.endswith(".png"))

    def test_gallery_upload_links_and_remove(self):
        files = [self._photo("a.jpg", (800, 600)), self._photo("b.jpg", (800, 600))]
        self.client.post(reverse("seller_add_product"), self._form(
            image_url="https://example.com/main.jpg", extra_files=files, extra_image_links="https://example.com/c.jpg"))
        p = Product.objects.get(name="Lamp")
        self.assertEqual(p.extra_images.count(), 3)
        page = self.client.get(reverse("product_detail", args=[p.id]))
        self.assertContains(page, "https://example.com/c.jpg")
        # remove one, keep the others
        first = p.extra_images.first()
        self.client.post(reverse("seller_edit_product", args=[p.id]), self._form(image_url=p.image_url, remove_extra=[first.id]))
        self.assertEqual(p.extra_images.count(), 2)
        self.assertFalse(p.extra_images.filter(pk=first.pk).exists())

    def test_too_many_gallery_pictures_rejected(self):
        links = "\n".join(f"https://example.com/{i}.jpg" for i in range(9))
        r = self.client.post(reverse("seller_add_product"), self._form(image_url="https://example.com/m.jpg", extra_image_links=links))
        self.assertContains(r, "up to 8")
        self.assertFalse(Product.objects.exists())

    def test_bad_links_and_heic_rejected(self):
        r = self.client.post(reverse("seller_add_product"), self._form(image_url="http://example.com/x.jpg"))
        self.assertContains(r, "must start with https://")
        heic = SimpleUploadedFile("IMG_0001.HEIC", b"\x00" * 100, content_type="image/heic")
        r = self.client.post(reverse("seller_add_product"), self._form(image_file=heic))
        self.assertContains(r, "Most Compatible")

    def test_admin_can_edit_product_with_uploaded_image_path(self):
        staff = User.objects.create_user("imgboss", "ib@example.com", "pass12345", is_staff=True)
        p = make_product(image_url="/media/products/abc.jpg")
        self.client.force_login(staff)
        page = self.client.get(reverse("manage_product_edit", args=[p.id]))
        self.assertNotContains(page, 'type="url" name="image_url"')  # a /media/ path must not be blocked by the browser
        r = self.client.post(reverse("manage_product_edit", args=[p.id]), {
            "name": p.name, "category": p.category, "price": "9.99", "stock": "3", "image_url": p.image_url,
            "approval_status": "approved",
        })
        self.assertEqual(r.status_code, 302)
        p.refresh_from_db()
        self.assertEqual(p.price, Decimal("9.99"))

    def test_store_logo_is_resized(self):
        r = self.client.post(reverse("seller_store_settings"), {"store_name": "Img Co", "store_logo": self._photo("logo.jpg", (3000, 3000))})
        self.assertEqual(r.status_code, 302)
        self.seller.refresh_from_db()
        from PIL import Image
        with Image.open(self.seller.store_logo.path) as img:
            self.assertLessEqual(max(img.size), 600)


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy")
class PaymentResilienceTests(TestCase):
    """Saved-card problems must never stop a customer from paying."""

    def setUp(self):
        self.user = User.objects.create_user("payer2", "payer2@example.com", "pass12345")
        self.product = make_product(price=Decimal("10.00"), stock=5)
        self.client.force_login(self.user)
        from .models import Profile
        Profile.objects.create(user=self.user, referral_code="PAYER2", stripe_customer_id="cus_stale")

    def _ok(self):
        fake = mock.MagicMock()
        fake.id, fake.url = "cs_ok", "https://checkout.stripe.com/c/pay/cs_ok"
        return fake

    def _pay(self, side_effect):
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        with mock.patch("stripe.checkout.Session.create", side_effect=side_effect) as create:
            r = self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card"})
        return r, create

    def test_stale_customer_falls_back_to_guest_payment(self):
        import stripe as _stripe
        from .models import Profile
        missing = _stripe.error.InvalidRequestError("No such customer: 'cus_stale'", "customer", code="resource_missing", http_status=400)
        r, create = self._pay([missing, self._ok()])
        self.assertEqual(r["Location"], "https://checkout.stripe.com/c/pay/cs_ok")
        second = create.call_args_list[1].kwargs
        self.assertNotIn("customer", second)
        self.assertEqual(second["customer_email"], "jane@example.com")
        self.assertEqual(Profile.objects.get(user=self.user).stripe_customer_id, "")
        self.assertTrue(AuditLog.objects.filter(action__contains="retried without saved cards").exists())

    def test_saved_card_option_rejected_falls_back(self):
        import stripe as _stripe
        bad = _stripe.error.InvalidRequestError("Invalid saved_payment_method_options", "saved_payment_method_options", http_status=400)
        r, create = self._pay([bad, self._ok()])
        self.assertEqual(r.status_code, 302)
        self.assertNotIn("saved_payment_method_options", create.call_args_list[1].kwargs)

    def test_network_errors_are_not_retried_as_guest(self):
        r, create = self._pay(Exception("network down"))
        self.assertEqual(create.call_count, 1)
        self.assertRedirects(r, reverse("cart"))

    def test_saved_cards_page_forgets_stale_customer(self):
        import stripe as _stripe
        from .models import Profile
        missing = _stripe.error.InvalidRequestError("No such customer: 'cus_stale'", "customer", code="resource_missing", http_status=400)
        with mock.patch("stripe.Customer.list_payment_methods", side_effect=missing):
            r = self.client.get(reverse("payment_methods"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "No saved cards yet")
        self.assertEqual(Profile.objects.get(user=self.user).stripe_customer_id, "")

    def test_new_customer_uses_fresh_idempotency_key(self):
        from . import payments
        from .models import Profile
        Profile.objects.filter(user=self.user).update(stripe_customer_id="")
        with mock.patch("stripe.Customer.create", return_value=mock.MagicMock(id="cus_a")) as create:
            payments.ensure_customer(self.user)
            Profile.objects.filter(user=self.user).update(stripe_customer_id="")
            payments.ensure_customer(self.user)
        keys = [c.kwargs["idempotency_key"] for c in create.call_args_list]
        self.assertNotEqual(keys[0], keys[1])


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy")
class PaymentSelfTestTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user("boss9", "boss9@example.com", "pass12345", is_staff=True)
        self.client.force_login(self.staff)

    def test_all_steps_pass(self):
        sess = mock.MagicMock(id="cs_1", url="https://checkout.stripe.com/x")
        with mock.patch("stripe.Balance.retrieve", return_value={}), \
                mock.patch("stripe.Customer.create", return_value=mock.MagicMock(id="cus_1")), \
                mock.patch("stripe.checkout.Session.create", return_value=sess) as create, \
                mock.patch("stripe.checkout.Session.expire") as expire, \
                mock.patch("stripe.Customer.delete") as delete:
            r = self.client.post(reverse("manage_system"), {"action": "stripe_full"}, follow=True)
        self.assertContains(r, "Full payment test passed")
        self.assertContains(r, "Add-a-card page")
        self.assertEqual(create.call_count, 3)
        self.assertEqual(expire.call_count, 3)
        delete.assert_called_once()

    def test_failing_step_is_explained(self):
        import stripe as _stripe
        err = _stripe.error.InvalidRequestError("The currency provided is not supported", "currency", http_status=400)
        sess = mock.MagicMock(id="cs_1")
        with mock.patch("stripe.Balance.retrieve", return_value={}), \
                mock.patch("stripe.Customer.create", return_value=mock.MagicMock(id="cus_1")), \
                mock.patch("stripe.checkout.Session.create", side_effect=[sess, sess, err]), \
                mock.patch("stripe.checkout.Session.expire"), mock.patch("stripe.Customer.delete"):
            r = self.client.post(reverse("manage_system"), {"action": "stripe_full"}, follow=True)
        self.assertContains(r, "found a problem")
        self.assertContains(r, "The currency provided is not supported")


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy")
class AddCardRequestTests(TestCase):
    def test_setup_page_uses_dashboard_payment_methods(self):
        """Stripe's current API rejects payment_method_types on Checkout."""
        from .models import Profile
        user = User.objects.create_user("cardq", "cq@example.com", "pass12345")
        Profile.objects.create(user=user, referral_code="CARDQ1", stripe_customer_id="cus_q")
        self.client.force_login(user)
        ok = mock.MagicMock(url="https://checkout.stripe.com/c/pay/cs_q")
        with mock.patch("stripe.checkout.Session.create", return_value=ok) as create:
            r = self.client.post(reverse("add_card"))
        self.assertEqual(r["Location"], "https://checkout.stripe.com/c/pay/cs_q")
        kwargs = create.call_args.kwargs
        self.assertNotIn("payment_method_types", kwargs)
        self.assertEqual(kwargs["mode"], "setup")
        self.assertEqual(kwargs["currency"], "usd")


# ---------------------------------------------------------------------------
# Growth pack: seller plans, setup wizard, AI writer, smart search, demo mode
# ---------------------------------------------------------------------------

import io  # noqa: E402
from datetime import timedelta  # noqa: E402

from django.utils import timezone  # noqa: E402

from .models import Notification, Order, PlanPayment, SellerAccount, SellerPlan  # noqa: E402


class SellerPlanTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user("planner", "planner@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.owner, account_type="individual", business_name="Plan Co")
        self.seller.status = "approved"
        self.seller.save()
        self.free = SellerPlan.objects.get(name="Starter")
        self.pro = SellerPlan.objects.get(name="Pro")
        self.client.force_login(self.owner)

    def test_default_plans_and_new_seller_starts_on_free(self):
        self.assertEqual(self.seller.active_plan, self.free)
        self.assertEqual(self.seller.product_limit, 50)
        self.assertEqual(self.seller.effective_commission_rate, 10)
        r = self.client.get(reverse("seller_plan"))
        self.assertContains(r, "Pro")
        self.assertContains(r, "Starter")

    def test_commission_defaults_come_from_store_settings(self):
        site = SiteSettings.load()
        site.commission_individual, site.commission_organization = Decimal("12"), Decimal("18")
        site.save()
        u = User.objects.create_user("org", "org@example.com", "pass12345")
        org = SellerAccount.objects.create(user=u, account_type="organization")
        self.assertEqual(org.commission_rate, Decimal("18"))

    def test_paid_plan_lowers_commission_and_expires(self):
        payment = PlanPayment.objects.create(seller=self.seller, plan=self.pro, plan_name="Pro", months=2, amount=Decimal("38"))
        self.assertTrue(payment.activate())
        self.assertFalse(payment.activate())  # second call does nothing
        seller = SellerAccount.objects.get(pk=self.seller.pk)
        self.assertEqual(seller.active_plan, self.pro)
        self.assertEqual(seller.effective_commission_rate, 7)
        self.assertTrue(seller.has_badge)
        self.assertGreater(seller.plan_expires_at, timezone.now() + timedelta(days=59))
        SellerAccount.objects.filter(pk=seller.pk).update(plan_expires_at=timezone.now() - timedelta(minutes=1))
        seller = SellerAccount.objects.get(pk=self.seller.pk)
        self.assertEqual(seller.active_plan, self.free)
        self.assertEqual(seller.effective_commission_rate, 10)

    def test_renewing_adds_to_time_left(self):
        PlanPayment.objects.create(seller=self.seller, plan=self.pro, plan_name="Pro", months=1, amount=19).activate()
        first = SellerAccount.objects.get(pk=self.seller.pk).plan_expires_at
        PlanPayment.objects.create(seller=self.seller, plan=self.pro, plan_name="Pro", months=1, amount=19).activate()
        second = SellerAccount.objects.get(pk=self.seller.pk).plan_expires_at
        self.assertAlmostEqual((second - first).days, 30, delta=1)

    def test_product_limit_blocks_adding(self):
        self.free.product_limit = 1
        self.free.save()
        make_product(name="One", seller_account=self.seller, seller_name="Plan Co")
        r = self.client.get(reverse("seller_add_product"))
        self.assertRedirects(r, reverse("seller_plan"), fetch_redirect_response=False)

    def test_buy_without_stripe_shows_message(self):
        with self.settings(STRIPE_SECRET_KEY=""):
            r = self.client.post(reverse("seller_buy_plan", args=[self.pro.id]), {"months": "3"}, follow=True)
        self.assertContains(r, "Card payments aren")
        self.assertFalse(PlanPayment.objects.exists())

    @override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy")
    def test_card_payment_activates_plan_via_webhook(self):
        fake = mock.Mock()
        fake.id, fake.url = "cs_plan_1", "https://checkout.stripe.com/c/pay/cs_plan_1"
        with mock.patch("stripe.checkout.Session.create", return_value=fake) as create:
            r = self.client.post(reverse("seller_buy_plan", args=[self.pro.id]), {"months": "3"})
        self.assertEqual(r.url, fake.url)
        kwargs = create.call_args.kwargs
        self.assertEqual(kwargs["line_items"][0]["quantity"], 3)
        self.assertEqual(kwargs["line_items"][0]["price_data"]["unit_amount"], 1900)
        payment = PlanPayment.objects.get()
        self.assertEqual(payment.amount, Decimal("57"))
        session = {"id": "cs_plan_1", "payment_status": "paid", "amount_total": 5700,
                   "metadata": {"plan_payment_id": str(payment.id)}}
        payload = json.dumps({"type": "checkout.session.completed", "data": {"object": session}})
        with mock.patch("stripe.Webhook.construct_event", return_value={}):
            self.client.post(reverse("stripe_webhook"), data=payload, content_type="application/json", HTTP_STRIPE_SIGNATURE="x")
        payment.refresh_from_db()
        self.assertEqual(payment.status, "paid")
        self.assertEqual(SellerAccount.objects.get(pk=self.seller.pk).active_plan, self.pro)

    @override_settings(STRIPE_SECRET_KEY="sk_test_dummy")
    def test_underpaid_session_does_not_activate(self):
        from . import payments
        payment = PlanPayment.objects.create(seller=self.seller, plan=self.pro, plan_name="Pro", months=1,
                                             amount=Decimal("19"), stripe_session_id="cs_x")
        payments.confirm_plan_payment(payment.id, {"id": "cs_x", "payment_status": "paid", "amount_total": 100})
        payment.refresh_from_db()
        self.assertEqual(payment.status, "pending")

    def test_staff_can_give_plan_and_manage_plans(self):
        admin = User.objects.create_superuser("boss", "boss@example.com", "pass12345")
        self.client.force_login(admin)
        r = self.client.post(reverse("manage_seller", args=[self.seller.id]),
                             {"action": "grant_plan", "plan": self.pro.id, "months": "1", "free": "on"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(SellerAccount.objects.get(pk=self.seller.pk).active_plan, self.pro)
        self.assertEqual(PlanPayment.objects.get().amount, 0)
        r = self.client.get(reverse("manage_plans"))
        self.assertContains(r, "Seller plans")
        self.client.post(reverse("manage_plans"), {"name": "Gold", "price": "99", "commission_discount": "6",
                                                   "product_limit": "", "position": "4", "active": "on"})
        self.assertTrue(SellerPlan.objects.filter(name="Gold", product_limit=None).exists())
        r = self.client.post(reverse("manage_plans"), {"name": "Bad", "price": "5", "commission_discount": "80", "position": "5"})
        self.assertFalse(SellerPlan.objects.filter(name="Bad").exists())

    def test_staff_team_member_cannot_open_plan(self):
        from .models import OrganizationMember
        staff = User.objects.create_user("helper", "h@example.com", "pass12345")
        self.seller.account_type = "organization"
        self.seller.save()
        OrganizationMember.objects.create(organization=self.seller, user=staff, role="staff")
        self.client.force_login(staff)
        r = self.client.get(reverse("seller_plan"))
        self.assertEqual(r.status_code, 302)

    def test_reminder_sent_once(self):
        from django.core.management import call_command
        SellerAccount.objects.filter(pk=self.seller.pk).update(plan=self.pro, plan_expires_at=timezone.now() + timedelta(days=2))
        call_command("plan_reminders", stdout=io.StringIO())
        call_command("plan_reminders", stdout=io.StringIO())
        self.assertEqual(Notification.objects.filter(user=self.owner, message__icontains="plan ends").count(), 1)


class SetupWizardTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser("owner", "owner@example.com", "pass12345")
        self.client.force_login(self.admin)

    def test_steps_save_and_finish(self):
        r = self.client.get(reverse("manage_setup"))
        self.assertContains(r, "Set up your store")
        r = self.client.post(reverse("manage_setup"), {"step": "basics", "site_name": "Nova Market", "tagline": "Hi",
                                                       "support_email": "help@nova.test"})
        self.assertRedirects(r, reverse("manage_setup") + "?step=look", fetch_redirect_response=False)
        self.assertEqual(SiteSettings.load().site_name, "Nova Market")
        r = self.client.post(reverse("manage_setup"), {"step": "selling", "commission_individual": "8", "commission_organization": "15"})
        self.assertEqual(SiteSettings.objects.get(pk=1).commission_individual, Decimal("8"))
        r = self.client.post(reverse("manage_setup"), {"step": "selling", "commission_individual": "120", "commission_organization": "15"})
        self.assertContains(r, "between 0 and 90")
        for step in ("look", "checkout", "launch"):
            self.assertEqual(self.client.get(reverse("manage_setup") + f"?step={step}").status_code, 200)
        self.client.post(reverse("manage_setup"), {"finish": "1"})
        self.assertTrue(SiteSettings.objects.get(pk=1).setup_completed)

    def test_dashboard_banner_and_redirect(self):
        SiteSettings.objects.update_or_create(pk=1, defaults={"setup_completed": False})
        from django.core.cache import cache as c
        c.clear()
        self.assertContains(self.client.get(reverse("manage_dashboard")), "Finish setting up your store")
        with self.settings(SETUP_WIZARD_REDIRECT=True):
            r = self.client.get(reverse("manage_dashboard"))
            self.assertRedirects(r, reverse("manage_setup"), fetch_redirect_response=False)
            self.assertEqual(self.client.get(reverse("manage_dashboard")).status_code, 200)  # only once

    def test_plain_staff_cannot_run_setup(self):
        staff = User.objects.create_user("clerk", "c@example.com", "pass12345", is_staff=True)
        self.client.force_login(staff)
        self.assertEqual(self.client.get(reverse("manage_setup")).status_code, 302)


class AIWriterAndSmartSearchTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user("writer", "w@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.owner, business_name="Write Co")
        self.seller.status = "approved"
        self.seller.save()

    def test_builtin_writer_uses_details_and_needs_a_name(self):
        self.client.force_login(self.owner)
        r = self.client.post(reverse("ai_write_listing"), {"name": "Cotton Hoodie", "category": "hoodies",
                                                           "notes": "100% cotton, machine washable"})
        data = r.json()
        self.assertEqual(data["source"], "builtin")
        self.assertIn("Cotton Hoodie", data["description"])
        self.assertIn("• 100% cotton", data["description"])
        self.assertEqual(self.client.post(reverse("ai_write_listing"), {"name": ""}).status_code, 400)

    def test_writer_falls_back_when_ai_unreachable(self):
        from . import ai
        with self.settings(AI_API_KEY="key"), mock.patch("urllib.request.urlopen", side_effect=OSError("offline")):
            text, source = ai.write_listing("Desk Lamp", "table-lamp", "")
        self.assertEqual(source, "builtin")
        self.assertIn("Desk Lamp", text)

    def test_writer_uses_ai_reply(self):
        from . import ai
        reply = mock.MagicMock()
        reply.__enter__.return_value.read.return_value = json.dumps(
            {"content": [{"type": "text", "text": "A **lovely** lamp for reading late into the night.\n\nWhy you'll love it:\n• Warm light"}]}).encode()
        with self.settings(AI_API_KEY="key"), mock.patch("urllib.request.urlopen", return_value=reply):
            text, source = ai.write_listing("Desk Lamp", "table-lamp", "")
        self.assertEqual(source, "ai")
        self.assertNotIn("**", text)

    def test_writer_only_for_sellers_and_staff(self):
        shopper = User.objects.create_user("shopper", "s@example.com", "pass12345")
        self.client.force_login(shopper)
        self.assertEqual(self.client.post(reverse("ai_write_listing"), {"name": "Lamp"}).status_code, 403)

    def test_parse_search(self):
        from . import ai
        self.assertEqual(ai.parse_search("red hoodie under 40"), {"words": ["red", "hoodie"], "min": None, "max": 40.0, "cheap": False})
        p = ai.parse_search("lamp between 20 and 50")
        self.assertEqual((p["min"], p["max"], p["words"]), (20.0, 50.0, ["lamp"]))
        p = ai.parse_search("sneakers $60-$20")
        self.assertEqual((p["min"], p["max"]), (20.0, 60.0))
        self.assertTrue(ai.parse_search("cheap earbuds")["cheap"])

    def test_smart_search_filters_price_and_matches_words_in_any_order(self):
        make_product(name="Red Cotton Hoodie", category="hoodies", price=Decimal("35"))
        make_product(name="Red Wool Hoodie", category="hoodies", price=Decimal("80"))
        make_product(name="Blue Lamp", category="table-lamp", price=Decimal("20"))
        r = self.client.get(reverse("search_products"), {"q": "hoodies red under 40"})
        self.assertContains(r, "Red Cotton Hoodie")
        self.assertNotContains(r, "Red Wool Hoodie")
        self.assertContains(r, "Under")
        r = self.client.get(reverse("search_products"), {"q": "red lamp"})  # nothing matches both words
        self.assertContains(r, "Close matches")
        self.assertContains(r, "Blue Lamp")


@override_settings(DEMO_MODE=True)
class DemoModeTests(TestCase):
    def setUp(self):
        from django.core.management import call_command
        make_product(name="Existing")
        call_command("demo_setup", stdout=io.StringIO())

    def test_demo_accounts_and_data(self):
        seller = SellerAccount.objects.get(user__username="demo_seller")
        self.assertEqual(seller.status, "approved")
        self.assertTrue(seller.has_badge)
        self.assertTrue(Order.objects.filter(user__username="demo_shopper").exists())
        self.assertFalse(User.objects.get(username="demo_admin").has_usable_password())

    def test_one_click_login_for_each_role(self):
        r = self.client.get(reverse("login"))
        self.assertContains(r, "Try the live demo")
        r = self.client.post(reverse("demo_login", args=["seller"]))
        self.assertRedirects(r, reverse("seller_dashboard"), fetch_redirect_response=False)
        self.assertEqual(self.client.get(reverse("seller_dashboard")).status_code, 200)
        self.assertEqual(self.client.get(reverse("login")).status_code, 200)  # can switch role
        self.client.post(reverse("demo_login", args=["admin"]))
        self.assertEqual(self.client.get(reverse("manage_dashboard")).status_code, 200)
        self.assertEqual(self.client.post(reverse("demo_login", args=["nobody"])).status_code, 404)

    def test_demo_admin_cannot_change_settings(self):
        self.client.post(reverse("demo_login", args=["admin"]))
        before = SiteSettings.load().site_name
        r = self.client.post(reverse("manage_settings"), {"site_name": "Hacked"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(SiteSettings.objects.get(pk=1).site_name, before)

    def test_reset_puts_accounts_back(self):
        from django.core.management import call_command
        SellerAccount.objects.filter(user__username="demo_seller").update(status="suspended", vacation_mode=True)
        call_command("demo_setup", "--reset", stdout=io.StringIO())
        seller = SellerAccount.objects.get(user__username="demo_seller")
        self.assertEqual((seller.status, seller.vacation_mode), ("approved", False))

    @override_settings(DEMO_MODE=False)
    def test_demo_login_off_by_default(self):
        self.assertEqual(self.client.post(reverse("demo_login", args=["admin"])).status_code, 404)
        self.assertNotContains(self.client.get(reverse("login")), "Try the live demo")


# ---------------------------------------------------------------------------
# Marketplace pro: currencies, languages, messages, deals, referrals,
# shipping/tracking, seller badges
# ---------------------------------------------------------------------------

from .models import Conversation, Coupon, Currency, Message, Profile, ShippingZone  # noqa: E402


class CurrencyAndLanguageTests(TestCase):
    def setUp(self):
        cache.clear()
        Currency.objects.all().delete()
        self.eur = Currency.objects.create(code="EUR", symbol="€", rate=Decimal("0.5"), active=True)
        self.product = make_product(name="Mug", price=Decimal("20.00"))

    def test_prices_shown_in_chosen_currency_but_cart_stays_in_store_currency(self):
        self.client.get(reverse("set_currency", args=["EUR"]))
        r = self.client.get(reverse("product_detail", args=[self.product.id]))
        self.assertContains(r, "€10.00")
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 1})
        r = self.client.get(reverse("cart"))
        self.assertContains(r, "$20.00")
        self.assertContains(r, "charged in")
        self.client.get(reverse("set_currency", args=["USD"]))
        self.assertContains(self.client.get(reverse("product_detail", args=[self.product.id])), "$20.00")

    def test_unknown_or_hidden_currency_ignored(self):
        self.eur.active = False
        self.eur.save()
        self.client.get(reverse("set_currency", args=["EUR"]))
        self.assertContains(self.client.get(reverse("product_detail", args=[self.product.id])), "$20.00")

    def test_rates_update(self):
        from . import currency
        reply = mock.MagicMock()
        reply.__enter__.return_value.read.return_value = json.dumps({"result": "success", "rates": {"EUR": 0.9}}).encode()
        with mock.patch("urllib.request.urlopen", return_value=reply):
            updated, error = currency.update_rates()
        self.assertEqual((updated, error), (1, None))
        self.eur.refresh_from_db()
        self.assertEqual(self.eur.rate, Decimal("0.9"))
        with mock.patch("urllib.request.urlopen", side_effect=OSError("blocked")):
            updated, error = currency.update_rates()
        self.assertEqual(updated, 0)
        self.assertIn("by hand", error)

    def test_admin_manages_currencies(self):
        admin = User.objects.create_superuser("cadmin", "c@example.com", "pass12345")
        self.client.force_login(admin)
        self.assertContains(self.client.get(reverse("manage_currencies")), "EUR")
        self.client.post(reverse("manage_currencies"), {"code": "jpy", "symbol": "¥", "rate": "150", "decimals": "0", "position": "2", "active": "on"})
        self.assertTrue(Currency.objects.filter(code="JPY", active=True).exists())
        r = self.client.post(reverse("manage_currencies"), {"code": "USD", "symbol": "$", "rate": "1", "decimals": "2", "position": "3"})
        self.assertContains(r, "store currency")

    def test_new_languages_and_rtl(self):
        site = SiteSettings.load()
        site.show_language_menu = True
        site.save()
        self.client.get(reverse("set_language", args=["ar"]))
        r = self.client.get(reverse("home"))
        self.assertContains(r, 'dir="rtl"')
        self.assertContains(r, "أضف إلى السلة")
        c = Client(HTTP_ACCEPT_LANGUAGE="es-ES,es;q=0.9")
        self.assertContains(c.get(reverse("home")), 'lang="es"')


class BuyerSellerMessageTests(TestCase):
    def setUp(self):
        from django.core import mail as _mail
        self.mail = _mail
        owner = User.objects.create_user("maker", "maker@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, business_name="Maker Co")
        self.seller.status = "approved"
        self.seller.save()
        self.product = make_product(name="Vase", seller_account=self.seller, seller_name="Maker Co")
        self.buyer = User.objects.create_user("asker", "asker@example.com", "pass12345")

    def _send(self, user, conv, body, **extra):
        self.client.force_login(user)
        return self.client.post(reverse("message_send", args=[conv.id]), {"body": body, **extra}, HTTP_X_REQUESTED_WITH="fetch")

    def test_full_conversation(self):
        self.client.force_login(self.buyer)
        r = self.client.get(reverse("message_seller", args=[self.product.id]))
        conv = Conversation.objects.get()
        self.assertIn(f"/messages/{conv.id}/?about={self.product.id}", r.url)
        self.assertContains(self.client.get(r.url), "About <strong>Vase</strong>")
        r = self._send(self.buyer, conv, "Is it hand made?", about=str(self.product.id))
        self.assertTrue(r.json()["message"]["mine"])
        conv.refresh_from_db()
        self.assertEqual(conv.seller_unread, 1)
        self.assertTrue(Notification.objects.filter(user=self.seller.user, message__icontains="New message").exists())
        self.assertTrue(any("maker@example.com" in m.to for m in self.mail.outbox))
        # seller sees and answers
        self.client.force_login(self.seller.user)
        r = self.client.get(reverse("seller_conversation", args=[conv.id]))
        self.assertContains(r, "Is it hand made?")
        conv.refresh_from_db()
        self.assertEqual(conv.seller_unread, 0)
        self._send(self.seller.user, conv, "Yes, by hand!")
        conv.refresh_from_db()
        self.assertEqual(conv.buyer_unread, 1)
        self.client.force_login(self.buyer)
        first = Message.objects.order_by("id").first()
        data = self.client.get(reverse("message_poll", args=[conv.id]) + f"?after={first.id}").json()
        self.assertEqual([m["body"] for m in data["messages"]], ["Yes, by hand!"])
        self.assertFalse(data["messages"][0]["mine"])

    def test_outsiders_cannot_read_or_post(self):
        conv = Conversation.objects.create(buyer=self.buyer, seller=self.seller)
        other = User.objects.create_user("snoop", "s@example.com", "pass12345")
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse("conversation", args=[conv.id])).status_code, 404)
        self.assertEqual(self.client.get(reverse("message_poll", args=[conv.id])).status_code, 404)
        self.assertEqual(self._send(other, conv, "hi").status_code, 404)

    def test_empty_and_rate_limited_messages(self):
        conv = Conversation.objects.create(buyer=self.buyer, seller=self.seller)
        self.assertEqual(self._send(self.buyer, conv, "   ").status_code, 400)
        cache.set(f"msg-rate:{self.buyer.pk}", 999, 60)
        self.assertEqual(self._send(self.buyer, conv, "hello").status_code, 400)
        cache.delete(f"msg-rate:{self.buyer.pk}")

    def test_report_reaches_admin(self):
        conv = Conversation.objects.create(buyer=self.buyer, seller=self.seller)
        self._send(self.buyer, conv, "Hello")
        self.client.post(reverse("message_report", args=[conv.id]))
        conv.refresh_from_db()
        self.assertTrue(conv.reported)
        admin = User.objects.create_superuser("mod", "mod@example.com", "pass12345")
        self.client.force_login(admin)
        self.assertContains(self.client.get(reverse("manage_conversations") + "?tab=reported"), "Reported")
        self.assertContains(self.client.get(reverse("manage_conversation", args=[conv.id])), "Hello")
        self.client.post(reverse("manage_conversation", args=[conv.id]), {"action": "resolve"})
        conv.refresh_from_db()
        self.assertFalse(conv.reported)

    def test_seller_cannot_message_own_store(self):
        self.client.force_login(self.seller.user)
        r = self.client.get(reverse("message_seller", args=[self.product.id]))
        self.assertRedirects(r, reverse("seller_messages"), fetch_redirect_response=False)
        self.assertFalse(Conversation.objects.exists())


class FulfilmentFlowTests(TestCase):
    """Multi-seller orders: each package is tracked, told and reviewable."""
    def setUp(self):
        from django.core import mail
        self.mail = mail
        self.owner_a = User.objects.create_user("sella", "a@example.com", "pass12345")
        self.a = SellerAccount.objects.create(user=self.owner_a, business_name="A Shop", status="approved")
        self.owner_b = User.objects.create_user("sellb", "b@example.com", "pass12345")
        self.b = SellerAccount.objects.create(user=self.owner_b, business_name="B Shop", status="approved")
        self.staff = User.objects.create_superuser("boss", "boss@example.com", "pass12345")
        self.buyer = User.objects.create_user("cust", "cust@example.com", "pass12345")
        self.pa = make_product(name="Alpha Mug", seller_account=self.a, seller_name="A Shop")
        self.pb = make_product(name="Beta Bowl", seller_account=self.b, seller_name="B Shop")
        self.pstore = make_product(name="Store Tea")
        self.order = Order.objects.create(user=self.buyer, email=self.buyer.email, full_name="Cus", address="1 St", city="X",
                                          country="GB", payment_method="card", payment_status="paid", status="confirmed")
        self.items = {}
        for p in (self.pa, self.pb, self.pstore):
            it = OrderItem(order=self.order, product=p, product_name=p.name, price=10, quantity=1)
            it.apply_commission(p.seller_account)
            it.save()
            self.items[p.name] = it

    def _seller_set(self, user, item, status):
        self.client.force_login(user)
        return self.client.post(reverse("update_fulfillment_status", args=[item.id]), {"fulfillment_status": status})

    def test_seller_delivery_lets_buyer_review_and_tells_everyone(self):
        self.mail.outbox.clear()
        with self.captureOnCommitCallbacks(execute=True):
            self._seller_set(self.owner_a, self.items["Alpha Mug"], "delivered")
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "confirmed")  # other packages still on their way
        item = OrderItem.objects.get(pk=self.items["Alpha Mug"].pk)
        self.assertIsNotNone(item.delivered_at)
        self.assertTrue(Notification.objects.filter(user=self.buyer, message__icontains="Alpha Mug").exists())
        self.assertTrue(Notification.objects.filter(user=self.staff, message__icontains="A Shop marked Alpha Mug").exists())
        self.assertTrue(any("Part of your order" in m.subject for m in self.mail.outbox))
        self.assertTrue(self.order.events.filter(text__icontains="Delivered: Alpha Mug").exists())
        # buyer can review the delivered item, not the others
        self.client.force_login(self.buyer)
        self.assertEqual(self.client.get(reverse("product_detail", args=[self.pa.id])).context["review_block"], "")
        self.assertNotEqual(self.client.get(reverse("product_detail", args=[self.pb.id])).context["review_block"], "")
        page = self.client.get(reverse("my_orders"))
        self.assertContains(page, "Rate this item")
        self.client.post(reverse("product_detail", args=[self.pa.id]), {"rating": "5", "comment": "Lovely mug"})
        self.assertContains(self.client.get(reverse("my_orders")), "Your review")

    def test_whole_order_delivered_once_all_packages_arrive(self):
        self._seller_set(self.owner_a, self.items["Alpha Mug"], "delivered")
        self._seller_set(self.owner_b, self.items["Beta Bowl"], "delivered")
        self.client.force_login(self.staff)
        self.client.post(reverse("manage_order", args=[self.order.id]), {"action": "item_status", "item": self.items["Store Tea"].id, "fulfillment_status": "delivered"})
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "delivered")
        self.assertTrue(Notification.objects.filter(user=self.buyer, message__icontains="is now Delivered").exists())

    def test_sellers_only_touch_their_own_items_and_cannot_undo_delivery(self):
        r = self._seller_set(self.owner_a, self.items["Beta Bowl"], "delivered")
        self.assertEqual(r.status_code, 404)
        self._seller_set(self.owner_a, self.items["Alpha Mug"], "delivered")
        self._seller_set(self.owner_a, self.items["Alpha Mug"], "pending")
        self.assertEqual(OrderItem.objects.get(pk=self.items["Alpha Mug"].pk).fulfillment_status, "delivered")
        self.client.force_login(self.owner_a)
        self.assertEqual(self.client.get(reverse("update_fulfillment_status", args=[self.items["Alpha Mug"].id])).status_code, 405)

    def test_unpaid_unconfirmed_and_pending_sellers_are_blocked(self):
        Order.objects.filter(pk=self.order.pk).update(payment_status="pending")
        self._seller_set(self.owner_a, self.items["Alpha Mug"], "handed_to_courier")
        self.assertEqual(OrderItem.objects.get(pk=self.items["Alpha Mug"].pk).fulfillment_status, "pending")
        Order.objects.filter(pk=self.order.pk).update(payment_status="not_applicable", payment_method="cod", cod_confirmed=False)
        self._seller_set(self.owner_a, self.items["Alpha Mug"], "handed_to_courier")
        self.assertEqual(OrderItem.objects.get(pk=self.items["Alpha Mug"].pk).fulfillment_status, "pending")
        Order.objects.filter(pk=self.order.pk).update(payment_status="paid", payment_method="card")
        SellerAccount.objects.filter(pk=self.a.pk).update(status="suspended")
        self._seller_set(self.owner_a, self.items["Alpha Mug"], "handed_to_courier")
        self.assertEqual(OrderItem.objects.get(pk=self.items["Alpha Mug"].pk).fulfillment_status, "pending")

    def test_customer_confirms_received(self):
        self._seller_set(self.owner_a, self.items["Alpha Mug"], "handed_to_courier")
        self.client.force_login(self.buyer)
        self.assertContains(self.client.get(reverse("my_orders")), "received it")
        self.client.post(reverse("confirm_received", args=[self.order.id]))
        self.assertEqual(OrderItem.objects.get(pk=self.items["Alpha Mug"].pk).fulfillment_status, "delivered")
        self.assertTrue(Notification.objects.filter(user=self.owner_a, message__icontains="confirmed they received").exists())
        self.assertTrue(Notification.objects.filter(user=self.staff, message__icontains="customer confirmed").exists())
        other = User.objects.create_user("nosy", "n@example.com", "pass12345")
        self.client.force_login(other)
        self.assertEqual(self.client.post(reverse("confirm_received", args=[self.order.id])).status_code, 404)

    def test_admin_cannot_move_delivered_order_back(self):
        self.client.force_login(self.staff)
        url = reverse("manage_order", args=[self.order.id])
        self.client.post(url, {"action": "update", "status": "delivered"})
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "delivered")
        self.assertTrue(all(i.delivered_at for i in self.order.items.all()))
        self.client.post(url, {"action": "update", "status": "pending"})
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "delivered")

    def test_removing_a_sold_product_keeps_history(self):
        self.client.force_login(self.owner_a)
        self.client.post(reverse("seller_delete_product", args=[self.pa.id]))
        self.pa.refresh_from_db()
        self.assertEqual(self.pa.approval_status, "archived")
        self.assertFalse(Product.objects.live().filter(pk=self.pa.pk).exists())
        self.client.force_login(self.buyer)
        self.assertContains(self.client.get(reverse("product_detail", args=[self.pa.id])), "no longer sold")
        stranger = User.objects.create_user("str", "str@example.com", "pass12345")
        self.client.force_login(stranger)
        self.assertEqual(self.client.get(reverse("product_detail", args=[self.pa.id])).status_code, 404)
        unsold = make_product(name="Never sold", seller_account=self.a)
        self.client.force_login(self.owner_a)
        self.client.post(reverse("seller_delete_product", args=[unsold.id]))
        self.assertFalse(Product.objects.filter(pk=unsold.pk).exists())

    def test_seller_team_gets_customer_messages(self):
        from bees.models import OrganizationMember
        helper = User.objects.create_user("helper", "h@example.com", "pass12345")
        OrganizationMember.objects.create(organization=self.a, user=helper, role="staff")
        conv = Conversation.objects.create(buyer=self.buyer, seller=self.a)
        self.client.force_login(self.buyer)
        self.client.post(reverse("message_send", args=[conv.id]), {"body": "Hi!"}, HTTP_X_REQUESTED_WITH="fetch")
        for u in (self.owner_a, helper):
            self.assertTrue(Notification.objects.filter(user=u, link=reverse("seller_conversation", args=[conv.id])).exists())
        self.client.force_login(helper)
        page = self.client.get(reverse("home"))
        self.assertContains(page, "Customer messages")
        self.assertContains(page, reverse("seller_messages"))
        self.assertContains(self.client.get(reverse("messages_inbox")), "Messages from your customers")
        self.assertContains(self.client.get(reverse("seller_conversation", args=[conv.id])), "Hi!")
        self.assertEqual(self.client.get(reverse("live")).json()["seller_messages"], 0)  # read just now

    def test_team_staff_cannot_reach_money_pages(self):
        from bees.models import OrganizationMember
        helper = User.objects.create_user("helper2", "h2@example.com", "pass12345")
        OrganizationMember.objects.create(organization=self.a, user=helper, role="staff")
        self.client.force_login(helper)
        for name in ("seller_earnings", "seller_team", "seller_plan", "seller_store_settings"):
            r = self.client.get(reverse(name))
            self.assertNotEqual(r.status_code, 200, name)
        self.assertEqual(self.client.get(reverse("seller_orders")).status_code, 200)


class TrustPackTests(TestCase):
    """Payout hold, item returns, problem reports, automation, review replies."""
    def setUp(self):
        from datetime import timedelta
        self.td = timedelta
        self.owner = User.objects.create_user("tpsell", "tps@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.owner, business_name="TP Shop", status="approved")
        self.staff = User.objects.create_superuser("tpboss", "tpb@example.com", "pass12345")
        self.buyer = User.objects.create_user("tpbuy", "tpbuy@example.com", "pass12345")
        self.p1 = make_product(name="Lamp", seller_account=self.seller, seller_name="TP Shop", price=Decimal("40.00"))
        self.p2 = make_product(name="Shade", seller_account=self.seller, seller_name="TP Shop", price=Decimal("10.00"))

    def _order(self, status="confirmed", method="card", pay="paid", age_days=0, items=None):
        o = Order.objects.create(user=self.buyer, email=self.buyer.email, full_name="Buyer", address="1 St", city="X", country="GB",
                                 payment_method=method, payment_status=pay, status=status,
                                 cod_confirmed=True)
        out = []
        for p in items or (self.p1,):
            it = OrderItem(order=o, product=p, product_name=p.name, price=p.price, quantity=1)
            it.apply_commission(self.seller)
            it.save()
            out.append(it)
        if age_days:
            Order.objects.filter(pk=o.pk).update(created_at=timezone.now() - self.td(days=age_days))
            o.refresh_from_db()
        return o, out

    def _deliver(self, item, days_ago=0):
        from bees import fulfilment
        fulfilment.set_items_status([item], "delivered", by_staff=True)
        if days_ago:
            OrderItem.objects.filter(pk=item.pk).update(delivered_at=timezone.now() - self.td(days=days_ago))

    def test_earnings_held_until_return_window_passes(self):
        from bees import seller_center as sc
        o, (item,) = self._order()
        self._deliver(item)
        self.assertEqual(sc.balance(self.seller)["available"], 0)
        OrderItem.objects.filter(pk=item.pk).update(delivered_at=timezone.now() - self.td(days=30))
        self.assertGreater(sc.balance(self.seller)["available"], 0)
        # an open problem report puts it back on hold
        Claim.objects.create(order_item=item, user=self.buyer, reason="damaged", details="Broken glass on arrival")
        self.assertEqual(sc.balance(self.seller)["available"], 0)

    def test_item_level_return_before_whole_order_is_delivered(self):
        o, (a, b) = self._order(items=(self.p1, self.p2))
        self._deliver(a)
        o.refresh_from_db()
        self.assertNotEqual(o.status, "delivered")
        self.client.force_login(self.buyer)
        page = self.client.get(reverse("my_orders")).content.decode()
        self.assertIn(reverse("request_return", args=[a.id]), page)
        self.assertNotIn(reverse("request_return", args=[b.id]), page)
        self.client.post(reverse("request_return", args=[a.id]), {"reason": "Too big"})
        self.assertTrue(ReturnRequest.objects.filter(order_item=a).exists())
        self.assertTrue(Notification.objects.filter(user=self.owner, message__icontains="Return requested").exists())
        self.assertTrue(Notification.objects.filter(user=self.staff, message__icontains="Return requested").exists())
        r = self.client.get(reverse("request_return", args=[b.id]))
        self.assertRedirects(r, reverse("my_orders"), fetch_redirect_response=False)

    def test_problem_report_full_flow_with_refund(self):
        o, (item,) = self._order()
        from bees import fulfilment
        fulfilment.set_items_status([item], "handed_to_courier", by_staff=True)
        self.client.force_login(self.buyer)
        self.assertContains(self.client.get(reverse("my_orders")), "Report a problem")
        form = self.client.get(reverse("open_claim", args=[item.id]))
        self.assertContains(form, "It didn&#x27;t arrive")
        self.assertNotContains(form, "Arrived damaged")  # not delivered yet
        r = self.client.post(reverse("open_claim", args=[item.id]), {"reason": "not_received", "details": "Tracking hasn't moved for 10 days."})
        claim = Claim.objects.get()
        self.assertRedirects(r, reverse("claim_detail", args=[claim.id]), fetch_redirect_response=False)
        self.assertTrue(Notification.objects.filter(user=self.owner, link=reverse("seller_claims")).exists())
        # second report blocked
        self.client.post(reverse("open_claim", args=[item.id]), {"reason": "other", "details": "Another report here"})
        self.assertEqual(Claim.objects.count(), 1)
        # seller answers and offers a refund -> goes to the store team
        self.client.force_login(self.owner)
        self.assertContains(self.client.get(reverse("seller_claims")), "Tracking hasn")
        self.client.post(reverse("seller_claims"), {"id": claim.id, "reply": "Sorry, courier lost it.", "offer_refund": "1"})
        claim.refresh_from_db()
        self.assertEqual(claim.status, "escalated")
        # another seller can't see or answer it
        other = User.objects.create_user("othersel", "os@example.com", "pass12345")
        SellerAccount.objects.create(user=other, business_name="Other", status="approved")
        self.client.force_login(other)
        self.assertNotContains(self.client.get(reverse("seller_claims")), "Tracking hasn")
        self.assertEqual(self.client.post(reverse("seller_claims"), {"id": claim.id, "reply": "x"}).status_code, 404)
        # store team refunds (card refund mocked)
        self.client.force_login(self.staff)
        self.assertContains(self.client.get(reverse("manage_claims")), "It didn")
        with mock.patch("bees.payments.refund_amount", return_value=True):
            Order.objects.filter(pk=o.pk).update(stripe_payment_intent="pi_test")
            self.client.post(reverse("manage_claim", args=[claim.id]), {"action": "refund", "note": "Courier lost it"})
        claim.refresh_from_db()
        self.assertEqual(claim.status, "refunded")
        self.assertEqual(claim.return_request.status, "refunded")
        self.assertTrue(Notification.objects.filter(user=self.buyer, message__icontains="was refunded").exists())
        # the refunded item no longer counts in the seller's earnings
        self.assertFalse(OrderItem.objects.counted().filter(pk=item.pk).exists())
        # strangers can't open the claim page
        self.client.force_login(other)
        self.assertRedirects(self.client.get(reverse("claim_detail", args=[claim.id])), reverse("my_orders"), fetch_redirect_response=False)

    def test_claim_reject_and_withdraw(self):
        o, (item,) = self._order()
        self._deliver(item)
        claim = Claim.objects.create(order_item=item, user=self.buyer, reason="not_as_described", details="Colour is different")
        self.client.force_login(self.staff)
        self.client.post(reverse("manage_claim", args=[claim.id]), {"action": "reject", "note": "Photos match the listing"})
        claim.refresh_from_db()
        self.assertEqual(claim.status, "rejected")
        o2, (item2,) = self._order()
        self._deliver(item2)
        c2 = Claim.objects.create(order_item=item2, user=self.buyer, reason="damaged", details="Cracked")
        self.client.force_login(self.buyer)
        self.client.post(reverse("claim_detail", args=[c2.id]), {"action": "withdraw"})
        c2.refresh_from_db()
        self.assertEqual(c2.status, "withdrawn")

    def test_automation_late_orders(self):
        from bees import automation
        o_new, _ = self._order(age_days=3)
        o_late, _ = self._order(age_days=5)
        o_dead, _ = self._order(age_days=8)
        cod_unconfirmed, _ = self._order(age_days=9, method="cod", pay="not_applicable")
        Order.objects.filter(pk=cod_unconfirmed.pk).update(cod_confirmed=False)
        with mock.patch("bees.payments.refund_order", return_value=True):
            summary = automation.run()
        self.assertEqual(summary["auto_cancelled"], 1)
        self.assertEqual(Order.objects.get(pk=o_dead.pk).status, "cancelled")
        self.assertEqual(Order.objects.get(pk=cod_unconfirmed.pk).status, "confirmed")
        self.assertIsNotNone(Order.objects.get(pk=o_new.pk).late_reminded_at)
        self.assertIsNotNone(Order.objects.get(pk=o_late.pk).late_escalated_at)
        self.assertTrue(Notification.objects.filter(user=self.owner, message__icontains=f"Order #{o_new.id} is waiting").exists())
        # running again doesn't repeat reminders
        before = Notification.objects.count()
        automation.run()
        self.assertEqual(Notification.objects.count(), before)

    def test_automation_arrival_reminder_and_auto_delivery(self):
        from bees import automation, fulfilment
        o, (item,) = self._order()
        fulfilment.set_items_status([item], "handed_to_courier", by_staff=True)
        OrderItem.objects.filter(pk=item.pk).update(shipped_at=timezone.now() - self.td(days=12))
        automation.run()
        self.assertTrue(Notification.objects.filter(user=self.buyer, message__icontains="arrived?").exists())
        self.assertEqual(OrderItem.objects.get(pk=item.pk).fulfillment_status, "handed_to_courier")
        OrderItem.objects.filter(pk=item.pk).update(shipped_at=timezone.now() - self.td(days=15))
        # an open "not received" report stops auto-delivery
        claim = Claim.objects.create(order_item=item, user=self.buyer, reason="not_received", details="Nothing yet")
        automation.run()
        self.assertEqual(OrderItem.objects.get(pk=item.pk).fulfillment_status, "handed_to_courier")
        claim.status = "withdrawn"
        claim.save()
        automation.run()
        self.assertEqual(OrderItem.objects.get(pk=item.pk).fulfillment_status, "delivered")
        self.assertEqual(Order.objects.get(pk=o.pk).status, "delivered")

    def test_unanswered_claims_go_to_store_team(self):
        from bees import automation
        o, (item,) = self._order()
        self._deliver(item)
        claim = Claim.objects.create(order_item=item, user=self.buyer, reason="damaged", details="Cracked")
        Claim.objects.filter(pk=claim.pk).update(created_at=timezone.now() - self.td(days=4))
        automation.run()
        claim.refresh_from_db()
        self.assertEqual(claim.status, "escalated")
        self.assertTrue(Notification.objects.filter(user=self.staff, message__icontains=f"problem #{claim.id}").exists())

    def test_review_alert_and_public_seller_reply(self):
        o, (item,) = self._order()
        self._deliver(item)
        self.client.force_login(self.buyer)
        self.client.post(reverse("product_detail", args=[self.p1.id]), {"rating": "4", "comment": "Nice lamp"})
        self.assertTrue(Notification.objects.filter(user=self.owner, message__icontains="4★ review").exists())
        review = Review.objects.get()
        self.client.force_login(self.owner)
        self.client.post(reverse("seller_reply_review", args=[review.id]), {"reply": "Thank you!"})
        self.assertTrue(Notification.objects.filter(user=self.buyer, message__icontains="replied to your review").exists())
        self.client.logout()
        self.assertContains(self.client.get(reverse("product_detail", args=[self.p1.id])), "Thank you!")
        other = User.objects.create_user("os2", "os2@example.com", "pass12345")
        SellerAccount.objects.create(user=other, business_name="Other2", status="approved")
        self.client.force_login(other)
        self.assertEqual(self.client.post(reverse("seller_reply_review", args=[review.id]), {"reply": "hack"}).status_code, 404)

    def test_order_rules_in_settings(self):
        self.client.force_login(self.staff)
        self.assertContains(self.client.get(reverse("manage_settings")), "auto_cancel_days")


class ProPackTests(TestCase):
    """Seller score, emails in the customer's language, price drops, statements."""
    def setUp(self):
        from datetime import timedelta
        self.td = timedelta
        self.owner = User.objects.create_user("ppsell", "pps@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.owner, business_name="PP Shop", status="approved")
        self.staff = User.objects.create_superuser("ppboss", "ppb@example.com", "pass12345")
        self.buyer = User.objects.create_user("ppbuy", "ppbuy@example.com", "pass12345")
        self.product = make_product(name="Kettle", seller_account=self.seller, seller_name="PP Shop", price=Decimal("50.00"))

    def _order(self, days_ago=10, shipped_after=None, lang=""):
        o = Order.objects.create(user=self.buyer, email=self.buyer.email, full_name="Buyer", address="1 St", city="X", country="GB",
                                 payment_method="card", payment_status="paid", status="confirmed", language=lang)
        it = OrderItem(order=o, product=self.product, product_name="Kettle", price=Decimal("50.00"), quantity=1)
        it.apply_commission(self.seller)
        it.save()
        created = timezone.now() - self.td(days=days_ago)
        Order.objects.filter(pk=o.pk).update(created_at=created)
        if shipped_after is not None:
            OrderItem.objects.filter(pk=it.pk).update(fulfillment_status="handed_to_courier", shipped_at=created + self.td(days=shipped_after))
        o.refresh_from_db()
        return o, it

    def test_score_levels_and_warning(self):
        from bees import badges, performance
        for _ in range(4):
            self._order(shipped_after=1)
        data = performance.measure(self.seller)
        self.assertEqual(data["score"], 100)
        self.assertEqual(data["level"], "excellent")
        for _ in range(6):
            self._order(shipped_after=5)
        data = performance.measure(self.seller)
        self.assertEqual(data["late_pct"], 60)
        self.assertLess(data["score"], 75)
        self.assertTrue(data["tips"])
        # lots of refunded problems -> at risk, warned once, no Top seller
        for it in OrderItem.objects.filter(seller_account=self.seller)[:4]:
            Claim.objects.create(order_item=it, user=self.buyer, reason="damaged", details="broken", status="refunded")
        self.assertEqual(performance.update_all(), 1)
        self.seller.refresh_from_db()
        self.assertEqual(self.seller.performance_level, "at_risk")
        self.assertTrue(Notification.objects.filter(user=self.owner, message__icontains="seller score").exists())
        self.assertTrue(Notification.objects.filter(user=self.staff, message__icontains="at risk").exists())
        before = Notification.objects.count()
        performance.update_all()
        self.assertEqual(Notification.objects.count(), before)
        self.assertFalse(badges.qualifies(self.seller))
        self.client.force_login(self.owner)
        self.assertContains(self.client.get(reverse("seller_dashboard")), "Seller score")
        self.client.force_login(self.staff)
        self.assertContains(self.client.get(reverse("manage_seller", args=[self.seller.id])), "Seller score")

    def test_new_sellers_are_not_judged(self):
        from bees import performance
        self._order(shipped_after=6)
        self.assertEqual(performance.measure(self.seller)["level"], "new")

    def test_emails_in_customer_language(self):
        from django.core import mail
        from bees import order_emails
        o, it = self._order(lang="ur")
        mail.outbox.clear()
        order_emails.send(o, "delivered")
        msg = mail.outbox[-1]
        self.assertIn("پہنچ گیا", msg.subject)
        html = msg.alternatives[0][0]
        self.assertIn('dir="rtl"', html)
        self.assertIn("آرڈر #", html)
        self.assertIn("Buyer", html)  # names stay as they are
        # language saved from the site switcher is used for other emails
        self.client.force_login(self.buyer)
        with self.settings(LANGUAGE_CODE="en"):
            SiteSettings.objects.filter(pk=1).update(show_language_menu=True)
            self.client.get(reverse("set_language", args=["es"]))
        self.assertEqual(Profile.objects.get(user=self.buyer).language, "es")
        o2, _ = self._order()
        mail.outbox.clear()
        order_emails.send(o2, "shipped")
        self.assertIn("en camino", mail.outbox[-1].subject)
        # English stays English
        Profile.objects.filter(user=self.buyer).update(language="")
        mail.outbox.clear()
        order_emails.send(Order.objects.get(pk=o2.pk), "shipped")
        self.assertIn("on its way", mail.outbox[-1].subject)

    def test_price_drop_alerts(self):
        from django.core import mail
        from bees import price_alerts
        from bees.models import Wishlist
        Profile.objects.get_or_create(user=self.buyer)
        w = Wishlist.objects.create(user=self.buyer, product=self.product)
        self.assertEqual(w.price_seen, Decimal("50.00"))
        Product.objects.filter(pk=self.product.pk).update(price=Decimal("48.00"))  # only 4% - no alert
        self.assertEqual(price_alerts.run(), 0)
        Product.objects.filter(pk=self.product.pk).update(price=Decimal("40.00"))
        mail.outbox.clear()
        self.assertEqual(price_alerts.run(), 1)
        self.assertTrue(Notification.objects.filter(user=self.buyer, message__icontains="20% cheaper").exists())
        self.assertTrue(any("Price dropped" in m.subject for m in mail.outbox))
        self.assertEqual(price_alerts.run(), 0)  # not twice for the same drop
        # opted out: notification yes, email no
        Profile.objects.filter(user=self.buyer).update(price_drop_alerts=False)
        Product.objects.filter(pk=self.product.pk).update(price=Decimal("30.00"))
        mail.outbox.clear()
        price_alerts.run()
        self.assertEqual(mail.outbox, [])
        self.client.force_login(self.buyer)
        self.client.post(reverse("account_privacy"), {"action": "emails", "price_drop_alerts": "1", "language": "fr"})
        p = Profile.objects.get(user=self.buyer)
        self.assertTrue(p.price_drop_alerts)
        self.assertEqual(p.language, "fr")

    def test_monthly_statement_pdf(self):
        o, it = self._order(days_ago=0)
        month = timezone.localdate().strftime("%Y-%m")
        self.client.force_login(self.owner)
        self.assertContains(self.client.get(reverse("seller_earnings")), "Monthly statement")
        r = self.client.get(reverse("seller_statement") + f"?month={month}")
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn(month, r["Content-Disposition"])
        from bees import statements
        d = statements.data(self.seller, month)
        self.assertEqual(d["sales"], Decimal("50.00"))
        self.assertEqual(d["net"], d["sales"] - d["commission"])
        self.assertEqual(self.client.get(reverse("seller_statement") + "?month=garbage").status_code, 200)
        # team staff (no money access) can't download it
        from bees.models import OrganizationMember
        helper = User.objects.create_user("pphelp", "pph@example.com", "pass12345")
        OrganizationMember.objects.create(organization=self.seller, user=helper, role="staff")
        self.client.force_login(helper)
        self.assertNotEqual(self.client.get(reverse("seller_statement")).status_code, 200)


class CourierAndKeysTests(TestCase):
    """Courier booking (EasyPost / Shippo, APIs mocked) and the .env editor."""
    def setUp(self):
        import os as _os
        import tempfile
        self.os = _os
        self.owner = User.objects.create_user("cksell", "cks@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.owner, business_name="CK Shop", status="approved",
                                                   business_address="1 Mill Rd", city="Lahore", country="Pakistan", phone="+923001234567")
        self.boss = User.objects.create_superuser("ckboss", "ckb@example.com", "pass12345")
        self.buyer = User.objects.create_user("ckbuy", "ckbuy@example.com", "pass12345")
        self.product = make_product(name="Rug", seller_account=self.seller, seller_name="CK Shop", price=Decimal("80.00"))
        Product.objects.filter(pk=self.product.pk).update(weight_grams=1200)
        self.order = Order.objects.create(user=self.buyer, email=self.buyer.email, full_name="Ann Buyer", address="5 High St", city="London",
                                          postal_code="E1 6AN", country="GB", phone="+447700900000",
                                          payment_method="card", payment_status="paid", status="confirmed")
        self.item = OrderItem(order=self.order, product=self.product, product_name="Rug", price=Decimal("80.00"), quantity=1)
        self.item.apply_commission(self.seller)
        self.item.save()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        for k in ("EASYPOST_API_KEY", "SHIPPO_API_TOKEN", "SHIPPING_WEBHOOK_TOKEN"):
            self.addCleanup(self.os.environ.pop, k, None)

    def _provider(self, name):
        site = SiteSettings.load()
        site.shipping_provider = name
        site.save()

    def _fake(self, routes):
        """Mocks urllib: routes = [(method, url-fragment, response dict)]."""
        import io
        calls = []

        def opener(req, timeout=None):
            calls.append((req.get_method(), req.full_url, json.loads(req.data) if req.data else None, dict(req.header_items())))
            for method, frag, payload in routes:
                if req.get_method() == method and frag in req.full_url:
                    resp = mock.MagicMock()
                    resp.__enter__.return_value = io.BytesIO(json.dumps(payload).encode())
                    return resp
            raise AssertionError(f"unexpected call {req.get_method()} {req.full_url}")
        return mock.patch("urllib.request.urlopen", side_effect=opener), calls

    def test_manual_mode_until_configured(self):
        self.client.force_login(self.owner)
        page = self.client.get(reverse("seller_order", args=[self.order.id]))
        self.assertNotContains(page, "Get courier rates")
        self._provider("easypost")
        self.assertNotContains(self.client.get(reverse("seller_order", args=[self.order.id])), "Get courier rates")  # no key yet
        from bees.couriers import status_line
        self.assertEqual(status_line()[0], "missing")

    def test_easypost_booking_and_tracking(self):
        self._provider("easypost")
        self.os.environ["EASYPOST_API_KEY"] = "EZTK_testkey123"
        self.os.environ["SHIPPING_WEBHOOK_TOKEN"] = "hook-secret-token-123456"
        patcher, calls = self._fake([
            ("POST", "/v2/shipments/shp_1/buy", {"id": "shp_1", "tracking_code": "1Z999", "postage_label": {"label_url": "https://labels.example/1.png"},
                                                  "tracker": {"id": "trk_1", "public_url": "https://track.example/1Z999"},
                                                  "selected_rate": {"carrier": "DHLExpress", "service": "ExpressWorldwide", "rate": "41.20", "currency": "USD"}}),
            ("POST", "/v2/shipments", {"id": "shp_1", "rates": [
                {"id": "rate_b", "carrier": "DHLExpress", "service": "ExpressWorldwide", "rate": "41.20", "currency": "USD", "delivery_days": 3},
                {"id": "rate_a", "carrier": "Aramex", "service": "PriorityParcel", "rate": "29.50", "currency": "USD", "delivery_days": 6}]}),
            ("GET", "/v2/trackers/trk_1", {"id": "trk_1", "status": "delivered"}),
        ])
        self.client.force_login(self.owner)
        with patcher:
            self.assertContains(self.client.get(reverse("seller_order", args=[self.order.id])), "Get courier rates")
            r = self.client.post(reverse("seller_courier_rates", args=[self.order.id]))
            shipment = Shipment.objects.get()
            self.assertIn(f"shipment={shipment.id}", r.url)
            self.assertEqual([x["id"] for x in shipment.rates], ["rate_a", "rate_b"])  # cheapest first
            body = calls[0][2]["shipment"]
            self.assertEqual(body["from_address"]["country"], "PK")
            self.assertEqual(body["to_address"]["country"], "GB")
            self.assertIn("customs_info", body)  # international
            self.assertAlmostEqual(body["parcel"]["weight"], 42.3, places=1)  # 1200 g in ounces
            self.assertTrue(calls[0][3]["Authorization"].startswith("Basic "))
            page = self.client.get(r.url)
            self.assertContains(page, "Aramex")
            self.client.post(reverse("seller_courier_buy", args=[shipment.id]), {"rate": "rate_b"})
        shipment.refresh_from_db()
        self.assertEqual(shipment.status, "booked")
        self.assertEqual(shipment.tracking_number, "1Z999")
        self.assertEqual(OrderItem.objects.get(pk=self.item.pk).fulfillment_status, "handed_to_courier")
        self.order.refresh_from_db()
        self.assertEqual(self.order.tracking_number, "1Z999")
        self.assertEqual(self.client.get(reverse("seller_courier_label", args=[shipment.id])).url, "https://labels.example/1.png")
        # webhook: wrong token -> 404, right token -> re-checks with the API and delivers
        self.assertEqual(self.client.post("/shipping/webhook/nope/", "{}", content_type="application/json").status_code, 404)
        patcher, calls = self._fake([("GET", "/v2/trackers/trk_1", {"id": "trk_1", "status": "delivered"})])
        with patcher:
            r = self.client.post("/shipping/webhook/hook-secret-token-123456/", json.dumps({"result": {"id": "trk_1", "status": "in_transit"}}),
                                 content_type="application/json")
        self.assertEqual(r.status_code, 200)
        shipment.refresh_from_db()
        self.assertEqual(shipment.status, "delivered")  # trusted the API, not the posted body
        self.assertEqual(OrderItem.objects.get(pk=self.item.pk).fulfillment_status, "delivered")
        # other sellers can't see the label
        other = User.objects.create_user("ckother", "cko@example.com", "pass12345")
        SellerAccount.objects.create(user=other, business_name="Other CK", status="approved")
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse("seller_courier_label", args=[shipment.id])).status_code, 404)

    def test_shippo_booking_and_refresh_job(self):
        from bees import couriers
        self._provider("shippo")
        self.os.environ["SHIPPO_API_TOKEN"] = "shippo_test_abc123"
        patcher, calls = self._fake([
            ("POST", "/customs/declarations/", {"object_id": "cust_1"}),
            ("POST", "/shipments/", {"object_id": "shp_9", "rates": [{"object_id": "r1", "provider": "UPS", "servicelevel": {"name": "Worldwide Saver"},
                                                                        "amount": "55.00", "currency": "USD", "estimated_days": 4}]}),
            ("POST", "/transactions/", {"status": "SUCCESS", "object_id": "tx_1", "tracking_number": "UPS123", "label_url": "https://l/1.pdf",
                                        "tracking_url_provider": "https://ups.example/UPS123"}),
            ("GET", "/transactions/tx_1", {"object_id": "tx_1", "tracking_status": "TRANSIT"}),
        ])
        with patcher:
            shipment = couriers.get_rates(self.order, [self.item], self.seller, self.owner)
            self.assertEqual(calls[0][3]["Authorization"], "Shippo-token".replace("-token", "Token") + " shippo_test_abc123")
            self.assertEqual(calls[1][2]["customs_declaration"], "cust_1")
            couriers.buy(shipment, "r1", user=self.owner)
            self.assertEqual(shipment.tracking_number, "UPS123")
            self.assertEqual(couriers.refresh_active(), 1)
        shipment.refresh_from_db()
        self.assertEqual(shipment.status, "in_transit")

    def test_prepaid_return_label(self):
        from django.core import mail as _mail
        from .models import ReturnRequest
        self._provider("easypost")
        self.os.environ["EASYPOST_API_KEY"] = "EZTK_testkey123"
        Order.objects.filter(pk=self.order.pk).update(status="delivered", delivered_at=timezone.now())
        rr = ReturnRequest.objects.create(order_item=self.item, user=self.buyer, reason="Wrong colour")
        patcher, calls = self._fake([
            ("POST", "/v2/shipments/shp_r/buy", {"id": "shp_r", "tracking_code": "RET1", "postage_label": {"label_pdf_url": "https://labels.example/r.pdf"},
                                                  "tracker": {"id": "trk_r"}, "selected_rate": {"carrier": "USPS", "service": "First", "rate": "9.10", "currency": "USD"}}),
            ("POST", "/v2/shipments", {"id": "shp_r", "rates": [
                {"id": "r_exp", "carrier": "DHL", "service": "X", "rate": "30.00", "currency": "USD"},
                {"id": "r_cheap", "carrier": "USPS", "service": "First", "rate": "9.10", "currency": "USD"}]}),
            ("GET", "/v2/trackers/trk_r", {"id": "trk_r", "status": "delivered"}),
        ])
        self.client.force_login(self.boss)
        with patcher:
            self.client.post(reverse("manage_returns"), {"id": rr.id, "action": "approve_label"})
            label = Shipment.objects.get(is_return=True)
            body = calls[0][2]["shipment"]
            self.assertEqual((body["from_address"]["country"], body["to_address"]["country"]), ("GB", "PK"))  # customer -> seller
            self.assertEqual(calls[1][2], {"rate": {"id": "r_cheap"}})
            self.assertEqual((label.return_request, label.label_url, label.status), (rr, "https://labels.example/r.pdf", "booked"))
            self.assertTrue(any("return label" in m.subject.lower() for m in _mail.outbox))
            self.client.force_login(self.buyer)
            self.assertContains(self.client.get(reverse("my_orders")), "Print return label")
            from bees import couriers
            couriers.refresh(label)
        self.item.refresh_from_db()
        self.assertNotEqual(self.item.fulfillment_status, "handed_to_courier")
        label.refresh_from_db()
        self.assertEqual(label.status, "delivered")

    def test_courier_errors_are_friendly(self):
        from bees import couriers
        self._provider("easypost")
        self.os.environ["EASYPOST_API_KEY"] = "EZTK_testkey123"
        SellerAccount.objects.filter(pk=self.seller.pk).update(business_address="")
        self.seller.refresh_from_db()
        with self.assertRaisesMessage(couriers.CourierError, "street address"):
            couriers.get_rates(self.order, [self.item], self.seller, self.owner)
        import urllib.error
        SellerAccount.objects.filter(pk=self.seller.pk).update(business_address="1 Mill Rd")
        self.seller.refresh_from_db()
        with mock.patch("urllib.request.urlopen", side_effect=urllib.error.URLError("blocked")):
            with self.assertRaisesMessage(couriers.CourierError, "Couldn't reach"):
                couriers.get_rates(self.order, [self.item], self.seller, self.owner)

    def test_keys_page_edits_env_safely(self):
        import os as _os
        from bees import envfile
        env = _os.path.join(self.tmp.name, ".env")
        with open(env, "w") as fh:
            fh.write("# my settings\nSECRET_KEY=abc\nSTRIPE_SECRET_KEY=sk_test_OLDKEY1234\nDEFAULT_FROM_EMAIL=Shop <a@b.com>\n")
        with mock.patch.object(envfile, "path", return_value=env), mock.patch.object(envfile, "reload_app", return_value=False):
            staff = User.objects.create_user("ckstaff", "ckst@example.com", "pass12345", is_staff=True)
            self.client.force_login(staff)
            self.assertRedirects(self.client.get(reverse("manage_keys")), reverse("manage_dashboard"), fetch_redirect_response=False)
            self.client.force_login(self.boss)
            page = self.client.get(reverse("manage_keys")).content.decode()
            self.assertIn("••••1234", page)
            self.assertNotIn("OLDKEY", page)  # secrets never shown
            # wrong password: nothing changes
            self.client.post(reverse("manage_keys"), {"confirm_password": "nope", "STRIPE_SECRET_KEY": "sk_test_NEW"})
            self.assertIn("OLDKEY1234", open(env).read())
            # bad format refused
            self.client.post(reverse("manage_keys"), {"confirm_password": "pass12345", "STRIPE_SECRET_KEY": "hello"})
            self.assertIn("OLDKEY1234", open(env).read())
            # good change + new key + webhook secret + removal
            self.client.post(reverse("manage_keys"), {"confirm_password": "pass12345", "STRIPE_SECRET_KEY": "sk_test_NEWKEY5678",
                                                      "EASYPOST_API_KEY": "EZAK_live_123456789", "new_webhook_token": "1",
                                                      "remove_DEFAULT_FROM_EMAIL": "1", "SECRET_KEY": "hacked"})
            text = open(env).read()
            self.assertIn("STRIPE_SECRET_KEY=sk_test_NEWKEY5678", text)
            self.assertIn("EASYPOST_API_KEY=EZAK_live_123456789", text)
            self.assertIn("SHIPPING_WEBHOOK_TOKEN=", text)
            self.assertNotIn("DEFAULT_FROM_EMAIL", text)
            self.assertIn("SECRET_KEY=abc", text)  # locked key untouched
            self.assertIn("# my settings", text)
            self.assertEqual(oct(_os.stat(env).st_mode & 0o777), "0o600")
            self.assertTrue(any(f.startswith(".env.bak-") for f in _os.listdir(self.tmp.name)))
            self.assertTrue(AuditLog.objects.filter(action__icontains="STRIPE_SECRET_KEY").exists())
            self.assertFalse(AuditLog.objects.filter(action__icontains="NEWKEY").exists())
            self.assertContains(self.client.get(reverse("manage_keys")), "/shipping/webhook/")
        for k in ("STRIPE_SECRET_KEY", "EASYPOST_API_KEY"):
            _os.environ.pop(k, None)

    def test_env_quoting_round_trip(self):
        import os as _os
        from bees import envfile
        env = _os.path.join(self.tmp.name, ".env")
        with mock.patch.object(envfile, "path", return_value=env):
            envfile.write({"DEFAULT_FROM_EMAIL": 'Lumen "Market" <hi@x.com>'})
            self.assertEqual(envfile.read()["DEFAULT_FROM_EMAIL"], 'Lumen "Market" <hi@x.com>')
            from dotenv import dotenv_values
            self.assertEqual(dotenv_values(env)["DEFAULT_FROM_EMAIL"], 'Lumen "Market" <hi@x.com>')
            with self.assertRaises(envfile.EnvError):
                envfile.write({"EMAIL_HOST_PASSWORD": "a\nSECRET_KEY=x"})
        _os.environ.pop("DEFAULT_FROM_EMAIL", None)


class GlobalPackTests(TestCase):
    """PayPal, per-country tax, cookie consent and SEO."""
    def setUp(self):
        import os as _os
        self.os = _os
        self.user = User.objects.create_user("glob", "glob@example.com", "pass12345")
        self.product = make_product(price=Decimal("25.00"), stock=50)
        self.client.force_login(self.user)
        for k in ("PAYPAL_CLIENT_ID", "PAYPAL_CLIENT_SECRET", "PAYPAL_MODE", "PAYPAL_WEBHOOK_ID"):
            self.addCleanup(self.os.environ.pop, k, None)

    def _site(self, **values):
        site = SiteSettings.load()
        for k, v in values.items():
            setattr(site, k, v)
        site.save()

    def _paypal_on(self):
        self.os.environ.update({"PAYPAL_CLIENT_ID": "A" * 30, "PAYPAL_CLIENT_SECRET": "B" * 30, "PAYPAL_MODE": "sandbox"})

    def _fake(self, routes):
        import io
        calls = []

        def opener(req, timeout=None):
            body = req.data.decode() if req.data else ""
            calls.append((req.get_method(), req.full_url, body, dict(req.header_items())))
            for method, frag, payload in routes:
                if req.get_method() == method and frag in req.full_url:
                    data = payload(order_amount=None) if callable(payload) else payload
                    resp = mock.MagicMock()
                    resp.__enter__.return_value = io.BytesIO(json.dumps(data).encode())
                    return resp
            raise AssertionError(f"unexpected {req.get_method()} {req.full_url}")
        return mock.patch("urllib.request.urlopen", side_effect=opener), calls

    def _paypal_order(self, value="50.00"):
        return {"id": "PP-1", "status": "COMPLETED", "purchase_units": [{"custom_id": str(Order.objects.get().id), "payments": {"captures": [
            {"id": "CAP-1", "status": "COMPLETED", "amount": {"currency_code": "USD", "value": value}}]}}]}

    def test_paypal_full_flow(self):
        self._paypal_on()
        self.assertContains(self.client.get(reverse("checkout")) if self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 2}) else None, 'value="paypal"')
        patcher, calls = self._fake([
            ("POST", "/v1/oauth2/token", {"access_token": "tok", "expires_in": 3600}),
            ("POST", "/v2/checkout/orders", {"id": "PP-1", "status": "PAYER_ACTION_REQUIRED", "links": [{"rel": "payer-action", "href": "https://www.sandbox.paypal.com/checkoutnow?token=PP-1"}]}),
        ])
        with patcher:
            r = self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "paypal"})
        self.assertTrue(r.url.startswith("https://www.sandbox.paypal.com/"))
        order = Order.objects.get()
        self.assertEqual((order.payment_method, order.payment_status, order.paypal_order_id), ("paypal", "pending", "PP-1"))
        body = json.loads(calls[1][2])
        self.assertEqual(body["purchase_units"][0]["amount"], {"currency_code": "USD", "value": "50.00"})
        self.assertEqual(body["purchase_units"][0]["custom_id"], str(order.id))
        self.assertTrue(calls[1][3]["Authorization"].startswith("Bearer "))
        # back from PayPal: approved -> captured -> paid
        patcher, calls = self._fake([
            ("GET", "/v2/checkout/orders/PP-1", {"id": "PP-1", "status": "APPROVED"}),
            ("POST", "/v2/checkout/orders/PP-1/capture", self._paypal_order()),
        ])
        with patcher:
            r = self.client.get(reverse("paypal_return", args=[order.id]) + "?token=PP-1")
        self.assertRedirects(r, reverse("order_success", args=[order.id]), fetch_redirect_response=False)
        order.refresh_from_db()
        self.assertEqual((order.payment_status, order.status, order.paypal_capture_id), ("paid", "confirmed", "CAP-1"))
        self.assertEqual(order.card_label, "PayPal")
        # cancelling refunds to PayPal
        patcher, calls = self._fake([("POST", "/v2/payments/captures/CAP-1/refund", {"id": "RF-1", "status": "COMPLETED"})])
        with patcher:
            self.client.post(reverse("cancel_order", args=[order.id]), {"reason": ""})
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "refunded")
        self.assertTrue(order.events.filter(text__icontains="PayPal").exists())

    def test_paypal_amount_mismatch_is_not_paid(self):
        self._paypal_on()
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 2})
        patcher, _ = self._fake([("POST", "/v1/oauth2/token", {"access_token": "tok", "expires_in": 3600}),
                                 ("POST", "/v2/checkout/orders", {"id": "PP-1", "links": [{"rel": "approve", "href": "https://p/x"}]})])
        with patcher:
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "paypal"})
        order = Order.objects.get()
        patcher, _ = self._fake([("GET", "/v2/checkout/orders/PP-1", self._paypal_order(value="1.00"))])
        with patcher:
            self.client.get(reverse("paypal_return", args=[order.id]))
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "pending")

    def test_paypal_cancel_returns_cart_and_webhook_needs_verification(self):
        self._paypal_on()
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 1})
        patcher, _ = self._fake([("POST", "/v1/oauth2/token", {"access_token": "tok", "expires_in": 3600}),
                                 ("POST", "/v2/checkout/orders", {"id": "PP-1", "links": [{"rel": "approve", "href": "https://p/x"}]})])
        with patcher:
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "paypal"})
        order = Order.objects.get()
        self.client.get(reverse("paypal_cancel", args=[order.id]))
        order.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.assertTrue(self.client.session.get("cart"))
        self.assertEqual(self.client.post(reverse("paypal_webhook"), "{}", content_type="application/json").status_code, 400)

    def test_paypal_hidden_for_unsupported_currency(self):
        from bees import paypal
        self._paypal_on()
        with self.settings(STORE_CURRENCY="pkr"):
            self.assertFalse(paypal.is_configured())
            self.assertEqual(paypal.status_line()[0], "missing")

    def test_tax_by_country_exclusive_and_inclusive(self):
        from bees.models import TaxRate
        TaxRate.objects.create(country="GB", name="VAT", rate=Decimal("20"), includes_shipping=True)
        self._site(tax_percent=Decimal("0"), shipping_flat_fee=Decimal("10"), ship_from_country="PK")
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 2})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "country": "GB", "payment_method": "cod", "customer_tax_id": "GB123456789"})
        order = Order.objects.get()
        self.assertEqual(order.tax_amount, Decimal("12.00"))  # 20% of 50 goods + 10 shipping
        self.assertEqual(order.total, Decimal("72.00"))
        self.assertEqual(order.tax_label, "VAT 20%")
        self.assertEqual(order.customer_tax_id, "GB123456789")
        # tax-inclusive store: shown, not added
        self._site(prices_include_tax=True)
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 2})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "country": "GB", "payment_method": "cod"})
        o2 = Order.objects.order_by("-id").first()
        self.assertEqual(o2.total, Decimal("60.00"))
        self.assertEqual(o2.tax_amount, Decimal("10.00"))  # 60 - 60/1.2
        self.assertTrue(o2.tax_included)
        # US has no rate -> no tax
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 1})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "country": "US", "payment_method": "cod"})
        self.assertEqual(Order.objects.order_by("-id").first().tax_amount, Decimal("0"))
        # checkout page carries the tax table and duties note
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 1})
        resp = self.client.get(reverse("checkout"))
        page = resp.content.decode()
        self.assertIn('"GB": {"name": "VAT", "rate": 20.0', page)
        self.assertIn("duties-note", page)
        r = self.client.get(reverse("invoice_pdf", args=[order.id]))
        self.assertEqual(r.status_code, 200)

    def test_tax_rates_admin(self):
        from bees.models import TaxRate
        boss = User.objects.create_superuser("taxboss", "tb@example.com", "pass12345")
        self.client.force_login(boss)
        self.client.post(reverse("manage_taxes"), {"action": "presets"})
        self.assertTrue(TaxRate.objects.filter(country="DE", rate=Decimal("19")).exists())
        self.client.post(reverse("manage_taxes"), {"action": "save", "country": "DE", "name": "MwSt", "rate": "7"})
        self.assertEqual(TaxRate.objects.get(country="DE").rate, Decimal("7"))
        self.client.post(reverse("manage_taxes"), {"action": "save", "country": "XX", "rate": "7"})
        self.assertFalse(TaxRate.objects.filter(country="XX").exists())
        self.assertContains(self.client.get(reverse("manage_taxes")), "Germany")

    def test_cookie_banner_and_analytics_only_after_consent(self):
        self._site(cookie_banner=True, ga4_id="G-ABC123XYZ")
        page = self.client.get(reverse("home")).content.decode()
        self.assertIn('id="cookie-bar"', page)
        self.assertNotIn("googletagmanager.com/gtag/js?id=G-", page)  # loaded by script only after consent
        self.assertIn("ABC123XYZ", page)
        self.assertContains(self.client.get(reverse("privacy_page")), "GDPR")
        boss = User.objects.create_superuser("cboss", "cb@example.com", "pass12345")
        self.client.force_login(boss)
        from bees.manage_views import SiteSettingsForm
        form = SiteSettingsForm(data={"ga4_id": "<script>"}, instance=SiteSettings.load())
        form.is_valid()
        self.assertIn("ga4_id", form.errors)

    def test_seo_tags_structured_data_and_sitemap(self):
        self._site(show_language_menu=True, google_site_verification="abc123def456ghi")
        r = self.client.get(reverse("product_detail", args=[self.product.id]))
        html = r.content.decode()
        self.assertIn('hreflang="ar"', html)
        self.assertIn('hreflang="x-default"', html)
        self.assertIn('google-site-verification', html)
        self.assertIn('max-image-preview:large', html)
        import re
        blocks = [json.loads(b) for b in re.findall(r'<script type="application/ld\+json">(.*?)</script>', html, re.S)]
        types = [b.get("@type") for b in blocks] + [g.get("@type") for b in blocks for g in b.get("@graph", [])]
        for t in ("Product", "BreadcrumbList", "Organization", "WebSite"):
            self.assertIn(t, types)
        product = next(b for b in blocks if b.get("@type") == "Product")
        self.assertEqual(product["offers"]["price"], "25.00")
        self.assertIn("hasMerchantReturnPolicy", product["offers"])
        # search / filter pages aren't indexed; ?lang shows that language
        self.assertIn("noindex", self.client.get(reverse("all_products") + "?sort=price").content.decode())
        self.assertIn('lang="es"', self.client.get(reverse("home") + "?lang=es").content.decode())
        sm = self.client.get("/sitemap.xml").content.decode()
        self.assertIn('hreflang="fr"', sm)
        self.assertIn("<image:image>", sm)
        # custom SEO title
        Product.objects.filter(pk=self.product.pk).update(seo_title="Best Mug Ever")
        self.assertContains(self.client.get(reverse("product_detail", args=[self.product.id])), "<title>Best Mug Ever</title>")


class LiveUpdatesTests(TestCase):
    def setUp(self):
        from django.core.cache import caches
        caches["live"].clear()
        owner = User.objects.create_user("livemaker", "lm@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, business_name="Live Co")
        self.seller.status = "approved"
        self.seller.save()
        self.buyer = User.objects.create_user("livebuyer", "lb@example.com", "pass12345")
        self.conv = Conversation.objects.create(buyer=self.buyer, seller=self.seller)

    def _send(self, user, body):
        self.client.force_login(user)
        return self.client.post(reverse("message_send", args=[self.conv.id]), {"body": body}, HTTP_X_REQUESTED_WITH="fetch")

    def test_live_counts_and_new_alerts(self):
        self._send(self.buyer, "Hello there")
        self.client.force_login(self.seller.user)
        first = self.client.get(reverse("live")).json()
        self.assertEqual(first["seller_messages"], 1)
        self.assertEqual(first["notifications"], 1)
        self.assertEqual(first["new"], [])
        Notification.objects.create(user=self.seller.user, message="New order #9", link="/seller/orders/")
        nxt = self.client.get(reverse("live") + f"?since={first['latest']}").json()
        self.assertEqual([n["text"] for n in nxt["new"]], ["New order #9"])
        self.assertIn("no-store", self.client.get(reverse("live")).headers["Cache-Control"])
        self.client.logout()
        self.assertEqual(self.client.get(reverse("live")).status_code, 302)

    def test_typing_presence_and_seen(self):
        self._send(self.buyer, "Is it in stock?")
        self.client.post(reverse("message_typing", args=[self.conv.id]))  # buyer is typing again
        self.client.force_login(self.seller.user)
        data = self.client.get(reverse("message_poll", args=[self.conv.id]) + "?after=0").json()
        self.assertTrue(data["typing"])
        self.assertEqual(data["presence"], "online")
        self.assertTrue(data["seen"])  # buyer has nothing unread
        self._send(self.seller.user, "Yes!")
        self.client.force_login(self.seller.user)
        self.assertFalse(self.client.get(reverse("message_poll", args=[self.conv.id]) + "?after=0").json()["seen"])
        self.client.force_login(self.buyer)
        self.client.get(reverse("message_poll", args=[self.conv.id]) + "?after=0")
        self.client.force_login(self.seller.user)
        self.assertTrue(self.client.get(reverse("message_poll", args=[self.conv.id]) + "?after=0").json()["seen"])

    def test_hidden_tab_does_not_mark_read(self):
        self._send(self.buyer, "Hi")
        self.client.force_login(self.seller.user)
        self.client.get(reverse("message_poll", args=[self.conv.id]) + "?after=0&visible=0")
        self.conv.refresh_from_db()
        self.assertEqual(self.conv.seller_unread, 1)

    def test_no_alert_when_chat_is_open(self):
        from django.core import mail
        self.client.force_login(self.seller.user)
        self.client.get(reverse("message_poll", args=[self.conv.id]) + "?after=0")  # seller has the chat open
        mail.outbox.clear()
        self._send(self.buyer, "Quick question")
        self.assertFalse(Notification.objects.filter(user=self.seller.user).exists())
        self.assertEqual(mail.outbox, [])

    def test_email_is_sent_after_the_reply(self):
        from django.core import mail
        from bees import background
        calls = []
        real = background._run
        with mock.patch("bees.background._run", side_effect=lambda fn, a, k: (calls.append(fn.__name__), real(fn, a, k))):
            self._send(self.buyer, "Hello")
        self.assertIn("_deliver_email", calls)
        self.assertTrue(any("lm@example.com" in m.to for m in mail.outbox))

    def _photo(self, name="p.jpg", fmt="JPEG"):
        import io as _io
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile
        buf = _io.BytesIO()
        img = Image.new("RGB", (2400, 1200), "red")
        exif = Image.Exif()
        exif[0x010F] = "SecretCam"  # camera maker - must be stripped
        img.save(buf, fmt, exif=exif) if fmt == "JPEG" else img.save(buf, fmt)
        return SimpleUploadedFile(name, buf.getvalue(), content_type="image/jpeg")

    def test_photo_message_is_private_and_cleaned(self):
        from PIL import Image
        self.addCleanup(lambda: [m.image.delete(save=False) for m in Message.objects.exclude(image="")])
        if True:
            self.client.force_login(self.buyer)
            r = self.client.post(reverse("message_send", args=[self.conv.id]), {"body": "", "photo": self._photo()}, HTTP_X_REQUESTED_WITH="fetch")
            self.assertEqual(r.status_code, 200, r.content)
            url = r.json()["message"]["image"]
            msg = Message.objects.get()
            with msg.image.open("rb") as fh:
                img = Image.open(fh)
                self.assertLessEqual(max(img.size), 1600)
                self.assertNotIn(0x010F, img.getexif())
            self.assertEqual(self.client.get(url).status_code, 200)
            self.client.force_login(self.seller.user)
            self.assertEqual(self.client.get(url).status_code, 200)
            outsider = User.objects.create_user("peek", "peek@example.com", "pass12345")
            self.client.force_login(outsider)
            self.assertEqual(self.client.get(url).status_code, 404)
            self.client.force_login(self.seller.user)
            self.assertContains(self.client.get(reverse("seller_messages")), "Photo")

    def test_bad_photo_is_refused(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.buyer)
        r = self.client.post(reverse("message_send", args=[self.conv.id]),
                             {"photo": SimpleUploadedFile("x.jpg", b"not an image", content_type="image/jpeg")}, HTTP_X_REQUESTED_WITH="fetch")
        self.assertEqual(r.status_code, 400)
        self.assertFalse(Message.objects.exists())

    def test_saved_replies(self):
        self.client.force_login(self.seller.user)
        self.client.post(reverse("seller_quick_replies"), {"quick_replies": "Hi there!\n\nShips tomorrow", "next": "https://evil.example/"})
        self.seller.refresh_from_db()
        self.assertEqual(self.seller.quick_replies, "Hi there!\nShips tomorrow")
        r = self.client.get(reverse("seller_conversation", args=[self.conv.id]))
        self.assertContains(r, "Ships tomorrow")
        r = self.client.post(reverse("seller_quick_replies"), {"quick_replies": "x", "next": "https://evil.example/"})
        self.assertNotIn("evil.example", r.url)
        self.client.force_login(self.buyer)
        self.assertEqual(self.client.post(reverse("seller_quick_replies"), {"quick_replies": "x"}).status_code, 404)

    def test_installable_app_files(self):
        m = self.client.get(reverse("pwa_manifest")).json()
        self.assertEqual(m["display"], "standalone")
        icon = self.client.get(m["icons"][0]["src"])
        self.assertEqual(icon["Content-Type"], "image/png")
        self.assertEqual(self.client.get("/app-icon-64.png").status_code, 404)
        sw = self.client.get(reverse("service_worker"))
        self.assertIn("navigate", sw.content.decode())
        self.assertContains(self.client.get(reverse("offline")), "offline")

    def test_inbox_shows_last_message_preview(self):
        self._send(self.buyer, "Do you ship to Dubai?")
        self.client.force_login(self.seller.user)
        self.assertContains(self.client.get(reverse("seller_messages")), "Do you ship to Dubai?")
        self.client.force_login(self.buyer)
        self.assertContains(self.client.get(reverse("messages_inbox")), "You: ")


class DealsAndOffersTests(TestCase):
    def test_flash_sale_ends(self):
        live = make_product(name="Live deal", is_flash_sale=True, flash_sale_ends=timezone.now() + timedelta(hours=3))
        make_product(name="Old deal", is_flash_sale=True, flash_sale_ends=timezone.now() - timedelta(minutes=1))
        names = list(Product.objects.flash().values_list("name", flat=True))
        self.assertEqual(names, ["Live deal"])
        r = self.client.get(reverse("product_detail", args=[live.id]))
        self.assertContains(r, "data-ends=")
        self.assertContains(r, "Deal ends in")

    def test_admin_sets_deal_end(self):
        admin = User.objects.create_superuser("deals", "d@example.com", "pass12345")
        p = make_product(name="Lamp")
        self.client.force_login(admin)
        self.client.post(reverse("manage_product_edit", args=[p.id]), {
            "name": "Lamp", "category": "skincare", "price": "10", "stock": "5", "discount_percent": "0",
            "image_url": "https://example.com/a.jpg", "is_flash_sale": "1", "flash_sale_ends": "2030-01-02T10:30",
            "approval_status": "approved", "bulk_min_qty": "3", "bulk_percent": "10"})
        p.refresh_from_db()
        self.assertEqual((p.flash_sale_ends.year, p.bulk_min_qty, p.bulk_percent), (2030, 3, 10))

    def test_quantity_offer_in_cart_and_order(self):
        self.client.force_login(User.objects.create_user("bulkbuyer", "bb@example.com", "pass12345"))
        p = make_product(name="Socks", price=Decimal("10.00"), stock=20, bulk_min_qty=3, bulk_percent=10)
        self.client.post(reverse("add_to_cart", args=[p.id]), {"quantity": 2})
        r = self.client.get(reverse("cart"))
        self.assertContains(r, "Buy 3 or more and save 10%")
        self.client.post(reverse("add_to_cart", args=[p.id]), {"quantity": 1})
        r = self.client.get(reverse("cart"))
        self.assertContains(r, "$27.00")
        self.assertContains(r, "Bundle discount applied")
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        item = OrderItem.objects.get(product=p)
        self.assertEqual((item.price, item.quantity), (Decimal("9.00"), 3))

    def test_offer_validation(self):
        owner = User.objects.create_user("offer", "o@example.com", "pass12345")
        seller = SellerAccount.objects.create(user=owner, business_name="Offer Co")
        seller.status = "approved"
        seller.save()
        self.client.force_login(owner)
        r = self.client.post(reverse("seller_add_product"), {"name": "X", "category": "skincare", "price": "5", "stock": "1",
                                                             "discount_percent": "0", "image_url": "https://example.com/x.jpg",
                                                             "bulk_min_qty": "3"}, follow=True)
        self.assertContains(r, "fill in both")


class ReferralProgramTests(TestCase):
    def setUp(self):
        self.friend_of = User.objects.create_user("sharer", "sharer@example.com", "pass12345")
        self.code = Profile.objects.get_or_create(user=self.friend_of)[0].referral_code
        site = SiteSettings.load()
        site.referral_friend_percent, site.referral_reward_percent = 15, 20
        site.save()

    def test_link_on_any_page_credits_signup(self):
        product = make_product(name="Shared thing")
        self.client.get(reverse("product_detail", args=[product.id]) + f"?ref={self.code}")
        self.client.post(reverse("signup"), {"username": "newbie", "email": "newbie@example.com",
                                             "password": "Str0ng-pass-123", "confirm_password": "Str0ng-pass-123",
                                             "password1": "Str0ng-pass-123", "password2": "Str0ng-pass-123"})
        newbie = User.objects.get(username="newbie")
        self.assertEqual(newbie.profile.referred_by, self.code)
        welcome = Coupon.objects.get(owner=newbie)
        self.assertEqual((welcome.purpose, welcome.percent_off), ("referral_welcome", 15))
        from .models import _reward_referrer_for
        _reward_referrer_for(newbie)
        reward = Coupon.objects.get(owner=self.friend_of)
        self.assertEqual((reward.purpose, reward.percent_off), ("referral_reward", 20))
        self.client.force_login(self.friend_of)
        r = self.client.get(reverse("referrals"))
        self.assertContains(r, reward.code)
        self.assertContains(r, "ne•••")
        self.assertContains(r, f"ref={self.code}")

    def test_admin_report_and_switch_off(self):
        admin = User.objects.create_superuser("radmin", "r@example.com", "pass12345")
        self.client.force_login(admin)
        self.assertContains(self.client.get(reverse("manage_referrals")), "Top referrers")
        site = SiteSettings.load()
        site.referral_enabled = False
        site.save()
        self.client.force_login(self.friend_of)
        self.assertRedirects(self.client.get(reverse("referrals")), reverse("profile"), fetch_redirect_response=False)


class ShippingAndTrackingTests(TestCase):
    def test_per_item_fee(self):
        from . import shipping
        ShippingZone.objects.create(name="US", countries="US", fee=Decimal("5"), per_item_fee=Decimal("2"))
        shipping.clear_cache()
        self.assertEqual(shipping.quote("US", Decimal("10"), 3)["fee"], Decimal("9"))
        self.assertEqual(shipping.quote("US", Decimal("10"))["fee"], Decimal("5"))

    def test_tracking_links(self):
        from .tracking import tracking_link
        self.assertIn("ups.com", tracking_link("ups", "1Z 999"))
        self.assertIn("1Z%20999", tracking_link("UPS", "1Z 999"))
        self.assertIn("17track", tracking_link("Some Local Courier", "AB12"))
        self.assertEqual(tracking_link("DHL", ""), "")
        self.assertEqual(tracking_link("DHL", "1", "https://my.courier/track/1"), "https://my.courier/track/1")

    def test_customer_sees_track_button(self):
        user = User.objects.create_user("trk", "trk@example.com", "pass12345")
        order = Order.objects.create(user=user, full_name="T", address="1", city="X", phone="1", country="US",
                                     status="shipped", courier_name="FedEx", tracking_number="123")
        self.client.force_login(user)
        self.assertContains(self.client.get(reverse("my_orders")), "fedex.com/fedextrack/?trknbr=123")
        admin = User.objects.create_superuser("tadmin", "t@example.com", "pass12345")
        self.client.force_login(admin)
        self.client.post(reverse("manage_order", args=[order.id]), {"action": "update", "status": "shipped", "courier_name": "Local",
                                                                   "tracking_number": "9", "tracking_url": "https://local.example/9"})
        order.refresh_from_db()
        self.assertEqual(order.tracking_link, "https://local.example/9")


class SellerBadgeTests(TestCase):
    def setUp(self):
        owner = User.objects.create_user("badge", "badge@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, business_name="Badge Co")
        self.seller.status = "approved"
        self.seller.save()
        self.product = make_product(name="Pot", seller_account=self.seller, seller_name="Badge Co")

    def test_staff_verifies_seller(self):
        admin = User.objects.create_superuser("vadmin", "v@example.com", "pass12345")
        self.client.force_login(admin)
        self.client.post(reverse("manage_seller", args=[self.seller.id]), {"action": "verify"})
        self.seller.refresh_from_db()
        self.assertTrue(self.seller.is_verified)
        self.assertContains(self.client.get(reverse("product_detail", args=[self.product.id])), "Verified seller")

    @override_settings(TOP_SELLER_MIN_ORDERS=2, TOP_SELLER_MIN_REVIEWS=1, TOP_SELLER_MIN_RATING=4.5)
    def test_top_seller_badge_is_earned_and_lost(self):
        from . import badges
        buyer = User.objects.create_user("b2", "b2@example.com", "pass12345")
        for _ in range(2):
            o = Order.objects.create(user=buyer, full_name="B", address="1", city="X", phone="1", country="US", status="delivered")
            item = OrderItem(order=o, product=self.product, product_name="Pot", price=Decimal("10"), quantity=1)
            item.apply_commission(self.seller)
            item.save()
        Review.objects.create(product=self.product, user=buyer, username="b2", rating=5, comment="Great")
        self.assertEqual(badges.update_all(), (1, 0))
        self.seller.refresh_from_db()
        self.assertTrue(self.seller.top_seller)
        self.client.force_login(self.seller.user)
        self.assertContains(self.client.get(reverse("seller_dashboard")), "Top seller")
        Review.objects.create(product=self.product, user=buyer, username="b2", rating=1, comment="Broke")
        Review.objects.create(product=self.product, user=buyer, username="b2", rating=1, comment="Broke")
        self.assertEqual(badges.update_all(), (0, 1))


# ---------------------------------------------------------------------------
# Ultra pack: COD confirmation, shipments, refunds, guides, translation,
# system health
# ---------------------------------------------------------------------------

from .models import OrderEvent, UIText  # noqa: E402


class CodConfirmationTests(TestCase):
    real_cod_confirmation = True

    def setUp(self):
        from django.core import mail as _mail
        self.mail = _mail
        owner = User.objects.create_user("codseller", "cs@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, business_name="Cod Co")
        self.seller.status = "approved"
        self.seller.save()
        self.product = make_product(name="Tea", price=Decimal("10.00"), stock=5, seller_account=self.seller, seller_name="Cod Co")
        self.buyer = User.objects.create_user("codbuyer", "cb@example.com", "pass12345")
        self.client.force_login(self.buyer)

    def _order(self):
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        return Order.objects.get()

    def test_unverified_buyer_confirms_by_email_then_seller_is_told(self):
        order = self._order()
        self.assertFalse(order.cod_confirmed)
        self.assertFalse(order.sellers_notified)
        self.assertTrue(any("confirm" in m.subject.lower() for m in self.mail.outbox))
        self.assertContains(self.client.get(reverse("my_orders")), "Please confirm this cash order")
        self.client.force_login(self.seller.user)
        self.client.post(reverse("seller_ship_order", args=[order.id]), {"status": "packed"})
        self.assertFalse(OrderItem.objects.filter(order=order, fulfillment_status="packed").exists())
        self.client.logout()
        with self.captureOnCommitCallbacks(execute=True):
            page = self.client.get(reverse("confirm_cod", args=[order.id, order.access_token]))
            self.assertContains(page, "Yes, confirm my order")  # a link scanner opening it doesn't confirm
            order.refresh_from_db()
            self.assertFalse(order.cod_confirmed)
            self.client.post(reverse("confirm_cod", args=[order.id, order.access_token]))
        order.refresh_from_db()
        self.assertTrue(order.cod_confirmed and order.sellers_notified)
        self.assertTrue(Notification.objects.filter(user=self.seller.user, message__icontains="New order").exists())
        self.assertEqual(self.client.get(reverse("confirm_cod", args=[order.id, "wrong"])).status_code, 404)

    def test_verified_buyer_skips_confirmation(self):
        from .models import Profile
        Profile.objects.update_or_create(user=self.buyer, defaults={"email_verified": True})
        order = self._order()
        self.assertTrue(order.cod_confirmed)

    def test_unconfirmed_orders_cancel_after_48_hours(self):
        from django.core.management import call_command
        order = self._order()
        Order.objects.filter(pk=order.pk).update(created_at=timezone.now() - timedelta(hours=49))
        call_command("release_unpaid_orders", stdout=io.StringIO())
        order.refresh_from_db()
        self.assertEqual((order.status, order.cancel_reason), ("cancelled", "not confirmed in time"))
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock, 5)

    def test_delivery_marks_cash_collected(self):
        from .models import Profile
        Profile.objects.update_or_create(user=self.buyer, defaults={"email_verified": True})
        order = self._order()
        order.status = "delivered"
        order.save()
        self.assertIsNotNone(order.cod_collected_at)
        self.assertTrue(OrderEvent.objects.filter(order=order, kind="cod").exists())


class ShipmentAndTrackingPageTests(TestCase):
    def test_every_order_gets_unique_shipment_number_and_courier(self):
        a = Order.objects.create(full_name="A", address="1", city="Lahore", phone="1", country="PK")
        b = Order.objects.create(full_name="B", address="1", city="Paris", phone="1", country="FR")
        self.assertRegex(a.shipment_number, r"^[A-Z]{2,4}-\d{4}-[2-9A-Z]{6}$")
        self.assertNotEqual(a.shipment_number, b.shipment_number)
        from .tracking import suggest_courier
        self.assertEqual((suggest_courier("PK"), suggest_courier("FR"), suggest_courier("BR")), ("TCS", "DPD", "DHL"))
        ShippingZone.objects.create(name="Pakistan", countries="PK", fee=Decimal("2"), courier="Leopards")
        self.assertEqual(suggest_courier("PK"), "Leopards")

    def test_checkout_picks_courier_and_public_tracking_page(self):
        user = User.objects.create_user("ship", "ship@example.com", "pass12345")
        self.client.force_login(user)
        p = make_product(name="Mat", price=Decimal("10"), stock=3)
        self.client.post(reverse("add_to_cart", args=[p.id]))
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get()
        self.assertEqual(order.courier_name, "USPS")
        order.status = "shipped"
        order.save()
        self.assertTrue(order.uses_store_tracking)
        self.client.logout()
        r = self.client.get(reverse("track_shipment", args=[order.shipment_number]))
        self.assertContains(r, order.shipment_number)
        self.assertContains(r, "Handed to USPS")
        self.assertNotContains(r, "1 Main St")  # no street address on the public page
        r = self.client.get(reverse("track_lookup") + f"?n={order.shipment_number.lower()}")
        self.assertRedirects(r, reverse("track_shipment", args=[order.shipment_number]), fetch_redirect_response=False)


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy")
class RefundRecordTests(TestCase):
    def test_cancel_paid_order_refunds_to_card_with_reason(self):
        user = User.objects.create_user("refundme", "r@example.com", "pass12345")
        p = make_product(name="Cup", price=Decimal("12"), stock=4)
        order = Order.objects.create(user=user, full_name="R", address="1", city="X", phone="1", country="US",
                                     payment_method="card", payment_status="paid", status="confirmed",
                                     stripe_payment_intent="pi_123", card_brand="visa", card_last4="4242")
        OrderItem.objects.create(order=order, product=p, product_name="Cup", price=Decimal("12"), quantity=1)
        self.client.force_login(user)
        refund = mock.Mock(id="re_999")
        with mock.patch("stripe.Refund.create", return_value=refund) as create:
            r = self.client.post(reverse("cancel_order", args=[order.id]), {"reason": "Changed my mind"}, follow=True)
        create.assert_called_once()
        order.refresh_from_db()
        self.assertEqual((order.payment_status, order.refund_id, order.cancel_reason), ("refunded", "re_999", "Changed my mind"))
        self.assertContains(r, "Visa •••• 4242")
        self.assertTrue(OrderEvent.objects.filter(order=order, kind="refunded").exists())

    def test_card_details_saved_after_payment(self):
        from . import payments
        order = Order.objects.create(full_name="R", address="1", city="X", phone="1", stripe_payment_intent="pi_1")
        intent = {"latest_charge": {"payment_method_details": {"card": {"brand": "mastercard", "last4": "4444"}}}}
        with self.settings(STRIPE_CARD_DETAILS=True), mock.patch("stripe.PaymentIntent.retrieve", return_value=intent):
            payments.remember_card(order)
        order.refresh_from_db()
        self.assertEqual(order.card_label, "Mastercard •••• 4444")


class BuyingGuideTests(TestCase):
    def setUp(self):
        owner = User.objects.create_user("guide", "g@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, business_name="Guide Co")
        self.seller.status = "approved"
        self.seller.save()
        self.client.force_login(owner)

    def test_seller_draft_and_save_guide(self):
        r = self.client.post(reverse("ai_buying_guide_draft"), {"name": "Yoga Mat", "category": "fashion"})
        self.assertEqual(r.json()["source"], "builtin")
        self.assertTrue(r.json()["guide"]["pros"])
        self.client.post(reverse("seller_add_product"), {"name": "Yoga Mat", "category": "fashion", "price": "20", "stock": "3",
                                                         "discount_percent": "0", "image_url": "https://example.com/m.jpg",
                                                         "guide_uses": "Yoga\\nPilates", "guide_pros": "Non-slip", "guide_cons": "Heavy"})
        p = Product.objects.get(name="Yoga Mat")
        p.approval_status = "approved"
        p.save()
        self.assertEqual(p.guide["pros"], ["Non-slip"])
        r = self.client.get(reverse("product_detail", args=[p.id]))
        self.assertContains(r, "Buying guide")
        self.assertContains(r, "Non-slip")
        self.assertContains(r, "message-seller" if False else "Message seller")

    def test_shopper_ai_guide_and_question(self):
        p = make_product(name="Kettle", seller_account=self.seller, seller_name="Guide Co")
        reply = mock.MagicMock()
        reply.__enter__.return_value.read.return_value = json.dumps({"content": [{"type": "text", "text":
            '{"uses": ["Boiling water fast"], "pros": ["Quick", "Easy"], "cons": ["Check the voltage"]}'}]}).encode()
        with self.settings(AI_API_KEY="k"), mock.patch("urllib.request.urlopen", return_value=reply):
            r = self.client.post(reverse("ai_product_guide", args=[p.id]))
        self.assertEqual(r.json()["guide"]["pros"], ["Quick", "Easy"])
        p.refresh_from_db()
        self.assertEqual(p.guide_source, "ai")
        answer = mock.MagicMock()
        answer.__enter__.return_value.read.return_value = json.dumps({"content": [{"type": "text", "text": "Yes, it suits travel."}]}).encode()
        with self.settings(AI_API_KEY="k"), mock.patch("urllib.request.urlopen", return_value=answer):
            r = self.client.post(reverse("ai_product_question", args=[p.id]), {"question": "Good for travel?"})
        self.assertEqual(r.json()["answer"], "Yes, it suits travel.")
        self.assertEqual(self.client.post(reverse("ai_product_question", args=[p.id]), {"question": "Hi"}).status_code, 400)


class WholePageTranslationTests(TestCase):
    def test_pages_translate_and_user_content_does_not(self):
        p = make_product(name="Plain Mug", description="Add to cart this mug")
        s = self.client.session
        s["site_lang"] = "fr"
        s.save()
        r = self.client.get(reverse("product_detail", args=[p.id]))
        html = r.content.decode()
        self.assertIn("Tous les produits", html)       # header (built-in pack)
        self.assertIn("Add to cart this mug", html)   # seller's description untouched
        r = self.client.get(reverse("help_support"))
        self.assertIn("Comment passer commande ?", r.content.decode())
        s = self.client.session
        s["site_lang"] = "ar"
        s.save()
        self.assertIn("قارن", self.client.get(reverse("all_products")).content.decode())

    def test_owner_translation_overrides_and_admin_page(self):
        from . import i18n
        UIText.objects.create(source="Plain Mug", fr="Tasse simple")
        i18n.clear_cache()
        p = make_product(name="Plain Mug")
        s = self.client.session
        s["site_lang"] = "fr"
        s.save()
        self.assertIn("Tasse simple", self.client.get(reverse("product_detail", args=[p.id])).content.decode())
        admin = User.objects.create_superuser("tr", "tr@example.com", "pass12345")
        self.client.force_login(admin)
        r = self.client.get(reverse("manage_translations") + "?lang=fr&show=all")
        self.assertContains(r, "Plain Mug")
        row = UIText.objects.get(source="Plain Mug")
        self.client.post(reverse("manage_translations"), {"action": "save", "id": row.id, "lang": "fr", "text": "Mug uni"})
        row.refresh_from_db()
        self.assertEqual(row.fr, "Mug uni")
        UIText.objects.create(source="Brand new text")

        def fake_claude(req, timeout=None):
            # Answer for exactly the texts that were sent: pages viewed above
            # may have added other untranslated texts to the batch.
            prompt = json.loads(req.data)["messages"][0]["content"]
            sent = json.loads(prompt.split("Texts:\n", 1)[1])
            answer = ["Texte tout neuf" if s == "Brand new text" else "" for s in sent]
            reply = mock.MagicMock()
            reply.__enter__.return_value.read.return_value = json.dumps(
                {"content": [{"type": "text", "text": json.dumps(answer)}]}).encode()
            return reply

        with self.settings(AI_API_KEY="k"), mock.patch("urllib.request.urlopen", side_effect=fake_claude):
            self.client.post(reverse("manage_translations"), {"action": "ai", "lang": "fr"})
        self.assertEqual(UIText.objects.get(source="Brand new text").fr, "Texte tout neuf")

    def test_keep_english_clears_the_system_check_warning(self):
        from .manage_views import operations_checks
        from .i18n_pack import PACK
        builtin = next(iter(PACK))
        UIText.objects.create(source=builtin)  # built-in texts never count as missing
        UIText.objects.create(source="Zuma Lamp")
        UIText.objects.create(source="Kora Vase", fr="Vase Kora")

        def check():
            return next(c for c in operations_checks() if c["name"] == "Untranslated shop text")
        self.assertEqual(check()["level"], "warn")
        self.assertIn("French 1", check()["detail"])
        self.assertNotIn(builtin, str(check()))
        admin = User.objects.create_superuser("kp", "kp@example.com", "pass12345")
        self.client.force_login(admin)
        row = UIText.objects.get(source="Zuma Lamp")
        self.client.post(reverse("manage_translations"), {"action": "keep", "id": row.id, "lang": "fr"})
        row.refresh_from_db()
        self.assertEqual(row.fr, "Zuma Lamp")
        for code in ("ur", "roman", "ar", "es"):
            self.client.post(reverse("manage_translations"), {"action": "keep_all", "lang": code})
        self.assertEqual(UIText.objects.get(source=builtin).ar, "")
        self.assertEqual(check()["level"], "ok")

    def test_placeholders_dates_and_store_name(self):
        from .i18n import translate_html
        out = translate_html("<p>Placed Oct 8, 2026</p><p>6 items</p><h2>Sell on Acme</h2>", "es", "Acme")
        self.assertIn("Realizado el oct 8, 2026", out)
        self.assertIn("6 artículos", out)
        self.assertIn("Vende en Acme", out)


class SystemHealthTests(TestCase):
    def test_system_check_lists_operations_and_dashboard_warns(self):
        admin = User.objects.create_superuser("health", "h@example.com", "pass12345")
        self.client.force_login(admin)
        from .alerts import SERVER_ERROR, log
        log(f"{SERVER_ERROR}: GET /boom/ - ValueError: bad")
        r = self.client.get(reverse("manage_system"))
        self.assertContains(r, "Page crashes (24 hours)")
        self.assertContains(r, "Daily tasks")
        self.assertContains(r, "/boom/")
        cache.clear()
        self.assertContains(self.client.get(reverse("manage_dashboard")), "Open System check")


class SecurityPackTests(TestCase):
    """Google/Apple sign-in, seller 2FA, SMS (Twilio) and cash-on-delivery protection."""
    real_cod_confirmation = True

    def setUp(self):
        import os as _os
        self.os = _os
        for k in ("GOOGLE_CLIENT_ID", "GOOGLE_CLIENT_SECRET", "TWILIO_ACCOUNT_SID", "TWILIO_AUTH_TOKEN", "TWILIO_FROM",
                  "APPLE_SERVICES_ID", "APPLE_TEAM_ID", "APPLE_KEY_ID", "APPLE_PRIVATE_KEY"):
            self.addCleanup(self.os.environ.pop, k, None)
        self.product = make_product(name="Lamp", price=Decimal("40.00"), stock=20)
        self.buyer = User.objects.create_user("secbuyer", "sb@example.com", "pass12345")

    def _site(self, **values):
        site = SiteSettings.load()
        for k, v in values.items():
            setattr(site, k, v)
        site.save()

    def _twilio_on(self):
        self.os.environ.update({"TWILIO_ACCOUNT_SID": "AC" + "a" * 32, "TWILIO_AUTH_TOKEN": "t" * 32, "TWILIO_FROM": "+15005550006"})

    def _sms_catch(self):
        sent = []

        def opener(req, timeout=None):
            from urllib.parse import parse_qs
            self.assertIn("api.twilio.com", req.full_url)
            sent.append({k: v[0] for k, v in parse_qs(req.data.decode()).items()})
            resp = mock.MagicMock()
            resp.__enter__.return_value.status = 201
            return resp
        return mock.patch("urllib.request.urlopen", side_effect=opener), sent

    # --- Google sign in -------------------------------------------------
    def _id_token(self, **claims):
        import base64 as b64
        body = {"iss": "https://accounts.google.com", "aud": "gid.apps.googleusercontent.com", "exp": 9999999999,
                "sub": "g-123", "email": "new.person@gmail.com", "email_verified": True, "given_name": "New", "family_name": "Person"}
        body.update(claims)
        enc = lambda d: b64.urlsafe_b64encode(json.dumps(d).encode()).decode().rstrip("=")
        return f"{enc({'alg': 'RS256'})}.{enc(body)}.sig"

    def _google(self, **claims):
        self.os.environ.update({"GOOGLE_CLIENT_ID": "gid.apps.googleusercontent.com", "GOOGLE_CLIENT_SECRET": "secret"})
        r = self.client.get(reverse("social_start", args=["google"]) + "?next=/my-orders/")
        self.assertTrue(r.url.startswith("https://accounts.google.com/"))
        from urllib.parse import parse_qs, urlparse
        q = {k: v[0] for k, v in parse_qs(urlparse(r.url).query).items()}
        token = self._id_token(nonce=q["nonce"], **claims)

        def opener(req, timeout=None):
            self.assertIn("oauth2.googleapis.com/token", req.full_url)
            resp = mock.MagicMock()
            resp.__enter__.return_value.read.return_value = json.dumps({"id_token": token}).encode()
            return resp
        with mock.patch("urllib.request.urlopen", side_effect=opener):
            return self.client.get(reverse("social_callback", args=["google"]), {"state": q["state"], "code": "abc"})

    def test_google_sign_in_creates_and_links_account(self):
        self.assertNotContains(self.client.get(reverse("login")), "Continue with Google")
        r = self._google()
        self.assertRedirects(r, "/my-orders/", fetch_redirect_response=False)
        from .models import SocialAccount
        link = SocialAccount.objects.get()
        self.assertEqual((link.provider, link.uid, link.user.email, link.user.first_name), ("google", "g-123", "new.person@gmail.com", "New"))
        self.assertFalse(link.user.has_usable_password())
        self.assertEqual(int(self.client.session["_auth_user_id"]), link.user.pk)
        self.client.logout()
        self.assertContains(self.client.get(reverse("login")), "Google")
        self._google()  # second time: same user, no duplicate
        self.assertEqual(User.objects.filter(email="new.person@gmail.com").count(), 1)

    def test_google_rejects_wrong_audience_and_reused_state(self):
        r = self._google(aud="someone-elses-app")
        self.assertRedirects(r, reverse("login"), fetch_redirect_response=False)
        self.assertNotIn("_auth_user_id", self.client.session)
        r = self.client.get(reverse("social_callback", args=["google"]), {"state": "made-up", "code": "x"})
        self.assertRedirects(r, reverse("login"), fetch_redirect_response=False)

    def test_google_unverified_email_does_not_take_over_account(self):
        r = self._google(email="sb@example.com", email_verified=False)
        self.assertNotIn("_auth_user_id", self.client.session)
        from .models import SocialAccount
        self.assertFalse(SocialAccount.objects.exists())

    def test_google_with_2fa_asks_for_code(self):
        from .models import Profile
        Profile.objects.update_or_create(user=self.buyer, defaults={"totp_enabled": True, "totp_secret": "JBSWY3DPEHPK3PXP"})
        r = self._google(email="sb@example.com", sub="g-9")
        self.assertRedirects(r, reverse("login_code"), fetch_redirect_response=False)
        self.assertNotIn("_auth_user_id", self.client.session)

    # --- Seller 2FA ----------------------------------------------------
    def test_sellers_must_set_up_2fa_when_required(self):
        seller_user = User.objects.create_user("twofa", "2fa@example.com", "pass12345")
        acc = SellerAccount.objects.create(user=seller_user, business_name="Safe Shop")
        acc.status = "approved"
        acc.save()
        self.client.force_login(seller_user)
        self.assertEqual(self.client.get("/seller/").status_code, 200)
        self._site(require_seller_2fa=True)
        self.assertRedirects(self.client.get("/seller/"), reverse("two_factor_setup"), fetch_redirect_response=False)

    # --- SMS -----------------------------------------------------------
    def test_phone_numbers_are_normalised(self):
        from . import sms
        self.assertEqual(sms.normalize("0300-1234567", "PK"), "+923001234567")
        self.assertEqual(sms.normalize("+44 7700 900123"), "+447700900123")
        self.assertEqual(sms.normalize("00971501234567"), "+971501234567")
        self.assertEqual(sms.normalize("555 0100", "ZZ"), "")

    def test_order_updates_are_texted_and_customers_can_opt_out(self):
        from .models import Profile
        Profile.objects.update_or_create(user=self.buyer, defaults={"email_verified": True})
        self._twilio_on()
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        patcher, sent = self._sms_catch()
        with patcher, self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get()
        with patcher, self.captureOnCommitCallbacks(execute=True):
            order.status = "shipped"
            order.save()
        self.assertEqual(sent[-1]["To"], "+15550100")
        self.assertIn(f"#{order.id} is on its way", sent[-1]["Body"])
        self.assertEqual(sent[-1]["From"], "+15005550006")
        count = len(sent)
        self.client.post(reverse("account_privacy"), {"action": "emails"})  # all boxes unticked
        self.assertFalse(Profile.objects.get(user=self.buyer).sms_updates)
        with patcher, self.captureOnCommitCallbacks(execute=True):
            order.status = "delivered"
            order.save()
        self.assertEqual(len(sent), count)

    def test_sellers_get_a_text_for_new_orders(self):
        from . import seller_center
        self._twilio_on()
        owner = User.objects.create_user("smsseller", "ss@example.com", "pass12345")
        acc = SellerAccount.objects.create(user=owner, business_name="Text Shop", phone="03001234567", country="Pakistan")
        product = make_product(name="Mug", price=Decimal("5.00"), stock=5, seller_account=acc, seller_name="Text Shop")
        order = Order.objects.create(user=self.buyer, full_name="B", address="a", city="c", postal_code="1", country="PK", phone="1",
                                     payment_method="cod", payment_status="not_applicable")
        OrderItem.objects.create(order=order, product=product, product_name="Mug", price=Decimal("5.00"), quantity=2, seller_account=acc)
        patcher, sent = self._sms_catch()
        with patcher, self.captureOnCommitCallbacks(execute=True):
            seller_center.notify_new_order(order)
        self.assertEqual(sent[0]["To"], "+923001234567")
        self.assertIn("2 items to ship", sent[0]["Body"])

    def test_no_keys_means_no_sms(self):
        from . import sms
        with mock.patch("urllib.request.urlopen") as urlopen:
            self.assertFalse(sms.send_now("+15550100", "hi"))
        urlopen.assert_not_called()

    # --- Cash on delivery protection ---------------------------------------
    def _cod_checkout(self, **extra):
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        patcher, sent = self._sms_catch()
        with patcher, self.captureOnCommitCallbacks(execute=True):
            r = self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod", **extra})
        return r, sent

    def test_cod_phone_code_confirms_order(self):
        from .models import Profile
        Profile.objects.update_or_create(user=self.buyer, defaults={"email_verified": True})
        self._twilio_on()
        self._site(cod_phone_check=True)
        r, sent = self._cod_checkout()
        order = Order.objects.get()
        self.assertFalse(order.cod_confirmed)  # even a verified email needs the phone code
        self.assertTrue(order.cod_otp_hash)
        import re as _re
        code = _re.search(r"is (\d{6})", sent[0]["Body"]).group(1)
        page = self.client.get(reverse("order_success", args=[order.id]))
        self.assertContains(page, "One last step")
        self.assertContains(page, 'autocomplete="one-time-code"')
        # confirm_cod email link can't bypass the phone check
        self.client.post(reverse("confirm_cod", args=[order.id, order.access_token]))
        order.refresh_from_db()
        self.assertFalse(order.cod_confirmed)
        wrong = "000000" if code != "000000" else "111111"
        self.client.post(reverse("cod_verify", args=[order.id]), {"code": wrong})
        order.refresh_from_db()
        self.assertEqual((order.cod_confirmed, order.cod_otp_tries), (False, 1))
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("cod_verify", args=[order.id]), {"code": code})
        order.refresh_from_db()
        self.assertTrue(order.cod_confirmed and order.sellers_notified is not None)
        self.assertTrue(order.events.filter(text__icontains="phone").exists())

    def test_cod_code_locks_after_five_wrong_tries_and_expires(self):
        from . import cod
        self._twilio_on()
        self._site(cod_phone_check=True)
        self._cod_checkout()
        order = Order.objects.get()
        for _ in range(5):
            cod.check_code(order, "999999" if order.cod_otp_hash != cod._hash(order, "999999") else "888888")
        ok, msg = cod.check_code(order, "123456")
        self.assertFalse(ok)
        self.assertIn("Too many", msg)
        patcher, sent = self._sms_catch()
        with patcher:
            self.assertTrue(cod.send_code(order))
        Order.objects.filter(pk=order.pk).update(cod_otp_sent_at=timezone.now() - timedelta(minutes=31))
        order.refresh_from_db()
        self.assertIn("expired", cod.check_code(order, "123456")[1])

    def test_cod_phone_check_needs_a_textable_number(self):
        self._twilio_on()
        self._site(cod_phone_check=True)
        r, _ = self._cod_checkout(phone="12", country="ZZ")
        self.assertFalse(Order.objects.exists())

    def test_cod_limit_and_block_list(self):
        from .models import CodBlock
        self._site(cod_max_order=Decimal("30.00"))
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[self.product.id]))
        page = self.client.get(reverse("checkout"))
        self.assertNotContains(page, 'value="cod"')
        self.assertContains(page, "Cash on delivery is available for orders up to")
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertFalse(Order.objects.exists())
        self._site(cod_max_order=Decimal("0"))
        from . import cod
        CodBlock.objects.create(kind="phone", value=cod.phone_key("+1 555-0100"))
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertFalse(Order.objects.exists())
        CodBlock.objects.all().delete()
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertTrue(Order.objects.exists())

    def test_refused_parcels_cancel_and_block_repeat_refusers(self):
        from . import cod
        from .models import CodBlock
        self._site(cod_refusal_limit=2)
        admin = User.objects.create_superuser("codadmin", "ca@example.com", "pass12345")
        orders = []
        for _ in range(2):
            o = Order.objects.create(user=self.buyer, email="sb@example.com", full_name="B", address="a", city="c", postal_code="1",
                                     country="US", phone="+1 (555) 0100", payment_method="cod", payment_status="not_applicable", status="shipped")
            OrderItem.objects.create(order=o, product=self.product, product_name="Lamp", price=Decimal("40"), quantity=1)
            orders.append(o)
        self.client.force_login(admin)
        self.client.post(reverse("manage_order", args=[orders[0].id]), {"action": "cod_refused"})
        orders[0].refresh_from_db()
        self.assertTrue(orders[0].cod_refused)
        self.assertEqual(orders[0].status, "cancelled")
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock, 21)
        self.assertFalse(CodBlock.objects.exists())
        cod.mark_refused(orders[1])
        self.assertEqual(set(CodBlock.objects.values_list("kind", flat=True)), {"email", "phone"})
        self.assertTrue(cod.blocked(phone="0015550100"))
        self.assertIn("pay online", cod.problem(Decimal("1"), email="SB@example.com"))
        page = self.client.get(reverse("manage_cod"))
        self.assertContains(page, "Refused on delivery")
        self.client.post(reverse("manage_cod"), {"action": "add", "kind": "email", "value": "Fraud@Example.com"})
        self.assertTrue(CodBlock.objects.filter(kind="email", value="fraud@example.com").exists())
        block = CodBlock.objects.get(value="fraud@example.com")
        self.client.post(reverse("manage_cod"), {"action": "delete", "id": block.id})
        self.assertFalse(CodBlock.objects.filter(pk=block.pk).exists())


class BulkProductTests(TestCase):
    def setUp(self):
        owner = User.objects.create_user("bulkseller", "bs@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, business_name="Bulk Co")
        self.seller.status = "approved"
        self.seller.save()
        self.mine = make_product(name="Old mug", price=Decimal("10.00"), stock=5, seller_account=self.seller, seller_name="Bulk Co")
        self.other = make_product(name="Not mine", price=Decimal("8.00"), stock=5)
        self.client.force_login(owner)

    def _upload(self, content, name="p.csv", url="seller_bulk", **extra):
        from django.core.files.uploadedfile import SimpleUploadedFile
        f = SimpleUploadedFile(name, content if isinstance(content, bytes) else content.encode())
        return self.client.post(reverse(url), {"action": "upload", "file": f, **extra})

    def test_csv_preview_then_save(self):
        csv_text = ("id,name,category,price,stock,image_url\n"
                    f"{self.mine.id},Old mug,{self.mine.category},12.00,9,\n"
                    ",New bowl,grocery,7.5,3,https://img.example.com/bowl.jpg\n"
                    f"{self.other.id},Hack,grocery,1,1,\n"
                    ",No photo,grocery,3,1,\n")
        r = self._upload(csv_text)
        self.assertContains(r, "Check before saving")
        self.assertContains(r, "isn&#x27;t yours")
        self.assertContains(r, "new products need image_url")
        self.mine.refresh_from_db()
        self.assertEqual(self.mine.price, Decimal("10.00"))  # nothing saved yet
        self.client.post(reverse("seller_bulk"), {"action": "confirm"})
        self.mine.refresh_from_db()
        self.assertEqual((self.mine.price, self.mine.stock, self.mine.approval_status), (Decimal("12.00"), 9, "approved"))
        new = Product.objects.get(name="New bowl")
        self.assertEqual((new.seller_account, new.approval_status, new.stock), (self.seller, "pending", 3))
        self.other.refresh_from_db()
        self.assertEqual(self.other.name, "Not mine")
        # confirming twice does nothing
        self.client.post(reverse("seller_bulk"), {"action": "confirm"})
        self.assertEqual(Product.objects.filter(name="New bowl").count(), 1)

    def test_name_change_goes_to_review(self):
        self._upload(f"id,name\n{self.mine.id},Better mug\n")
        self.client.post(reverse("seller_bulk"), {"action": "confirm"})
        self.mine.refresh_from_db()
        self.assertEqual((self.mine.name, self.mine.approval_status), ("Better mug", "pending"))

    def test_excel_round_trip_and_formula_safety(self):
        from openpyxl import load_workbook
        Product.objects.filter(pk=self.mine.pk).update(description="=HYPERLINK(\"http://evil\")")
        r = self.client.get(reverse("seller_bulk") + "?download=xlsx")
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb["products"]
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][:4], ("id", "name", "category", "price"))
        self.assertEqual(len(rows), 2)  # only my product
        self.assertTrue(rows[1][6].startswith("'="))
        ws.cell(row=2, column=4, value=15)
        ws.cell(row=2, column=7, value="Plain text")
        buf = io.BytesIO()
        wb.save(buf)
        self._upload(buf.getvalue(), name="p.xlsx")
        self.client.post(reverse("seller_bulk"), {"action": "confirm"})
        self.mine.refresh_from_db()
        self.assertEqual(self.mine.price, Decimal("15.00"))
        csv_resp = self.client.get(reverse("seller_bulk") + "?download=csv&what=template")
        self.assertIn(b"id,name,category,price", csv_resp.content)

    def test_sizes_and_colours_from_excel_and_csv(self):
        from openpyxl import Workbook, load_workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "products"
        ws.append(["name", "category", "price", "image_url"])
        ws.append(["Tee shirt", "fashion", 15, "https://img.example.com/t.jpg"])
        opt = wb.create_sheet("options")
        opt.append(["product", "size", "colour", "price", "stock"])
        opt.append(["Tee shirt", "S", "Red", "", 3])
        opt.append(["Tee shirt", "M", "Red", 17, 4])
        opt.append(["Tee shirt", "M", "red", "", 1])  # duplicate
        opt.append(["Nope", "S", "", "", 1])
        buf = io.BytesIO()
        wb.save(buf)
        r = self._upload(buf.getvalue(), name="p.xlsx")
        self.assertContains(r, "2 size/colour options")
        self.assertContains(r, "listed twice")
        self.client.post(reverse("seller_bulk"), {"action": "confirm"})
        tee = Product.objects.get(name="Tee shirt")
        self.assertTrue(tee.has_variants)
        self.assertEqual(tee.stock, 7)
        self.assertEqual(sorted((v.size, v.color, v.price) for v in tee.variants.all()), [("M", "Red", Decimal("17.00")), ("S", "Red", None)])
        v_s = tee.variants.get(size="S")
        # CSV for an existing product: keeps S (same id), drops M, adds L
        self._upload(f"product,size,color,stock\n{tee.id},S,Red,9\n{tee.id},L,Red,2\n")
        self.client.post(reverse("seller_bulk"), {"action": "confirm"})
        tee.refresh_from_db()
        self.assertEqual(tee.stock, 11)
        self.assertEqual(tee.variants.get(size="S").pk, v_s.pk)
        self.assertFalse(tee.variants.filter(size="M").exists())
        out = load_workbook(io.BytesIO(self.client.get(reverse("seller_bulk") + "?download=xlsx").content))
        self.assertIn("options", out.sheetnames)
        self.assertEqual(len(list(out["options"].iter_rows(values_only=True))), 3)

    def test_bulk_price_change(self):
        self.client.post(reverse("seller_bulk"), {"action": "change", "mode": "price_pct", "value": "-10"})
        self.mine.refresh_from_db()
        self.other.refresh_from_db()
        self.assertEqual((self.mine.price, self.other.price), (Decimal("9.00"), Decimal("8.00")))
        self.client.post(reverse("seller_bulk"), {"action": "change", "mode": "stock_set", "value": "2.5"})
        self.mine.refresh_from_db()
        self.assertEqual(self.mine.stock, 5)

    def test_admin_imports_store_products_live(self):
        admin = User.objects.create_superuser("bulkadmin", "ba@example.com", "pass12345")
        self.client.force_login(admin)
        self._upload("name,category,price,stock,image_url\nStore lamp,Table lamp,30,4,https://img.example.com/l.jpg\n", url="manage_bulk")
        self.client.post(reverse("manage_bulk"), {"action": "confirm"})
        p = Product.objects.get(name="Store lamp")
        self.assertEqual((p.seller_account, p.approval_status, p.category), (None, "approved", "table-lamp"))
        self.assertEqual(self.client.get(reverse("manage_bulk")).status_code, 200)


class GiftCardTests(TestCase):
    def setUp(self):
        import os as _os
        self.os = _os
        self.buyer = User.objects.create_user("gcbuyer", "gb@example.com", "pass12345")
        self.friend = User.objects.create_user("gcfriend", "friend@example.com", "pass12345")

    def _stripe_on(self):
        return mock.patch("bees.payments.is_configured", return_value=True)

    def test_buy_with_card_then_redeem(self):
        from django.core import mail as _mail
        from .models import GiftCard, Profile
        session = mock.MagicMock(id="cs_gift", url="https://checkout.stripe.com/c/pay/cs_gift")
        fake = mock.MagicMock()
        fake.checkout.Session.create.return_value = session
        self.client.force_login(self.buyer)
        with self._stripe_on(), mock.patch("bees.payments._stripe", return_value=fake):
            self.assertContains(self.client.get(reverse("gift_cards")), 'value="card"')
            r = self.client.post(reverse("gift_cards"), {"amount": "custom", "custom": "75", "email": "gb@example.com", "from_name": "Ali",
                                                          "recipient_name": "Sara", "recipient_email": "friend@example.com", "message": "Happy birthday", "method": "card"})
        self.assertEqual(r.url, "https://checkout.stripe.com/c/pay/cs_gift")
        kwargs = fake.checkout.Session.create.call_args.kwargs
        self.assertEqual(kwargs["line_items"][0]["price_data"]["unit_amount"], 7500)
        card = GiftCard.objects.get()
        self.assertEqual((card.status, card.stripe_session_id), ("pending", "cs_gift"))
        # underpaid session doesn't activate
        from . import giftcards
        giftcards.confirm_stripe(card.id, {"id": "cs_gift", "payment_status": "paid", "amount_total": 100})
        card.refresh_from_db()
        self.assertEqual(card.status, "pending")
        with self.captureOnCommitCallbacks(execute=True):
            with mock.patch("bees.payments.retrieve_session", return_value={"id": "cs_gift", "payment_status": "paid", "amount_total": 7500}):
                page = self.client.get(reverse("gift_card_done", args=[card.id]) + "?session_id=cs_gift")
        self.assertContains(page, "Gift card sent")
        card.refresh_from_db()
        self.assertEqual(card.status, "active")
        sent = [m for m in _mail.outbox if "friend@example.com" in m.to]
        self.assertTrue(sent and card.code in sent[0].alternatives[0][0])
        self.assertIn("Happy birthday", sent[0].alternatives[0][0])
        # the friend redeems (code typed loosely)
        self.client.force_login(self.friend)
        self.client.post(reverse("gift_cards_mine"), {"code": card.code.lower().replace("-", " ")})
        self.assertEqual(Profile.objects.get(user=self.friend).store_credit, Decimal("75.00"))
        card.refresh_from_db()
        self.assertEqual((card.status, card.redeemed_by), ("redeemed", self.friend))
        # can't be used twice
        self.client.force_login(self.buyer)
        r = self.client.post(reverse("gift_cards_mine"), {"code": card.code}, follow=True)
        self.assertContains(r, "already been used")
        self.assertEqual(Profile.objects.get_or_create(user=self.buyer)[0].store_credit, Decimal("0"))

    def test_strangers_cannot_see_purchase_and_bad_amounts_rejected(self):
        from .models import GiftCard
        card = GiftCard.objects.create(code="GIFT-AAAA-BBBB-CCCC", amount=Decimal("20"), purchaser=self.buyer, purchaser_email="gb@example.com")
        self.client.force_login(self.friend)
        self.assertEqual(self.client.get(reverse("gift_card_done", args=[card.id])).status_code, 404)
        with self._stripe_on():
            self.client.post(reverse("gift_cards"), {"amount": "custom", "custom": "5000", "email": "a@b.com", "method": "card"})
        self.assertEqual(GiftCard.objects.count(), 1)
        r = self.client.post(reverse("gift_cards_mine"), {"code": "GIFT-AAAA-BBBB-CCCC"}, follow=True)
        self.assertContains(r, "couldn")  # unpaid cards can't be redeemed

    def test_paypal_purchase(self):
        from .models import GiftCard
        self.os.environ.update({"PAYPAL_CLIENT_ID": "A" * 30, "PAYPAL_CLIENT_SECRET": "B" * 30, "PAYPAL_MODE": "sandbox"})
        for k in ("PAYPAL_CLIENT_ID", "PAYPAL_CLIENT_SECRET", "PAYPAL_MODE"):
            self.addCleanup(self.os.environ.pop, k, None)
        state = {}

        def opener(req, timeout=None):
            url, method = req.full_url, req.get_method()
            if "/v1/oauth2/token" in url:
                data = {"access_token": "tok", "expires_in": 3600}
            elif method == "POST" and url.endswith("/v2/checkout/orders"):
                body = json.loads(req.data)
                state["amount"] = body["purchase_units"][0]["amount"]
                data = {"id": "PPG-1", "links": [{"rel": "payer-action", "href": "https://www.sandbox.paypal.com/checkoutnow?token=PPG-1"}]}
            elif method == "GET":
                data = {"id": "PPG-1", "status": "APPROVED"}
            else:
                card = GiftCard.objects.get()
                data = {"id": "PPG-1", "status": "COMPLETED", "purchase_units": [{"custom_id": f"gift-{card.id}", "payments": {"captures": [
                    {"id": "CAP", "status": "COMPLETED", "amount": {"currency_code": "USD", "value": "50.00"}}]}}]}
            resp = mock.MagicMock()
            resp.__enter__.return_value = io.BytesIO(json.dumps(data).encode())
            return resp
        with mock.patch("urllib.request.urlopen", side_effect=opener):
            r = self.client.post(reverse("gift_cards"), {"amount": "50", "email": "guest@example.com", "method": "paypal"})
            self.assertTrue(r.url.startswith("https://www.sandbox.paypal.com/"))
            self.assertEqual(state["amount"], {"currency_code": "USD", "value": "50.00"})
            card = GiftCard.objects.get()
            with self.captureOnCommitCallbacks(execute=True):
                self.client.get(reverse("gift_card_paypal_return", args=[card.id]) + "?token=PPG-1")
        card.refresh_from_db()
        self.assertEqual((card.status, card.payment_method), ("active", "paypal"))

    def test_admin_issue_and_void(self):
        from .models import GiftCard
        admin = User.objects.create_superuser("gcadmin", "ga@example.com", "pass12345")
        self.client.force_login(admin)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("manage_gift_cards"), {"action": "issue", "amount": "15", "email": "Lucky@Example.com", "note": "Sorry for the delay"})
        card = GiftCard.objects.get()
        self.assertEqual((card.status, card.recipient_email, card.payment_method, card.issued_by), ("active", "lucky@example.com", "staff", admin))
        self.assertContains(self.client.get(reverse("manage_gift_cards")), card.code)
        self.client.post(reverse("manage_gift_cards"), {"action": "void", "id": card.id})
        card.refresh_from_db()
        self.assertEqual(card.status, "void")
        self.client.force_login(self.friend)
        r = self.client.post(reverse("gift_cards_mine"), {"code": card.code}, follow=True)
        self.assertContains(r, "cancelled")


class ReportsTests(TestCase):
    UA = {"HTTP_USER_AGENT": "Mozilla/5.0 (Windows NT 10.0) Firefox/130.0"}

    def test_funnel_counts_humans_not_bots_and_report_numbers(self):
        from .models import DailyStat
        product = make_product(name="Report lamp", price=Decimal("20.00"), stock=10)
        c = self.client
        buyer = User.objects.create_user("repbuyer", "rb@example.com", "pass12345")
        c.force_login(buyer)
        c.get("/", **self.UA)
        c.get("/", **self.UA)  # same visitor, same day: one visit
        c.get(reverse("product_detail", args=[product.id]), **self.UA)
        c.get(reverse("product_detail", args=[product.id]), HTTP_USER_AGENT="Googlebot/2.1")
        c.get(reverse("search_products") + "?q=Unicorn+Kettle", **self.UA)
        c.post(reverse("add_to_cart", args=[product.id]), **self.UA)
        c.get(reverse("checkout"), **self.UA)
        c.get(reverse("checkout"), **self.UA)  # refresh doesn't double count
        from . import analytics

        def stat(kind, key=""):
            analytics.flush()
            return DailyStat.objects.filter(kind=kind, key=key).values_list("count", flat=True).first()
        self.assertEqual(stat("visit"), 1)
        self.assertEqual(stat("view", str(product.id)), 1)
        self.assertEqual(stat("search", "unicorn kettle"), 1)
        self.assertEqual(stat("search_zero", "unicorn kettle"), 1)
        self.assertEqual(stat("cart"), 1)
        self.assertEqual(stat("checkout"), 1)
        for country in ("GB", "GB", "DE"):
            o = Order.objects.create(user=buyer, email="rb@example.com", full_name="R", address="a", city="c", postal_code="1", country=country,
                                     phone="1", payment_method="card", payment_status="paid", status="confirmed", shipping_amount=Decimal("5"))
            OrderItem.objects.create(order=o, product=product, product_name="Report lamp", price=Decimal("20.00"), quantity=1)
        Order.objects.create(user=buyer, email="rb@example.com", full_name="R", address="a", city="c", postal_code="1", country="GB",
                             phone="1", payment_method="card", payment_status="refunded", status="cancelled")
        admin = User.objects.create_superuser("repadmin", "ra@example.com", "pass12345")
        c.force_login(admin)
        page = c.get(reverse("manage_reports") + "?range=7")
        r = page.context["r"]
        self.assertEqual((r["orders"], r["revenue"], r["aov"]), (3, Decimal("75.00"), Decimal("25.00")))
        self.assertEqual(r["by_country"][0]["name"], "United Kingdom")
        self.assertEqual(r["repeat_rate"], 100.0)
        self.assertEqual(r["funnel"][0]["n"], 1)
        self.assertContains(page, "unicorn kettle")
        csv_resp = c.get(reverse("manage_reports") + "?range=7&export=orders")
        self.assertEqual(csv_resp["Content-Type"], "text/csv; charset=utf-8")
        self.assertEqual(csv_resp.content.decode("utf-8-sig").count("\n"), 5)
        for kind in ("summary", "products", "sellers", "searches"):
            self.assertEqual(c.get(reverse("manage_reports") + f"?range=30&export={kind}").status_code, 200)
        self.assertEqual(c.get(reverse("manage_reports") + "?from=2026-01-01&to=2026-02-01").status_code, 200)
        # staff browsing doesn't count
        c.get(reverse("product_detail", args=[product.id]), **self.UA)
        self.assertEqual(stat("view", str(product.id)), 1)


class AutoPayoutTests(TestCase):
    def setUp(self):
        import os as _os
        self.os = _os
        owner = User.objects.create_user("payseller", "ps@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, business_name="Pay Co", country="United Kingdom")
        self.seller.status = "approved"
        self.seller.save()
        self.owner = owner

    def _available(self, amount):
        from . import seller_center as sc
        return mock.patch.object(sc, "balance", return_value={"available": Decimal(amount), "owed": Decimal(amount), "pending": Decimal("0"),
                                                               "earned": Decimal(amount), "paid": Decimal("0"), "requested": Decimal("0")})

    def _paypal(self, item_status="SUCCESS"):
        self.os.environ.update({"PAYPAL_CLIENT_ID": "A" * 30, "PAYPAL_CLIENT_SECRET": "B" * 30, "PAYPAL_MODE": "sandbox"})
        for k in ("PAYPAL_CLIENT_ID", "PAYPAL_CLIENT_SECRET", "PAYPAL_MODE"):
            self.addCleanup(self.os.environ.pop, k, None)
        calls = []

        def opener(req, timeout=None):
            url, method = req.full_url, req.get_method()
            calls.append((method, url, req.data))
            if "/v1/oauth2/token" in url:
                data = {"access_token": "tok", "expires_in": 3600}
            elif method == "POST":
                data = {"batch_header": {"payout_batch_id": "BATCH1", "batch_status": "PENDING"}}
            else:
                data = {"batch_header": {"batch_status": "SUCCESS"}, "items": [{"transaction_status": item_status, "transaction_id": "TX9",
                                                                                 "errors": {"message": "Receiver unregistered"}}]}
            resp = mock.MagicMock()
            resp.__enter__.return_value = io.BytesIO(json.dumps(data).encode())
            return resp
        return mock.patch("urllib.request.urlopen", side_effect=opener), calls

    def test_seller_sets_paypal_and_withdraws(self):
        from .models import Payout
        patcher, calls = self._paypal()
        self.client.force_login(self.owner)
        with patcher:
            self.client.post(reverse("seller_payout_settings"), {"payout_method": "paypal", "paypal_email": "Me@PayPal.com"})
        self.seller.refresh_from_db()
        self.assertEqual((self.seller.payout_method, self.seller.paypal_email), ("paypal", "me@paypal.com"))
        with patcher, self._available("40.00"), self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("seller_request_payout"), {"amount": "30"})
        payout = Payout.objects.get()
        self.assertEqual((payout.status, payout.amount, payout.method, payout.automatic), ("paid", Decimal("30.00"), "paypal", True))
        self.assertEqual(payout.reference, "PayPal TX9")
        body = json.loads([c for c in calls if c[0] == "POST" and "payouts" in c[1]][0][2])
        self.assertEqual(body["items"][0]["receiver"], "me@paypal.com")
        self.assertEqual(body["items"][0]["amount"], {"value": "30.00", "currency": "USD"})
        self.seller.refresh_from_db()
        self.assertEqual(self.seller.total_paid_out, Decimal("30.00"))

    def test_cannot_withdraw_more_than_available(self):
        from .models import Payout
        patcher, _ = self._paypal()
        self.seller.payout_method, self.seller.paypal_email = "paypal", "me@paypal.com"
        self.seller.save()
        self.client.force_login(self.owner)
        with patcher, self._available("10.00"):
            self.client.post(reverse("seller_request_payout"), {"amount": "50"})
        self.assertFalse(Payout.objects.exists())

    def test_failed_paypal_payout_gives_balance_back(self):
        from . import payouts
        patcher, _ = self._paypal(item_status="FAILED")
        self.seller.payout_method, self.seller.paypal_email = "paypal", "me@paypal.com"
        self.seller.save()
        with patcher, self._available("25.00"):
            payout = payouts.send(self.seller)
        self.assertEqual(payout.status, "failed")
        self.assertIn("Receiver unregistered", payout.error)
        self.seller.refresh_from_db()
        self.assertEqual(self.seller.total_paid_out, Decimal("0.00"))

    def test_stripe_connect_and_admin_pay_all(self):
        from . import payouts
        from .models import Payout
        fake = mock.MagicMock()
        fake.Account.create.return_value = {"id": "acct_123"}
        fake.AccountLink.create.return_value = {"url": "https://connect.stripe.com/setup/e/acct_123"}
        fake.Account.retrieve.return_value = {"id": "acct_123", "payouts_enabled": True, "capabilities": {"transfers": "active"}}
        fake.Transfer.create.return_value = {"id": "tr_1"}
        self.client.force_login(self.owner)
        with mock.patch("bees.payments.is_configured", return_value=True), mock.patch("bees.payments._stripe", return_value=fake):
            r = self.client.post(reverse("seller_payout_settings"), {"payout_method": "stripe"})
            self.assertRedirects(r, reverse("seller_stripe_connect"), fetch_redirect_response=False)
            r = self.client.get(reverse("seller_stripe_connect"))
            self.assertEqual(r.url, "https://connect.stripe.com/setup/e/acct_123")
            self.assertEqual(fake.Account.create.call_args.kwargs["country"], "GB")
            self.client.get(reverse("seller_stripe_return"))
            self.seller.refresh_from_db()
            self.assertEqual((self.seller.payout_method, self.seller.stripe_account_id, self.seller.stripe_payouts_ready), ("stripe", "acct_123", True))
            admin = User.objects.create_superuser("payadmin", "pa@example.com", "pass12345")
            self.client.force_login(admin)
            with self._available("60.00"), self.captureOnCommitCallbacks(execute=True):
                self.client.post(reverse("manage_sellers"), {"action": "pay_all"})
        payout = Payout.objects.get()
        self.assertEqual((payout.status, payout.reference, payout.method), ("paid", "Stripe tr_1", "stripe"))
        self.assertEqual(fake.Transfer.create.call_args.kwargs["destination"], "acct_123")
        self.assertEqual(fake.Transfer.create.call_args.kwargs["amount"], 6000)
        # weekly schedule
        site = SiteSettings.load()
        site.auto_payouts = "weekly"
        site.save()
        from datetime import date
        with mock.patch("django.utils.timezone.localdate", return_value=date(2026, 10, 12)):
            self.assertTrue(payouts.due_today())  # a Monday
        with mock.patch("django.utils.timezone.localdate", return_value=date(2026, 10, 13)):
            self.assertFalse(payouts.due_today())


@override_settings(STRIPE_SECRET_KEY="sk_test_dummy", STRIPE_WEBHOOK_SECRET="whsec_dummy")
class PresentmentCurrencyTests(TestCase):
    """Charging in the shopper's currency."""
    def setUp(self):
        cache.clear()
        Currency.objects.all().delete()
        Currency.objects.create(code="EUR", symbol="€", rate=Decimal("0.9"), active=True)
        Currency.objects.create(code="KWD", symbol="KD", rate=Decimal("0.31"), active=True)
        self.user = User.objects.create_user("eurbuyer", "eb@example.com", "pass12345")
        self.product = make_product(price=Decimal("25.00"), stock=5)
        self.client.force_login(self.user)

    def _on(self, value=True):
        site = SiteSettings.load()
        site.charge_local_currency = value
        site.save()

    def _checkout(self, code="EUR"):
        self.client.get(reverse("set_currency", args=[code]))
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 2})
        fake = mock.MagicMock(id="cs_eur", url="https://checkout.stripe.com/c/pay/cs_eur")
        with mock.patch("stripe.checkout.Session.create", return_value=fake) as create:
            self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "card"})
        return Order.objects.get(), create.call_args.kwargs

    def test_off_by_default_charges_store_currency(self):
        order, kwargs = self._checkout()
        self.assertEqual(order.charge_currency, "")
        self.assertEqual(kwargs["line_items"][0]["price_data"]["currency"], "usd")

    def test_card_charged_in_euros_and_refunds_converted(self):
        from . import payments
        self._on()
        order, kwargs = self._checkout()
        self.assertEqual((order.charge_currency, order.charge_rate), ("eur", Decimal("0.900000")))
        expected = (order.total * Decimal("0.9")).quantize(Decimal("0.01"))
        self.assertEqual(order.charge_amount, expected)
        self.assertEqual(len(kwargs["line_items"]), 1)
        self.assertEqual(kwargs["line_items"][0]["price_data"]["currency"], "eur")
        self.assertEqual(kwargs["line_items"][0]["price_data"]["unit_amount"], int(expected * 100))
        # a USD-sized payment no longer counts; the EUR one does
        order.stripe_session_id = "cs_eur"
        self.assertEqual(payments.mark_order_paid(order.id, {"id": "cs_eur", "payment_status": "paid", "amount_total": order.total_cents, "currency": "usd"})[1], False)
        _o, paid = payments.mark_order_paid(order.id, {"id": "cs_eur", "payment_status": "paid", "amount_total": order.pay_cents, "currency": "eur", "payment_intent": "pi_1"})
        self.assertTrue(paid)
        order.refresh_from_db()
        with mock.patch("stripe.Refund.create") as refund:
            payments.refund_amount(order, Decimal("10.00"), "r1")
        self.assertEqual(refund.call_args.kwargs["amount"], 900)
        page = self.client.get(reverse("order_success", args=[order.id]))
        self.assertContains(page, f"{order.charge_amount} EUR")

    def test_three_decimal_currency_falls_back_to_store_currency(self):
        self._on()
        order, kwargs = self._checkout("KWD")
        self.assertEqual(order.charge_currency, "")
        self.assertEqual(kwargs["line_items"][0]["price_data"]["currency"], "usd")

    def test_paypal_amount_in_euros(self):
        from . import paypal
        self._on()
        self.client.get(reverse("set_currency", args=["EUR"]))
        order = Order.objects.create(user=self.user, full_name="E", address="a", city="c", postal_code="1", country="DE", phone="1",
                                     payment_method="paypal", payment_status="pending", shipping_amount=Decimal("10"))
        from .currency import apply_presentment, _display
        cur = Currency.objects.get(code="EUR")
        token = _display.set(cur)
        try:
            apply_presentment(order, "paypal")
        finally:
            _display.reset(token)
        self.assertEqual(paypal._amount(order), {"currency_code": "EUR", "value": "9.00"})


class StaffRoleTests(TestCase):
    def setUp(self):
        self.owner = User.objects.create_superuser("owner1", "o@example.com", "pass12345")

    def test_owner_adds_support_agent_with_limited_access(self):
        from django.core import mail as _mail
        from .models import Profile
        self.client.force_login(self.owner)
        self.client.post(reverse("manage_team"), {"action": "add", "email": "agent@example.com", "role": "support"})
        agent = User.objects.get(email="agent@example.com")
        self.assertTrue(agent.is_staff and not agent.has_usable_password())
        self.assertEqual(Profile.objects.get(user=agent).staff_role, "support")
        self.assertTrue(any("agent@example.com" in m.to for m in _mail.outbox))
        self.client.force_login(agent)
        self.assertEqual(self.client.get(reverse("manage_orders")).status_code, 200)
        self.assertEqual(self.client.get(reverse("manage_support")).status_code, 200)
        for blocked in ("manage_settings", "manage_reports", "manage_sellers", "manage_keys", "manage_team", "manage_gift_cards"):
            r = self.client.get(reverse(blocked))
            self.assertRedirects(r, reverse("manage_dashboard"), fetch_redirect_response=False, msg_prefix=blocked)
        page = self.client.get(reverse("manage_dashboard"))
        self.assertNotContains(page, reverse("manage_settings"))
        self.assertContains(page, reverse("manage_support"))
        # POSTs are blocked too
        self.client.post(reverse("manage_settings"), {"site_name": "Hacked"})
        self.assertNotEqual(SiteSettings.load().site_name, "Hacked")

    def test_change_role_and_remove(self):
        from .models import Profile
        member = User.objects.create_user("m1", "m1@example.com", "pass12345", is_staff=True)
        self.client.force_login(self.owner)
        self.client.post(reverse("manage_team"), {"action": "role", "user": member.id, "role": "finance"})
        self.assertEqual(Profile.objects.get(user=member).staff_role, "finance")
        self.client.force_login(member)
        self.assertEqual(self.client.get(reverse("manage_reports")).status_code, 200)
        self.client.force_login(self.owner)
        self.client.post(reverse("manage_team"), {"action": "remove", "user": member.id})
        member.refresh_from_db()
        self.assertFalse(member.is_staff)
        # managers (default) can't manage the team
        manager = User.objects.create_user("mg", "mg@example.com", "pass12345", is_staff=True)
        self.client.force_login(manager)
        self.assertEqual(self.client.get(reverse("manage_settings")).status_code, 200)
        self.assertRedirects(self.client.get(reverse("manage_team")), reverse("manage_dashboard"), fetch_redirect_response=False)


class SellerPromotionsAndAdsTests(TestCase):
    def setUp(self):
        owner = User.objects.create_user("promoseller", "pr@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=owner, business_name="Promo Co")
        self.seller.status = "approved"
        self.seller.save()
        self.owner = owner
        self.mine = make_product(name="Promo lamp", price=Decimal("40.00"), stock=10, seller_account=self.seller, seller_name="Promo Co")
        self.other = make_product(name="Other lamp", price=Decimal("20.00"), stock=10)
        self.buyer = User.objects.create_user("promobuyer", "pb@example.com", "pass12345")

    def test_seller_coupon_only_discounts_their_products_and_seller_pays(self):
        self.client.force_login(self.owner)
        self.client.post(reverse("seller_promotions"), {"action": "coupon", "code": "promo25", "percent_off": "25", "per_user_limit": "1"})
        coupon = Coupon.objects.get(code="PROMO25")
        self.assertEqual(coupon.seller_account, self.seller)
        self.client.force_login(self.buyer)
        self.client.post(reverse("add_to_cart", args=[self.other.id]))
        r = self.client.post(reverse("apply_coupon"), {"coupon_code": "PROMO25"}, follow=True)
        self.assertContains(r, "for products from Promo Co only")
        self.client.post(reverse("add_to_cart", args=[self.mine.id]))
        self.client.post(reverse("apply_coupon"), {"coupon_code": "PROMO25"})
        page = self.client.get(reverse("checkout"))
        self.assertEqual(page.context["pricing"]["discount"], Decimal("10.00"))
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get()
        self.assertEqual((order.discount_amount, order.coupon_code), (Decimal("0"), "PROMO25"))
        prices = {i.product_name: i.price for i in order.items.all()}
        self.assertEqual(prices, {"Promo lamp": Decimal("30.00"), "Other lamp": Decimal("20.00")})

    def test_timed_sale_starts_and_restores_prices(self):
        from . import promotions
        from .models import Promotion
        self.client.force_login(self.owner)
        now = timezone.localtime()
        self.client.post(reverse("seller_promotions"), {"action": "sale", "name": "Flash", "percent_off": "20",
                                                        "starts_at": now.strftime("%Y-%m-%dT%H:%M"), "ends_at": (now + timedelta(days=1)).strftime("%Y-%m-%dT%H:%M")})
        promo = Promotion.objects.get()
        self.assertEqual(promo.status, "active")
        self.mine.refresh_from_db()
        self.other.refresh_from_db()
        self.assertEqual((self.mine.price, self.mine.old_price, self.other.price), (Decimal("32.00"), Decimal("40.00"), Decimal("20.00")))
        Promotion.objects.filter(pk=promo.pk).update(ends_at=timezone.now() - timedelta(minutes=1))
        promotions.tick()
        self.mine.refresh_from_db()
        self.assertEqual((self.mine.price, self.mine.old_price), (Decimal("40.00"), None))
        self.assertEqual(Promotion.objects.get().status, "ended")

    def test_sponsored_product_click_is_charged_from_earnings(self):
        from . import ads
        from .models import AdCampaign, AdStat
        self.client.force_login(self.owner)
        self.client.post(reverse("seller_ads"), {"action": "create", "product": self.mine.id, "bid": "0.50", "daily_budget": "1.00"})
        campaign = AdCampaign.objects.get()
        # no earnings yet: not shown
        self.assertEqual(ads.pick(), [])
        with mock.patch.object(SellerAccount, "amount_owed", new_callable=mock.PropertyMock, return_value=100.0):
            cache.delete("ads:eligible")  # the eligible list is cached for a minute
            shopper = self.client_class()
            shopper.force_login(self.buyer)
            ua = {"HTTP_USER_AGENT": "Mozilla/5.0 Firefox/130"}
            page = shopper.get(reverse("search_products") + "?q=lamp", **ua)
            self.assertContains(page, "Sponsored")
            url = ads.click_url(campaign)
            self.assertRedirects(shopper.get(url, **ua), reverse("product_detail", args=[self.mine.id]), fetch_redirect_response=False)
            shopper.get(url, **ua)  # same visitor, same day: free
            self.assertEqual(shopper.get(reverse("ad_click", args=[campaign.id]) + "?s=bad", **ua).status_code, 404)
            sneaky = self.client_class()
            sneaky.get(url, **ua)  # new session, same address: free (click-fraud guard)
            other = self.client_class()
            other.get(url, REMOTE_ADDR="203.0.113.7", **ua)
            third = self.client_class()
            third.get(url, REMOTE_ADDR="203.0.113.8", **ua)  # over the daily budget: not charged
        self.seller.refresh_from_db()
        campaign.refresh_from_db()
        self.assertEqual((self.seller.ad_spend, campaign.spent), (Decimal("1.00"), Decimal("1.00")))
        ads.flush_impressions()
        stat = AdStat.objects.get()
        self.assertEqual((stat.spend, stat.clicks), (Decimal("1.00"), 3))  # free clicks aren't counted as paid clicks
        self.assertGreaterEqual(stat.impressions, 1)
        # seller's own clicks are free
        self.client.get(url, HTTP_USER_AGENT="Mozilla/5.0 Firefox/130")
        self.seller.refresh_from_db()
        self.assertEqual(self.seller.ad_spend, Decimal("1.00"))


class EmailCampaignTests(TestCase):
    def test_campaign_to_subscribers_with_unsubscribe_and_tracking(self):
        import re as _re
        from django.core import mail as _mail
        from .models import EmailCampaign, NewsletterSubscriber
        admin = User.objects.create_superuser("mkadmin", "mk@example.com", "pass12345")
        for e in ("a@example.com", "b@example.com"):
            NewsletterSubscriber.objects.create(email=e)
        p = make_product(name="Campaign mug", price=Decimal("9.00"))
        self.client.force_login(admin)
        self.client.post(reverse("manage_campaigns"), {"action": "save", "subject": "Big sale", "heading": "Hello!", "body": "Line one.\n\nLine two.",
                                                        "button_text": "Shop", "button_url": "/", "product_ids": str(p.id), "audience": "all", "coupon_code": "save10"})
        c = EmailCampaign.objects.get()
        self.assertContains(self.client.get(reverse("manage_campaign", args=[c.id])), "srcdoc")
        self.client.post(reverse("manage_campaign", args=[c.id]), {"action": "test", "subject": "Big sale", "heading": "Hello!", "body": "Line one.", "audience": "all"})
        self.assertEqual(_mail.outbox[-1].to, ["mk@example.com"])
        _mail.outbox.clear()
        self.client.post(reverse("manage_campaign", args=[c.id]), {"action": "send", "subject": "Big sale", "heading": "Hello!", "body": "Line one.\n\nLine two.",
                                                                    "button_text": "Shop", "button_url": "/", "product_ids": str(p.id), "audience": "all", "coupon_code": "SAVE10"})
        c.refresh_from_db()
        self.assertEqual((c.status, c.sent_count), ("sent", 2))
        self.assertEqual(sorted(m.to[0] for m in _mail.outbox), ["a@example.com", "b@example.com"])
        msg = _mail.outbox[0]
        html = msg.alternatives[0][0]
        self.assertIn("SAVE10", html)
        self.assertIn("Campaign mug", html)
        self.assertIn("List-Unsubscribe", msg.extra_headers)
        path = lambda pat: _re.search(r'(?:href|src)="https?://[^/"]+(' + pat + r')"', html).group(1).replace("&amp;", "&")
        unsub = path(r'/unsubscribe/\?[^"]+')
        click = path(r'/e/\d+/c/\?[^"]+')
        pixel = path(r'/e/\d+/o\.gif\?[^"]+')
        anon = self.client_class()
        self.assertEqual(anon.get(pixel).status_code, 200)
        self.assertEqual(anon.get(click).status_code, 302)
        self.assertEqual(anon.get(click.replace("t=", "t=x")).status_code, 404)
        c.refresh_from_db()
        self.assertEqual((c.opens, c.clicks), (1, 1))
        email = msg.to[0]
        self.assertContains(anon.get(unsub), "Stop news and offers")
        anon.post(unsub, {"List-Unsubscribe": "One-Click"})
        self.assertFalse(NewsletterSubscriber.objects.filter(email=email).exists())
        self.assertEqual(anon.get(unsub.replace("t=", "t=bad")).status_code, 404)


class WebPushTests(TestCase):
    def setUp(self):
        import os as _os
        from . import webpush
        pub, priv = webpush.generate_keys()
        _os.environ.update({"VAPID_PUBLIC_KEY": pub, "VAPID_PRIVATE_KEY": priv})
        for k in ("VAPID_PUBLIC_KEY", "VAPID_PRIVATE_KEY"):
            self.addCleanup(_os.environ.pop, k, None)
        self.user = User.objects.create_user("pushy", "push@example.com", "pass12345")

    def _browser_keys(self):
        import os as _os
        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric import ec
        from . import webpush
        key = ec.generate_private_key(ec.SECP256R1())
        pub = key.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint)
        auth = _os.urandom(16)
        return key, webpush.b64u(pub), webpush.b64u(auth), auth, pub

    def test_encryption_round_trip_rfc8291(self):
        import struct
        from cryptography.hazmat.primitives import serialization
        from cryptography.hazmat.primitives.asymmetric import ec
        from cryptography.hazmat.primitives.ciphers.aead import AESGCM
        from . import webpush
        key, p256dh, auth_b64, auth, ua_pub = self._browser_keys()
        body = webpush.encrypt(b'{"title":"Hi"}', p256dh, auth_b64)
        salt, rs, idlen = body[:16], struct.unpack("!L", body[16:20])[0], body[20]
        as_pub = body[21:21 + idlen]
        self.assertEqual((rs, idlen), (4096, 65))
        shared = key.exchange(ec.ECDH(), ec.EllipticCurvePublicKey.from_encoded_point(ec.SECP256R1(), as_pub))
        ikm = webpush._hkdf(auth, shared, b"WebPush: info\x00" + ua_pub + as_pub, 32)
        cek = webpush._hkdf(salt, ikm, b"Content-Encoding: aes128gcm\x00", 16)
        nonce = webpush._hkdf(salt, ikm, b"Content-Encoding: nonce\x00", 12)
        plain = AESGCM(cek).decrypt(nonce, body[21 + idlen:], None)
        self.assertEqual(plain, b'{"title":"Hi"}\x02')

    def test_subscribe_and_notifications_are_pushed(self):
        import jwt
        from .models import Notification, PushSubscription
        _key, p256dh, auth_b64, _auth, _pub = self._browser_keys()
        self.client.force_login(self.user)
        r = self.client.post(reverse("push_subscribe"), json.dumps({"endpoint": "https://fcm.googleapis.com/fcm/send/abc", "keys": {"p256dh": p256dh, "auth": auth_b64}}),
                             content_type="application/json")
        self.assertEqual(r.json(), {"ok": True})
        self.assertEqual(PushSubscription.objects.count(), 1)
        bad = self.client.post(reverse("push_subscribe"), json.dumps({"endpoint": "https://169.254.169.254/latest/meta-data", "keys": {"p256dh": p256dh, "auth": auth_b64}}),
                               content_type="application/json")
        self.assertEqual(bad.status_code, 400)  # no requests to internal addresses
        self.assertEqual(PushSubscription.objects.count(), 1)
        calls = []

        def opener(req, timeout=None):
            calls.append(req)
            resp = mock.MagicMock()
            resp.__enter__.return_value.status = 201
            return resp
        with mock.patch("urllib.request.urlopen", side_effect=opener), self.captureOnCommitCallbacks(execute=True):
            Notification.objects.create(user=self.user, message="Order #5 is on its way.", link="/my-orders/")
        self.assertEqual(len(calls), 1)
        req = calls[0]
        self.assertEqual(req.full_url, "https://fcm.googleapis.com/fcm/send/abc")
        self.assertEqual(req.get_header("Content-encoding"), "aes128gcm")
        auth_header = req.get_header("Authorization")
        token = auth_header.split("t=")[1].split(",")[0]
        claims = jwt.decode(token, options={"verify_signature": False})
        self.assertEqual(claims["aud"], "https://fcm.googleapis.com")
        # expired subscriptions are removed
        import urllib.error
        with mock.patch("urllib.request.urlopen", side_effect=urllib.error.HTTPError("u", 410, "Gone", {}, None)), self.captureOnCommitCallbacks(execute=True):
            Notification.objects.create(user=self.user, message="Another", link="/")
        self.assertFalse(PushSubscription.objects.exists())


class SearchAutocompleteTests(TestCase):
    def setUp(self):
        from django.core.cache import cache as _c
        _c.clear()
        self.p = make_product(name="Hydrating Moisturizer Cream", price=Decimal("12.00"), stock=5)
        make_product(name="Wireless Headphones", price=Decimal("50.00"), stock=5)

    def test_suggest_returns_products_categories_and_fixes_typos(self):
        r = self.client.get(reverse("search_suggest") + "?q=moistur").json()
        self.assertEqual(r["products"][0]["name"], "Hydrating Moisturizer Cream")
        self.assertTrue(r["products"][0]["url"].endswith(f"/product/{self.p.id}/"))
        self.assertEqual(r["results"], ["Hydrating Moisturizer Cream"])
        typo = self.client.get(reverse("search_suggest") + "?q=headfones").json()
        self.assertEqual(typo["did_you_mean"], "headphones")
        self.assertEqual(typo["products"][0]["name"], "Wireless Headphones")

    def test_search_page_autocorrects_when_nothing_found(self):
        page = self.client.get(reverse("search_products") + "?q=moisturiser")
        self.assertContains(page, "Showing results for")
        self.assertContains(page, "Hydrating Moisturizer Cream")
        exact = self.client.get(reverse("search_products") + "?q=moisturiser&exact=1")
        self.assertNotContains(exact, "Hydrating Moisturizer Cream")


class RecommendationTests(TestCase):
    def test_customers_also_bought(self):
        from django.core.cache import cache as _c
        _c.clear()
        a = make_product(name="Tea pot", price=Decimal("20"), stock=9)
        b = make_product(name="Tea cups", price=Decimal("10"), stock=9)
        c = make_product(name="Unrelated bike", price=Decimal("99"), stock=9)
        buyer = User.objects.create_user("recbuyer", "rb2@example.com", "pass12345")
        for _ in range(2):
            o = Order.objects.create(user=buyer, full_name="R", address="a", city="c", postal_code="1", country="GB", phone="1",
                                     payment_method="card", payment_status="paid", status="confirmed")
            OrderItem.objects.create(order=o, product=a, product_name=a.name, price=a.price, quantity=1)
            OrderItem.objects.create(order=o, product=b, product_name=b.name, price=b.price, quantity=1)
        page = self.client.get(reverse("product_detail", args=[a.id]))
        self.assertEqual([p.id for p in page.context["also_bought"]][:1], [b.id])
        self.assertNotIn(c.id, [p.id for p in page.context["also_bought"]])
        self.client.post(reverse("add_to_cart", args=[a.id]))
        cart = self.client.get(reverse("cart"))
        self.assertContains(cart, "Goes well with your cart")
        self.assertEqual(cart.context["recommended"][0].id, b.id)


class StagingSiteTests(TestCase):
    def test_staging_protects_and_redirects_everything(self):
        import base64 as b64
        from django.core import mail as _mail
        from django.core.management import call_command, CommandError
        from django.test import override_settings
        with self.assertRaises(CommandError):
            call_command("anonymize_data", "--yes")  # never on the live store
        with override_settings(STAGING=True, STAGING_PASSWORD="letmein", STAGING_EMAIL_TO="qa@example.com",
                               STAGING_REAL_EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
                               EMAIL_BACKEND="bees.staging.RedirectEmailBackend"):
            self.assertEqual(self.client.get("/").status_code, 401)
            auth = {"HTTP_AUTHORIZATION": "Basic " + b64.b64encode(b"team:letmein").decode()}
            page = self.client.get("/", **auth)
            self.assertContains(page, "TEST SITE")
            self.assertEqual(page["X-Robots-Tag"], "noindex, nofollow")
            self.assertContains(self.client.get("/robots.txt", **auth), "Disallow: /")
            from django.core.mail import send_mail
            send_mail("Your order", "hi", None, ["real.customer@gmail.com"])
            self.assertEqual(_mail.outbox[-1].to, ["qa@example.com"])
            self.assertIn("[TEST] Your order (to real.customer@gmail.com)", _mail.outbox[-1].subject)
            import os as _os
            from bees import sms
            with mock.patch.dict(_os.environ, {"TWILIO_ACCOUNT_SID": "AC" + "a" * 32, "TWILIO_AUTH_TOKEN": "t", "TWILIO_FROM": "+15005550006"}):
                self.assertFalse(sms.is_configured())
            u = User.objects.create_user("realname", "real@gmail.com", "pass12345", first_name="Real")
            call_command("anonymize_data", "--yes", stdout=io.StringIO())
            u.refresh_from_db()
            self.assertEqual((u.email, u.first_name, u.has_usable_password()), (f"user{u.pk}@example.test", "Test", False))


class HardeningTests(TestCase):
    def test_health_and_security_headers(self):
        r = self.client.get("/health/")
        self.assertEqual(r.json(), {"ok": True})
        page = self.client.get("/")
        self.assertIn("object-src 'none'", page["Content-Security-Policy"])
        self.assertIn("frame-ancestors 'none'", page["Content-Security-Policy"])
        self.assertIn("camera=()", page["Permissions-Policy"])

    def test_coupon_guessing_is_rate_limited(self):
        p = make_product(name="RL mug", price=Decimal("10"), stock=5)
        self.client.post(reverse("add_to_cart", args=[p.id]))
        for i in range(15):
            self.client.post(reverse("apply_coupon"), {"coupon_code": f"GUESS{i}"})
        r = self.client.post(reverse("apply_coupon"), {"coupon_code": "GUESS99"}, follow=True)
        self.assertContains(r, "Too many coupon tries")


class SellerApplicationTests(TestCase):
    def test_application_is_validated_and_terms_recorded(self):
        SellerAccount.objects.create(user=User.objects.create_user("taken", "t@example.com", "pass12345"), business_name="Sunny Shop")
        u = User.objects.create_user("applicant", "ap@example.com", "pass12345")
        self.client.force_login(u)
        base = {"account_type": "individual", "business_name": "sunny shop", "phone": "+92 300 1234567", "country": "Pakistan", "agree": "1"}
        r = self.client.post(reverse("become_seller"), base)
        self.assertContains(r, "already taken")
        r = self.client.post(reverse("become_seller"), {**base, "business_name": "Moon Shop", "agree": ""})
        self.assertContains(r, "accept the seller terms")
        r = self.client.post(reverse("become_seller"), {**base, "business_name": "Moon Shop", "country": "Atlantis"})
        self.assertContains(r, "choose your country")
        self.client.post(reverse("become_seller"), {**base, "business_name": "  Moon   Shop "})
        acc = SellerAccount.objects.get(user=u)
        self.assertEqual((acc.business_name, acc.status), ("Moon Shop", "pending"))
        self.assertIsNotNone(acc.terms_accepted_at)


class ListingReportTests(TestCase):
    def test_report_and_take_down(self):
        from django.core import mail as _mail
        from .models import ListingReport
        owner = User.objects.create_user("lrseller", "lr@example.com", "pass12345")
        acc = SellerAccount.objects.create(user=owner, business_name="LR Co", status="approved")
        p = make_product(name="Fake watch", seller_account=acc, seller_name="LR Co")
        self.client.post(reverse("report_listing", args=[p.id]), {"reason": "counterfeit", "details": "Logo is fake", "email": "Spotter@Example.com"})
        rep = ListingReport.objects.get()
        self.assertEqual((rep.reason, rep.email, rep.status), ("counterfeit", "spotter@example.com", "open"))
        admin = User.objects.create_superuser("lradmin", "lra@example.com", "pass12345")
        self.client.force_login(admin)
        self.assertContains(self.client.get(reverse("manage_listing_reports")), "Fake watch")
        self.client.post(reverse("manage_listing_reports"), {"id": rep.id, "action": "remove", "note": "Brand confirmed it's fake."})
        p.refresh_from_db()
        rep.refresh_from_db()
        self.assertEqual((p.approval_status, rep.status), ("rejected", "removed"))
        self.assertTrue(any(m.to == ["spotter@example.com"] for m in _mail.outbox))
        self.assertTrue(Notification.objects.filter(user=owner, message__icontains="taken down").exists())


# ---------------------------------------------------------------------------
# REST API (bees/api)
# ---------------------------------------------------------------------------

class ApiTests(TestCase):
    """The REST API follows the same rules as the website."""

    def setUp(self):
        from .models import Profile
        self.buyer = User.objects.create_user("apibuyer", "apibuyer@example.com", "pass12345")
        Profile.objects.get_or_create(user=self.buyer)
        self.seller_user = User.objects.create_user("apiseller", "apiseller@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.seller_user, status="approved", business_name="Api Studio")
        self.live = make_product(name="Rose Serum", price=Decimal("25.00"), seller_account=self.seller, seller_name="Api Studio")
        self.cheap = make_product(name="Plain Soap", price=Decimal("3.00"), category="grocery")
        self.hidden = make_product(name="Secret Serum", approval_status="pending", seller_account=self.seller, seller_name="Api Studio")

    # --- helpers -----------------------------------------------------------
    def token_for(self, username, password="pass12345", **extra):
        r = self.client.post("/api/v1/auth/token/", {"username": username, "password": password, **extra},
                             content_type="application/json")
        self.assertEqual(r.status_code, 201, r.content)
        return r.json()["token"]

    def get(self, path, token=None, **params):
        headers = {"HTTP_AUTHORIZATION": f"Bearer {token}"} if token else {}
        return self.client.get(path, params, **headers)

    def send(self, method, path, data=None, token=None):
        headers = {"HTTP_AUTHORIZATION": f"Bearer {token}"} if token else {}
        return getattr(self.client, method)(path, data or {}, content_type="application/json", **headers)

    def order_for(self, user, product, seller=None, **kw):
        order = Order.objects.create(user=user, email=user.email, full_name="Api Buyer", address="1 St", city="Austin",
                                     country="US", phone="+15550100", payment_method="cod", status="confirmed", **kw)
        item = OrderItem.objects.create(order=order, product=product, product_name=product.name, price=product.price,
                                        quantity=2, seller_account=seller)
        return order, item

    # --- catalog -----------------------------------------------------------
    def test_index_and_docs(self):
        r = self.get("/api/v1/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("/api/v1/products/", r.json()["catalog"]["products"])
        self.assertEqual(self.get("/api/v1/schema/").status_code, 200)
        self.assertEqual(self.get("/api/v1/docs/").status_code, 200)

    def test_products_only_live(self):
        r = self.get("/api/v1/products/")
        self.assertEqual(r.status_code, 200)
        names = [p["name"] for p in r.json()["results"]]
        self.assertIn("Rose Serum", names)
        self.assertNotIn("Secret Serum", names)
        self.seller.status = "suspended"
        self.seller.save()
        names = [p["name"] for p in self.get("/api/v1/products/").json()["results"]]
        self.assertNotIn("Rose Serum", names)
        self.assertEqual(self.get(f"/api/v1/products/{self.hidden.id}/").status_code, 404)

    def test_product_fields(self):
        p = self.get(f"/api/v1/products/{self.live.id}/").json()
        self.assertEqual(p["price"], "25.00")
        self.assertTrue(p["url"].startswith("http"))
        self.assertEqual(p["seller"]["name"], "Api Studio")
        self.assertEqual(p["category"]["code"], "skincare")
        self.assertTrue(p["in_stock"])

    def test_search_filters_and_sort(self):
        r = self.get("/api/v1/products/", search="serum").json()
        self.assertEqual([p["name"] for p in r["results"]], ["Rose Serum"])
        r = self.get("/api/v1/products/", category="grocery").json()
        self.assertEqual([p["name"] for p in r["results"]], ["Plain Soap"])
        r = self.get("/api/v1/products/", max_price="10").json()
        self.assertEqual([p["name"] for p in r["results"]], ["Plain Soap"])
        r = self.get("/api/v1/products/", sort="price_asc").json()
        self.assertEqual(r["results"][0]["name"], "Plain Soap")
        r = self.get("/api/v1/products/", seller_id=self.seller.id).json()
        self.assertEqual([p["name"] for p in r["results"]], ["Rose Serum"])
        self.assertEqual(self.get("/api/v1/products/", seller_id="x").status_code, 400)

    def test_categories_and_reviews_and_seller(self):
        cats = {c["code"]: c["product_count"] for c in self.get("/api/v1/categories/").json()}
        self.assertEqual(cats["skincare"], 1)
        Review.objects.create(product=self.live, username="Ann", rating=5, comment="Lovely")
        r = self.get(f"/api/v1/products/{self.live.id}/reviews/").json()
        self.assertEqual(r["results"][0]["comment"], "Lovely")
        self.assertEqual(self.get(f"/api/v1/sellers/{self.seller.id}/").json()["name"], "Api Studio")
        self.seller.status = "pending"
        self.seller.save()
        self.assertEqual(self.get(f"/api/v1/sellers/{self.seller.id}/").status_code, 404)

    # --- sign-in ------------------------------------------------------------
    def test_token_sign_in_and_out(self):
        from .models import ApiToken
        r = self.send("post", "/api/v1/auth/token/", {"username": "apibuyer", "password": "wrong"})
        self.assertEqual(r.status_code, 400)
        token = self.token_for("apibuyer@example.com")  # email works too
        self.assertTrue(token.startswith("lm_"))
        stored = ApiToken.objects.get(user=self.buyer)
        self.assertNotEqual(stored.key_hash, token)  # only the hash is stored
        self.assertEqual(stored.key_hash, ApiToken.hash(token))
        me = self.get("/api/v1/me/", token)
        self.assertEqual(me.status_code, 200)
        self.assertEqual(me.json()["username"], "apibuyer")
        self.assertEqual(self.get("/api/v1/me/").status_code, 401)
        self.assertEqual(self.get("/api/v1/me/", "lm_not-a-real-token").status_code, 401)
        self.assertEqual(self.send("post", "/api/v1/auth/logout/", token=token).status_code, 204)
        self.assertEqual(self.get("/api/v1/me/", token).status_code, 401)

    def test_expired_and_inactive_tokens_refused(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models import ApiToken
        token = self.token_for("apibuyer")
        ApiToken.objects.update(expires_at=timezone.now() - timedelta(seconds=1))
        self.assertEqual(self.get("/api/v1/me/", token).status_code, 401)
        token = self.token_for("apibuyer")
        self.buyer.is_active = False
        self.buyer.save()
        self.assertEqual(self.get("/api/v1/me/", token).status_code, 401)

    def test_two_step_accounts_need_a_code(self):
        from . import twofactor
        from .models import Profile
        secret = twofactor.new_secret()
        Profile.objects.filter(user=self.buyer).update(totp_secret=secret, totp_enabled=True)
        r = self.send("post", "/api/v1/auth/token/", {"username": "apibuyer", "password": "pass12345"})
        self.assertEqual(r.status_code, 400)
        self.assertTrue(r.json()["code_required"])
        r = self.send("post", "/api/v1/auth/token/", {"username": "apibuyer", "password": "pass12345", "code": "000000"})
        self.assertEqual(r.status_code, 400)
        code = twofactor.code_at(secret, twofactor.current_step())
        self.token_for("apibuyer", code=code)
        # the same code can't be used twice
        r = self.send("post", "/api/v1/auth/token/", {"username": "apibuyer", "password": "pass12345", "code": code})
        self.assertEqual(r.status_code, 400)

    def test_sign_in_is_rate_limited(self):
        codes = [self.send("post", "/api/v1/auth/token/", {"username": "apibuyer", "password": "nope"}).status_code
                 for _ in range(12)]
        self.assertIn(429, codes)

    # --- customer -----------------------------------------------------------
    def test_orders_are_private(self):
        mine, _ = self.order_for(self.buyer, self.cheap)
        other = User.objects.create_user("other", "other@example.com", "pass12345")
        theirs, _ = self.order_for(other, self.cheap)
        token = self.token_for("apibuyer")
        r = self.get("/api/v1/orders/", token).json()
        self.assertEqual([o["id"] for o in r["results"]], [mine.id])
        detail = self.get(f"/api/v1/orders/{mine.id}/", token).json()
        self.assertEqual(detail["totals"]["subtotal"], "6.00")
        self.assertEqual(detail["items"][0]["quantity"], 2)
        self.assertEqual(self.get(f"/api/v1/orders/{theirs.id}/", token).status_code, 404)

    def test_wishlist(self):
        token = self.token_for("apibuyer")
        self.assertEqual(self.send("post", "/api/v1/wishlist/", {"product_id": self.live.id}, token).status_code, 201)
        self.assertEqual(self.send("post", "/api/v1/wishlist/", {"product_id": self.hidden.id}, token).status_code, 404)
        self.assertEqual([p["id"] for p in self.get("/api/v1/wishlist/", token).json()], [self.live.id])
        self.assertEqual(self.send("delete", f"/api/v1/wishlist/{self.live.id}/", token=token).status_code, 204)
        self.assertEqual(self.get("/api/v1/wishlist/", token).json(), [])

    # --- seller ---------------------------------------------------------------
    def test_seller_endpoints_need_a_seller(self):
        token = self.token_for("apibuyer")
        self.assertEqual(self.get("/api/v1/seller/products/", token).status_code, 403)

    def test_seller_products_create_and_review_rules(self):
        token = self.token_for("apiseller")
        r = self.get("/api/v1/seller/products/", token).json()
        self.assertEqual({p["name"] for p in r["results"]}, {"Rose Serum", "Secret Serum"})
        self.assertEqual(self.send("post", "/api/v1/seller/products/", {"name": "No price"}, token).status_code, 400)
        r = self.send("post", "/api/v1/seller/products/", {"name": "Lip Balm", "category": "skincare", "price": "4.50",
                                                           "stock": 7, "image_url": "https://example.com/balm.jpg"}, token)
        self.assertEqual(r.status_code, 201, r.content)
        self.assertEqual(r.json()["status"], "pending")  # new products go to review
        balm = Product.objects.get(name="Lip Balm")
        self.assertEqual(balm.seller_account, self.seller)
        self.assertEqual(balm.stock, 7)
        # price and stock are live at once, a new name goes back to review
        r = self.send("patch", f"/api/v1/seller/products/{self.live.id}/", {"price": "19.00", "stock": 3}, token)
        self.assertEqual(r.status_code, 200)
        self.live.refresh_from_db()
        self.assertEqual((self.live.price, self.live.stock, self.live.approval_status), (Decimal("19.00"), 3, "approved"))
        self.send("patch", f"/api/v1/seller/products/{self.live.id}/", {"name": "Rose Serum XL"}, token)
        self.live.refresh_from_db()
        self.assertEqual(self.live.approval_status, "pending")

    def test_sellers_only_touch_their_own_products(self):
        other = SellerAccount.objects.create(user=User.objects.create_user("o2", "o2@example.com", "pass12345"),
                                             status="approved", business_name="Other")
        theirs = make_product(name="Not yours", seller_account=other, seller_name="Other")
        token = self.token_for("apiseller")
        self.assertEqual(self.get(f"/api/v1/seller/products/{theirs.id}/", token).status_code, 404)
        self.assertEqual(self.send("patch", f"/api/v1/seller/products/{theirs.id}/", {"price": "1.00"}, token).status_code, 404)
        theirs.refresh_from_db()
        self.assertEqual(theirs.price, Decimal("500.00"))

    def test_unapproved_seller_cannot_change_things(self):
        self.seller.status = "pending"
        self.seller.save()
        token = self.token_for("apiseller")
        r = self.send("patch", f"/api/v1/seller/products/{self.live.id}/", {"price": "1.00"}, token)
        self.assertEqual(r.status_code, 403)

    def test_seller_orders_and_fulfilment(self):
        order, item = self.order_for(self.buyer, self.live, seller=self.seller)
        token = self.token_for("apiseller")
        r = self.get("/api/v1/seller/orders/", token, status="to_pack").json()
        self.assertEqual([o["id"] for o in r["results"]], [order.id])
        self.assertEqual(r["results"][0]["your_sale"], "50.00")
        r = self.send("post", f"/api/v1/seller/orders/{order.id}/fulfil/",
                      {"status": "handed_to_courier", "courier_name": "DHL", "tracking_number": "JD0001"}, token)
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(r.json()["customer_notified"])
        item.refresh_from_db()
        order.refresh_from_db()
        self.assertEqual(item.fulfillment_status, "handed_to_courier")
        self.assertEqual((order.courier_name, order.tracking_number, order.status), ("DHL", "JD0001", "shipped"))

    def test_fulfilment_blocked_for_unconfirmed_cash_orders(self):
        order, item = self.order_for(self.buyer, self.live, seller=self.seller, cod_confirmed=False)
        token = self.token_for("apiseller")
        r = self.send("post", f"/api/v1/seller/orders/{order.id}/fulfil/", {"status": "handed_to_courier"}, token)
        self.assertEqual(r.status_code, 400)
        item.refresh_from_db()
        self.assertEqual(item.fulfillment_status, "pending")

    def test_team_staff_see_no_money(self):
        staff = User.objects.create_user("teamstaff", "teamstaff@example.com", "pass12345")
        OrganizationMember.objects.create(organization=self.seller, user=staff, role="staff")
        token = self.token_for("teamstaff")
        data = self.get("/api/v1/seller/", token).json()
        self.assertEqual(data["role"], "staff")
        self.assertNotIn("money", data)
        owner = self.get("/api/v1/seller/", self.token_for("apiseller")).json()
        self.assertIn("balance_available", owner["money"])


class ApiTokenAccountPageTests(TestCase):
    def test_tokens_listed_revoked_and_cleared_on_password_change(self):
        from .models import ApiToken
        user = User.objects.create_user("tok", "tok@example.com", "pass12345")
        t1, _k = ApiToken.issue(user, name="Phone app")
        ApiToken.issue(user, name="Script")
        self.client.force_login(user)
        page = self.client.get(reverse("account_security"))
        self.assertContains(page, "Phone app")
        self.assertContains(page, t1.prefix)
        self.client.post(reverse("account_security"), {"action": "revoke_token", "token": t1.pk})
        self.assertEqual(list(ApiToken.objects.filter(user=user).values_list("name", flat=True)), ["Script"])
        self.client.post(reverse("account_security"), {"old_password": "pass12345", "new_password1": "N3w-strong-pass!",
                                                       "new_password2": "N3w-strong-pass!"})
        self.assertFalse(ApiToken.objects.filter(user=user).exists())


class DockerSupportTests(TestCase):
    def test_env_file_location_and_gunicorn_reload(self):
        import os
        import tempfile
        from unittest import mock
        from django.test import override_settings
        from . import envfile
        with tempfile.TemporaryDirectory() as tmp:
            target = os.path.join(tmp, ".env")
            with override_settings(ENV_FILE=target), mock.patch.dict(os.environ, {}, clear=False):
                envfile.write({"STORE_CURRENCY": "eur"})
                self.assertEqual(envfile.read()["STORE_CURRENCY"], "eur")
                self.assertTrue(os.path.exists(target))
        with mock.patch.dict(os.environ, {"RELOAD_WITH_HUP": "1"}), mock.patch("os.kill") as kill:
            self.assertTrue(envfile.reload_app())
            kill.assert_called_once()

    def test_database_ssl_choice(self):
        import os
        from unittest import mock
        from backend import settings as st
        cases = [
            ("postgresql://u:p@db:5432/x", "", False),
            ("postgresql://u:p@abc.pooler.supabase.com:5432/postgres", "", True),
            ("postgresql://u:p@abc.example.com:5432/x?sslmode=disable", "", False),
            ("postgresql://u:p@db:5432/x", "True", True),
            ("postgresql://u:p@abc.example.com:5432/x", "False", False),
        ]
        for url, choice, expected in cases:
            with mock.patch.dict(os.environ, {"DATABASE_SSL": choice}), mock.patch.object(st, "DATABASE_URL", url):
                self.assertEqual(st._database_ssl(), expected, (url, choice))


# ---------------------------------------------------------------------------
# Competitor-parity features: categories, pages, specs, digital, quotes...
# ---------------------------------------------------------------------------

from .models import Category, HomeBanner, Page, QuoteRequest, StoreFollow  # noqa: E402


class CategoryAdminTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user("catboss", "catboss@example.com", "pass12345", is_staff=True, is_superuser=True)
        self.client.force_login(self.staff)

    def test_starting_categories_are_grouped(self):
        beauty = Category.objects.get(slug="beauty")
        self.assertEqual(Category.objects.get(slug="skincare").parent, beauty)
        self.assertIn(("beauty", "Beauty"), __import__("bees.categories", fromlist=["x"]).top_level())

    def test_add_subcategory_and_parent_page_lists_its_products(self):
        r = self.client.post(reverse("manage_categories"), {"action": "add", "name": "Yoga mats", "parent": Category.objects.get(slug="beauty").pk, "is_active": "1"})
        self.assertRedirects(r, reverse("manage_categories"))
        self.assertTrue(Category.objects.filter(slug="yoga-mats", parent__slug="beauty").exists())
        make_product(name="Cork Mat", category="yoga-mats")
        page = self.client.get(reverse("category_products", args=["beauty"]))
        self.assertContains(page, "Cork Mat")
        self.assertContains(page, "Yoga mats")  # subcategory chip
        child = self.client.get(reverse("category_products", args=["yoga-mats"]))
        self.assertContains(child, "Cork Mat")
        self.assertContains(child, 'aria-label="Breadcrumb"')

    def test_no_nesting_deeper_than_one_level(self):
        skincare = Category.objects.get(slug="skincare")
        self.client.post(reverse("manage_categories"), {"action": "add", "name": "Deep", "parent": skincare.pk})
        self.assertFalse(Category.objects.filter(name="Deep").exists())

    def test_delete_moves_products(self):
        cat = Category.objects.create(slug="old-cat", name="Old")
        p = make_product(category="old-cat")
        self.client.post(reverse("manage_categories"), {"action": "delete", "id": cat.pk})
        self.assertTrue(Category.objects.filter(pk=cat.pk).exists())  # products need a new home first
        self.client.post(reverse("manage_categories"), {"action": "delete", "id": cat.pk, "move_to": "grocery"})
        self.assertFalse(Category.objects.filter(pk=cat.pk).exists())
        p.refresh_from_db()
        self.assertEqual(p.category, "grocery")

    def test_hidden_category_is_404_and_not_in_menu(self):
        cat = Category.objects.get(slug="charity")
        cat.is_active = False
        cat.save()
        self.assertEqual(self.client.get(reverse("category_products", args=["charity"])).status_code, 404)
        self.assertEqual(self.client.get(reverse("category_products", args=["donate-education"])).status_code, 404)
        self.assertNotContains(self.client.get("/"), "/category/charity/")

    def test_catalog_staff_can_open_categories_marketing_cannot(self):
        from .models import Profile
        u = User.objects.create_user("mk", "mk@example.com", "pass12345", is_staff=True)
        Profile.objects.update_or_create(user=u, defaults={"staff_role": "marketing"})
        self.client.force_login(u)
        self.assertRedirects(self.client.get(reverse("manage_categories")), reverse("manage_dashboard"), fetch_redirect_response=False)
        self.assertEqual(self.client.get(reverse("manage_pages")).status_code, 200)


class PagesAndBannersTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user("pgboss", "pgboss@example.com", "pass12345", is_staff=True, is_superuser=True)

    def test_store_text_replaces_built_in_page(self):
        self.assertContains(self.client.get(reverse("terms_page")), "Using ")
        self.client.force_login(self.staff)
        form = self.client.get(reverse("manage_page_builtin", args=["terms"]))
        self.assertContains(form, "## Using")  # starts from the built-in text
        self.client.post(reverse("manage_page_builtin", args=["terms"]), {"title": "Our terms", "body": "## Rules\nBe nice.\n- one\n- two\n<script>x</script>", "is_published": "1"})
        page = self.client.get(reverse("terms_page"))
        self.assertContains(page, "<h2>Rules</h2>", html=True)
        self.assertContains(page, "<li>two</li>", html=True)
        self.assertNotContains(page, "<script>x</script>")
        self.assertNotContains(page, "Using ")

    def test_custom_page_in_footer_and_drafts_hidden(self):
        self.client.force_login(self.staff)
        self.client.post(reverse("manage_page_new"), {"title": "Shipping policy", "body": "We ship fast.", "is_published": "1", "show_in_footer": "1"})
        self.client.post(reverse("manage_page_new"), {"title": "Secret", "body": "Draft text."})
        self.assertContains(self.client.get("/pages/shipping-policy/"), "We ship fast.")
        self.assertEqual(self.client.get("/pages/secret/").status_code, 404)
        self.assertContains(self.client.get("/"), "/pages/shipping-policy/")
        self.assertEqual(self.client.get("/pages/terms/").status_code, 404)  # built-ins live at their own address

    def test_home_banner_respects_schedule(self):
        from datetime import timedelta
        from django.utils import timezone
        HomeBanner.objects.create(title="Summer sale", link="/products/")
        HomeBanner.objects.create(title="Future deal", link="/products/", starts_at=timezone.now() + timedelta(days=2))
        home = self.client.get("/")
        self.assertContains(home, "Summer sale")
        self.assertNotContains(home, "Future deal")

    def test_banner_link_must_be_safe(self):
        self.client.force_login(self.staff)
        self.client.post(reverse("manage_pages"), {"action": "add_banner", "title": "X", "link": "javascript:alert(1)", "is_active": "1"})
        self.assertFalse(HomeBanner.objects.exists())


class ProductDetailsAndFiltersTests(TestCase):
    def test_brand_size_colour_filters(self):
        from .models import ProductVariant
        a = make_product(name="Alpha Tee", brand="Acme", category="fashion")
        b = make_product(name="Beta Tee", brand="Bolt", category="fashion", has_variants=True)
        ProductVariant.objects.create(product=b, size="M", color="Red", stock=3)
        ProductVariant.objects.create(product=b, size="L", color="Blue", stock=0)
        Review.objects.create(product=b, user=User.objects.create_user("rv", "rv@example.com", "x"), rating=5, comment="ok")
        page = self.client.get(reverse("category_products", args=["fashion"]), {"brand": "Acme"})
        self.assertContains(page, "Alpha Tee")
        self.assertNotContains(page, "Beta Tee")
        page = self.client.get(reverse("category_products", args=["fashion"]), {"size": "M"})
        self.assertContains(page, "Beta Tee")
        self.assertNotContains(page, "Alpha Tee")
        self.assertEqual(page.context["products"][0].rating_count, 1)  # no double counting
        # Sold-out options don't match.
        self.assertNotContains(self.client.get(reverse("category_products", args=["fashion"]), {"color": "Blue"}), "Beta Tee")
        self.assertContains(self.client.get(reverse("category_products", args=["fashion"])), 'name="brand" value="Bolt"')
        _ = a

    def test_product_page_shows_specs_safety_video_and_size_guide(self):
        p = make_product(name="Wool Hoodie", category="hoodies", brand="Nordic", specs="Material: Merino wool\nFit: Regular",
                         manufacturer="Nordic Knit AB, Storgatan 1, Stockholm", video_url="https://youtu.be/dQw4w9WgXcQ")
        page = self.client.get(reverse("product_detail", args=[p.id]))
        self.assertContains(page, "Merino wool")
        self.assertContains(page, "Nordic Knit AB")
        self.assertContains(page, "youtube-nocookie.com/embed/dQw4w9WgXcQ")
        self.assertContains(page, "Size guide")  # from the Fashion > Hoodies category
        self.assertIn("player.vimeo.com", page["Content-Security-Policy"])

    def test_seller_form_validates_new_fields(self):
        from .views import _extra_product_fields
        from django.core.exceptions import ValidationError
        from django.test import RequestFactory
        rf = RequestFactory()
        with self.assertRaises(ValidationError):
            _extra_product_fields(rf.post("/", {"video_url": "https://evil.example/x.html"}))
        with self.assertRaises(ValidationError):
            _extra_product_fields(rf.post("/", {"min_order_qty": "5", "max_order_qty": "2"}))
        with self.assertRaises(ValidationError):
            _extra_product_fields(rf.post("/", {"specs": "no colon here"}))
        ok = _extra_product_fields(rf.post("/", {"brand": "  Acme  ", "specs": "A: b", "min_order_qty": "2"}))
        self.assertEqual((ok["brand"], ok["min_order_qty"]), ("Acme", 2))


class OrderLimitTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("lim", "lim@example.com", "pass12345")
        self.client.force_login(self.user)

    def test_minimum_added_and_maximum_enforced(self):
        p = make_product(price=Decimal("5.00"), stock=50, min_order_qty=3, max_order_qty=4)
        self.client.post(reverse("add_to_cart", args=[p.id]), {"quantity": 1})
        self.assertEqual(self.client.session["cart"][str(p.id)], 3)
        self.client.post(reverse("add_to_cart", args=[p.id]), {"quantity": 5})
        self.assertEqual(self.client.session["cart"][str(p.id)], 4)
        self.client.post(reverse("update_cart_item", args=[str(p.id)]), {"action": "increase"})
        self.assertEqual(self.client.session["cart"][str(p.id)], 4)
        Product.objects.filter(pk=p.pk).update(max_order_qty=2)
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertFalse(Order.objects.exists())


class DigitalProductTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("dig", "dig@example.com", "pass12345")
        self.client.force_login(self.user)
        self.p = make_product(name="Recipe e-book", price=Decimal("9.00"), stock=999, is_digital=True,
                              digital_file=SimpleUploadedFile("book.pdf", b"%PDF-1.4 test", content_type="application/pdf"))
        SiteSettings.objects.filter(pk=1).update(shipping_flat_fee=Decimal("5.00"))

    def tearDown(self):
        self.p.digital_file.delete(save=False)

    def test_no_cod_no_shipping_and_download_after_payment(self):
        from .views import _price_cart
        self.assertEqual(_price_cart(Decimal("9"), country="US", digital_only=True)["shipping"], Decimal("0"))
        self.client.post(reverse("add_to_cart", args=[self.p.id]))
        page = self.client.get(reverse("checkout"))
        self.assertNotIn("cod", page.context["payment_methods"])
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertFalse(Order.objects.exists())

        order = Order.objects.create(user=self.user, email="dig@example.com", full_name="D", address="1", city="X", country="US",
                                     phone="+15550100", payment_method="card", payment_status="pending", status="pending")
        item = OrderItem.objects.create(order=order, product=self.p, product_name=self.p.name, price=self.p.price, quantity=1)
        self.assertFalse(item.can_return)
        self.assertRedirects(self.client.get(reverse("digital_download", args=[item.id])), reverse("my_orders"), fetch_redirect_response=False)
        order.payment_status, order.status = "paid", "confirmed"
        order.save(update_fields=["payment_status", "status"])
        item.refresh_from_db()
        order.refresh_from_db()
        self.assertEqual(item.fulfillment_status, "delivered")
        self.assertEqual(order.status, "delivered")
        r = self.client.get(reverse("digital_download", args=[item.id]))
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r["Content-Disposition"])
        item.refresh_from_db()
        self.assertEqual(item.downloads, 1)
        # Someone else can't download it.
        other = User.objects.create_user("oth", "oth@example.com", "pass12345")
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse("digital_download", args=[item.id])).status_code, 404)
        # Limit
        OrderItem.objects.filter(pk=item.pk).update(downloads=5)
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("digital_download", args=[item.id])).status_code, 302)

    def test_digital_product_needs_a_file(self):
        from .views import check_digital
        from django.core.exceptions import ValidationError
        with self.assertRaises(ValidationError):
            check_digital({"is_digital": True})
        check_digital({"is_digital": True}, self.p)  # keeps the current file


class QuoteTests(TestCase):
    def setUp(self):
        self.buyer = User.objects.create_user("qb", "qb@example.com", "pass12345")
        self.seller_user = User.objects.create_user("qs", "qs@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.seller_user, status="approved", business_name="Bulk Co")
        self.p = make_product(name="Mug", price=Decimal("10.00"), stock=500, quote_only=True, seller_account=self.seller, seller_name="Bulk Co")

    def test_request_offer_accept_and_order_at_agreed_price(self):
        self.client.force_login(self.buyer)
        self.assertContains(self.client.get(reverse("product_detail", args=[self.p.id])), "Request a quote")
        self.client.post(reverse("request_quote", args=[self.p.id]), {"quantity": "100", "target_price": "7", "message": "Logo print?"})
        q = QuoteRequest.objects.get()
        self.assertEqual((q.status, q.quantity, q.seller_account), ("open", 100, self.seller))
        self.assertTrue(Notification.objects.filter(user=self.seller_user, message__contains="quote").exists())
        # Buyer can't answer their own quote as a seller.
        self.client.force_login(self.seller_user)
        self.client.post(reverse("seller_quotes"), {"action": "offer", "id": q.id, "offer_price": "8.00", "valid_days": "5"})
        q.refresh_from_db()
        self.assertEqual((q.status, q.offer_price), ("offered", Decimal("8.00")))
        # Another user can't accept it.
        self.client.force_login(User.objects.create_user("x9", "x9@example.com", "pass12345"))
        self.assertEqual(self.client.post(reverse("my_quotes"), {"action": "accept", "id": q.id}).status_code, 404)
        self.client.force_login(self.buyer)
        self.client.post(reverse("my_quotes"), {"action": "accept", "id": q.id})
        q.refresh_from_db()
        self.assertEqual(q.status, "accepted")
        cart = self.client.get(reverse("cart"))
        self.assertContains(cart, "Quote #")
        self.client.post(reverse("update_cart_item", args=[f"q{q.id}"]), {"action": "increase"})
        self.assertEqual(self.client.session["cart"][f"q{q.id}"], 100)
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "email": "qb@example.com", "payment_method": "cod"})
        order = Order.objects.get()
        item = order.items.get()
        self.assertEqual((item.price, item.quantity, item.quote_id), (Decimal("8.00"), 100, q.id))
        q.refresh_from_db()
        self.assertEqual(q.status, "ordered")

    def test_quote_not_available_unless_turned_on(self):
        self.client.force_login(self.buyer)
        Product.objects.filter(pk=self.p.pk).update(quote_only=False)
        self.assertEqual(self.client.get(reverse("request_quote", args=[self.p.id])).status_code, 404)

    def test_counter_offer_reopens(self):
        q = QuoteRequest.objects.create(product=self.p, buyer=self.buyer, seller_account=self.seller, quantity=50,
                                        offer_price=Decimal("9"), status="offered")
        self.client.force_login(self.buyer)
        self.client.post(reverse("my_quotes"), {"action": "counter", "id": q.id, "target_price": "8.5", "quantity": "60"})
        q.refresh_from_db()
        self.assertEqual((q.status, q.quantity, q.target_price), ("open", 60, Decimal("8.50")))


class FollowStoreTests(TestCase):
    def test_follow_and_daily_notification(self):
        from django.core.management import call_command
        seller = SellerAccount.objects.create(user=User.objects.create_user("fs", "fs@example.com", "x"), status="approved", business_name="Fern Shop")
        fan = User.objects.create_user("fan", "fan@example.com", "pass12345")
        self.client.force_login(fan)
        self.client.post(reverse("follow_store", args=[seller.id]), {"next": "/"})
        self.assertTrue(StoreFollow.objects.filter(user=fan, seller_account=seller).exists())
        make_product(name="Fern pot", seller_account=seller, seller_name="Fern Shop")
        call_command("notify_followers", stdout=open(os.devnull, "w"))
        self.assertTrue(Notification.objects.filter(user=fan, message__contains="Fern Shop added 1 new product").exists())
        self.client.post(reverse("follow_store", args=[seller.id]), {"next": "/"})
        self.assertFalse(StoreFollow.objects.exists())


class ApiNewFieldsTests(TestCase):
    def test_product_detail_and_categories(self):
        p = make_product(name="Spec Lamp", brand="Lumo", specs="Watts: 9", min_order_qty=2)
        data = self.client.get(f"/api/v1/products/{p.id}/").json()
        self.assertEqual(data["brand"], "Lumo")
        self.assertEqual(data["specifications"], [{"name": "Watts", "value": "9"}])
        self.assertEqual(data["order_limits"], {"min": 2, "max": None})
        cats = {c["code"]: c for c in self.client.get("/api/v1/categories/").json()}
        self.assertEqual(cats["skincare"]["parent"], "beauty")


# ---------------------------------------------------------------------------
# Big-marketplace features: Amazon / eBay / Etsy / Daraz parity
# ---------------------------------------------------------------------------

from .models import (Bid, CouponClaim, MembershipPayment, PickupPoint, PriceHistory, ReviewVote,  # noqa: E402
                     Subscription, Wishlist)


class ConditionOffersAndSellersTests(TestCase):
    def setUp(self):
        self.s1 = SellerAccount.objects.create(user=User.objects.create_user("s1", "s1@example.com", "x"), status="approved", business_name="One Shop")
        self.s2_user = User.objects.create_user("s2", "s2@example.com", "pass12345")
        self.s2 = SellerAccount.objects.create(user=self.s2_user, status="approved", business_name="Two Shop")
        self.p = make_product(name="Retro Camera", price=Decimal("100"), seller_account=self.s1, seller_name="One Shop")

    def test_used_condition_shows_and_other_seller_offer_is_grouped(self):
        used = make_product(name="Retro Camera", price=Decimal("70"), seller_account=self.s2, seller_name="Two Shop",
                            condition="used_good", offer_of=self.p)
        page = self.client.get(reverse("product_detail", args=[self.p.id]))
        self.assertContains(page, "Other sellers on")
        self.assertContains(page, "Two Shop")
        self.assertContains(page, "Used - good")
        # The duplicate listing doesn't appear twice in search.
        found = self.client.get(reverse("search_products"), {"q": "retro camera"}).context["products"]
        self.assertEqual([x.id for x in found], [self.p.id])
        _ = used

    def test_sell_one_like_this_prefills_form(self):
        self.client.force_login(self.s2_user)
        form = self.client.get(reverse("seller_add_product"), {"offer_of": self.p.id})
        self.assertContains(form, 'name="offer_of"')
        self.assertContains(form, "Retro Camera")

    def test_make_offer_and_accept_buyer_price(self):
        Product.objects.filter(pk=self.p.pk).update(accepts_offers=True, offer_floor=Decimal("60"))
        buyer = User.objects.create_user("ob", "ob@example.com", "pass12345")
        self.client.force_login(buyer)
        self.client.post(reverse("make_offer", args=[self.p.id]), {"price": "50", "quantity": "1"})
        self.assertEqual(QuoteRequest.objects.get().status, "declined")  # below the floor
        QuoteRequest.objects.all().delete()
        self.client.post(reverse("make_offer", args=[self.p.id]), {"price": "85", "quantity": "1"})
        q = QuoteRequest.objects.get()
        self.assertEqual((q.kind, q.status, q.target_price), ("offer", "open", Decimal("85.00")))
        self.client.force_login(self.s1.user)
        self.client.post(reverse("seller_quotes"), {"action": "accept_target", "id": q.id})
        q.refresh_from_db()
        self.assertEqual((q.status, q.offer_price), ("offered", Decimal("85.00")))


class AuctionTests(TestCase):
    def setUp(self):
        from datetime import timedelta
        from django.utils import timezone
        seller = SellerAccount.objects.create(user=User.objects.create_user("au", "au@example.com", "x"), status="approved", business_name="Bid Co")
        self.p = make_product(name="Vintage Watch", price=Decimal("1"), stock=1, seller_account=seller, seller_name="Bid Co", is_auction=True,
                              auction_start_price=Decimal("10"), auction_reserve=Decimal("15"),
                              auction_ends_at=timezone.now() + timedelta(hours=2))
        self.a = User.objects.create_user("ba", "ba@example.com", "pass12345")
        self.b = User.objects.create_user("bb", "bb@example.com", "pass12345")

    def test_bidding_outbid_and_win(self):
        from datetime import timedelta
        from django.utils import timezone
        from .market import close_auctions
        self.client.force_login(self.a)
        self.client.post(reverse("place_bid", args=[self.p.id]), {"amount": "9"})
        self.assertFalse(Bid.objects.exists())  # below the start
        self.client.post(reverse("place_bid", args=[self.p.id]), {"amount": "12"})
        self.client.force_login(self.b)
        self.client.post(reverse("place_bid", args=[self.p.id]), {"amount": "12.10"})
        self.assertEqual(Bid.objects.count(), 1)  # needs the increment
        self.client.post(reverse("place_bid", args=[self.p.id]), {"amount": "16"})
        self.assertTrue(Notification.objects.filter(user=self.a, message__contains="outbid").exists())
        # Can't add an auction to the cart.
        self.client.post(reverse("add_to_cart", args=[self.p.id]))
        self.assertEqual(self.client.session.get("cart", {}), {})
        Product.objects.filter(pk=self.p.pk).update(auction_ends_at=timezone.now() - timedelta(minutes=1))
        self.assertEqual(close_auctions(), 1)
        win = QuoteRequest.objects.get(kind="auction")
        self.assertEqual((win.buyer, win.offer_price, win.status), (self.b, Decimal("16.00"), "accepted"))
        self.client.post(reverse("my_quotes"), {"action": "readd", "id": win.id})
        self.assertIn(f"q{win.id}", self.client.session["cart"])

    def test_reserve_not_met_no_sale_and_anti_sniping(self):
        from datetime import timedelta
        from django.utils import timezone
        from .market import close_auctions
        Product.objects.filter(pk=self.p.pk).update(auction_ends_at=timezone.now() + timedelta(seconds=30))
        self.client.force_login(self.a)
        self.client.post(reverse("place_bid", args=[self.p.id]), {"amount": "11"})
        self.p.refresh_from_db()
        self.assertGreater(self.p.auction_ends_at, timezone.now() + timedelta(seconds=90))  # extended
        Product.objects.filter(pk=self.p.pk).update(auction_ends_at=timezone.now() - timedelta(minutes=1))
        close_auctions()
        self.assertFalse(QuoteRequest.objects.exists())  # 11 < reserve 15


class PersonalizationGiftPickupTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("pg", "pg@example.com", "pass12345")
        self.client.force_login(self.user)

    def test_personalization_required_and_saved_on_order(self):
        p = make_product(name="Engraved Pen", price=Decimal("20"), personalization_prompt="Name to engrave", personalization_required=True,
                         personalization_max=10)
        self.client.post(reverse("add_to_cart", args=[p.id]))
        self.assertEqual(self.client.session.get("cart", {}), {})
        self.client.post(reverse("add_to_cart", args=[p.id]), {"personalization": "Samad Mangrio Long"})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod", "gift_wrap": "1", "gift_message": "Happy day"})
        order = Order.objects.get()
        self.assertEqual(order.items.get().personalization, "Samad Mang")
        self.assertTrue(order.gift_wrap)
        self.assertEqual(order.gift_message, "Happy day")

    def test_pickup_point_sets_fee_and_address(self):
        SiteSettings.objects.filter(pk=1).update(shipping_flat_fee=Decimal("7"))
        point = PickupPoint.objects.create(name="Locker 1", address="Main Sq 1", city="Lisbon", country="PT", fee=Decimal("2"))
        p = make_product(price=Decimal("30"))
        self.client.post(reverse("add_to_cart", args=[p.id]))
        page = self.client.get(reverse("checkout"))
        self.assertContains(page, "Locker 1")
        form = {**CHECKOUT_FORM, "address": "", "city": "", "postal_code": "", "country": "",
                "payment_method": "cod", "delivery": "pickup", "pickup_point": point.id}
        self.client.post(reverse("checkout"), form)
        order = Order.objects.get()
        self.assertEqual((order.pickup_point, order.shipping_amount, order.country), (point, Decimal("2"), "PT"))


class VoucherAndMembershipTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("vm", "vm@example.com", "pass12345")
        self.client.force_login(self.user)
        SiteSettings.objects.filter(pk=1).update(shipping_flat_fee=Decimal("5"))

    def test_collect_public_free_shipping_voucher(self):
        c = Coupon.objects.create(code="SHIPFREE", percent_off=0, free_shipping=True, is_public=True, title="Free delivery")
        Coupon.objects.create(code="HIDDEN", percent_off=50)
        page = self.client.get(reverse("vouchers"))
        self.assertContains(page, "Free delivery")
        self.assertNotContains(page, "HIDDEN")
        self.client.post(reverse("collect_voucher", args=[c.id]))
        self.assertTrue(CouponClaim.objects.filter(user=self.user, coupon=c).exists())
        p = make_product(price=Decimal("20"))
        self.client.post(reverse("add_to_cart", args=[p.id]))
        self.assertContains(self.client.get(reverse("checkout")), "SHIPFREE")
        self.client.post(reverse("apply_coupon"), {"coupon_code": "SHIPFREE"})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        self.assertEqual(Order.objects.get().shipping_amount, Decimal("0"))

    def test_member_gets_free_shipping_and_discount(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models import Profile
        SiteSettings.objects.filter(pk=1).update(membership_enabled=True, membership_discount=10)
        Profile.objects.update_or_create(user=self.user, defaults={"member_until": timezone.now() + timedelta(days=5)})
        p = make_product(price=Decimal("50"))
        self.client.post(reverse("add_to_cart", args=[p.id]))
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get()
        self.assertEqual((order.shipping_amount, order.discount_amount), (Decimal("0"), Decimal("5.00")))
        self.assertEqual(self.client.get(reverse("membership")).status_code, 200)

    def test_membership_payment_confirmed(self):
        from unittest import mock
        SiteSettings.objects.filter(pk=1).update(membership_enabled=True)
        pay = MembershipPayment.objects.create(user=self.user, amount=Decimal("4.99"), days=30, stripe_session_id="cs_1")
        with mock.patch("bees.payments._stripe") as st:
            st.return_value.checkout.Session.retrieve.return_value = {"id": "cs_1", "payment_status": "paid", "amount_total": 499}
            self.client.get(reverse("membership_done", args=[pay.id]), {"session_id": "cs_1"})
            self.client.get(reverse("membership_done", args=[pay.id]), {"session_id": "cs_1"})  # twice: no double days
        self.user.profile.refresh_from_db()
        from .market import is_member
        self.assertTrue(is_member(self.user))
        from django.utils import timezone
        self.assertLess((self.user.profile.member_until - timezone.now()).days, 31)


class CoinsSubscriptionsBundlesTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("cs", "cs@example.com", "pass12345")
        self.client.force_login(self.user)

    def test_daily_check_in_once_a_day_and_streak(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models import Profile
        self.client.post(reverse("check_in"))
        self.client.post(reverse("check_in"))
        profile = Profile.objects.get(user=self.user)
        self.assertEqual((profile.loyalty_points, profile.checkin_streak), (5, 1))
        Profile.objects.filter(pk=profile.pk).update(last_checkin=timezone.localdate() - timedelta(days=1), checkin_streak=6)
        self.client.post(reverse("check_in"))
        profile.refresh_from_db()
        self.assertEqual((profile.loyalty_points, profile.checkin_streak), (15, 7))  # day 7 = double

    def test_subscribe_and_save_order_and_reminder(self):
        from django.core.management import call_command
        from django.utils import timezone
        p = make_product(name="Coffee Beans", price=Decimal("20"), subscribe_percent=10)
        self.client.post(reverse("subscribe", args=[p.id]), {"every_weeks": "4", "quantity": "2"})
        sub = Subscription.objects.get()
        self.assertIn(f"s{sub.id}", self.client.session["cart"])
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        item = Order.objects.get().items.get()
        self.assertEqual((item.price, item.quantity), (Decimal("18.00"), 2))
        Subscription.objects.filter(pk=sub.pk).update(next_date=timezone.localdate())
        call_command("subscription_reminders", stdout=open(os.devnull, "w"))
        self.assertTrue(Notification.objects.filter(user=self.user, message__contains="refill").exists())
        sub.refresh_from_db()
        self.assertGreater(sub.next_date, timezone.localdate())

    def test_bought_together_add_many_and_save_for_later(self):
        a, b = make_product(name="Tea Pot", price=Decimal("10")), make_product(name="Tea Cups", price=Decimal("6"))
        self.client.post(reverse("add_many"), {"ids": [a.id, b.id]})
        self.assertEqual(set(self.client.session["cart"]), {str(a.id), str(b.id)})
        self.client.post(reverse("save_for_later", args=[str(b.id)]))
        self.assertNotIn(str(b.id), self.client.session["cart"])
        self.assertTrue(Wishlist.objects.filter(user=self.user, product=b, list_name="Saved for later").exists())

    def test_wishlist_share_and_move_to_cart(self):
        p = make_product(name="Lamp")
        Wishlist.objects.create(user=self.user, product=p)
        self.client.post(reverse("wishlist_manage"), {"action": "share"})
        token = self.user.profile.__class__.objects.get(user=self.user).wishlist_token
        self.client.logout()
        self.assertContains(self.client.get(reverse("shared_wishlist", args=[token])), "Lamp")
        self.assertEqual(self.client.get(reverse("shared_wishlist", args=["wrongtoken123"])).status_code, 404)


class ReviewsPriceAndTraderTests(TestCase):
    def test_review_photos_helpful_votes_and_filter(self):
        from io import BytesIO
        from PIL import Image
        p = make_product(name="Desk Lamp")
        u = User.objects.create_user("rp", "rp@example.com", "pass12345")
        order = Order.objects.create(user=u, email="rp@example.com", full_name="R", address="1", city="X", country="US", phone="+15550100",
                                     payment_method="cod", status="delivered")
        OrderItem.objects.create(order=order, product=p, product_name=p.name, price=p.price, quantity=1, fulfillment_status="delivered")
        self.client.force_login(u)
        buf = BytesIO()
        Image.new("RGB", (40, 40), "red").save(buf, "JPEG")
        photo = SimpleUploadedFile("p.jpg", buf.getvalue(), content_type="image/jpeg")
        self.client.post(reverse("product_detail", args=[p.id]), {"rating": "4", "comment": "Bright and solid", "photos": [photo]})
        review = Review.objects.get()
        self.assertEqual(review.photos.count(), 1)
        if review.photos.exists():
            self.assertContains(self.client.get(reverse("product_detail", args=[p.id]), {"with_photos": "1"}), "Bright and solid")
        voter = User.objects.create_user("vt", "vt@example.com", "pass12345")
        self.client.force_login(voter)
        self.client.post(reverse("review_helpful", args=[review.id]))
        review.refresh_from_db()
        self.assertEqual(review.helpful_count, 1)
        self.client.post(reverse("review_helpful", args=[review.id]))  # toggles off
        review.refresh_from_db()
        self.assertEqual(review.helpful_count, 0)
        self.assertNotContains(self.client.get(reverse("product_detail", args=[p.id]), {"stars": "5"}), "Bright and solid")
        for f in review.photos.all():
            from django.core.files.storage import default_storage
            name = f.image_url.split("/media/", 1)[-1]
            if default_storage.exists(name):
                default_storage.delete(name)

    def test_lowest_price_in_30_days_shown_with_discount(self):
        p = make_product(name="Sale Kettle", price=Decimal("40"))
        p.price = Decimal("30")
        p.save()
        p.price, p.old_price, p.discount_percent = Decimal("35"), Decimal("50"), 30
        p.save()
        self.assertEqual(PriceHistory.objects.filter(product=p).count(), 3)
        self.assertContains(self.client.get(reverse("product_detail", args=[p.id])), "Lowest price in the last 30 days")

    def test_trader_information_shown(self):
        seller = SellerAccount.objects.create(user=User.objects.create_user("tr", "tr@example.com", "x"), status="approved",
                                              business_name="Trader Ltd", account_type="organization", company_registration="HRB 12345",
                                              trader_email="legal@trader.example", business_address="Main 1", city="Berlin", country="Germany")
        p = make_product(seller_account=seller, seller_name="Trader Ltd")
        page = self.client.get(reverse("product_detail", args=[p.id]))
        self.assertContains(page, "Seller information")
        self.assertContains(page, "HRB 12345")

    def test_bnpl_message(self):
        SiteSettings.objects.filter(pk=1).update(bnpl_enabled=True, bnpl_instalments=4)
        p = make_product(price=Decimal("100"))
        self.assertContains(self.client.get(reverse("product_detail", args=[p.id])), "4 interest-free payments of")

    def test_photo_search(self):
        from io import BytesIO
        from unittest import mock
        from PIL import Image
        make_product(name="Red Running Shoes", category="sneakers")
        buf = BytesIO()
        Image.new("RGB", (40, 40), "red").save(buf, "JPEG")
        with self.settings(AI_API_KEY="test"), mock.patch("bees.market.describe_image", return_value="red running shoes"):
            r = self.client.post(reverse("photo_search"), {"photo": SimpleUploadedFile("s.jpg", buf.getvalue(), content_type="image/jpeg")})
        self.assertIn("q=red+running+shoes", r["Location"])


class StoreBestTests(TestCase):
    """Round 3: signals, alerts, tiers, affiliates, deals and the assistant."""

    def _sell(self, product, qty, user=None):
        user = user or User.objects.create_user(f"b{Order.objects.count()}x", f"b{Order.objects.count()}@example.com", "x")
        order = Order.objects.create(user=user, email=user.email, full_name="B", address="1", city="X", country="US", phone="+15550100",
                                     payment_method="cod", status="confirmed")
        OrderItem.objects.create(order=order, product=product, product_name=product.name, price=product.price, quantity=qty)
        return order

    def test_bought_count_best_seller_and_sort(self):
        a = make_product(name="Hot Hoodie", category="hoodies")
        b = make_product(name="Cool Hoodie", category="hoodies")
        self._sell(a, 60)
        self._sell(b, 4)
        page = self.client.get(reverse("product_detail", args=[a.id]))
        self.assertContains(page, "50+ bought in the past month")
        self.assertContains(page, "#1 Best Seller")
        listing = self.client.get(reverse("category_products", args=["hoodies"]), {"sort": "popular"})
        self.assertEqual(listing.context["products"][0].id, a.id)
        self.assertContains(page, "Arrives")

    def test_price_alert_notifies_once(self):
        u = User.objects.create_user("pa", "pa@example.com", "pass12345")
        p = make_product(name="Fancy Kettle", price=Decimal("80"))
        self.client.force_login(u)
        self.client.post(reverse("price_alert", args=[p.id]), {"target_price": "90"})  # not below the price
        from .models import PriceAlert
        self.assertFalse(PriceAlert.objects.exists())
        self.client.post(reverse("price_alert", args=[p.id]), {"target_price": "60"})
        p.price = Decimal("70")
        p.save()
        self.assertFalse(Notification.objects.filter(user=u, message__contains="Price alert").exists())
        p.price = Decimal("59")
        p.save()
        p.price = Decimal("58")
        p.save()
        self.assertEqual(Notification.objects.filter(user=u, message__contains="Price alert").count(), 1)

    def test_loyalty_tier_multiplies_points(self):
        from .models import Profile
        from . import loyalty
        u = User.objects.create_user("lt", "lt@example.com", "pass12345")
        p = make_product(price=Decimal("300"))
        old = self._sell(p, 2, u)
        Order.objects.filter(pk=old.pk).update(status="delivered")
        self.assertEqual(loyalty.tier_for(u)["key"], "gold")  # 600 spent
        new = self._sell(make_product(price=Decimal("10")), 1, u)
        new.status = "delivered"
        new.save()
        self.assertEqual(Profile.objects.get(user=u).loyalty_points, 15)  # 10 x 1.5
        self.client.force_login(u)
        self.assertContains(self.client.get(reverse("profile")), "Gold level")

    def test_affiliate_earns_on_delivery(self):
        from django.db import transaction
        from .models import Profile
        SiteSettings.objects.filter(pk=1).update(affiliate_percent=5)
        aff = User.objects.create_user("af", "af@example.com", "pass12345")
        Profile.objects.update_or_create(user=aff, defaults={"referral_code": "AFFCODE1"})
        buyer = User.objects.create_user("abuy", "abuy@example.com", "pass12345")
        p = make_product(price=Decimal("40"))
        self.client.force_login(buyer)
        self.client.get(reverse("product_detail", args=[p.id]), {"ref": "AFFCODE1"})
        self.client.post(reverse("add_to_cart", args=[p.id]))
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "email": "abuy@example.com", "payment_method": "cod"})
        order = Order.objects.get(user=buyer)
        self.assertEqual(order.affiliate_code, "AFFCODE1")
        with self.captureOnCommitCallbacks(execute=True):
            order.status = "delivered"
            order.save()
        self.assertEqual(Profile.objects.get(user=aff).store_credit, Decimal("2.00"))
        from . import loyalty
        self.assertIsNone(loyalty.pay_affiliate(Order.objects.get(pk=order.pk)))  # only once
        _ = transaction

    def test_deals_assistant_and_protection_pages(self):
        make_product(name="Big Discount Mug", discount_percent=40, old_price=Decimal("900"))
        self.assertContains(self.client.get(reverse("deals")), "Big Discount Mug")
        r = self.client.post(reverse("shop_assistant"), {"q": "discount mug"})
        self.assertContains(r, "Big Discount Mug")
        self.assertContains(self.client.get(reverse("buyer_protection")), "money back")

    def test_recommended_for_you_from_viewed_category(self):
        seen = make_product(name="Seen Sneaker", category="sneakers")
        other = make_product(name="Other Sneaker", category="sneakers")
        self.client.get(reverse("product_detail", args=[seen.id]))
        home = self.client.get("/")
        self.assertIn(other.id, [p.id for p in home.context["recommended"]])
        self.assertNotIn(seen.id, [p.id for p in home.context["recommended"]])


class EasyModeTests(TestCase):
    """Time savers: imports from other shops, trusted sellers, digests,
    admin search, WhatsApp, photo listings and price tips."""

    def test_shopify_and_woocommerce_exports_are_understood(self):
        from .bulk import _to_dicts, _clean
        shopify = [["Handle", "Title", "Body (HTML)", "Vendor", "Variant Price", "Variant Compare At Price", "Variant Inventory Qty", "Image Src", "Type"],
                   ["tee", "Cotton Tee", "<p>Soft <b>cotton</b></p>", "Acme", "19.00", "25.00", "7", "https://cdn.example.com/t.jpg", "Hoodie"],
                   ["tee", "", "", "", "", "", "", "https://cdn.example.com/t2.jpg", ""]]
        rows = _to_dicts(shopify)
        self.assertEqual(len(rows), 1)  # the extra image row is skipped
        clean = _clean(rows[0])
        self.assertEqual((clean["name"], clean["price"], clean["old_price"], clean["stock"], clean["brand"]),
                         ("Cotton Tee", Decimal("19.00"), Decimal("25.00"), 7, "Acme"))
        self.assertEqual(clean["description"], "Soft cotton")
        self.assertEqual(clean["category"], "hoodies")
        woo = [["Name", "Regular price", "Sale price", "Stock", "Images", "Categories", "Weight (kg)"],
               ["Desk Lamp", "40", "30", "3", "https://x.example.com/a.jpg, https://x.example.com/b.jpg", "Home & Kitchen > Table lamp", "1.2"]]
        clean = _clean(_to_dicts(woo)[0])
        self.assertEqual((clean["price"], clean["old_price"], clean["image_url"], clean["category"], clean["weight_grams"]),
                         (Decimal("30"), Decimal("40"), "https://x.example.com/a.jpg", "table-lamp", 1200))

    def test_trusted_seller_products_go_live(self):
        from django.utils import timezone
        user = User.objects.create_user("ts", "ts@example.com", "pass12345")
        seller = SellerAccount.objects.create(user=user, status="approved", business_name="Trusty", verified_at=timezone.now(),
                                              performance_level="excellent")
        self.client.force_login(user)
        form = {"name": "Trusted Mug", "category": "casserole-pot", "price": "10", "stock": "5", "discount_percent": "0",
                "image_url": "https://example.com/m.jpg"}
        self.client.post(reverse("seller_add_product"), form)
        self.assertEqual(Product.objects.get(name="Trusted Mug").approval_status, "pending")  # setting is off
        st = SiteSettings.load(); st.auto_approve_trusted = True; st.save()
        self.client.post(reverse("seller_add_product"), dict(form, name="Trusted Mug 2"))
        self.assertEqual(Product.objects.get(name="Trusted Mug 2").approval_status, "approved")
        _ = seller

    def test_daily_digests(self):
        from django.core import mail
        from django.core.management import call_command
        User.objects.create_user("own", "owner@example.com", "x", is_staff=True, is_superuser=True)
        seller = SellerAccount.objects.create(user=User.objects.create_user("ds", "ds@example.com", "x"), status="approved", business_name="Digest Shop")
        make_product(name="Last One", stock=1, seller_account=seller, seller_name="Digest Shop")
        quiet = SellerAccount.objects.create(user=User.objects.create_user("qq", "qq@example.com", "x"), status="approved", business_name="Quiet")
        with self.captureOnCommitCallbacks(execute=True):
            call_command("send_digests", stdout=open(os.devnull, "w"))
        to = [m.to[0] for m in mail.outbox]
        self.assertIn("owner@example.com", to)
        self.assertIn("ds@example.com", to)
        self.assertNotIn("qq@example.com", to)  # nothing to tell them
        _ = quiet

    def test_admin_search(self):
        staff = User.objects.create_user("ad", "ad@example.com", "pass12345", is_staff=True, is_superuser=True)
        self.client.force_login(staff)
        order = Order.objects.create(email="find@example.com", full_name="Fin Der", address="1", city="X", country="US", phone="+15550100")
        r = self.client.get(reverse("manage_search"), {"q": f"#{order.id}"})
        self.assertRedirects(r, reverse("manage_order", args=[order.id]), fetch_redirect_response=False)
        make_product(name="Searchable Lamp")
        make_product(name="Searchable Lamp Two")
        self.assertContains(self.client.get(reverse("manage_search"), {"q": "searchable"}), "Searchable Lamp Two")

    def test_whatsapp_link_and_channel(self):
        from unittest import mock
        from .templatetags.bees_extras import wa_link
        self.assertEqual(wa_link("+92 300 1234567", "Hi"), "https://wa.me/923001234567?text=Hi")
        st = SiteSettings.load(); st.whatsapp_number = "+923001234567"; st.message_channel = "whatsapp"; st.save()
        self.assertContains(self.client.get(reverse("help_support")), "wa.me/923001234567")
        from . import sms
        env = {"TWILIO_ACCOUNT_SID": "AC" + "0" * 32, "TWILIO_AUTH_TOKEN": "0" * 32, "TWILIO_FROM": "+15550001111"}
        with mock.patch.dict(os.environ, env), mock.patch("urllib.request.urlopen") as op:
            op.return_value.__enter__.return_value.status = 201
            sms.send_now("+923001234567", "Your order shipped")
            body = op.call_args[0][0].data.decode()
        self.assertIn("To=whatsapp%3A%2B923001234567", body)
        self.assertIn("From=whatsapp%3A%2B15550001111", body)

    def test_listing_from_photo_and_price_tips(self):
        from io import BytesIO
        from unittest import mock
        from PIL import Image
        user = User.objects.create_user("lp", "lp@example.com", "pass12345")
        SellerAccount.objects.create(user=user, status="approved", business_name="Photo Shop")
        for price in ("10", "20", "30", "40"):
            make_product(category="sneakers", price=Decimal(price))
        self.client.force_login(user)
        tips = self.client.get(reverse("price_tips"), {"category": "sneakers"}).json()["tips"]
        self.assertEqual(tips["count"], 4)
        buf = BytesIO()
        Image.new("RGB", (40, 40), "white").save(buf, "JPEG")
        draft = {"name": "White Sneakers", "category": "sneakers", "description": "Clean white sneakers.", "brand": "", "specs": "Colour: White"}
        with self.settings(AI_API_KEY="k"), mock.patch("bees.market.listing_from_photo", return_value=draft):
            r = self.client.post(reverse("listing_from_photo"), {"photo": SimpleUploadedFile("s.jpg", buf.getvalue(), content_type="image/jpeg")})
        self.assertEqual(r.json()["draft"]["name"], "White Sneakers")
        self.client.logout()
        self.assertEqual(self.client.post(reverse("listing_from_photo")).status_code, 403)


from .models import Conversation, Message, Profile, SiteSettings  # noqa: E402


class Round5Tests(TestCase):
    """Shopping channels, AI agents (UCP), review requests, passkeys,
    chat translation and import-duty estimates."""

    UCP = {"HTTP_UCP_AGENT": 'profile="https://agent.example/.well-known/ucp"'}

    # ---------------------------------------------------------- barcode & feeds
    def test_gtin_check_digit(self):
        from .views import valid_gtin
        for good in ("5012345678900", "036000291452", "96385074", "9780306406157"):
            self.assertTrue(valid_gtin(good), good)
        for bad in ("5012345678901", "12345", "abcdefghijklm", ""):
            self.assertFalse(valid_gtin(bad), bad)
        from django.core.exceptions import ValidationError
        from django.test import RequestFactory
        from .views import _extra_product_fields
        rf = RequestFactory()
        self.assertEqual(_extra_product_fields(rf.post("/", {"gtin": "5012 3456 78900", "mpn": " AB-1 "}))["gtin"], "5012345678900")
        with self.assertRaises(ValidationError):
            _extra_product_fields(rf.post("/", {"gtin": "5012345678901"}))

    def test_bulk_reads_barcodes_and_export_lines_up(self):
        from .bulk import COLUMNS, EXPORT_ONLY, _clean, _rows_for, _to_dicts
        rows = _to_dicts([["Title", "Variant Price", "Variant Barcode", "Vendor"], ["Mug", "9", "5012345678900.0", "Acme"]])
        self.assertEqual(_clean(rows[0])["gtin"], "5012345678900")
        p = make_product(name="Exported", brand="Acme", gtin="5012345678900", mpn="M-1")
        row = list(_rows_for(Product.objects.filter(pk=p.pk)))[0]
        header = COLUMNS + EXPORT_ONLY
        self.assertEqual(len(row), len(header))
        named = dict(zip(header, row))
        self.assertEqual((named["brand"], named["gtin"], named["mpn"], named["status"]), ("Acme", "5012345678900", "M-1", "Approved"))

    def test_feeds(self):
        make_product(name="Feed Lamp", brand="Lumo", gtin="5012345678900", price=Decimal("30"), old_price=Decimal("40"))
        make_product(name="Feed Auction", is_auction=True)
        make_product(name="Feed Quote", quote_only=True)
        xml = self.client.get(reverse("feed_google")).content.decode()
        self.assertIn("<g:gtin>5012345678900</g:gtin>", xml)
        self.assertIn("<g:price>40.00 USD</g:price>", xml)
        self.assertIn("<g:sale_price>30.00 USD</g:sale_price>", xml)
        self.assertNotIn("Feed Auction", xml)
        self.assertNotIn("Feed Quote", xml)
        csv_body = self.client.get(reverse("feed_csv")).content.decode()
        self.assertTrue(csv_body.startswith("id,title,description,availability"))
        self.assertIn("in stock", csv_body)
        from .feeds import readiness
        self.assertEqual(readiness()["count"], 1)

    def test_sales_channels_page(self):
        staff = User.objects.create_user("ch", "ch@example.com", "pass12345", is_staff=True, is_superuser=True)
        self.client.force_login(staff)
        page = self.client.get(reverse("manage_channels"))
        self.assertContains(page, "/feeds/google.xml")
        self.assertContains(page, "/.well-known/ucp")
        self.assertContains(self.client.get(reverse("manage_settings")), 'name="ai_agents"')

    # ---------------------------------------------------------- AI agents (UCP)
    def test_ucp_profile_and_llms_txt(self):
        data = self.client.get("/.well-known/ucp").json()
        self.assertEqual(data["ucp"]["version"], "2026-04-08")
        self.assertIn("dev.ucp.shopping.checkout", data["ucp"]["capabilities"])
        self.assertTrue(data["ucp"]["services"]["dev.ucp.shopping"][0]["endpoint"].endswith("/ucp/v1"))
        self.assertContains(self.client.get("/llms.txt"), "/feeds/google.xml")
        st = SiteSettings.load(); st.ai_agents = False; st.save()
        self.assertEqual(self.client.get("/.well-known/ucp").status_code, 404)
        self.assertNotContains(self.client.get("/llms.txt"), "/.well-known/ucp")

    def _create(self, items, **headers):
        import json
        return self.client.post("/ucp/v1/checkout-sessions", json.dumps({"line_items": items, "buyer": {"email": "ai@example.com"}}),
                                content_type="application/json", **{**self.UCP, **headers})

    def test_ucp_checkout_flow(self):
        import json
        p = make_product(name="Agent Mug", price=Decimal("12.99"), stock=5)
        self.assertEqual(self.client.post("/ucp/v1/checkout-sessions", "{}", content_type="application/json").status_code, 400)
        r = self._create([{"item": {"id": str(p.id)}, "quantity": 2}], HTTP_IDEMPOTENCY_KEY="k1")
        self.assertEqual(r.status_code, 201)
        body = r.json()
        self.assertEqual(body["status"], "requires_escalation")
        self.assertEqual(body["line_items"][0]["item"]["price"], 1299)
        self.assertEqual(body["totals"][-1], {"type": "total", "amount": 2598})
        self.assertTrue(any(m.get("severity") == "requires_buyer_input" for m in body["messages"]))
        self.assertIn("/agent-checkout/", body["continue_url"])
        again = self._create([{"item": {"id": str(p.id)}, "quantity": 2}], HTTP_IDEMPOTENCY_KEY="k1").json()
        self.assertEqual(again["id"], body["id"])
        url = f"/ucp/v1/checkout-sessions/{body['id']}"
        self.assertEqual(self.client.get(url, **self.UCP).json()["line_items"][0]["quantity"], 2)
        updated = self.client.put(url, json.dumps({"line_items": [{"item": {"id": str(p.id)}, "quantity": 9}]}),
                                  content_type="application/json", **self.UCP).json()
        self.assertEqual(updated["line_items"][0]["quantity"], 5)  # only 5 in stock
        self.assertTrue(any(m["code"] == "quantity_adjusted" for m in updated["messages"]))
        done = self.client.post(url + "/complete", **self.UCP).json()
        self.assertEqual(done["status"], "requires_escalation")
        # The shopper opens the link: items land in the cart, checkout opens.
        page = self.client.get(done["continue_url"].replace("http://testserver", ""))
        self.assertRedirects(page, reverse("checkout"), fetch_redirect_response=False)
        self.assertEqual(self.client.session["cart"], {str(p.id): 5})
        self.assertEqual(self.client.session["agent_email"], "ai@example.com")
        self.assertEqual(self.client.post(url + "/cancel", **self.UCP).json()["status"], "canceled")
        self.assertEqual(self.client.put(url, "{}", content_type="application/json", **self.UCP).status_code, 409)

    def test_ucp_rejects_bad_items(self):
        sold_out = make_product(name="Gone", stock=0)
        self.assertEqual(self._create([{"item": {"id": "nonsense"}, "quantity": 1}]).status_code, 400)
        self.assertEqual(self._create([{"item": {"id": "1"}, "quantity": 500}]).status_code, 400)
        r = self._create([{"item": {"id": str(sold_out.id)}, "quantity": 1}])
        self.assertEqual(r.status_code, 200)
        self.assertNotIn("id", r.json())
        from .models import AgentCheckout
        self.assertFalse(AgentCheckout.objects.exists())

    # ---------------------------------------------------------- review requests
    def _delivered_order(self, user, product, days_ago):
        from django.utils import timezone
        from datetime import timedelta
        when = timezone.now() - timedelta(days=days_ago)
        order = Order.objects.create(user=user, email=user.email, full_name="R", address="1", city="X", country="US",
                                     phone="+15550100", status="delivered", delivered_at=when)
        OrderItem.objects.create(order=order, product=product, product_name=product.name, price=product.price, quantity=1)
        return order

    def test_review_request_emails(self):
        from django.core import mail
        from .review_requests import send
        buyer = User.objects.create_user("rv", "rv@example.com", "x")
        p1, p2 = make_product(name="Review Me"), make_product(name="Already Reviewed")
        order = self._delivered_order(buyer, p1, 5)
        OrderItem.objects.create(order=order, product=p2, product_name=p2.name, price=p2.price, quantity=1)
        Review.objects.create(product=p2, user=buyer, username="rv", rating=5, comment="ok")
        too_new = self._delivered_order(buyer, p1, 1)
        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(send(), 1)
        self.assertEqual(len(mail.outbox), 1)
        html = mail.outbox[0].alternatives[0][0]
        self.assertIn("Review Me", html)
        self.assertNotIn("Already Reviewed", html)
        self.assertIn("#write-review", html)
        self.assertEqual(send(), 0)  # never twice
        too_new.refresh_from_db()
        self.assertFalse(too_new.review_request_sent)
        other = User.objects.create_user("rv2", "rv2@example.com", "x")
        Profile.objects.update_or_create(user=other, defaults={"review_emails": False})
        self._delivered_order(other, p1, 5)
        self.assertEqual(send(), 0)

    def test_photo_review_points_once(self):
        from io import BytesIO
        from unittest import mock
        from PIL import Image
        buyer = User.objects.create_user("ph", "ph@example.com", "pass12345")
        p = make_product(name="Photo Product")
        self._delivered_order(buyer, p, 5)
        self.client.force_login(buyer)

        def photo():
            buf = BytesIO(); Image.new("RGB", (30, 30), "red").save(buf, "JPEG")
            return SimpleUploadedFile("r.jpg", buf.getvalue(), content_type="image/jpeg")
        with mock.patch("bees.images.save_public", return_value="/media/reviews/r.jpg"):
            self.client.post(reverse("product_detail", args=[p.id]), {"rating": "5", "comment": "Great", "photos": photo()})
            self.client.post(reverse("product_detail", args=[p.id]), {"rating": "4", "comment": "Still great", "photos": photo()})
        self.assertEqual(Profile.objects.get(user=buyer).loyalty_points, 50)
        self.assertTrue(Review.objects.get(user=buyer).photo_bonus)

    # ---------------------------------------------------------- passkeys
    def test_passkey_register_login_and_remove(self):
        import json
        from types import SimpleNamespace
        from unittest import mock
        from .models import Passkey
        user = User.objects.create_user("pk", "pk@example.com", "pass12345")
        self.assertNotEqual(self.client.post(reverse("passkey_register_options")).status_code, 200)  # signed out
        self.client.force_login(user)
        opts = self.client.post(reverse("passkey_register_options")).json()
        self.assertEqual(opts["rp"]["id"], "testserver")
        self.assertIn("challenge", opts)
        reg = SimpleNamespace(credential_id=b"\x01\x02cred", credential_public_key=b"pubkey", sign_count=0)
        cred = {"id": "AQJjcmVk", "rawId": "AQJjcmVk", "type": "public-key", "response": {"transports": ["internal"]}}
        with mock.patch("webauthn.verify_registration_response", return_value=reg):
            r = self.client.post(reverse("passkey_register"), json.dumps({"credential": cred}), content_type="application/json")
        self.assertTrue(r.json()["ok"])
        key = Passkey.objects.get(user=user)
        self.assertEqual(key.credential_id, "AQJjcmVk")
        self.assertContains(self.client.get(reverse("account_security")), "Remove")
        self.client.logout()
        self.client.post(reverse("passkey_login_options"))
        with mock.patch("webauthn.verify_authentication_response", return_value=SimpleNamespace(new_sign_count=3)):
            ok = self.client.post(reverse("passkey_login"), json.dumps({"credential": cred, "next": "/my-orders/"}),
                                  content_type="application/json").json()
        self.assertEqual(ok, {"ok": True, "next": "/my-orders/"})
        self.assertEqual(int(self.client.session["_auth_user_id"]), user.pk)
        key.refresh_from_db()
        self.assertEqual(key.sign_count, 3)
        # Without a fresh challenge, or with an unknown key, nothing happens.
        self.client.logout()
        bad = self.client.post(reverse("passkey_login"), json.dumps({"credential": cred}), content_type="application/json")
        self.assertEqual(bad.status_code, 400)
        self.client.force_login(user)
        self.client.post(reverse("passkey_delete", args=[key.id]))
        self.assertFalse(Passkey.objects.exists())

    # ---------------------------------------------------------- chat translation
    def test_translate_chat_message(self):
        from unittest import mock
        buyer = User.objects.create_user("tb", "tb@example.com", "x")
        seller_user = User.objects.create_user("ts2", "ts2@example.com", "x")
        seller = SellerAccount.objects.create(user=seller_user, status="approved", business_name="Tr Shop")
        conv = Conversation.objects.create(buyer=buyer, seller=seller)
        msg = Message.objects.create(conversation=conv, sender=seller_user, from_seller=True, body="Hola, ¿qué talla?")
        stranger = User.objects.create_user("st", "st@example.com", "x")
        self.client.force_login(stranger)
        self.assertEqual(self.client.post(reverse("translate_message", args=[msg.id])).status_code, 404)
        self.client.force_login(buyer)
        with mock.patch("bees.round5_views.translate_text", return_value="Hi, what size?") as tr:
            r = self.client.post(reverse("translate_message", args=[msg.id]), {"lang": "en"}).json()
            self.client.post(reverse("translate_message", args=[msg.id]), {"lang": "en"})
        self.assertEqual(r["text"], "Hi, what size?")
        self.assertEqual(tr.call_count, 1)  # cached

    # ---------------------------------------------------------- import duty
    def test_duty_estimate(self):
        from .models import TaxRate
        from .taxes import duty_estimate, table
        st = SiteSettings.load(); st.ship_from_country = "PK"; st.save()
        TaxRate.objects.create(country="GB", name="VAT", rate=Decimal("20"), duty_rate=Decimal("12"), duty_free_up_to=Decimal("150"))
        self.assertIsNone(duty_estimate(Decimal("500"), "PK"))
        self.assertEqual(duty_estimate(Decimal("100"), "GB"), Decimal("0"))
        self.assertEqual(duty_estimate(Decimal("500"), "GB"), Decimal("60.00"))
        self.assertEqual(table()["countries"]["GB"]["duty"], 12.0)
        st.duties_policy = "included"; st.save()
        self.assertIsNone(duty_estimate(Decimal("500"), "GB"))
        staff = User.objects.create_user("tx", "tx@example.com", "pass12345", is_staff=True, is_superuser=True)
        self.client.force_login(staff)
        self.client.post(reverse("manage_taxes"), {"action": "save", "country": "DE", "name": "VAT", "rate": "19",
                                                    "duty_rate": "8", "duty_free_up_to": "150"})
        self.assertEqual(TaxRate.objects.get(country="DE").duty_rate, Decimal("8"))


from .models import GiftRegistry, GroupBuy, LiveShow, RegistryItem  # noqa: E402


class ShopperPackTests(TestCase):
    """Live shows, videos, group buys, prize wheel, pre-orders, filters,
    blog, gift registries, 'still looking?' emails and address lookup."""

    def _buy(self, user, product, **post):
        self.client.force_login(user)
        self.client.post(reverse("add_to_cart", args=[product.id]), {"quantity": 1, **post})
        return self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "email": user.email, "payment_method": "cod"})

    def test_group_buy_fills_and_pays_back_after_delivery(self):
        from django.utils import timezone
        from .shopper import settle_groups
        p = make_product(name="Team Kettle", price=Decimal("50"), stock=10, group_size=2, group_percent=10)
        a = User.objects.create_user("ga", "ga@example.com", "x")
        b = User.objects.create_user("gb", "gb@example.com", "x")
        self.client.force_login(a)
        self.assertRedirects(self.client.post(reverse("group_start", args=[p.id])), reverse("cart"), fetch_redirect_response=False)
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "email": a.email, "payment_method": "cod"})
        group = GroupBuy.objects.get(product=p)
        self.assertEqual((group.leader, group.joined, group.completed_at), (a, 1, None))
        self.assertNotIn("group_for", self.client.session)
        page = self.client.get(reverse("group_page", args=[group.code]))
        self.assertContains(page, "1 more person needed")
        self.client.force_login(b)
        self.client.post(reverse("group_join", args=[group.code]))
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "email": b.email, "payment_method": "cod"})
        group.refresh_from_db()
        self.assertIsNotNone(group.completed_at)
        # A third person can't join a full group.
        c = User.objects.create_user("gc", "gc@example.com", "x")
        self.client.force_login(c)
        self.client.post(reverse("group_join", args=[group.code]))
        self.assertEqual(group.members.count(), 2)
        self.assertEqual(settle_groups(), 0)  # nothing delivered yet
        Order.objects.filter(user__in=[a, b]).update(status="delivered", delivered_at=timezone.now())
        self.assertEqual(settle_groups(), 2)
        self.assertEqual(Profile.objects.get(user=a).store_credit, Decimal("5.00"))
        self.assertEqual(settle_groups(), 0)  # only once

    def test_spin_wheel(self):
        import random
        from .shopper import prizes, spin
        self.assertEqual(prizes("10% off:5\nnonsense\nFree shipping:2\n500% off:1\nTry again:0"), [("10% off", 5), ("Free shipping", 2)])
        st = SiteSettings.load(); st.spin_enabled = True; st.spin_prizes = "10% off:1"; st.save()
        user = User.objects.create_user("sp", "sp@example.com", "x")
        result, index = spin(user, rng=random.Random(1))
        self.assertEqual((result.prize, index, result.coupon.percent_off), ("10% off", 0, 10))
        other = User.objects.create_user("sp2", "sp2@example.com", "x")
        self.assertFalse(result.coupon.is_valid_for(other, Decimal("100"))[0])
        self.assertTrue(result.coupon.is_valid_for(user, Decimal("100"))[0])
        with self.assertRaises(ValueError):
            spin(user)  # once a day
        st.spin_prizes = "75 points:1"; st.save()
        self.client.force_login(other)
        j = self.client.post(reverse("spin")).json()
        self.assertEqual(j["points"], 75)
        self.assertEqual(Profile.objects.get(user=other).loyalty_points, 75)
        st.spin_enabled = False; st.save()
        self.assertEqual(self.client.get(reverse("spin")).status_code, 404)

    def test_preorder(self):
        from datetime import timedelta
        from django.utils import timezone
        from .insights import delivery_window
        ships = timezone.localdate() + timedelta(days=30)
        p = make_product(name="Future Phone", preorder_ships_on=ships, stock=5)
        self.assertGreaterEqual(delivery_window(p)[0], ships)
        self.assertContains(self.client.get(reverse("product_detail", args=[p.id])), "Pre-order")
        user = User.objects.create_user("po", "po@example.com", "x")
        self._buy(user, p)
        self.assertGreater(Order.objects.get(user=user).estimated_delivery, ships)

    def test_preorder_and_group_fields_validated(self):
        from django.core.exceptions import ValidationError
        from django.test import RequestFactory
        from .views import _selling_fields
        rf = RequestFactory()
        with self.assertRaises(ValidationError):
            _selling_fields(rf.post("/", {"group_size": "2"}))  # no % back
        with self.assertRaises(ValidationError):
            _selling_fields(rf.post("/", {"preorder_ships_on": "2001-01-01"}))
        ok = _selling_fields(rf.post("/", {"group_size": "3", "group_percent": "5"}))
        self.assertEqual((ok["group_size"], ok["group_percent"], ok["preorder_ships_on"]), (3, 5, None))

    def test_material_and_condition_filters(self):
        make_product(name="Cotton Tee", specs="Material: 100% Cotton\nFit: Regular")
        make_product(name="Wool Hat", specs="Material: Wool")
        make_product(name="Used Lamp", condition="used_good")
        page = self.client.get(reverse("all_products"))
        self.assertContains(page, 'value="Cotton"')
        names = [p.name for p in self.client.get(reverse("all_products"), {"material": "Cotton"}).context["products"]]
        self.assertEqual(names, ["Cotton Tee"])
        names = [p.name for p in self.client.get(reverse("all_products"), {"condition": "used"}).context["products"]]
        self.assertEqual(names, ["Used Lamp"])

    def test_live_show_flow(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models import LiveReminder, Notification, StoreFollow
        seller_user = User.objects.create_user("lv", "lv@example.com", "x")
        seller = SellerAccount.objects.create(user=seller_user, status="approved", business_name="Live Shop")
        mine = make_product(name="Live Product", seller_account=seller, seller_name="Live Shop")
        fan = User.objects.create_user("fan", "fan@example.com", "x")
        StoreFollow.objects.create(user=fan, seller_account=seller)
        self.client.force_login(seller_user)
        start = (timezone.now() + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M")
        self.client.post(reverse("seller_live"), {"title": "Spring try-on", "starts_at": start,
                                                  "stream_url": "https://www.youtube.com/live/abcdefghijk", "products": [mine.id]})
        show = LiveShow.objects.get()
        self.assertEqual(list(show.products.all()), [mine])
        self.assertEqual(show.embed()["kind"], "youtube")
        watcher = User.objects.create_user("wt", "wt@example.com", "x")
        self.client.force_login(watcher)
        self.client.post(reverse("live_remind", args=[show.id]))
        self.assertTrue(LiveReminder.objects.filter(show=show, user=watcher).exists())
        self.assertEqual(self.client.post(reverse("live_comments", args=[show.id]), {"body": "hi"}).status_code, 400)  # not live yet
        self.client.force_login(seller_user)
        self.client.post(reverse("seller_live"), {"action": "start", "id": show.id})
        show.refresh_from_db()
        self.assertEqual(show.status, "live")
        self.assertTrue(Notification.objects.filter(user=fan, message__contains="Live now").exists())
        self.assertTrue(Notification.objects.filter(user=watcher, message__contains="Live now").exists())
        self.assertContains(self.client.get(reverse("live_list")), "Spring try-on")
        self.client.force_login(watcher)
        data = self.client.post(reverse("live_comments", args=[show.id]), {"body": "Is it waterproof?"}).json()
        self.assertEqual(data["comments"][-1]["body"], "Is it waterproof?")
        self.assertContains(self.client.get(reverse("live_detail", args=[show.id])), "youtube-nocookie.com/embed/abcdefghijk")
        other_seller = User.objects.create_user("os", "os@example.com", "x")
        SellerAccount.objects.create(user=other_seller, status="approved", business_name="Other")
        self.client.force_login(other_seller)
        self.client.post(reverse("seller_live"), {"action": "end", "id": show.id})
        show.refresh_from_db()
        self.assertEqual(show.status, "live")  # not theirs

    def test_videos_page(self):
        make_product(name="Video Lamp", video_url="https://youtu.be/abcdefghijk")
        make_product(name="No Video")
        page = self.client.get(reverse("videos"))
        self.assertContains(page, "Video Lamp")
        self.assertNotContains(page, "No Video")

    def test_blog(self):
        staff = User.objects.create_user("bl", "bl@example.com", "pass12345", is_staff=True, is_superuser=True)
        p = make_product(name="Blog Shoe")
        self.client.force_login(staff)
        self.client.post(reverse("manage_blog"), {"title": "How to choose running shoes", "excerpt": "A guide",
                                                  "body": "## Fit\n\nTry them on <b>late</b>.\n\n- Light\n- Soft", "products": [p.id]})
        from .models import BlogPost
        post = BlogPost.objects.get()
        self.assertFalse(post.is_published)
        self.client.logout()
        self.assertEqual(self.client.get(post.get_absolute_url()).status_code, 404)
        post.is_published = True; post.published_at = post.created_at; post.save()
        page = self.client.get(post.get_absolute_url())
        self.assertContains(page, "<h2>Fit</h2>")
        self.assertContains(page, "&lt;b&gt;late&lt;/b&gt;")  # escaped
        self.assertContains(page, "<li>Light</li>")
        self.assertContains(page, "Blog Shoe")
        self.assertContains(self.client.get(reverse("blog_rss")), "How to choose running shoes")
        self.assertContains(self.client.get("/sitemap.xml"), post.get_absolute_url())

    def test_gift_registry(self):
        from .models import Notification
        owner = User.objects.create_user("ro", "ro@example.com", "x")
        guest = User.objects.create_user("rg", "rg@example.com", "x")
        p = make_product(name="Toaster", stock=5)
        self.client.force_login(owner)
        r = self.client.post(reverse("registries"), {"title": "Sara & Ali", "event": "wedding", "event_date": "2027-05-01"})
        reg = GiftRegistry.objects.get(owner=owner)
        self.assertRedirects(r, reverse("registry_manage", args=[reg.id]), fetch_redirect_response=False)
        self.client.post(reverse("registry_add", args=[p.id]), {"registry": reg.id})
        self.client.post(reverse("registry_add", args=[p.id]), {"registry": reg.id})
        line = RegistryItem.objects.get(registry=reg)
        self.assertEqual(line.wanted, 2)
        self.client.force_login(guest)
        self.assertEqual(self.client.post(reverse("registry_add", args=[p.id]), {"registry": reg.id}).status_code, 404)  # not theirs
        self.assertContains(self.client.get(reverse("registry_public", args=[reg.token])), "2 of 2 still wanted")
        self.client.post(reverse("registry_buy", args=[line.id]))
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "email": guest.email, "payment_method": "cod"})
        line.refresh_from_db()
        self.assertEqual(line.bought, 1)
        self.assertEqual(OrderItem.objects.get(order__user=guest).registry_item, line)
        self.assertTrue(Notification.objects.filter(user=owner, message__contains="from your registry").exists())

    def test_browse_reminders(self):
        from datetime import timedelta
        from django.core import mail
        from django.utils import timezone
        from .models import ProductView
        from .shopper import send_browse_reminders
        user = User.objects.create_user("bw", "bw@example.com", "x")
        p = make_product(name="Looked At Lamp")
        self.client.force_login(user)
        self.client.get(reverse("product_detail", args=[p.id]))
        view = ProductView.objects.get(user=user)
        ProductView.objects.filter(pk=view.pk).update(viewed_at=timezone.now() - timedelta(days=1))
        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(send_browse_reminders(), 1)
        self.assertIn("Looked At Lamp", mail.outbox[0].alternatives[0][0])
        self.assertEqual(send_browse_reminders(), 0)  # at most weekly

    def test_address_lookup(self):
        import json
        import os
        from unittest import mock
        from .shopper_views import parse_components
        parts = [{"longText": "10", "types": ["street_number"]}, {"longText": "Downing Street", "types": ["route"]},
                 {"longText": "London", "types": ["postal_town"]}, {"longText": "SW1A 2AA", "types": ["postal_code"]},
                 {"longText": "United Kingdom", "shortText": "GB", "types": ["country"]}]
        self.assertEqual(parse_components(parts), {"address": "10 Downing Street", "city": "London", "state": "",
                                                   "postal_code": "SW1A 2AA", "country": "GB"})
        self.assertEqual(self.client.get(reverse("address_suggest"), {"q": "10 Down"}).json(), {"suggestions": []})  # no key
        reply = mock.MagicMock()
        reply.__enter__.return_value.read.return_value = json.dumps(
            {"suggestions": [{"placePrediction": {"placeId": "ChIJabc123xyz", "text": {"text": "10 Downing St, London"}}}]}).encode()
        with mock.patch.dict(os.environ, {"GOOGLE_MAPS_API_KEY": "AIza" + "x" * 35}), mock.patch("urllib.request.urlopen", return_value=reply) as op:
            data = self.client.get(reverse("address_suggest"), {"q": "10 Down", "country": "GB"}).json()
            sent = json.loads(op.call_args[0][0].data)
        self.assertEqual(data["suggestions"][0]["id"], "ChIJabc123xyz")
        self.assertEqual(sent["includedRegionCodes"], ["gb"])


from .models import Company, CompanyMember, StockLevel, Warehouse, WebhookDelivery, WebhookEndpoint  # noqa: E402


class OpsB2BTests(TestCase):
    """Webhooks, stock locations, business accounts and invoices, customer
    item cancel / address change, AI help chat, accounting export, ID
    checks and 3D models."""

    def _order(self, user, product, method="cod", **post):
        self.client.force_login(user)
        self.client.post(reverse("add_to_cart", args=[product.id]), {"quantity": post.pop("qty", 1)})
        return self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "email": user.email, "payment_method": method, **post})

    def test_webhooks_signed_and_retried(self):
        import hashlib
        import hmac
        import json
        from unittest import mock
        from .ops import retry_due
        ep = WebhookEndpoint.objects.create(url="https://hooks.example.com/x", events=["order.created"], secret="s3cret")
        user = User.objects.create_user("wh", "wh@example.com", "x")
        p = make_product(name="Hooked Lamp", stock=5)
        ok = mock.MagicMock(); ok.__enter__.return_value.status = 200
        with mock.patch("urllib.request.urlopen", return_value=ok) as op, self.captureOnCommitCallbacks(execute=True):
            self._order(user, p)
        req = op.call_args[0][0]
        body = req.data
        self.assertEqual(req.headers["X-lumen-signature"], "sha256=" + hmac.new(b"s3cret", body, hashlib.sha256).hexdigest())
        data = json.loads(body)
        self.assertEqual(data["event"], "order.created")
        self.assertEqual(data["data"]["items"][0]["name"], "Hooked Lamp")
        self.assertIsNotNone(WebhookDelivery.objects.get().delivered_at)
        # A failing endpoint is retried later.
        WebhookDelivery.objects.all().delete()
        with mock.patch("urllib.request.urlopen", side_effect=OSError("down")), self.captureOnCommitCallbacks(execute=True):
            self._order(user, p)
        d = WebhookDelivery.objects.get()
        self.assertIsNone(d.delivered_at)
        self.assertEqual(d.attempts, 1)
        WebhookDelivery.objects.filter(pk=d.pk).update(next_try_at=d.created_at)
        with mock.patch("urllib.request.urlopen", return_value=ok):
            self.assertEqual(retry_due(), 1)
        self.assertIsNotNone(WebhookDelivery.objects.get().delivered_at)
        staff = User.objects.create_user("wa", "wa@example.com", "pass12345", is_staff=True, is_superuser=True)
        self.client.force_login(staff)
        self.assertContains(self.client.get(reverse("manage_integrations")), "hooks.example.com")

    def test_stock_locations(self):
        from .payments import restock
        seller_user = User.objects.create_user("loc", "loc@example.com", "x")
        seller = SellerAccount.objects.create(user=seller_user, status="approved", business_name="Loc Shop")
        p = make_product(name="Located Chair", stock=0, seller_account=seller, seller_name="Loc Shop")
        self.client.force_login(seller_user)
        self.client.post(reverse("seller_locations"), {"action": "add", "name": "Lisbon", "country": "PT"})
        self.client.post(reverse("seller_locations"), {"action": "add", "name": "Austin", "country": "US"})
        pt, us = Warehouse.objects.order_by("id")
        self.client.post(reverse("seller_locations"), {"action": "stock", f"q_{pt.id}_{p.id}_0": "5", f"q_{us.id}_{p.id}_0": "2"})
        p.refresh_from_db()
        self.assertEqual(p.stock, 7)
        buyer = User.objects.create_user("lb", "lb@example.com", "x")
        self._order(buyer, p, qty=2)  # ships to the US
        item = OrderItem.objects.get(product=p)
        self.assertEqual(item.warehouse, us)
        self.assertEqual(StockLevel.objects.get(warehouse=us).quantity, 0)
        p.refresh_from_db()
        self.assertEqual(p.stock, 5)
        restock(item.order)
        self.assertEqual(StockLevel.objects.get(warehouse=us).quantity, 2)
        other = User.objects.create_user("lo", "lo@example.com", "x")
        SellerAccount.objects.create(user=other, status="approved", business_name="Other")
        self.client.force_login(other)
        self.client.post(reverse("seller_locations"), {"action": "stock", f"q_{pt.id}_{p.id}_0": "99"})
        self.assertEqual(StockLevel.objects.get(warehouse=pt).quantity, 5)  # not their location

    def test_business_account_invoice_flow(self):
        from datetime import timedelta
        from django.core import mail
        from django.utils import timezone
        from .fulfilment import blocked_reason
        from .models import TaxRate
        from .ops import remind_invoices
        from .seller_center import on_hold_items
        TaxRate.objects.create(country="US", name="Sales tax", rate=Decimal("8"))
        buyer = User.objects.create_user("biz", "biz@example.com", "x")
        self.client.force_login(buyer)
        self.client.post(reverse("company_account"), {"name": "Acme Ltd", "billing_email": "ap@acme.example", "tax_id": "US123",
                                                      "tax_exempt_requested": "1", "exemption_note": "Cert 42"})
        company = Company.objects.get()
        self.assertEqual((company.status, CompanyMember.objects.get().role), ("pending", "admin"))
        staff = User.objects.create_user("bs", "bs@example.com", "pass12345", is_staff=True, is_superuser=True)
        self.client.force_login(staff)
        self.client.post(reverse("manage_company", args=[company.id]), {"action": "save", "status": "approved", "payment_terms": "30",
                                                                         "credit_limit": "100", "tax_exempt": "1"})
        company.refresh_from_db()
        self.assertEqual((company.status, company.payment_terms, company.tax_exempt), ("approved", 30, True))
        seller_user = User.objects.create_user("bsel", "bsel@example.com", "x")
        seller = SellerAccount.objects.create(user=seller_user, status="approved", business_name="B Shop")
        p = make_product(name="Office Chair", price=Decimal("40"), stock=10, seller_account=seller, seller_name="B Shop")
        self.client.force_login(buyer)
        self.client.post(reverse("add_to_cart", args=[p.id]), {"quantity": 1})
        self.assertContains(self.client.get(reverse("checkout")), 'value="invoice"')
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "email": buyer.email, "payment_method": "invoice", "po_number": "PO-777"})
        order = Order.objects.get(user=buyer)
        self.assertEqual((order.payment_status, order.po_number, order.company, order.tax_amount), ("invoiced", "PO-777", company, Decimal("0")))
        self.assertEqual(order.payment_due, timezone.localdate() + timedelta(days=30))
        self.assertEqual(blocked_reason(order), "")  # sellers can ship it
        self.assertTrue(on_hold_items(seller).exists())  # but aren't paid out until the invoice is paid
        # Over the credit limit: invoice isn't offered.
        self.client.post(reverse("add_to_cart", args=[p.id]), {"quantity": 2})
        self.assertNotContains(self.client.get(reverse("checkout")), 'value="invoice"')
        Order.objects.filter(pk=order.pk).update(payment_due=timezone.localdate() - timedelta(days=2))
        mail.outbox.clear()
        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(remind_invoices(), 1)
        self.assertEqual(mail.outbox[0].to, ["ap@acme.example"])
        self.client.force_login(staff)
        self.client.post(reverse("manage_company", args=[company.id]), {"action": "paid", "order": order.id})
        order.refresh_from_db()
        self.assertEqual(order.payment_status, "paid")

    def test_customer_cancels_one_item_and_changes_address(self):
        from unittest import mock
        from .fulfilment import FulfilmentError, set_items_status
        buyer = User.objects.create_user("ci", "ci@example.com", "x")
        a, b = make_product(name="Keep Me", price=Decimal("20"), stock=5), make_product(name="Drop Me", price=Decimal("10"), stock=5)
        order = Order.objects.create(user=buyer, email=buyer.email, full_name="C I", address="1 St", city="X", country="US", phone="+15550100",
                                     payment_method="card", payment_status="paid", stripe_payment_intent="pi_1")
        OrderItem.objects.create(order=order, product=a, product_name=a.name, price=a.price, quantity=1)
        drop = OrderItem.objects.create(order=order, product=b, product_name=b.name, price=b.price, quantity=1)
        self.client.force_login(buyer)
        with mock.patch("bees.payments.refund_amount", return_value=True) as refund:
            self.client.post(reverse("cancel_item", args=[drop.id]))
        drop.refresh_from_db()
        self.assertIsNotNone(drop.cancelled_at)
        self.assertEqual(refund.call_args[0][1], Decimal("10.00"))
        b.refresh_from_db()
        self.assertEqual(b.stock, 6)
        with self.assertRaises(FulfilmentError):
            set_items_status([drop], "packed")
        keep = order.items.get(product=a)
        self.assertFalse(keep.customer_can_cancel)  # last item: cancel the order instead
        self.client.post(reverse("change_address", args=[order.id]), {"full_name": "C I", "address": "2 New St", "city": "Austin",
                                                                       "state": "TX", "postal_code": "73301", "phone": "+15550100"})
        order.refresh_from_db()
        self.assertEqual(order.address, "2 New St")
        OrderItem.objects.filter(pk=keep.pk).update(fulfillment_status="handed_to_courier")
        self.client.post(reverse("change_address", args=[order.id]), {"full_name": "C I", "address": "3 Late St", "city": "Austin",
                                                                       "postal_code": "73301", "phone": "+15550100"})
        order.refresh_from_db()
        self.assertEqual(order.address, "2 New St")
        cod = Order.objects.create(user=buyer, email=buyer.email, full_name="C", address="1", city="X", country="US", phone="1", payment_method="cod")
        i1 = OrderItem.objects.create(order=cod, product=a, product_name=a.name, price=a.price, quantity=1)
        OrderItem.objects.create(order=cod, product=b, product_name=b.name, price=b.price, quantity=1)
        self.assertFalse(i1.customer_can_cancel)

    def test_ai_help_chat(self):
        import json
        from unittest import mock
        reply = mock.MagicMock()
        reply.__enter__.return_value.read.return_value = json.dumps(
            {"content": [{"type": "text", "text": '{"answer": "Your order #5 ships tomorrow.", "handoff": false}'}]}).encode()
        with self.settings(AI_API_KEY="k"), mock.patch("urllib.request.urlopen", return_value=reply) as op:
            r = self.client.post(reverse("chat_send"), {"message": "Is the blue lamp good for reading at night?"}).json()
            prompt = json.loads(op.call_args[0][0].data)["messages"][0]["content"]
        self.assertIn("ships tomorrow", r["reply"]["message"])
        self.assertIn("Use ONLY the facts", prompt)
        with self.settings(AI_API_KEY=""):
            r = self.client.post(reverse("chat_send"), {"message": "Is the blue lamp good for reading at night?"}).json()
        self.assertIn("member of our team", r["reply"]["message"])

    def test_accounting_export(self):
        staff = User.objects.create_user("ae", "ae@example.com", "pass12345", is_staff=True, is_superuser=True)
        p = make_product(name="Booked Lamp", price=Decimal("25"))
        o = Order.objects.create(email="x@example.com", full_name="Book Keeper", address="1", city="X", country="GB", phone="1",
                                 payment_method="card", payment_status="paid", shipping_amount=Decimal("5"))
        OrderItem.objects.create(order=o, product=p, product_name=p.name, price=p.price, quantity=2)
        self.client.force_login(staff)
        xero = self.client.get(reverse("accounting_export"), {"format": "xero"}).content.decode()
        self.assertTrue(xero.startswith("*ContactName"))
        self.assertIn(f"LM-{o.id}", xero)
        self.assertIn("Shipping", xero)
        qb = self.client.get(reverse("accounting_export"), {"format": "quickbooks"}).content.decode()
        self.assertTrue(qb.startswith("InvoiceNo,Customer"))
        self.assertIn("Booked Lamp,2,25.00,50.00", qb)

    def test_identity_check_and_3d(self):
        from types import SimpleNamespace
        from unittest import mock
        from .ops import identity_event
        user = User.objects.create_user("idc", "idc@example.com", "x")
        seller = SellerAccount.objects.create(user=user, status="approved", business_name="ID Shop", identity_session="vs_1")
        identity_event("identity.verification_session.verified", {"id": "vs_1"})
        seller.refresh_from_db()
        self.assertEqual(seller.identity_status, "verified")
        self.assertIsNotNone(seller.verified_at)
        st = SiteSettings.load(); st.identity_checks = True; st.save()
        fake_stripe = SimpleNamespace(identity=SimpleNamespace(VerificationSession=SimpleNamespace(
            create=mock.MagicMock(return_value={"id": "vs_2", "url": "https://verify.stripe.com/x"}))))
        other = User.objects.create_user("id2", "id2@example.com", "x")
        SellerAccount.objects.create(user=other, status="approved", business_name="ID2")
        self.client.force_login(other)
        with mock.patch("bees.payments.is_configured", return_value=True), mock.patch("bees.payments._stripe", return_value=fake_stripe):
            r = self.client.post(reverse("seller_identity_start"))
        self.assertEqual(r["Location"], "https://verify.stripe.com/x")
        p = make_product(name="3D Vase", model_3d_url="https://models.example.com/vase.glb")
        page = self.client.get(reverse("product_detail", args=[p.id]))
        self.assertContains(page, "<model-viewer")
        self.assertIn("connect-src 'self' https://models.example.com", page["Content-Security-Policy"])
        self.assertIn("script-src 'self' 'wasm-unsafe-eval'", page["Content-Security-Policy"])
        self.assertNotIn("models.example.com", self.client.get(reverse("home"))["Content-Security-Policy"])


class MultiStoreTests(TestCase):
    """Multi-store mode: every store has its own database, cache space,
    files, keys and on/off switch."""

    def setUp(self):
        import tempfile
        from pathlib import Path
        self.tmp = Path(tempfile.mkdtemp())
        folder = self.tmp / "alpha"
        folder.mkdir()
        (folder / "store.json").write_text('{"name": "Alpha", "hosts": ["alpha.test"], "active": true}')
        (folder / ".env").write_text("STORE_CURRENCY=eur\nSTRIPE_SECRET_KEY=sk_test_alpha123456789012345\n")

    def tearDown(self):
        import shutil
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _store(self):
        from . import multistore
        with self.settings(MULTISTORE=True, MULTISTORE_DIR=str(self.tmp)):
            return multistore.load(force=True)["alpha"]

    def test_values_are_per_store(self):
        from . import multistore
        store = self._store()
        token = multistore._current.set(store)
        try:
            self.assertEqual(multistore.store_settings.STORE_CURRENCY, "EUR")
            self.assertTrue(multistore.cache_key("k", "", 1).startswith("alpha:"))
            self.assertEqual(multistore.media_prefix("reviews/a.jpg"), "stores/alpha/reviews/a.jpg")
            self.assertEqual(multistore.media_prefix("stores/alpha/x.jpg"), "stores/alpha/x.jpg")
            self.assertEqual(multistore.env("STRIPE_SECRET_KEY"), "sk_test_alpha123456789012345")
        finally:
            multistore._current.reset(token)
        self.assertNotEqual(multistore.store_settings.STORE_CURRENCY, "EUR")
        self.assertTrue(multistore.cache_key("k", "", 1).startswith("_:"))
        self.assertEqual(multistore.media_prefix("reviews/a.jpg"), "reviews/a.jpg")

    def test_keys_page_edits_the_store_file_only(self):
        import os
        from . import envfile, multistore
        store = self._store()
        token = multistore._current.set(store)
        try:
            envfile.write({"PAYPAL_CLIENT_ID": "A" * 40})
            with self.assertRaises(envfile.EnvError):
                envfile.write({"SENTRY_DSN": "https://x@y.ingest.sentry.io/1"})
        finally:
            multistore._current.reset(token)
        self.assertIn("PAYPAL_CLIENT_ID", (self.tmp / "alpha" / ".env").read_text())
        self.assertNotEqual(os.environ.get("PAYPAL_CLIENT_ID"), "A" * 40)

    def test_stripe_key_sent_per_call(self):
        from unittest import mock
        from . import multistore, payments
        store = self._store()
        token = multistore._current.set(store)
        try:
            with self.settings(MULTISTORE=True), mock.patch("stripe.Refund.create") as create:
                payments._stripe().Refund.create(payment_intent="pi_1")
        finally:
            multistore._current.reset(token)
        self.assertEqual(create.call_args.kwargs["api_key"], "sk_test_alpha123456789012345")

    def test_paused_store_shows_break_page(self):
        import json
        path = self.tmp / "alpha" / "store.json"
        path.write_text(json.dumps({"name": "Alpha", "hosts": ["alpha.test"], "active": False}))
        from . import multistore
        from django.conf import settings as dj
        with self.settings(MULTISTORE=True, MULTISTORE_DIR=str(self.tmp), ALLOWED_HOSTS=["*"],
                           MIDDLEWARE=["bees.multistore.StoreMiddleware"] + list(dj.MIDDLEWARE)):
            multistore.load(force=True)
            r = self.client.get("/", HTTP_HOST="alpha.test")
        self.assertEqual(r.status_code, 503)
        self.assertContains(r, "short break", status_code=503)

    def test_settings_memo_is_per_store(self):
        from . import multistore
        from .models import _SETTINGS_MEMO
        store = self._store()
        SiteSettings.load()
        token = multistore._current.set(store)
        try:
            SiteSettings.load()
        finally:
            multistore._current.reset(token)
        self.assertIn("", _SETTINGS_MEMO)
        self.assertIn("alpha", _SETTINGS_MEMO)

    def test_two_real_stores_are_separate(self):
        """Creates two stores with their own databases in a separate process
        and checks nothing crosses between them."""
        import os
        import subprocess
        import sys
        import tempfile
        from django.conf import settings as dj
        root = tempfile.mkdtemp()
        env = {**os.environ, "MULTISTORE": "True", "MULTISTORE_DIR": os.path.join(root, "stores"), "DEBUG": "True",
               "DATABASE_URL": f"sqlite:///{os.path.join(root, 'main.sqlite3')}", "SETUP_WIZARD_REDIRECT": "False",
               "ALLOWED_HOSTS": "*", "PYTHONPATH": str(dj.BASE_DIR)}
        env.pop("REDIS_URL", None)
        manage = [sys.executable, str(dj.BASE_DIR / "manage.py")]

        def run(*args):
            out = subprocess.run(manage + list(args), env=env, capture_output=True, text=True, timeout=300)
            self.assertEqual(out.returncode, 0, out.stdout[-800:] + out.stderr[-1500:])
            return out.stdout
        run("migrate", "--verbosity", "0")
        run("store_create", "alpha", "--host", "alpha.test", "--name", "Alpha Shop", "--owner-email", "a@alpha.test", "--currency", "EUR")
        run("store_create", "beta", "--host", "beta.test", "--name", "Beta Bazaar", "--owner-email", "b@beta.test")
        script = """
import django, re
django.setup()
from django.test import Client
from django.contrib.auth import get_user_model
from bees import multistore
from bees.models import Product
with multistore.use(multistore.get('alpha')):
    Product.objects.create(name='Alpha Only Teapot', category='skincare', price='12.50', stock=3, image_url='https://example.com/t.jpg')
    u = get_user_model().objects.get(email='a@alpha.test'); u.set_password('Alpha-pass-123'); u.save()
c = Client()
title = lambda h: re.search(r'<title>(.*?)</title>', c.get('/', HTTP_HOST=h).content.decode(), re.S).group(1)
print('TITLES', 'Alpha Shop' in title('alpha.test'), 'Beta Bazaar' in title('beta.test'), 'Alpha' not in title('localhost'))
print('SEARCH', 'Alpha Only Teapot' in c.get('/search/?q=teapot', HTTP_HOST='alpha.test').content.decode(),
      'Alpha Only Teapot' in c.get('/search/?q=teapot', HTTP_HOST='beta.test').content.decode())
r = Client().post('/login/', {'username': 'a@alpha.test', 'password': 'Alpha-pass-123'}, HTTP_HOST='alpha.test')
r2 = Client().post('/login/', {'username': 'a@alpha.test', 'password': 'Alpha-pass-123'}, HTTP_HOST='beta.test')
print('LOGIN', r.status_code, r2.status_code)
"""
        out = subprocess.run([sys.executable, "-c", script], env={**env, "DJANGO_SETTINGS_MODULE": "backend.settings"},
                             capture_output=True, text=True, timeout=300, cwd=str(dj.BASE_DIR))
        self.assertEqual(out.returncode, 0, out.stderr[-1500:])
        self.assertIn("TITLES True True True", out.stdout)
        self.assertIn("SEARCH True False", out.stdout)
        self.assertIn("LOGIN 302 200", out.stdout)
        self.assertIn("alpha", run("stores"))
        import shutil
        shutil.rmtree(root, ignore_errors=True)


class LaunchReadyTests(TestCase):
    """Mail guard, daily-task catch-up, 'mark resolved' and crash-safe menus."""

    def test_mail_guard_drops_reserved_addresses_and_retries(self):
        import smtplib
        from django.core.mail import EmailMessage
        from .mailguard import GuardedBackend, deliverable
        self.assertFalse(deliverable("shopper@demo.invalid"))
        self.assertFalse(deliverable("a@example.com"))
        self.assertTrue(deliverable("real@gmail.com"))
        with override_settings(MAILGUARD_INNER_BACKEND="django.core.mail.backends.locmem.EmailBackend"):
            backend = GuardedBackend()
            self.assertEqual(backend.send_messages([EmailMessage("Hi", "x", None, ["shopper@demo.invalid"])]), 0)
            self.assertEqual(len(mail.outbox), 0)
            msg = EmailMessage("Hi", "x", None, ["real@gmail.com", "b@demo.invalid"])
            calls = []

            def flaky(messages):
                calls.append(1)
                if len(calls) == 1:
                    raise smtplib.SMTPServerDisconnected("Connection unexpectedly closed")
                return len(messages)
            with mock.patch.object(backend.inner, "send_messages", side_effect=flaky):
                self.assertEqual(backend.send_messages([msg]), 1)
            self.assertEqual(len(calls), 2)
            self.assertEqual(msg.to, ["real@gmail.com"])

    def test_console_detection_looks_through_wrappers(self):
        from .mailguard import is_console
        with override_settings(EMAIL_BACKEND="bees.mailguard.GuardedBackend",
                               EMAIL_REAL_BACKEND="django.core.mail.backends.console.EmailBackend"):
            self.assertTrue(is_console())
        with override_settings(EMAIL_BACKEND="django.core.mail.backends.smtp.EmailBackend"):
            self.assertFalse(is_console())

    def test_overdue_daily_tasks_start_once(self):
        from datetime import timedelta
        from django.utils import timezone
        from . import catchup
        from .alerts import DAILY_TASKS
        from .models import AuditLog
        entry = AuditLog.objects.create(action=f"{DAILY_TASKS}: ok")
        AuditLog.objects.filter(pk=entry.pk).update(created_at=timezone.now() - timedelta(hours=2))
        with override_settings(DAILY_CATCHUP=True), mock.patch("bees.catchup.threading.Thread") as thread:
            self.assertFalse(catchup.maybe_start())
            AuditLog.objects.filter(pk=entry.pk).update(created_at=timezone.now() - timedelta(days=3))
            cache.clear()
            self.assertTrue(catchup.maybe_start())
            self.assertFalse(catchup.maybe_start())  # checked recently
            cache.delete("daily_catchup_checked:-")
            self.assertFalse(catchup.maybe_start())  # already running
        self.assertEqual(thread.call_count, 1)

    def test_catchup_run_skips_demo_reset_and_is_flagged(self):
        from django.core.management import call_command
        from .alerts import DAILY_TASKS
        from .models import AuditLog
        ran = []
        real = call_command

        def fake(name, *args, **kwargs):
            if name == "daily_tasks":
                return real(name, *args, **kwargs)
            ran.append(name)
        with override_settings(DEMO_MODE=True), mock.patch("bees.management.commands.daily_tasks.call_command", side_effect=fake):
            call_command("daily_tasks", "--catchup", stdout=io.StringIO())
        self.assertNotIn("demo_setup", ran)
        self.assertIn("backup_data", ran)
        self.assertIn("catch-up", AuditLog.objects.get(action__startswith=DAILY_TASKS).action)

    def test_mark_resolved_and_demo_staff_two_step(self):
        from .alerts import EMAIL_FAILED, SERVER_ERROR
        from .models import AuditLog
        boss = User.objects.create_user("boss2", "boss2@example.com", "pass12345", is_staff=True, is_superuser=True)
        User.objects.create_user("demo_admin", "admin@demo.invalid", "x", is_staff=True)
        AuditLog.objects.create(action=f"{EMAIL_FAILED}: 'Your order' to a@b.com - SMTPServerDisconnected: Connection unexpectedly closed")
        AuditLog.objects.create(action=f"{SERVER_ERROR}: GET /ask/ - ImportError: cannot import name 'LiveShow'")
        self.client.force_login(boss)
        page = self.client.get(reverse("manage_system")).content.decode()
        self.assertIn("Mark errors as resolved", page)
        self.assertIn("Gmail hiccup", page)
        self.assertIn("crashed again after the last Reload", page)
        with override_settings(DEMO_MODE=True):
            page = self.client.get(reverse("manage_system")).content.decode()
        self.assertNotIn("demo_admin", page.split("Staff two-step sign-in", 1)[1][:400])
        self.client.post(reverse("manage_system"), {"action": "resolve"})
        page = self.client.get(reverse("manage_system")).content.decode()
        self.assertNotIn("Mark errors as resolved", page)
        self.assertNotIn("Gmail hiccup", page)

    def test_menu_extras_never_break_a_page(self):
        with mock.patch("bees.context_processors._live_on", side_effect=RuntimeError("table missing")):
            self.assertEqual(self.client.get("/").status_code, 200)

    def test_page_translation_handles_icons_names_dates_and_codes(self):
        from .i18n import NAME_KEYS, _drop_names, lookup_text, table
        lk = table("fr")
        missing = set()
        self.assertEqual(lookup_text("🔥 Deals", lk, missing), "🔥 Bons plans")
        self.assertEqual(lookup_text("Search Demo Studio", lk, missing), "Rechercher dans Demo Studio")
        self.assertEqual(lookup_text("This item: Sunblock SPF 50", lk, missing), "Cet article : Sunblock SPF 50")
        self.assertEqual(lookup_text("Thu 15 Oct 2026", lk, missing), "jeu. 15 oct. 2026")
        self.assertEqual(lookup_text("Pakistan", table("ar"), missing), "باكستان")
        self.assertIsNone(lookup_text("· TCS", lk, missing))
        self.assertEqual(lookup_text("Asked for 5 pcs: please wrap them", lk, missing), "Demandé : 5 pièces: please wrap them")
        self.assertEqual(missing, set())
        self.assertIn(NAME_KEYS, lk)
        make_product(name="Sunblock SPF 50")
        self.assertEqual(_drop_names(["Sunblock SPF {0}", "Brand new text"]), ["Brand new text"])

    def test_prune_untranslated_rows(self):
        from .i18n import prune_missing
        from .models import UIText
        make_product(name="Linen Throw")
        for src in ("Pakistan", "Linen Throw", "A brand new sentence nobody translated", "Search {0} Studio"):
            UIText.objects.create(source=src)
        UIText.objects.create(source="Typed by owner", fr="Saisi")
        self.assertEqual(prune_missing(), 3)  # the pack, a product name, and a text the engine now handles
        self.assertEqual(set(UIText.objects.values_list("source", flat=True)),
                         {"A brand new sentence nobody translated", "Typed by owner"})


class SmartToolsTests(TestCase):
    def setUp(self):
        from .models import SellerAccount
        self.boss = User.objects.create_user("owner9", "owner9@example.com", "pass12345", is_staff=True, is_superuser=True)
        self.seller_user = User.objects.create_user("seller9", "seller9@example.com", "pass12345")
        self.seller = SellerAccount.objects.create(user=self.seller_user, status="approved", business_name="Nine Shop")

    def _order(self, user=None, email="buyer@example.com", product=None, qty=1, price="10.00", days_ago=0, **extra):
        from datetime import timedelta
        from django.utils import timezone
        from .models import Order, OrderItem
        o = Order.objects.create(user=user, email=email, full_name="Buyer One", address="1 Road", city="Town", country=extra.pop("country", "US"),
                                 phone=extra.pop("phone", "+15551234567"), payment_method=extra.pop("payment_method", "cod"),
                                 status=extra.pop("status", "confirmed"), payment_status=extra.pop("payment_status", "not_applicable"), **extra)
        Order.objects.filter(pk=o.pk).update(created_at=timezone.now() - timedelta(days=days_ago))
        o.refresh_from_db()
        if product:
            OrderItem.objects.create(order=o, product=product, product_name=product.name, price=Decimal(price), quantity=qty,
                                     seller_account=product.seller_account)
        return o

    def test_restock_forecast(self):
        from .smart import restock_plan
        fast = make_product(name="Fast seller", stock=4)
        idle = make_product(name="Sits on shelf", stock=20)
        self._order(product=fast, qty=30, days_ago=3)
        rows = {r["product"].name: r for r in restock_plan([fast, idle])}
        self.assertEqual(rows["Fast seller"]["status"], "urgent")
        self.assertGreater(rows["Fast seller"]["reorder"], 30)
        self.assertEqual(rows["Sits on shelf"]["status"], "slow")
        self.client.force_login(self.boss)
        page = self.client.get(reverse("manage_restock"))
        self.assertContains(page, "Fast seller")
        self.assertContains(page, "Reorder now")
        self.assertIn("text/csv", self.client.get(reverse("manage_restock") + "?export=csv")["Content-Type"])

    def test_seller_sees_only_own_restock(self):
        mine = make_product(name="Mine", stock=1)
        mine.seller_account = self.seller
        mine.save()
        make_product(name="Someone else's", stock=1)
        self.client.force_login(self.seller_user)
        page = self.client.get(reverse("seller_restock") + "?show=all")
        self.assertContains(page, "Mine")
        self.assertNotContains(page, "Someone else&#x27;s")

    def test_order_risk_flags_and_page(self):
        from .smart import order_risk
        p = make_product(name="Gadget", stock=50, price="20.00")
        for i in range(3):
            self._order(email=f"n{i}@example.com", product=p, days_ago=10)
        normal = self._order(email="calm@example.com", product=p)
        self.assertEqual(order_risk(normal)["level"], "low")
        self._order(email="x@mailinator.com", phone="", product=p, status="cancelled", cod_refused=True, days_ago=5)
        risky = self._order(email="x@mailinator.com", phone="", product=p, qty=40, price="20.00")
        r = order_risk(risky)
        self.assertEqual(r["level"], "high")
        self.assertIn("Throwaway email address", r["reasons"])
        self.client.force_login(self.boss)
        self.assertContains(self.client.get(reverse("manage_order", args=[risky.pk])), "High risk")
        self.assertContains(self.client.get(reverse("manage_orders")), f"Risk {r['score']}")

    def test_price_drop_protection(self):
        from .models import Notification, Profile, SiteSettings
        from .smart import price_protection
        buyer = User.objects.create_user("pp", "pp@example.com", "x")
        p = make_product(name="Lamp", price="50.00")
        self._order(user=buyer, email="pp@example.com", product=p, qty=2, price="50.00", payment_status="paid", payment_method="card", days_ago=3)
        self.assertEqual(price_protection(), 0)  # off by default
        site = SiteSettings.load()
        site.price_protection_days = 14
        site.save()
        p.price = Decimal("45.00")
        p.save()
        self.assertEqual(price_protection(), 1)
        self.assertEqual(Profile.objects.get(user=buyer).store_credit, Decimal("10.00"))
        self.assertEqual(price_protection(), 0)  # paid once
        p.price = Decimal("40.00")
        p.save()
        self.assertEqual(price_protection(), 1)  # a further drop pays the rest
        self.assertEqual(Profile.objects.get(user=buyer).store_credit, Decimal("20.00"))
        self.assertTrue(Notification.objects.filter(user=buyer, message__startswith="Price drop protection").exists())
        self.assertContains(self.client.get(reverse("product_detail", args=[p.pk])), "You get the difference back")

    def test_segments_campaign_audience_and_copilot(self):
        from .campaigns import audience_emails
        from .models import NewsletterSubscriber
        from .smart import customers
        p = make_product(name="Tea", price="30.00")
        for d in (1, 10, 20):
            self._order(email="vip@example.com", product=p, payment_status="paid", payment_method="card", days_ago=d, price="90.00")
        for d in (120, 200):
            self._order(email="gone@example.com", product=p, payment_status="paid", payment_method="card", days_ago=d)
        self._order(email="fresh@example.com", product=p, payment_status="paid", payment_method="card", days_ago=2)
        people = customers()
        self.assertEqual(people["vip@example.com"]["segment"], "vip")
        self.assertEqual(people["gone@example.com"]["segment"], "at_risk")
        self.assertEqual(people["fresh@example.com"]["segment"], "new")
        NewsletterSubscriber.objects.create(email="gone@example.com")
        self.assertEqual(audience_emails("at_risk"), ["gone@example.com"])
        self.client.force_login(self.boss)
        page = self.client.get(reverse("manage_segments"))
        self.assertContains(page, "At risk")
        self.assertContains(page, "Email them")
        page = self.client.get(reverse("manage_copilot"))
        self.assertContains(page, "Today's insights")
        self.assertContains(page, "haven&#x27;t ordered for 90+ days")
        with override_settings(AI_API_KEY=""):
            r = self.client.post(reverse("manage_copilot"), {"question": "What should I do?"})
        self.assertContains(r, "Quick insights")
        with mock.patch("bees.ai._call_claude", return_value="Promote Tea to your VIPs.") as call, \
                mock.patch("bees.multistore.store_settings", mock.Mock(AI_API_KEY="k", STORE_CURRENCY="usd")):
            r = self.client.post(reverse("manage_copilot"), {"question": "What should I promote?"})
        self.assertContains(r, "Promote Tea to your VIPs.")
        self.assertNotIn("vip@example.com", call.call_args[0][0])  # no personal data goes to the AI


JC_ENV = {"JAZZCASH_MERCHANT_ID": "MC12345", "JAZZCASH_PASSWORD": "pw0rd123", "JAZZCASH_INTEGRITY_SALT": "s4ltSalt", "JAZZCASH_MODE": "sandbox"}


class PakistanPaymentsTests(TestCase):
    def setUp(self):
        from .models import Currency
        Currency.objects.update_or_create(code="PKR", defaults={"symbol": "Rs", "rate": Decimal("280"), "active": True})
        self.user = User.objects.create_user("ali", "ali@example.com", "pass12345")
        self.product = make_product(price=Decimal("10.00"), stock=5)
        self.client.force_login(self.user)

    def _checkout(self, method):
        self.client.post(reverse("add_to_cart", args=[self.product.id]), {"quantity": 2})
        return self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": method})

    def test_jazzcash_full_flow_and_tampering(self):
        import os
        from . import pk_payments
        with mock.patch.dict(os.environ, JC_ENV):
            page = self._checkout("jazzcash")
            self.assertContains(page, "Taking you to JazzCash")
            order = Order.objects.get(user=self.user)
            self.assertEqual(order.payment_status, "pending")
            self.assertEqual(order.charge_currency, "pkr")
            self.assertEqual(order.charge_amount, Decimal("5600.00"))
            fields = page.context["fields"]
            self.assertEqual(fields["pp_Amount"], "560000")
            self.assertEqual(fields["pp_SecureHash"], pk_payments.secure_hash(fields, "s4ltSalt"))
            reply = {"pp_TxnRefNo": order.gateway_ref, "pp_Amount": "560000", "pp_ResponseCode": "000", "pp_ResponseMessage": "OK",
                     "pp_MerchantID": "MC12345", "pp_TxnCurrency": "PKR"}
            bad = {**reply, "pp_Amount": "100"}
            bad["pp_SecureHash"] = pk_payments.secure_hash(bad, "s4ltSalt")
            r = self.client.post(reverse("jazzcash_return"), bad)
            self.assertEqual(Order.objects.get(pk=order.pk).payment_status, "pending")  # amount mismatch never confirms
            forged = {**reply, "pp_SecureHash": "ABC"}
            self.assertEqual(self.client.post(reverse("jazzcash_return"), forged).status_code, 400)
            reply["pp_SecureHash"] = pk_payments.secure_hash(reply, "s4ltSalt")
            with self.captureOnCommitCallbacks(execute=True):
                r = self.client.post(reverse("jazzcash_return"), reply)
            self.assertEqual(r.status_code, 302)
            r = self.client.get(r["Location"])
            self.assertRedirects(r, reverse("order_success", args=[order.id]))
            order.refresh_from_db()
            self.assertEqual((order.payment_status, order.status), ("paid", "confirmed"))
            self.assertTrue(any("received" in m.subject or "confirmed" in m.subject for m in mail.outbox))

    def test_jazzcash_failure_releases_and_restores_cart(self):
        import os
        from . import pk_payments
        with mock.patch.dict(os.environ, JC_ENV):
            self._checkout("jazzcash")
            order = Order.objects.get(user=self.user)
            reply = {"pp_TxnRefNo": order.gateway_ref, "pp_Amount": "560000", "pp_ResponseCode": "199"}
            reply["pp_SecureHash"] = pk_payments.secure_hash(reply, "s4ltSalt")
            done = self.client.post(reverse("jazzcash_return"), reply)["Location"]
            self.assertRedirects(self.client.get(done), reverse("cart"))
            self.client.get(done)  # reloading doesn't add the items twice
        order.refresh_from_db()
        self.assertEqual(order.status, "cancelled")
        self.assertEqual(sum(self.client.session["cart"].values()), 2)

    def test_bank_transfer_flow(self):
        from .models import SiteSettings
        site = SiteSettings.load()
        site.transfer_enabled, site.transfer_accounts = True, "Easypaisa: 0345 1234567 (Lumen)\nMeezan (Raast): PK00MEZN0001"
        site.save()
        r = self._checkout("transfer")
        order = Order.objects.get(user=self.user)
        self.assertRedirects(r, reverse("pay_transfer", args=[order.id]))
        page = self.client.get(reverse("pay_transfer", args=[order.id]))
        self.assertContains(page, "0345 1234567")
        self.assertContains(page, "5600.00 PKR")
        self.client.post(reverse("pay_transfer", args=[order.id]), {"account": "Easypaisa", "reference": "TX99887766"})
        order.refresh_from_db()
        self.assertEqual(order.transfer_reference, "TX99887766")
        boss = User.objects.create_user("boss5", "boss5@example.com", "x", is_staff=True, is_superuser=True)
        self.client.force_login(boss)
        self.assertContains(self.client.get(reverse("manage_transfers")), "TX99887766")
        self.client.post(reverse("manage_transfers"), {"order": order.id, "action": "approve"})
        order.refresh_from_db()
        self.assertEqual((order.payment_status, order.status), ("paid", "confirmed"))
        # The same transaction ID can't pay for a second order.
        other = Order.objects.create(user=self.user, email="ali@example.com", full_name="A", address="x", city="y", country="PK",
                                     phone="+923001234567", payment_method="transfer", payment_status="pending")
        from . import pk_payments
        with self.assertRaises(ValueError):
            pk_payments.submit_transfer(other, "Easypaisa", "TX99887766")

    def test_unsent_transfer_is_released(self):
        from datetime import timedelta
        from django.utils import timezone
        from .models import SiteSettings
        from .pk_payments import release_unpaid_transfers
        site = SiteSettings.load()
        site.transfer_enabled, site.transfer_accounts = True, "Easypaisa: 0345 1234567"
        site.save()
        self._checkout("transfer")
        order = Order.objects.get(user=self.user)
        self.assertEqual(release_unpaid_transfers(), 0)
        Order.objects.filter(pk=order.pk).update(created_at=timezone.now() - timedelta(hours=49))
        self.assertEqual(release_unpaid_transfers(), 1)
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock, 5)


class DeliveryCodeTests(TestCase):
    def test_cod_needs_the_customer_code(self):
        from .models import SellerAccount
        buyer = User.objects.create_user("b1", "b1@example.com", "x")
        su = User.objects.create_user("s1", "s1@example.com", "x")
        seller = SellerAccount.objects.create(user=su, status="approved", business_name="S1")
        p = make_product(price=Decimal("10"), stock=5, seller_account=seller)
        self.client.force_login(buyer)
        self.client.post(reverse("add_to_cart", args=[p.id]), {"quantity": 1})
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod"})
        order = Order.objects.get(user=buyer)
        Order.objects.filter(pk=order.pk).update(cod_confirmed=True, status="confirmed")
        self.assertRegex(order.delivery_code, r"^\d{4}$")
        self.assertContains(self.client.get(reverse("my_orders")), order.delivery_code)
        self.client.force_login(su)
        self.client.post(reverse("seller_ship_order", args=[order.id]), {"status": "handed_to_courier"})
        r = self.client.post(reverse("seller_ship_order", args=[order.id]), {"status": "delivered"}, follow=True)
        self.assertContains(r, "delivery code")
        wrong = "0000" if order.delivery_code != "0000" else "1111"
        self.client.post(reverse("seller_ship_order", args=[order.id]), {"status": "delivered", "delivery_code": wrong})
        order.refresh_from_db()
        self.assertEqual(order.delivery_code_tries, 1)
        self.assertNotEqual(order.items.first().fulfillment_status, "delivered")
        self.client.post(reverse("seller_ship_order", args=[order.id]), {"status": "delivered", "delivery_code": order.delivery_code})
        self.assertEqual(order.items.first().fulfillment_status, "delivered")


class WhatsAppTests(TestCase):
    def _post(self, body, phone="+923001234567"):
        import base64, hashlib, hmac, os
        url = "http://testserver" + reverse("whatsapp_webhook")
        data = {"From": f"whatsapp:{phone}", "Body": body}
        sig = base64.b64encode(hmac.new(b"tok123", (url + "".join(f"{k}{v}" for k, v in sorted(data.items()))).encode(), hashlib.sha1).digest()).decode()
        with mock.patch.dict(os.environ, {"TWILIO_AUTH_TOKEN": "tok123"}):
            return self.client.post(reverse("whatsapp_webhook"), data, HTTP_X_TWILIO_SIGNATURE=sig)

    def test_bot_confirms_status_and_builds_cart_link(self):
        import os
        with mock.patch.dict(os.environ, {"TWILIO_AUTH_TOKEN": "tok123"}):
            self.assertEqual(self.client.post(reverse("whatsapp_webhook"), {"From": "whatsapp:+1", "Body": "YES"},
                                              HTTP_X_TWILIO_SIGNATURE="bad").status_code, 403)
        p = make_product(name="Mug", price=Decimal("9"), stock=4)
        order = Order.objects.create(email="w@example.com", full_name="W", address="x", city="y", country="PK", phone="0300-1234567",
                                     payment_method="cod", payment_status="not_applicable", cod_confirmed=False)
        r = self._post("yes")
        self.assertContains(r, f"Order #{order.id} is confirmed")
        order.refresh_from_db()
        self.assertTrue(order.cod_confirmed)
        self.assertNotContains(self._post("yes", phone="+923339999999"), "is confirmed")  # someone else's number
        self.assertContains(self._post("status"), f"Order #{order.id} is")
        r = self._post(f"Hi! I'd like to order:\n• Mug (#{p.pk}) x2")
        import re
        link = re.search(r"https?://\S+/whatsapp/cart/\S+", r.content.decode()).group(0)
        self.client.get(link.split("testserver")[-1])
        self.assertEqual(list(self.client.session["cart"].values()), [2])


class TrustCheckTests(TestCase):
    def test_copied_listing_reviews_and_networks(self):
        import io
        from datetime import timedelta
        from django.core.files.base import ContentFile
        from django.core.files.storage import default_storage
        from django.utils import timezone
        from PIL import Image, ImageDraw
        from .models import Review, SellerAccount, TrustFlag
        from . import trustcheck
        img = Image.new("RGB", (300, 300), "white")
        ImageDraw.Draw(img).ellipse((60, 60, 240, 240), fill=(200, 40, 40))
        buf = io.BytesIO(); img.save(buf, "JPEG")
        name = default_storage.save("products/test-copy.jpg", ContentFile(buf.getvalue()))
        url = default_storage.url(name)
        sellers = [SellerAccount.objects.create(user=User.objects.create_user(f"sx{i}", f"sx{i}@e.com", "x"), status="approved",
                                                business_name=f"Shop{i}") for i in range(2)]
        a = make_product(name="Original lamp", image_url=url, seller_account=sellers[0])
        b = make_product(name="Copy lamp", image_url=url, seller_account=sellers[1])
        result = trustcheck.run_checks()
        self.assertGreaterEqual(result["copied"], 1)
        flag = TrustFlag.objects.get(kind="copied_listing")
        self.assertEqual(flag.product_id, b.pk)
        for i in range(3):
            u = User.objects.create_user(f"fake{i}", f"fake{i}@e.com", "x")
            Review.objects.create(product=a, user=u, username=u.username, rating=5, comment=f"Amazing product number {i}")
        self.assertGreaterEqual(trustcheck.find_review_rings(), 1)
        for i in range(3):
            Order.objects.create(email=f"n{i}@e.com", full_name="N", address="1 Lane", city="Lahore", country="PK", phone="0300 7654321",
                                 payment_method="cod", status="cancelled", cod_refused=i < 2)
        self.assertGreaterEqual(trustcheck.find_order_networks(), 1)
        fresh = Order.objects.create(email="new@e.com", full_name="N", address="1 Lane", city="Lahore", country="PK", phone="+92 300 7654321",
                                     payment_method="cod")
        reasons = " ".join(r for _p, r in trustcheck.network_signals(fresh))
        self.assertIn("refused", reasons)
        boss = User.objects.create_user("tb", "tb@e.com", "x", is_staff=True, is_superuser=True)
        self.client.force_login(boss)
        self.assertContains(self.client.get(reverse("manage_trust")), "Copy lamp")
        self.client.post(reverse("manage_trust"), {"flag": flag.pk, "action": "unlist"})
        b.refresh_from_db()
        self.assertEqual(b.approval_status, "pending")


class SellerPowerToolsTests(TestCase):
    def setUp(self):
        from .models import SellerAccount
        self.su = User.objects.create_user("pw", "pw@e.com", "x")
        self.seller = SellerAccount.objects.create(user=self.su, status="approved", business_name="PW Shop")
        self.client.force_login(self.su)

    def test_photo_studio(self):
        import io
        from PIL import Image, ImageDraw
        from .models import Product
        p = make_product(name="Bottle", seller_account=self.seller)
        img = Image.new("RGB", (800, 600), (230, 228, 222))
        ImageDraw.Draw(img).rectangle((300, 150, 500, 500), fill=(20, 120, 60))
        buf = io.BytesIO(); img.save(buf, "JPEG"); buf.seek(0); buf.name = "phone.jpg"
        self.client.post(reverse("seller_photo_studio", args=[p.pk]), {"photo": buf, "background": "white", "shadow": "1"})
        page = self.client.get(reverse("seller_photo_studio", args=[p.pk]))
        self.assertContains(page, "Background removed")
        self.client.post(reverse("seller_photo_studio", args=[p.pk]), {"action": "main"})
        p = Product.objects.get(pk=p.pk)
        self.assertIn("/products/studio/", p.image_url)
        self.assertTrue(p.extra_images.filter(image_url="https://example.com/img.jpg").exists())
        # A forged result URL is never accepted.
        s = self.client.session; s["studio_product"] = p.pk; s["studio_result"] = {"url": "https://evil.example/x.jpg"}; s.save()
        self.client.post(reverse("seller_photo_studio", args=[p.pk]), {"action": "main"})
        self.assertNotIn("evil", Product.objects.get(pk=p.pk).image_url)

    def test_repricer(self):
        from .models import SellerAccount
        other = SellerAccount.objects.create(user=User.objects.create_user("o2", "o2@e.com", "x"), status="approved", business_name="O2")
        master = make_product(name="Headphones", price=Decimal("50"), seller_account=other)
        mine = make_product(name="Headphones", price=Decimal("55"), seller_account=self.seller, offer_of=master)
        self.client.post(reverse("seller_repricer"), {"product": mine.pk, "min_price": "45", "max_price": "60", "mode": "beat", "step": "0.50"})
        mine.refresh_from_db()
        self.assertEqual(mine.price, Decimal("49.50"))
        master.price = Decimal("44"); master.save()  # competitor drops below my floor
        mine.refresh_from_db()
        self.assertEqual(mine.price, Decimal("45.00"))
        master.stock = 0; master.save(update_fields=["stock"])
        from .repricer import run_all
        run_all()
        mine.refresh_from_db()
        self.assertEqual(mine.price, Decimal("60.00"))  # alone again: back to the normal price
        self.assertContains(self.client.get(reverse("seller_repricer")), "only seller")

    def test_price_test(self):
        from datetime import timedelta
        from django.utils import timezone
        from . import pricetest
        from .models import DailyStat, OrderItem, PriceTest
        p = make_product(name="Candle", price=Decimal("20"), seller_account=self.seller)
        self.client.post(reverse("seller_price_tests"), {"action": "start", "product": p.pk, "price_b": "17", "days": "7"})
        test = PriceTest.objects.get(product=p)
        pricetest.tick()
        p.refresh_from_db()
        self.assertEqual(p.price, Decimal("17.00"))
        today = timezone.localdate()
        block_b = test.blocks.get(arm="b")
        DailyStat.objects.create(date=today, kind="view", key=str(p.pk), count=100)
        for i in range(3):
            o = Order.objects.create(email=f"c{i}@e.com", full_name="C", address="x", city="y", country="US", phone="1",
                                     payment_method="card", payment_status="paid")
            OrderItem.objects.create(order=o, product=p, product_name="Candle", price=Decimal("17"), quantity=2, seller_account=self.seller)
        PriceTest.objects.filter(pk=test.pk).update(started_at=timezone.now() - timedelta(days=8))
        test.blocks.filter(arm="a").update(start=timezone.now() - timedelta(days=8), end=timezone.now() - timedelta(days=3))
        test.blocks.filter(arm="b").update(start=timezone.now() - timedelta(days=3))
        DailyStat.objects.create(date=today - timedelta(days=5), kind="view", key=str(p.pk), count=100)
        pricetest.tick()
        test.refresh_from_db()
        p.refresh_from_db()
        self.assertEqual((test.status, test.applied, p.price), ("finished", "b", Decimal("17.00")))
        self.assertContains(self.client.get(reverse("seller_price_tests")), "🏆")

    def test_reply_draft(self):
        from .models import Conversation, Message, OrderItem
        buyer = User.objects.create_user("cb", "cb@e.com", "x", first_name="Sana")
        p = make_product(name="Scarf", stock=3, seller_account=self.seller)
        o = Order.objects.create(user=buyer, email="cb@e.com", full_name="Sana", address="x", city="y", country="PK", phone="1",
                                 status="shipped", tracking_number="TCS123", courier_name="TCS", payment_method="cod")
        OrderItem.objects.create(order=o, product=p, product_name="Scarf", price=Decimal("5"), quantity=1, seller_account=self.seller)
        conv = Conversation.objects.create(buyer=buyer, seller=self.seller, product=p)
        Message.objects.create(conversation=conv, sender=buyer, body="Where is my order?")
        with override_settings(AI_API_KEY=""):
            r = self.client.post(reverse("message_draft", args=[conv.pk]))
        self.assertIn("TCS123", r.json()["draft"])
        self.assertIn("Sana", r.json()["draft"])
        self.client.force_login(buyer)
        self.assertEqual(self.client.post(reverse("message_draft", args=[conv.pk])).status_code, 404)


class MapsAndAppTests(TestCase):
    def test_checkout_pin_signup_location_and_links(self):
        user = User.objects.create_user("mp", "mp@e.com", "x")
        p = make_product(price=Decimal("10"), stock=5)
        self.client.force_login(user)
        self.client.post(reverse("add_to_cart", args=[p.id]), {"quantity": 1})
        page = self.client.get(reverse("checkout"))
        self.assertContains(page, "data-map-picker")
        self.client.post(reverse("checkout"), {**CHECKOUT_FORM, "payment_method": "cod", "latitude": "24.860734", "longitude": "67.001137"})
        order = Order.objects.get(user=user)
        self.assertEqual((order.latitude, order.longitude), (Decimal("24.860734"), Decimal("67.001137")))
        from .templatetags.bees_extras import gmaps_directions
        self.assertIn("destination=24.860734,67.001137", gmaps_directions(order))
        from .maps import parse
        self.assertEqual(parse("200", "10"), (None, None))
        self.assertEqual(parse("abc", "1"), (None, None))
        self.client.post(reverse("become_seller"), {"account_type": "individual", "business_name": "Map Shop", "phone": "+923001112223",
                                                    "country": "Pakistan", "agree": "1", "pickup_latitude": "31.52", "pickup_longitude": "74.35"})
        self.assertEqual(user.seller_account.pickup_latitude, Decimal("31.520000"))

    def test_assetlinks_and_keys_page(self):
        import os
        r = self.client.get("/.well-known/assetlinks.json")
        self.assertEqual(r.json(), [])
        fp = ":".join(["AB"] * 32)
        with mock.patch.dict(os.environ, {"ANDROID_APP_PACKAGE": "com.lumen.app", "ANDROID_APP_SHA256": fp}):
            data = self.client.get("/.well-known/assetlinks.json").json()
        self.assertEqual(data[0]["target"]["package_name"], "com.lumen.app")
        boss = User.objects.create_user("kb", "kb@e.com", "x", is_staff=True, is_superuser=True)
        self.client.force_login(boss)
        page = self.client.get(reverse("manage_keys"))
        for text in ("JazzCash (Pakistan)", "GOOGLE_MAPS_BROWSER_KEY", "Android app (Play Store)", "REMOVE_BG_API_KEY"):
            self.assertContains(page, text)
        self.assertContains(self.client.get(reverse("manage_mobile_app")), "PWABuilder")


class MailRetryAndUpdateCrashTests(TestCase):
    def _backend(self, failures):
        import smtplib
        from .mailretry import RetrySMTPBackend
        backend = RetrySMTPBackend(host="smtp.example.com", port=587, fail_silently=False)
        sent = []

        class FakeSMTP:
            def sendmail(self, frm, to, msg):
                if failures:
                    raise failures.pop(0)
                sent.append(to)

            def quit(self):
                pass

            def close(self):
                pass

        def fake_open(self=backend):
            if backend.connection is None:
                backend.connection = FakeSMTP()
                return True
            return False
        backend.open = fake_open
        return backend, sent, smtplib

    def test_dropped_connection_is_retried(self):
        import smtplib
        from django.core.mail import EmailMessage
        backend, sent, _ = self._backend([smtplib.SMTPServerDisconnected("Connection unexpectedly closed"),
                                          smtplib.SMTPResponseException(421, b"try again later")])
        with mock.patch("bees.mailretry.time.sleep") as nap:
            n = backend.send_messages([EmailMessage("Hi", "Body", "shop@example.com", ["a@example.com"])])
        self.assertEqual(n, 1)
        self.assertEqual(sent, [["a@example.com"]])
        self.assertEqual(nap.call_count, 2)

    def test_permanent_errors_are_not_retried(self):
        import smtplib
        from django.core.mail import EmailMessage
        backend, sent, _ = self._backend([smtplib.SMTPAuthenticationError(535, b"bad password")])
        with mock.patch("bees.mailretry.time.sleep") as nap, self.assertRaises(smtplib.SMTPAuthenticationError):
            backend.send_messages([EmailMessage("Hi", "Body", "shop@example.com", ["a@example.com"])])
        self.assertEqual(nap.call_count, 0)
        self.assertEqual(sent, [])

    def test_gives_up_after_retries(self):
        import smtplib
        from django.core.mail import EmailMessage
        backend, sent, _ = self._backend([smtplib.SMTPServerDisconnected("x")] * 3)
        with mock.patch("bees.mailretry.time.sleep"), self.assertRaises(smtplib.SMTPServerDisconnected):
            backend.send_messages([EmailMessage("Hi", "Body", "shop@example.com", ["a@example.com"])])

    def test_crashes_from_before_the_reload_are_ignored(self):
        from datetime import timedelta
        from django.utils import timezone
        from . import alerts
        from .alerts import SERVER_ERROR
        from .models import AuditLog
        boss = User.objects.create_user("boss9", "boss9@example.com", "pass12345", is_staff=True, is_superuser=True)
        old = AuditLog.objects.create(action=f"{SERVER_ERROR}: GET /ask/ - ImportError: cannot import name 'LiveShow' from 'bees.models'")
        AuditLog.objects.filter(pk=old.pk).update(created_at=timezone.now() - timedelta(minutes=5))
        self.client.force_login(boss)
        with mock.patch.object(alerts, "PROCESS_STARTED", timezone.now() - timedelta(minutes=1)):
            page = self.client.get(reverse("manage_system")).content.decode()
            row = page.split("Page crashes (24 hours)", 1)[1][:900]
            self.assertIn("already fixed, ignored", row)
            self.assertNotIn("Problem", row[:300])
            # a different, real crash still shows
            AuditLog.objects.create(action=f"{SERVER_ERROR}: GET /cart/ - KeyError: 'x'")
            row = self.client.get(reverse("manage_system")).content.decode().split("Page crashes (24 hours)", 1)[1][:900]
            self.assertIn("KeyError", row)
